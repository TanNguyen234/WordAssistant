import os
import sys
import re
import json

if sys.platform.startswith('win'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except AttributeError:
        pass
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=120, bottom=120, left=120, right=120):
    """Set inner padding of a table cell in DXA (1440 DXA = 1 inch)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def apply_text_formatting(run, font_name="Times New Roman", size_pt=13, bold=False, italic=False, color_rgb=None):
    """Applies font styles to a run."""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor(*color_rgb)

def set_para_formatting(paragraph, style_name='Normal', align='justify', line_spacing=1.5, space_after=6, space_before=0, first_line_indent=18):
    """Formats a paragraph to meet Vietnamese academic standards."""
    paragraph.style = style_name
    p_format = paragraph.paragraph_format
    
    # Alignment
    if align == 'justify':
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == 'center':
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'right':
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
    p_format.line_spacing = line_spacing
    p_format.space_after = Pt(space_after)
    p_format.space_before = Pt(space_before)
    
    if first_line_indent > 0 and style_name == 'Normal':
        p_format.first_line_indent = Pt(first_line_indent)
    else:
        p_format.first_line_indent = Pt(0)

def set_page_setup(section, top_cm=2.0, bottom_cm=2.0, left_cm=3.0, right_cm=2.0):
    """Configures page setup for A4 size and margins."""
    section.page_width = Inches(8.27)  # A4 width (~21.0cm)
    section.page_height = Inches(11.69) # A4 height (~29.7cm)
    
    # Convert cm to inches (1 cm = 0.3937 inches)
    section.top_margin = Inches(top_cm * 0.3937)
    section.bottom_margin = Inches(bottom_cm * 0.3937)
    section.left_margin = Inches(left_cm * 0.3937)
    section.right_margin = Inches(right_cm * 0.3937)

def parse_markdown_table(lines):
    """Parses markdown table lines into headers and rows."""
    headers = []
    rows = []
    
    # Find table header
    header_line = lines[0].strip()
    if not (header_line.startswith('|') and header_line.endswith('|')):
        return None, None
        
    headers = [cell.strip() for cell in header_line.split('|')[1:-1]]
    
    # Check separator line
    separator_line = lines[1].strip()
    if not (separator_line.startswith('|') and '-' in separator_line):
        return None, None
        
    # Parse rows
    for line in lines[2:]:
        line = line.strip()
        if line.startswith('|') and line.endswith('|'):
            row_cells = [cell.strip() for cell in line.split('|')[1:-1]]
            rows.append(row_cells)
        else:
            break
            
    return headers, rows

def build_docx_from_markdown(md_path, references_json_path, output_docx_path):
    if not os.path.exists(md_path):
        print(f"Error: Markdown file '{md_path}' does not exist.")
        return False
        
    # Load references database
    ref_db = {"citation_style": "IEEE", "sources": {}}
    if os.path.exists(references_json_path):
        with open(references_json_path, 'r', encoding='utf-8') as f:
            ref_db = json.load(f)
            
    citation_style = ref_db.get("citation_style", "IEEE")
    sources = ref_db.get("sources", {})
    
    # Parse markdown lines
    with open(md_path, 'r', encoding='utf-8') as f:
        md_content = f.read()
        
    lines = md_content.split('\n')
    doc = Document()
    
    # Initialize page layout
    section = doc.sections[0]
    set_page_setup(section)
    
    # Track citation mappings (key -> numeric index or APA string)
    citations_tracker = []
    citations_keys = {} # key -> index (1-based)
    
    # Pre-scan document for citations to build the reference list order (IEEE)
    all_citations = re.findall(r'\[@(\w+)\]', md_content)
    for key in all_citations:
        if key in sources and key not in citations_keys:
            citations_tracker.append(key)
            citations_keys[key] = len(citations_tracker)

    # Markdown processing loop variables
    i = 0
    table_lines = []
    in_table = False
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # 1. Handle table block
        if stripped.startswith('|'):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            # Table ended, process table
            in_table = False
            headers, rows = parse_markdown_table(table_lines)
            if headers and rows:
                table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                table.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                # Apply header
                hdr_cells = table.rows[0].cells
                for idx, text in enumerate(headers):
                    hdr_cells[idx].text = text
                    hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
                    hdr_cells[idx].paragraphs[0].runs[0].font.name = "Times New Roman"
                    hdr_cells[idx].paragraphs[0].runs[0].font.size = Pt(12)
                    hdr_cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_cell_margins(hdr_cells[idx])
                    
                # Apply rows
                for r_idx, row in enumerate(rows):
                    row_cells = table.rows[r_idx + 1].cells
                    for c_idx, val in enumerate(row):
                        if c_idx < len(row_cells):
                            row_cells[c_idx].text = val
                            run = row_cells[c_idx].paragraphs[0].runs[0] if row_cells[c_idx].paragraphs[0].runs else row_cells[c_idx].paragraphs[0].add_run(val)
                            apply_text_formatting(run, size_pt=12)
                            row_cells[c_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                            set_cell_margins(row_cells[c_idx])
                            
                # Add spacing after table
                p = doc.add_paragraph()
                set_para_formatting(p, line_spacing=1.5, space_after=6, space_before=6)
                
            table_lines = []
            
        # 2. Handle Heading levels
        if stripped.startswith('#'):
            h_match = re.match(r'^(#+)\s+(.*)$', stripped)
            if h_match:
                level = len(h_match.group(1))
                h_text = h_match.group(2)
                
                # We limit headings to 1, 2, 3
                docx_level = min(level, 3)
                style_name = f'Heading {docx_level}'
                
                # Heading 1 starts on a new page (except first H1)
                if docx_level == 1 and len(doc.paragraphs) > 1:
                    doc.add_page_break()
                    
                p = doc.add_paragraph()
                p.style = style_name
                
                # Set formatting for Heading
                p_format = p.paragraph_format
                p_format.space_before = Pt(12)
                p_format.space_after = Pt(6)
                p_format.keep_with_next = True
                
                # Heading font details
                run = p.add_run(h_text)
                if docx_level == 1:
                    apply_text_formatting(run, size_pt=16, bold=True)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif docx_level == 2:
                    apply_text_formatting(run, size_pt=13, bold=True)
                else:
                    apply_text_formatting(run, size_pt=13, bold=True, italic=True)
            i += 1
            continue
            
        # 3. Handle images
        img_match = re.match(r'^!\[(.*)\]\((.*)\)$', stripped)
        if img_match:
            img_caption = img_match.group(1)
            img_path = img_match.group(2)
            
            # Insert image (if path is valid)
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run()
                r.add_picture(img_path, width=Inches(5.5))
            else:
                p = doc.add_paragraph(f"[Image Missing: {img_path}]")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
            # Insert caption below image
            caption_p = doc.add_paragraph()
            set_para_formatting(caption_p, align='center', first_line_indent=0, space_before=6, space_after=12)
            c_run = caption_p.add_run(f"Hình: {img_caption}")
            apply_text_formatting(c_run, size_pt=11, italic=True)
            i += 1
            continue
            
        # 4. Handle captions for tables/figures written explicitly
        caption_match = re.match(r'^\*(Hình|Bảng)\s+(\d+\.\d+):\s+(.*)\*$', stripped)
        if caption_match:
            cap_type = caption_match.group(1)
            cap_num = caption_match.group(2)
            cap_text = caption_match.group(3)
            
            p = doc.add_paragraph()
            if cap_type == 'Hình':
                set_para_formatting(p, align='center', first_line_indent=0, space_before=6, space_after=12)
            else: # Bảng
                set_para_formatting(p, align='left', first_line_indent=0, space_before=12, space_after=6)
                
            run = p.add_run(f"{cap_type} {cap_num}: {cap_text}")
            apply_text_formatting(run, size_pt=11, italic=True)
            i += 1
            continue

        # 5. Handle standard paragraph text
        if stripped:
            p = doc.add_paragraph()
            set_para_formatting(p, style_name='Normal')
            
            # Parse in-line formatting (bold, italic, citations)
            # Find all citation tags [@cite_key]
            tokens = re.split(r'(\[@\w+\]|\*\*.*?\*\*|\*.*?\*)', line)
            
            for token in tokens:
                if not token:
                    continue
                    
                # Handle Citation token
                cite_match = re.match(r'^\[@(\w+)\]$', token)
                if cite_match:
                    key = cite_match.group(1)
                    if key in sources:
                        if citation_style == "IEEE":
                            num = citations_keys.get(key, 1)
                            r = p.add_run(f"[{num}]")
                            apply_text_formatting(r, bold=True)
                        else: # APA
                            source = sources[key]
                            authors = source.get("author", "Anon").split(" and ")
                            last_name = authors[0].split(",")[-1].strip() if "," in authors[0] else authors[0].split(" ")[-1]
                            year = source.get("year", "n.d.")
                            
                            citation_text = f"({last_name}, {year})"
                            if len(authors) > 1:
                                citation_text = f"({last_name} et al., {year})"
                            r = p.add_run(citation_text)
                            apply_text_formatting(r)
                    else:
                        r = p.add_run(token)
                        apply_text_formatting(r)
                        
                # Handle Bold token
                elif token.startswith('**') and token.endswith('**'):
                    r = p.add_run(token[2:-2])
                    apply_text_formatting(r, bold=True)
                # Handle Italic token
                elif token.startswith('*') and token.endswith('*'):
                    r = p.add_run(token[1:-1])
                    apply_text_formatting(r, italic=True)
                # Plain Text
                else:
                    r = p.add_run(token)
                    apply_text_formatting(r)
                    
        i += 1

    # 6. Add Bibliography / References Section at the end
    if citations_tracker or (citation_style == "APA" and sources):
        doc.add_page_break()
        p_ref = doc.add_paragraph()
        p_ref.style = 'Heading 1'
        p_ref.paragraph_format.keep_with_next = True
        run_ref = p_ref.add_run("TÀI LIỆU THAM KHẢO")
        apply_text_formatting(run_ref, size_pt=16, bold=True)
        p_ref.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Sort sources depending on style
        sources_to_render = []
        if citation_style == "IEEE":
            # Sort by order of appearance
            for key in citations_tracker:
                sources_to_render.append((citations_keys[key], key, sources[key]))
        else: # APA
            # Sort alphabetically by author name
            all_keys = list(sources.keys())
            sorted_keys = sorted(all_keys, key=lambda k: sources[k].get("author", ""))
            for idx, key in enumerate(sorted_keys):
                sources_to_render.append((idx + 1, key, sources[key]))
                
        for num, key, src in sources_to_render:
            p = doc.add_paragraph()
            set_para_formatting(p, style_name='Normal', align='left', first_line_indent=0)
            
            # Hanging indent of 1.0cm (~0.39 inches / 28 pt)
            p.paragraph_format.left_indent = Pt(28)
            p.paragraph_format.first_line_indent = Pt(-28)
            
            author = src.get("author", "Unknown Author")
            title = src.get("title", "Untitled Work")
            journal = src.get("journal", "")
            year = src.get("year", "")
            volume = src.get("volume", "")
            number = src.get("number", "")
            pages = src.get("pages", "")
            url = src.get("url", "")
            
            if citation_style == "IEEE":
                p.add_run(f"[{num}]\t").bold = True
                p.runs[-1].font.name = "Times New Roman"
                p.runs[-1].font.size = Pt(13)
                
                # Authors
                p.add_run(f"{author}, ")
                p.runs[-1].font.name = "Times New Roman"
                p.runs[-1].font.size = Pt(13)
                
                # Title (in quotes)
                p.add_run(f'"{title}", ')
                p.runs[-1].font.name = "Times New Roman"
                p.runs[-1].font.size = Pt(13)
                
                # Journal / Conference (Italic)
                if journal:
                    p.add_run(f"{journal}, ")
                    p.runs[-1].font.name = "Times New Roman"
                    p.runs[-1].font.size = Pt(13)
                    p.runs[-1].italic = True
                    
                if volume:
                    p.add_run(f"vol. {volume}, ")
                if number:
                    p.add_run(f"no. {number}, ")
                if pages:
                    p.add_run(f"pp. {pages}, ")
                p.add_run(f"{year}.")
                p.runs[-1].font.name = "Times New Roman"
                p.runs[-1].font.size = Pt(13)
                
            else: # APA
                # Author (Year).
                p.add_run(f"{author} ({year}). ")
                p.runs[-1].font.name = "Times New Roman"
                p.runs[-1].font.size = Pt(13)
                
                # Title (Italic)
                p.add_run(f"{title}. ")
                p.runs[-1].font.name = "Times New Roman"
                p.runs[-1].font.size = Pt(13)
                p.runs[-1].italic = True
                
                # Journal / Publisher
                if journal:
                    p.add_run(f"{journal}. ")
                    p.runs[-1].font.name = "Times New Roman"
                    p.runs[-1].font.size = Pt(13)
                    
                if url:
                    p.add_run(f"Retrieved from {url}")
                    p.runs[-1].font.name = "Times New Roman"
                    p.runs[-1].font.size = Pt(13)
                    
    # Save the generated document
    doc.save(output_docx_path)
    print(f"Document successfully generated at '{output_docx_path}'")
    return True

if __name__ == '__main__':
    if len(sys.argv) < 4:
        print("Usage: python markdown_to_docx.py <md_path> <references_json_path> <output_docx_path>")
        sys.exit(1)
        
    build_docx_from_markdown(sys.argv[1], sys.argv[2], sys.argv[3])
