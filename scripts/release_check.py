from __future__ import annotations
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from docx import Document

import os

try:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    if hasattr(sys.stderr, 'reconfigure'):
        sys.stderr.reconfigure(encoding='utf-8')
except Exception:
    pass

ROOT = Path(__file__).resolve().parents[1]


def verify_offline_build_toolchain() -> None:
    missing = []
    for module_name in ("build", "setuptools", "wheel"):
        try:
            __import__(module_name)
        except ImportError:
            missing.append(module_name)
    if missing:
        raise RuntimeError(
            "Offline wheel build prerequisites are missing: "
            + ", ".join(missing)
            + ". Install the project dev extra before running release_check.py."
        )

OFFICIAL_ASSETS = {
    ROOT / 'use_cases/hcmunre_dacn_2026_2027/assets/huong_dan_trinh_bay_dacn_2026_2027.docx':
        '0f3622dbcd74009fb0867e0d3fa4a1624a38983af502e55d48868edf33f74e78',
    ROOT / 'use_cases/hcmunre_dacn_2026_2027/assets/mau_de_cuong_do_an_chuyen_nganh.docx':
        'bb291a6de459e33267170c3dd9087903a0d9f1911086bbfc935cb9f7781904e6'
}

def run(
    cmd: list[str],
    *,
    extra_env: dict[str, str] | None = None,
    cwd: Path = ROOT,
) -> None:
    env = dict(os.environ)
    env['PYTHONIOENCODING'] = 'utf-8'
    env['PYTHONUTF8'] = '1'
    if extra_env:
        env.update(extra_env)
    try:
        print('+', ' '.join(str(x) for x in cmd))
    except Exception:
        print('+', str(cmd[0]))
    subprocess.run(cmd, cwd=cwd, check=True, env=env)

def run_expect_failure(
    cmd: list[str],
    *,
    expected_issue_codes: set[str] | None = None,
    expected_values: dict[str, object | set[object]] | None = None,
) -> None:
    env = dict(os.environ)
    env['PYTHONIOENCODING'] = 'utf-8'
    env['PYTHONUTF8'] = '1'
    try:
        print('+ [expect-fail]', ' '.join(str(x) for x in cmd))
    except Exception:
        print('+ [expect-fail]', str(cmd[0]))
    res = subprocess.run(
        cmd,
        cwd=ROOT,
        env=env,
        capture_output=True,
        text=True,
        encoding='utf-8',
        errors='replace',
    )
    if res.stdout:
        print(res.stdout, end='')
    if res.stderr:
        print(res.stderr, end='', file=sys.stderr)
    assert res.returncode != 0, f"Expected non-zero exit code, got {res.returncode}"
    try:
        payload = json.loads(res.stdout)
    except (TypeError, json.JSONDecodeError) as exc:
        raise AssertionError(
            "Expected a machine-readable JSON failure payload, not a harness/import/parser crash"
        ) from exc
    if expected_issue_codes:
        actual_codes = {
            str(issue.get('code'))
            for issue in payload.get('issues', [])
            if isinstance(issue, dict)
        }
        missing = expected_issue_codes - actual_codes
        assert not missing, (
            f"Expected failure issue code(s) {sorted(missing)}, observed {sorted(actual_codes)}"
        )
    for key, expected in (expected_values or {}).items():
        actual = payload.get(key)
        if isinstance(expected, set):
            assert actual in expected, f"Expected {key} in {expected!r}, got {actual!r}"
        else:
            assert actual == expected, f"Expected {key}={expected!r}, got {actual!r}"

def verify_source_hashes() -> None:
    for path, expected_hash in OFFICIAL_ASSETS.items():
        if not path.exists():
            raise FileNotFoundError(f"Official asset missing: {path}")
        actual_hash = hashlib.sha256(path.read_bytes()).hexdigest().lower()
        if actual_hash != expected_hash.lower():
            raise ValueError(f"Asset byte mismatch for {path.name}: expected {expected_hash}, got {actual_hash}")
    print("source-asset-hashes: VERIFIED")

def main() -> None:
    verify_offline_build_toolchain()
    print("=== 1. Checking Package Data Synchronization ===")
    run([sys.executable, 'scripts/sync_package_data.py', '--check'])

    print("\n=== 2. Verifying Official Source Asset Hashes ===")
    verify_source_hashes()

    print("\n=== 3. Running Pytest Suite ===")
    run([sys.executable, '-m', 'pytest', '-v'])

    print("\n=== 4. Verifying Python Compilation ===")
    run([sys.executable, '-m', 'compileall', '-q', 'src', 'tests', 'scripts'])

    print("\n=== 4b. Verifying generated feature matrix and public command help ===")
    run([sys.executable, 'scripts/generate_feature_matrix.py', '--check'])
    command_families = json.loads((ROOT / 'wordassistant-agent.json').read_text(encoding='utf-8'))['commands']
    for command in command_families:
        run([sys.executable, '-m', 'wordassistant.cli', command, '--help'])

    print("\n=== 5. Running Document Indexing & Context Hard Budget Check ===")
    context_cmd = [
        sys.executable, '-c',
        "from wordassistant.document.structure import index_docx; "
        "from wordassistant.context.retrieve import build_context; "
        "from wordassistant.document.model import DocumentSection; "
        "from pathlib import Path; "
        "r = Path('use_cases/hcmunre_dacn_2026_2027/assets'); "
        "a = index_docx(r / 'huong_dan_trinh_bay_dacn_2026_2027.docx'); "
        "b = index_docx(r / 'mau_de_cuong_do_an_chuyen_nganh.docx'); "
        "assert 10 <= len(a) <= 45 and 8 <= len(b) <= 30, f'Unexpected section counts: {len(a)}, {len(b)}'; "
        "ctx = build_context(a, 'trích dẫn nguồn tài liệu gốc IEEE', max_tokens=1500); "
        "assert ctx.estimated_tokens <= 1500, f'Hard budget exceeded: {ctx.estimated_tokens} > 1500'; "
        "assert any('trích dẫn' in s.title.lower() or 'tài liệu' in s.title.lower() or 'ieee' in s.text.lower() for s in ctx.sections), 'Context missing relevant citation guidance'; "
        "stress_secs = [DocumentSection(f'sec_{i}', f'Section {i} Topic Title', f'Text content of paragraph {i} in stress document.', 1, None, i) for i in range(200)]; "
        "stress_ctx = build_context(stress_secs, 'Topic Title 150', max_tokens=2000); "
        "assert stress_ctx.estimated_tokens <= 2000, 'Stress token budget exceeded'; "
        "assert any(s.section_id == 'sec_150' for s in stress_ctx.sections), 'Stress query missed exact target'; "
        "print({'guidance_sections': len(a), 'proposal_sections': len(b), 'estimated_tokens': ctx.estimated_tokens, 'stress_status': 'PASS'})"
    ]
    run(context_cmd)

    print("\n=== 6. Running CLI End-to-End Workflow Smoke & Verification Gates ===")
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        report_docx = tmp / 'test_report.docx'
        proposal_docx = tmp / 'test_proposal.docx'
        generic_docx = tmp / 'test_generic.docx'
        neg_report_docx = tmp / 'test_neg_report.docx'
        evidence_file = tmp / 'evidence.json'
        
        # doctor
        run([sys.executable, '-m', 'wordassistant.cli', 'doctor', '--json'])
        # capabilities & advanced
        run([sys.executable, '-m', 'wordassistant.cli', 'capabilities', '--json'])
        run([sys.executable, '-m', 'wordassistant.cli', 'advanced', 'status', '--json'])
        # profiles
        run([sys.executable, '-m', 'wordassistant.cli', 'profiles', '--json'])
        run([sys.executable, '-m', 'wordassistant.cli', 'profile', 'validate', 'generic', '--json'])
        run([sys.executable, '-m', 'wordassistant.cli', 'profile', 'validate', 'hcmunre_dacn_2026_2027', '--json'])
        
        # A. Draft scaffold smoke & negative verify proof
        run([
            sys.executable, '-m', 'wordassistant.cli', 'scaffold', str(report_docx),
            '--profile', 'hcmunre_dacn_2026_2027', '--variant', 'report',
            '--title', 'XAY DUNG HE THONG RAG',
            '--student', 'NGUYEN VAN A - 2351010001',
            '--supervisor', 'TS. NGUYEN VAN B',
            '--json'
        ])
        # Draft validate on fresh scaffold must succeed
        run([
            sys.executable, '-m', 'wordassistant.cli', 'validate', str(report_docx),
            '--profile', 'hcmunre_dacn_2026_2027', '--variant', 'report', '--json'
        ])
        # Unedited scaffold MUST FAIL final verify because placeholders remain
        run_expect_failure([
            sys.executable, '-m', 'wordassistant.cli', 'verify', str(report_docx),
            '--profile', 'hcmunre_dacn_2026_2027', '--variant', 'report', '--json'
        ], expected_issue_codes={'draft_placeholder_detected'})
        print("draft-scaffold-negative-verify: OK (rejected as expected)")

        # B. Completed structural fixture. Dynamic fields still require a
        # real Word refresh and must not be presented as final-delivery PASS.
        d = Document(report_docx)
        for p in d.paragraphs:
            if "[Front matter content]" in p.text:
                p.text = "Lời cảm ơn chi tiết và đầy đủ gửi đến thầy cô và gia đình."
            elif "[Section content...]" in p.text:
                p.text = "Nội dung nghiên cứu chi tiết của đồ án giải quyết trọn vẹn bài toán đặt ra."
            elif "[List of figures]" in p.text:
                p.text = "Hình 1.1: Sơ đồ kiến trúc\nHình 1.2: Lưu đồ xử lý"
            elif "[List of tables]" in p.text:
                p.text = "Bảng 1.1: Danh sách tham số\nBảng 1.2: Kết quả đo đạc"
            elif "[Document content]" in p.text:
                p.text = "Nội dung báo cáo đã hoàn thiện đầy đủ."
                
        # Add 11 references under bibliography
        for idx in range(1, 12):
            d.add_paragraph(f"[{idx}] Tác giả {idx}, Tên tài liệu {idx}, Nhà xuất bản {idx}, 202{idx % 5}.")
        d.save(report_docx)
        
        run_expect_failure([
            sys.executable, '-m', 'wordassistant.cli', 'verify', str(report_docx),
            '--profile', 'hcmunre_dacn_2026_2027', '--variant', 'report', '--json'
        ], expected_values={'field_freshness_status': {'FAIL', 'STALE_OR_UNKNOWN'}})
        print("completed-synthetic-report-verify: INCOMPLETE_AS_EXPECTED (native field refresh required)")

        # C. Negative final-delivery fixture (one placeholder left)
        shutil.copyfile(report_docx, neg_report_docx)
        d_neg = Document(neg_report_docx)
        d_neg.add_paragraph("[Section content...]")
        d_neg.save(neg_report_docx)
        run_expect_failure([
            sys.executable, '-m', 'wordassistant.cli', 'verify', str(neg_report_docx),
            '--profile', 'hcmunre_dacn_2026_2027', '--variant', 'report', '--json'
        ], expected_issue_codes={'draft_placeholder_detected'})
        print("negative-delivery-fixture-verify: OK (rejected as expected)")

        # D. Same-surname supervisor mismatch negative check
        sup_mismatch_docx = tmp / 'sup_mismatch.docx'
        shutil.copyfile(report_docx, sup_mismatch_docx)
        d_sup = Document(sup_mismatch_docx)
        d_sup.sections[-1].header.paragraphs[0].text = "XAY DUNG HE THONG RAG CHO VAN BAN PHAP LUAT\tTS. NGUYEN VAN C"
        d_sup.save(sup_mismatch_docx)
        run_expect_failure([
            sys.executable, '-m', 'wordassistant.cli', 'verify', str(sup_mismatch_docx),
            '--profile', 'hcmunre_dacn_2026_2027', '--variant', 'report', '--json'
        ], expected_issue_codes={'header_supervisor_mismatch'})
        print("same-surname-supervisor-negative-verify: OK (rejected as expected)")

        # E. Same-surname student mismatch negative check
        stu_mismatch_docx = tmp / 'stu_mismatch.docx'
        shutil.copyfile(report_docx, stu_mismatch_docx)
        d_stu = Document(stu_mismatch_docx)
        for p in d_stu.sections[-1].footer.paragraphs:
            if "NGUYEN VAN A" in p.text:
                p.text = "NGUYEN VAN B - 2351010002"
        d_stu.save(stu_mismatch_docx)
        run_expect_failure([
            sys.executable, '-m', 'wordassistant.cli', 'verify', str(stu_mismatch_docx),
            '--profile', 'hcmunre_dacn_2026_2027', '--variant', 'report', '--json'
        ], expected_issue_codes={'footer_student_mismatch'})
        print("same-surname-student-negative-verify: OK (rejected as expected)")

        # F. Missing inner cover negative check
        missing_inner_docx = tmp / 'missing_inner.docx'
        shutil.copyfile(report_docx, missing_inner_docx)
        d_inner = Document(missing_inner_docx)
        found_first = False
        for p in d_inner.paragraphs:
            if "BÁO CÁO ĐỒ ÁN" in p.text.upper() or "CÔNG NGHỆ THÔNG TIN" in p.text.upper():
                if not found_first:
                    found_first = True
                else:
                    p.text = ""
        d_inner.save(missing_inner_docx)
        run_expect_failure([
            sys.executable, '-m', 'wordassistant.cli', 'verify', str(missing_inner_docx),
            '--profile', 'hcmunre_dacn_2026_2027', '--variant', 'report', '--json'
        ], expected_issue_codes={'missing_inner_cover'})
        print("missing-inner-cover-negative-verify: OK (rejected as expected)")

        # scaffold proposal clone & generic
        run([
            sys.executable, '-m', 'wordassistant.cli', 'scaffold', str(proposal_docx),
            '--profile', 'hcmunre_dacn_2026_2027', '--variant', 'proposal', '--json'
        ])
        run([
            sys.executable, '-m', 'wordassistant.cli', 'scaffold', str(generic_docx),
            '--profile', 'generic', '--variant', 'default', '--title', 'GENERIC DOC', '--json'
        ])
        
        # replace dry-run
        run([
            sys.executable, '-m', 'wordassistant.cli', 'replace', str(report_docx),
            '--old', 'XAY DUNG HE THONG', '--new', 'PHAT TRIEN HE THONG',
            '--output', str(tmp / 'replace-preview.docx'), '--dry-run', '--json'
        ])
        # evidence
        run([
            sys.executable, '-m', 'wordassistant.cli', 'evidence', 'add', str(evidence_file),
            '--id', 'smoke1', '--title', 'Smoke Title', '--authors', 'Author A', '--year', '2024',
            '--url', 'https://example.com/smoke', '--verified-by', 'CI',
            '--verified-real-source',
            '--json'
        ])
        run([sys.executable, '-m', 'wordassistant.cli', 'evidence', 'validate', str(evidence_file), '--json'])

    print("\n=== 7. Verifying LibreOffice Render Availability & Direct Render CLI ===")
    from wordassistant.doctor import find_libreoffice_executable
    lo_exe = find_libreoffice_executable()
    if lo_exe:
        print(f"libreoffice-render: AVAILABLE ({lo_exe})")
        with tempfile.TemporaryDirectory() as render_tmpdir:
            r_docx = Path(render_tmpdir) / 'render_test.docx'
            r_pdf_dir = Path(render_tmpdir) / 'pdf_out'
            run([
                sys.executable, '-m', 'wordassistant.cli', 'scaffold', str(r_docx),
                '--profile', 'generic', '--variant', 'default', '--title', 'RENDER TEST', '--json'
            ])
            run([
                sys.executable, '-m', 'wordassistant.cli', 'render', str(r_docx),
                '--outdir', str(r_pdf_dir), '--timeout', '180', '--json'
            ])
            expected_pdf = r_pdf_dir / f"{r_docx.stem}.pdf"
            assert expected_pdf.exists() and expected_pdf.stat().st_size > 0, "Direct render CLI did not produce expected non-empty PDF"
            print(f"direct-render-cli: OK ({expected_pdf.name}, {expected_pdf.stat().st_size} bytes)")
    else:
        print("libreoffice-render: SKIPPED_WITH_REASON (LibreOffice binary not on PATH/default locations)")

    print("\n=== 8. Offline Wheel Build and Isolated Install Smoke Test ===")
    dist_dir = ROOT / 'dist'
    if dist_dir.exists():
        shutil.rmtree(dist_dir)
    # Release verification must not depend on live package-index access. Build
    # against the already-provisioned, version-checked toolchain.
    run([sys.executable, '-m', 'build', '--wheel', '--no-isolation'])
    wheel_files = list(dist_dir.glob('*.whl'))
    assert len(wheel_files) == 1, f"Expected 1 wheel file in dist/, found {len(wheel_files)}"
    wheel_file = wheel_files[0]
    print(f"wheel-built: {wheel_file.name}")

    with tempfile.TemporaryDirectory() as clean_env_dir:
        venv_path = Path(clean_env_dir) / 'test_env'
        # Inherit only third-party runtime dependencies from the provisioned
        # interpreter; the WordAssistant wheel itself is installed explicitly.
        # This still catches missing package data, entry points, and import
        # errors without weakening TLS or downloading during a release gate.
        run([sys.executable, '-m', 'venv', '--system-site-packages', str(venv_path)])
        pip_exe = str(venv_path / ('Scripts/pip.exe' if sys.platform == 'win32' else 'bin/pip'))
        py_exe = str(venv_path / ('Scripts/python.exe' if sys.platform == 'win32' else 'bin/python'))
        wa_exe = str(venv_path / ('Scripts/wordassistant.exe' if sys.platform == 'win32' else 'bin/wordassistant'))
        ag_exe = str(venv_path / ('Scripts/wordassistant-antigravity.exe' if sys.platform == 'win32' else 'bin/wordassistant-antigravity'))
        clean_cwd = Path(clean_env_dir)
        clean_execution_env = {'PYTHONPATH': '', 'PYTHONNOUSERSITE': '1'}

        def run_installed(command: list[str]) -> None:
            run(command, extra_env=clean_execution_env, cwd=clean_cwd)

        run([pip_exe, 'install', '--quiet', '--no-deps', str(wheel_file)])

        run_installed([
            py_exe,
            '-c',
            "from pathlib import Path; import wordassistant; "
            f"venv=Path({str(venv_path)!r}).resolve(); "
            "origin=Path(wordassistant.__file__).resolve(); "
            "assert origin.is_relative_to(venv), origin; "
            "print('installed-wheel-origin: VERIFIED', origin)",
        ])

        # Test CLI invocation via python -m and entrypoint executable
        run_installed([py_exe, '-m', 'wordassistant.cli', 'doctor', '--json'])
        run_installed([py_exe, '-m', 'wordassistant.cli', 'profiles', '--json'])
        run_installed([py_exe, '-m', 'wordassistant.cli', 'profile', 'validate', 'generic', '--json'])
        installed_commands = json.loads(
            (ROOT / 'wordassistant-agent.json').read_text(encoding='utf-8')
        )['commands']
        for command in installed_commands:
            run_installed([py_exe, '-m', 'wordassistant.cli', command, '--help'])
        
        clean_gen_docx = Path(clean_env_dir) / 'clean_generic.docx'
        clean_prop_docx = Path(clean_env_dir) / 'clean_proposal.docx'
        run_installed([py_exe, '-m', 'wordassistant.cli', 'scaffold', str(clean_gen_docx), '--profile', 'generic', '--variant', 'default', '--json'])
        run_installed([py_exe, '-m', 'wordassistant.cli', 'inspect', str(clean_gen_docx), '--json'])
        run_installed([py_exe, '-m', 'wordassistant.cli', 'scaffold', str(clean_prop_docx), '--profile', 'hcmunre_dacn_2026_2027', '--variant', 'proposal', '--json'])
        
        # Verify packaged official docx hashes inside installed package
        verify_pkg_code = (
            "import hashlib, json; from importlib.resources import files; "
            "hd = files('wordassistant').joinpath('data', 'use_cases', 'hcmunre_dacn_2026_2027', 'assets', 'huong_dan_trinh_bay_dacn_2026_2027.docx').read_bytes(); "
            "dc = files('wordassistant').joinpath('data', 'use_cases', 'hcmunre_dacn_2026_2027', 'assets', 'mau_de_cuong_do_an_chuyen_nganh.docx').read_bytes(); "
            "assert hashlib.sha256(hd).hexdigest().lower() == '0f3622dbcd74009fb0867e0d3fa4a1624a38983af502e55d48868edf33f74e78'; "
            "assert hashlib.sha256(dc).hexdigest().lower() == 'bb291a6de459e33267170c3dd9087903a0d9f1911086bbfc935cb9f7781904e6'; "
            "pkg = files('wordassistant').joinpath('data', 'agent'); "
            f"root = __import__('pathlib').Path({str(ROOT)!r}); "
            "assert pkg.joinpath('wordassistant-agent.json').read_bytes() == root.joinpath('wordassistant-agent.json').read_bytes(); "
            "assert pkg.joinpath('word_feature_matrix.json').read_bytes() == root.joinpath('docs', 'word_feature_matrix.json').read_bytes(); "
            "assert pkg.joinpath('antigravity-workflow.json').read_bytes() == root.joinpath('antigravity-workflow.json').read_bytes(); "
            "assert pkg.joinpath('run_antigravity_workflow.py').is_file(); "
            "commands = json.loads(pkg.joinpath('wordassistant-agent.json').read_text(encoding='utf-8'))['commands']; "
            "assert {'doctor','capabilities','inspect','context','create','edit','verify','render'} <= set(commands); "
            "print('installed-package-asset-hashes: VERIFIED')"
        )
        run_installed([py_exe, '-c', verify_pkg_code])
        run_installed([
            py_exe, '-c',
            "from importlib.metadata import requires; "
            "reqs = requires('wordassistant') or []; "
            "assert any(r.lower().startswith('python-docx') for r in reqs), reqs; "
            "assert any(r.lower().startswith('pypdfium2') for r in reqs), reqs; "
            "assert any(r.lower().startswith('pillow') for r in reqs), reqs; "
            "print('installed-wheel-dependency-metadata: VERIFIED')"
        ])
        fallback_pdf = Path(clean_env_dir) / 'fallback-smoke.pdf'
        fallback_pages = Path(clean_env_dir) / 'fallback-smoke-pages'
        installed_site = subprocess.check_output(
            [
                py_exe,
                '-c',
                "import importlib.util; from pathlib import Path; "
                "spec=importlib.util.find_spec('wordassistant'); "
                "assert spec is not None and spec.origin is not None; "
                "print(Path(spec.origin).resolve().parents[1])",
            ],
            cwd=clean_cwd,
            text=True,
            env={
                **os.environ,
                **clean_execution_env,
                'PYTHONIOENCODING': 'utf-8',
                'PYTHONUTF8': '1',
            },
        ).strip()
        installed_site_path = Path(installed_site).resolve()
        assert installed_site_path.is_relative_to(venv_path.resolve()), installed_site_path
        fallback_code = (
            "import sys; from pathlib import Path; "
            f"sys.path.insert(0, {installed_site!r}); "
            "from PIL import Image; from unittest.mock import patch; import wordassistant; "
            f"installed=Path({installed_site!r}).resolve(); "
            "assert Path(wordassistant.__file__).resolve().is_relative_to(installed), wordassistant.__file__; "
            "from wordassistant.document.render import render_pdf_pages; "
            f"pdf=Path({str(fallback_pdf)!r}); out=Path({str(fallback_pages)!r}); "
            "im=Image.new('RGB',(200,300),'white'); im.save(pdf,'PDF'); im.close(); "
            "ctx=patch('wordassistant.document.render.shutil.which', return_value=None); ctx.start(); "
            "pages=render_pdf_pages(pdf,out); ctx.stop(); "
            "assert len(pages)==1 and pages[0].read_bytes().startswith(b'\\x89PNG\\r\\n\\x1a\\n'); "
            "print('installed-wheel-pdfium-page-render: VERIFIED')"
        )
        run(
            [sys.executable, '-c', fallback_code],
            extra_env=clean_execution_env,
            cwd=clean_cwd,
        )

        if Path(wa_exe).exists():
            run_installed([wa_exe, 'doctor', '--json'])
        assert Path(ag_exe).exists(), "wordassistant-antigravity entry point is missing"
        run_installed([ag_exe, '--help'])
        print('isolated-wheel-smoke: OK (installed wheel code, release dependencies)')

    print('\n========================================')
    print('release-check: PASS')
    print('========================================')

if __name__ == '__main__':
    main()
