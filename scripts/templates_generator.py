import os
import sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_page_setup(section, top_cm=2.0, bottom_cm=2.0, left_cm=3.0, right_cm=2.0):
    """Configures page setup for A4 size and margins."""
    section.page_width = Inches(8.27)  # A4 width (~21.0cm)
    section.page_height = Inches(11.69) # A4 height (~29.7cm)
    section.top_margin = Inches(top_cm * 0.3937)
    section.bottom_margin = Inches(bottom_cm * 0.3937)
    section.left_margin = Inches(left_cm * 0.3937)
    section.right_margin = Inches(right_cm * 0.3937)

def configure_styles(doc):
    """Configures default Styles to meet Vietnamese Academic Standards."""
    # Configure Normal Style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    
    p_format = style_normal.paragraph_format
    p_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_format.line_spacing = 1.5
    p_format.space_after = Pt(6)
    p_format.space_before = Pt(0)
    
    # Configure Heading 1
    style_h1 = doc.styles['Heading 1']
    h1_font = style_h1.font
    h1_font.name = 'Times New Roman'
    h1_font.size = Pt(16)
    h1_font.bold = True
    style_h1.paragraph_format.space_before = Pt(12)
    style_h1.paragraph_format.space_after = Pt(6)
    style_h1.paragraph_format.keep_with_next = True
    
    # Configure Heading 2
    style_h2 = doc.styles['Heading 2']
    h2_font = style_h2.font
    h2_font.name = 'Times New Roman'
    h2_font.size = Pt(13)
    h2_font.bold = True
    style_h2.paragraph_format.space_before = Pt(12)
    style_h2.paragraph_format.space_after = Pt(6)
    style_h2.paragraph_format.keep_with_next = True

    # Configure Heading 3
    style_h3 = doc.styles['Heading 3']
    h3_font = style_h3.font
    h3_font.name = 'Times New Roman'
    h3_font.size = Pt(13)
    h3_font.bold = True
    h3_font.italic = True
    style_h3.paragraph_format.space_before = Pt(12)
    style_h3.paragraph_format.space_after = Pt(6)
    style_h3.paragraph_format.keep_with_next = True

def create_template(name, title_desc, output_path):
    doc = Document()
    set_page_setup(doc.sections[0])
    configure_styles(doc)
    
    # --- PAGE 1: COVER PAGE ---
    # School / University headers
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_univ = p_univ.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC BÁCH KHOA\nKHOA CÔNG NGHỆ THÔNG TIN")
    r_univ.font.name = "Times New Roman"
    r_univ.font.size = Pt(13)
    r_univ.bold = True
    p_univ.paragraph_format.space_before = Pt(0)
    p_univ.paragraph_format.space_after = Pt(72) # Large space before title
    
    # Document Type (e.g. ĐỒ ÁN TỐT NGHIỆP)
    p_type = doc.add_paragraph()
    p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_type = p_type.add_run(name.upper())
    r_type.font.name = "Times New Roman"
    r_type.font.size = Pt(16)
    r_type.bold = True
    p_type.paragraph_format.space_after = Pt(36)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"ĐỀ TÀI: [TÊN ĐỀ TÀI {title_desc.upper()}]")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(20)
    r_title.bold = True
    p_title.paragraph_format.space_after = Pt(96)
    
    # Student / Supervisor Info
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Inches(2.5) # Asymmetric layout
    r_info = p_info.add_run(
        "Sinh viên thực hiện: [Họ và tên SV]\n"
        "Mã số sinh viên:      [MSSV]\n"
        "Lớp:                         [Lớp học]\n"
        "Giảng viên hướng dẫn: [Họ và tên GV]"
    )
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(13)
    p_info.paragraph_format.space_after = Pt(72)
    
    # Location and Date
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_date = p_date.add_run("Hà Nội, Năm 2026")
    r_date.font.name = "Times New Roman"
    r_date.font.size = Pt(13)
    r_date.bold = True
    
    # --- PAGE 2: LỜI CAM ĐOAN / CẢM ƠN ---
    doc.add_page_break()
    p_camdoan = doc.add_paragraph()
    p_camdoan.style = 'Heading 1'
    run_cd = p_camdoan.add_run("LỜI CẢM ƠN")
    run_cd.font.size = Pt(16)
    run_cd.bold = True
    p_camdoan.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p_cd_content = doc.add_paragraph()
    r_cd_content = p_cd_content.add_run("Em xin gửi lời cảm ơn chân thành tới các thầy cô giáo trong Khoa Công nghệ thông tin...")
    r_cd_content.font.name = "Times New Roman"
    r_cd_content.font.size = Pt(13)
    p_cd_content.paragraph_format.first_line_indent = Pt(36)
    
    # --- PAGE 3: MỤC LỤC ---
    doc.add_page_break()
    p_toc = doc.add_paragraph()
    p_toc.style = 'Heading 1'
    run_toc = p_toc.add_run("MỤC LỤC")
    run_toc.font.size = Pt(16)
    run_toc.bold = True
    p_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p_toc_desc = doc.add_paragraph()
    r_toc_desc = p_toc_desc.add_run("[Mục lục tự động sinh ở đây]")
    r_toc_desc.font.name = "Times New Roman"
    r_toc_desc.font.size = Pt(13)
    p_toc_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # --- PAGE 4: MỞ ĐẦU ---
    doc.add_page_break()
    p_modau = doc.add_paragraph()
    p_modau.style = 'Heading 1'
    run_md = p_modau.add_run("MỞ ĐẦU")
    run_md.font.size = Pt(16)
    run_md.bold = True
    p_modau.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    p_md_content = doc.add_paragraph()
    r_md_content = p_md_content.add_run("Ngày nay, sự phát triển vượt bậc của công nghệ thông tin đã mang lại nhiều thay đổi...")
    r_md_content.font.name = "Times New Roman"
    r_md_content.font.size = Pt(13)
    p_md_content.paragraph_format.first_line_indent = Pt(36)
    
    # Save document
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Created template: {output_path}")

if __name__ == '__main__':
    create_template("Tiểu Luận", "tiểu luận", "../templates/tieu_luan_template.docx")
    create_template("Đồ Án Chuyên Ngành", "đồ án chuyên ngành", "../templates/do_an_chuyen_nganh_template.docx")
    create_template("Đồ Án Tốt Nghiệp", "đồ án tốt nghiệp", "../templates/do_an_tot_nghiep_template.docx")
