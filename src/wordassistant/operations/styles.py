from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor

from ..document.fonts import set_word_font
from ..selectors import Selector, resolve_paragraphs
from .common import atomic_save
from .models import Fidelity, OperationResult


def _result(operation, output, data, dry_run=False):
    return OperationResult(
        status="DRY_RUN" if dry_run else "SUCCESS",
        operation=operation,
        backend_used="python-docx",
        fidelity=Fidelity.HIGH,
        output_path=None if dry_run else Path(output).resolve(),
        data=data,
    )


def create_style(
    input_path: str | Path,
    *,
    output_path: str | Path,
    name: str,
    base_style: str = "Normal",
    dry_run: bool = False,
    **properties: Any,
) -> OperationResult:
    document = Document(input_path)
    if name in document.styles:
        raise ValueError(f"Style already exists: {name}")
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles[base_style]
    if "font_family" in properties:
        set_word_font(style.font, properties["font_family"])
    if "font_size_pt" in properties:
        style.font.size = Pt(float(properties["font_size_pt"]))
    for key in ("bold", "italic"):
        if key in properties:
            setattr(style.font, key, bool(properties[key]))
    if not dry_run:
        atomic_save(document, output_path)
    return _result("style.create", output_path, {"style": name}, dry_run)


def apply_style(
    input_path: str | Path,
    *,
    output_path: str | Path,
    name: str,
    selector=None,
    dry_run: bool = False,
) -> OperationResult:
    document = Document(input_path)
    if name not in document.styles:
        raise ValueError(f"Style not found: {name}")
    paragraphs = resolve_paragraphs(
        document, Selector.from_value(selector or {"type": "document"})
    )
    for paragraph in paragraphs:
        paragraph.style = name
    if not dry_run:
        atomic_save(document, output_path)
    return _result(
        "style.apply", output_path, {"style": name, "paragraphs": len(paragraphs)}, dry_run
    )


def delete_style(
    input_path: str | Path,
    *,
    output_path: str | Path,
    name: str,
    dry_run: bool = False,
) -> OperationResult:
    document = Document(input_path)
    if name not in document.styles:
        raise ValueError(f"Style not found: {name}")
    style = document.styles[name]
    used = sum(1 for paragraph in document.paragraphs if paragraph.style == style)
    if style.builtin or used:
        return OperationResult(
            status="UNSAFE_EDIT_BLOCKED",
            operation="style.delete",
            backend_used="python-docx",
            fidelity=Fidelity.HIGH,
            warnings=[
                f"Style {name!r} is {'built-in' if style.builtin else f'used by {used} paragraph(s)'}"
            ],
            manual_review_required=True,
            data={"style": name, "references": used},
        )
    style.delete()
    if not dry_run:
        atomic_save(document, output_path)
    return _result("style.delete", output_path, {"style": name}, dry_run)


def update_style(input_path, *, output_path, name, dry_run=False, **properties):
    document = Document(input_path)
    if name not in document.styles:
        raise ValueError(f"Style not found: {name}")
    style = document.styles[name]
    if "font_family" in properties:
        set_word_font(style.font, properties["font_family"])
    if "font_size_pt" in properties:
        style.font.size = Pt(float(properties["font_size_pt"]))
    for key in ("bold", "italic"):
        if key in properties:
            setattr(style.font, key, bool(properties[key]))
    if "font_color" in properties:
        value = str(properties["font_color"]).strip().lstrip("#")
        if not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
            raise ValueError("font_color must be a six-digit RGB value")
        style.font.color.rgb = RGBColor.from_string(value.upper())
    if "left_indent_pt" in properties:
        style.paragraph_format.left_indent = Pt(float(properties["left_indent_pt"]))
    if "first_line_indent_pt" in properties:
        style.paragraph_format.first_line_indent = Pt(float(properties["first_line_indent_pt"]))
    if not dry_run:
        atomic_save(document, output_path)
    return _result("style.update", output_path, {"style": name}, dry_run)


def clone_style(input_path, *, output_path, source_name, name, dry_run=False):
    document = Document(input_path)
    if source_name not in document.styles:
        raise ValueError(f"Style not found: {source_name}")
    if name in document.styles:
        raise ValueError(f"Style already exists: {name}")
    source = document.styles[source_name]
    clone = document.styles.add_style(name, source.type)
    clone.base_style = source
    if not dry_run:
        atomic_save(document, output_path)
    return _result(
        "style.clone", output_path, {"source": source_name, "style": name}, dry_run
    )
