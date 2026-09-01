from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .common import atomic_save
from .models import Fidelity, OperationResult


_FORMATS = {"decimal", "lowerRoman", "upperRoman", "lowerLetter", "upperLetter"}
_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def _add_field(paragraph, code: str):
    begin_run = paragraph.add_run()._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    instruction_run = paragraph.add_run()._r
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {code} "
    instruction_run.append(instruction)
    separate_run = paragraph.add_run()._r
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)
    paragraph.add_run("1")
    end_run = paragraph.add_run()._r
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)


def _remove_existing_fields(paragraph):
    """Remove complex/simple fields without deleting surrounding footer text."""
    in_field = False
    for child in list(paragraph._p):
        if child.tag == qn("w:fldSimple"):
            paragraph._p.remove(child)
            continue
        field_chars = child.xpath(".//w:fldChar/@w:fldCharType")
        starts_field = "begin" in field_chars
        ends_field = "end" in field_chars
        if in_field or starts_field or child.xpath(".//w:instrText"):
            paragraph._p.remove(child)
        if starts_field:
            in_field = True
        if ends_field:
            in_field = False


def _set_page_number_type(section, number_format, start):
    element = section._sectPr.find(qn("w:pgNumType"))
    if element is None:
        element = OxmlElement("w:pgNumType")
        section._sectPr.append(element)
    if number_format is not None:
        if number_format not in _FORMATS:
            raise ValueError(f"Unsupported page number format: {number_format}")
        element.set(qn("w:fmt"), number_format)
    if start is None:
        element.attrib.pop(qn("w:start"), None)
    else:
        if start < 0:
            raise ValueError("Page number start must be non-negative")
        element.set(qn("w:start"), str(start))


def _story(section, position, variant="primary"):
    if position not in {"header", "footer"}:
        raise ValueError("position must be header or footer")
    attribute = position if variant == "primary" else f"{variant}_page_{position}"
    return getattr(section, attribute)


def apply_page_numbers_to_section(
    document,
    section_index: int,
    *,
    position="footer",
    alignment="right",
    number_format="decimal",
    start=1,
    field="PAGE",
    page_x_of_y=False,
    unlink_previous=False,
    variant="primary",
):
    section = document.sections[section_index]
    _set_page_number_type(section, number_format, start)
    container = _story(section, position, variant)
    if unlink_previous:
        container.is_linked_to_previous = False
    paragraph = container.paragraphs[0]
    if alignment not in _ALIGNMENTS:
        raise ValueError(f"Unsupported page-number alignment: {alignment}")
    _remove_existing_fields(paragraph)
    has_surrounding_text = bool(paragraph.text.strip())
    if not has_surrounding_text:
        paragraph.clear()
        paragraph.alignment = _ALIGNMENTS[alignment]
    elif alignment == "right" and not paragraph.text.endswith("\t"):
        paragraph.add_run("\t")
    if page_x_of_y:
        paragraph.add_run("Page ")
        _add_field(paragraph, "PAGE")
        paragraph.add_run(" of ")
        _add_field(paragraph, "NUMPAGES")
    else:
        _add_field(paragraph, field)


def configure_page_numbers(
    input_path,
    *,
    output_path,
    section_index: int,
    position="footer",
    alignment="right",
    number_format="decimal",
    start=1,
    field="PAGE",
    page_x_of_y=False,
    hide_first_page=False,
    different_odd_even=None,
    unlink_previous=False,
    continue_from_previous=False,
    dry_run=False,
):
    document = Document(input_path)
    apply_page_numbers_to_section(
        document,
        section_index,
        position=position,
        alignment=alignment,
        number_format=number_format,
        start=None if continue_from_previous else start,
        field=field,
        page_x_of_y=page_x_of_y,
        unlink_previous=unlink_previous,
    )
    section = document.sections[section_index]
    section.different_first_page_header_footer = hide_first_page
    settings = document.settings.element
    even = settings.find(qn("w:evenAndOddHeaders"))
    if different_odd_even:
        if even is None:
            settings.append(OxmlElement("w:evenAndOddHeaders"))
        apply_page_numbers_to_section(
            document,
            section_index,
            position=position,
            alignment=alignment,
            number_format=number_format,
            start=None if continue_from_previous else start,
            field=field,
            page_x_of_y=page_x_of_y,
            unlink_previous=unlink_previous,
            variant="even",
        )
    elif different_odd_even is False and even is not None:
        active_even_sections = []
        for candidate_index, candidate_section in enumerate(document.sections):
            if any(
                _field_codes(_story(candidate_section, candidate_position, "even"))
                for candidate_position in ("header", "footer")
            ):
                active_even_sections.append(candidate_index)
        if len(document.sections) > 1 and active_even_sections:
            raise ValueError(
                "different_odd_even is a document-level setting and cannot be disabled "
                "through a section-scoped operation while even-page fields exist in "
                f"sections: {active_even_sections}"
            )
        settings.remove(even)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success(
        "page-number.configure",
        "ooxml",
        Fidelity.HIGH,
        {
            "section_index": section_index,
            "format": number_format,
            "start": None if continue_from_previous else start,
        },
        output_path=None if dry_run else output,
    )


def _section_headings(document):
    headings: dict[int, str] = {}
    section_index = 0
    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if section_index not in headings and re.match(r"Heading [1-9]$", style, re.IGNORECASE) and paragraph.text.strip():
            headings[section_index] = paragraph.text.strip()
        ppr = paragraph._p.pPr
        if ppr is not None and ppr.sectPr is not None:
            section_index += 1
    return headings


def _field_codes(container):
    codes = []
    for value in container._element.xpath(".//w:instrText/text()"):
        code = str(value).strip().split()[0].upper() if str(value).strip() else ""
        if code and code not in codes:
            codes.append(code)
    for value in container._element.xpath(".//w:fldSimple/@w:instr"):
        code = str(value).strip().split()[0].upper() if str(value).strip() else ""
        if code and code not in codes:
            codes.append(code)
    return codes


def inspect_page_numbers(input_path):
    document = Document(input_path)
    odd_even_enabled = (
        document.settings.element.find(qn("w:evenAndOddHeaders")) is not None
    )
    headings = _section_headings(document)
    result = []
    for index, section in enumerate(document.sections):
        pg_type = section._sectPr.find(qn("w:pgNumType"))
        number_format = pg_type.get(qn("w:fmt")) if pg_type is not None else None
        start_value = pg_type.get(qn("w:start")) if pg_type is not None else None
        header_fields = _field_codes(section.header)
        footer_fields = _field_codes(section.footer)
        if footer_fields:
            position = "footer"
            fields = footer_fields
            paragraph = section.footer.paragraphs[0]
        elif header_fields:
            position = "header"
            fields = header_fields
            paragraph = section.header.paragraphs[0]
        else:
            position = None
            fields = []
            paragraph = None
        alignment = None
        if paragraph is not None:
            alignment = {
                WD_ALIGN_PARAGRAPH.LEFT: "left",
                WD_ALIGN_PARAGRAPH.CENTER: "center",
                WD_ALIGN_PARAGRAPH.RIGHT: "right",
            }.get(paragraph.alignment)
        variants = {}
        for variant in ("primary", "even", "first"):
            variant_position = None
            variant_fields = []
            variant_paragraph = None
            for candidate_position in ("footer", "header"):
                container = _story(section, candidate_position, variant)
                candidate_fields = _field_codes(container)
                if candidate_fields:
                    variant_position = candidate_position
                    variant_fields = candidate_fields
                    variant_paragraph = container.paragraphs[0]
                    break
            variant_alignment = None
            if variant_paragraph is not None:
                variant_alignment = {
                    WD_ALIGN_PARAGRAPH.LEFT: "left",
                    WD_ALIGN_PARAGRAPH.CENTER: "center",
                    WD_ALIGN_PARAGRAPH.RIGHT: "right",
                }.get(variant_paragraph.alignment)
            variants[variant] = {
                "fields": variant_fields,
                "position": variant_position,
                "alignment": variant_alignment,
                "linked_to_previous": (
                    _story(section, variant_position, variant).is_linked_to_previous
                    if variant_position is not None
                    else None
                ),
            }
        result.append(
            {
                "section_index": index,
                "first_heading": headings.get(index),
                "format": number_format,
                "start": int(start_value) if start_value is not None else None,
                "fields": fields,
                "position": position,
                "alignment": alignment,
                "odd_even_enabled": odd_even_enabled,
                "variants": variants,
            }
        )
    return result
