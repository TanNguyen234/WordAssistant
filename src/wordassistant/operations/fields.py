from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..selectors import add_paragraph_at_boundary, resolve_single_paragraph
from .common import atomic_save
from .models import Fidelity, OperationResult


FIELD_TYPES = {
    "PAGE",
    "NUMPAGES",
    "SECTIONPAGES",
    "TOC",
    "SEQ",
    "REF",
    "PAGEREF",
    "NOTEREF",
    "DATE",
    "TIME",
    "AUTHOR",
    "TITLE",
    "SUBJECT",
    "KEYWORDS",
    "FILENAME",
    "DOCPROPERTY",
}


def normalize_instruction(field_type: str, argument: str | None = None, switches: str | None = None) -> str:
    field = field_type.upper().strip()
    if field not in FIELD_TYPES:
        raise ValueError(f"Unsupported field type: {field_type}")
    values = [field]
    if argument:
        if any(character in argument for character in "\r\n\0"):
            raise ValueError("Field argument contains control characters")
        values.append(argument)
    if switches:
        if any(character in switches for character in "\r\n\0"):
            raise ValueError("Field switches contain control characters")
        values.append(switches)
    return " ".join(values)


def append_field(paragraph, instruction: str, display_text: str = "0", locked: bool = False):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    if locked:
        field.set(qn("w:fldLock"), "true")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = display_text
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    return field


def _all_story_roots(document):
    yield "document", document.element
    seen_parts = {str(document.part.partname)}
    for index, section in enumerate(document.sections):
        for name in (
            "header",
            "footer",
            "first_page_header",
            "first_page_footer",
            "even_page_header",
            "even_page_footer",
        ):
            container = getattr(section, name)
            part_name = str(container.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            yield f"section:{index}:{name}", container._element


def _parse_complex_fields(root) -> tuple[list[dict], list[dict]]:
    """Parse complex fields as ordered begin/instruction/separate/result/end units."""
    completed: list[dict] = []
    issues: list[dict] = []
    stack: list[dict] = []
    for element in root.iter():
        if element.tag == qn("w:fldChar"):
            kind = element.get(qn("w:fldCharType"))
            if kind == "begin":
                stack.append(
                    {
                        "begin": element,
                        "instruction_parts": [],
                        "result_parts": [],
                        "separated": False,
                    }
                )
            elif kind == "separate":
                if not stack:
                    issues.append({"code": "orphan_complex_field_separator"})
                elif stack[-1]["separated"]:
                    issues.append({"code": "duplicate_complex_field_separator"})
                else:
                    stack[-1]["separated"] = True
            elif kind == "end":
                if not stack:
                    issues.append({"code": "orphan_complex_field_end"})
                    continue
                context = stack.pop()
                if not context["separated"]:
                    issues.append({"code": "complex_field_missing_separator"})
                instruction = "".join(context["instruction_parts"]).strip()
                completed.append(
                    {
                        "type": instruction.split()[0].upper() if instruction else "",
                        "instruction": " ".join(instruction.split()),
                        "locked": context["begin"].get(qn("w:fldLock"))
                        in {"1", "true", "on"},
                        "kind": "complex",
                        "cached_result": "".join(context["result_parts"]).strip(),
                    }
                )
        elif element.tag == qn("w:instrText") and stack:
            if not stack[-1]["separated"]:
                stack[-1]["instruction_parts"].append(element.text or "")
        elif element.tag == qn("w:t") and stack:
            for context in stack:
                if context["separated"]:
                    context["result_parts"].append(element.text or "")
    for _ in stack:
        issues.append({"code": "unclosed_complex_field"})
    return completed, issues


def inspect_fields(input_path) -> list[dict]:
    document = Document(input_path)
    result: list[dict] = []
    index = 0
    for story, root in _all_story_roots(document):
        for field in root.xpath(".//w:fldSimple"):
            instruction = field.get(qn("w:instr"), "").strip()
            if not instruction:
                continue
            display = "".join(
                element.text or "" for element in field.iter(qn("w:t"))
            ).strip()
            result.append(
                {
                    "index": index,
                    "story": story,
                    "type": instruction.split()[0].upper(),
                    "instruction": instruction,
                    "locked": field.get(qn("w:fldLock")) in {"1", "true", "on"},
                    "kind": "simple",
                    "cached_result": display,
                    "freshness": _cached_freshness(display),
                }
            )
            index += 1
        complex_fields, _ = _parse_complex_fields(root)
        for field in complex_fields:
            if not field["instruction"]:
                continue
            result.append(
                {
                    "index": index,
                    "story": story,
                    **field,
                    "freshness": _cached_freshness(field["cached_result"]),
                }
            )
            index += 1
    return result


def inspect_field_structure(input_path) -> dict:
    document = Document(input_path)
    issues: list[dict] = []
    bookmarks = {
        element.get(qn("w:name"))
        for element in document.element.body.xpath(".//w:bookmarkStart")
    }
    for story, root in _all_story_roots(document):
        _, structural_issues = _parse_complex_fields(root)
        issues.extend({"story": story, **issue} for issue in structural_issues)
    for field in inspect_fields(input_path):
        if field["type"] not in {"REF", "PAGEREF", "NOTEREF"}:
            continue
        parts = field["instruction"].split()
        if len(parts) < 2 or parts[1] not in bookmarks:
            issues.append(
                {
                    "story": field["story"],
                    "code": "dangling_reference",
                    "instruction": field["instruction"],
                }
            )
    return {"status": "PASS" if not issues else "FAIL", "issues": issues}


def _cached_freshness(value: str) -> str:
    normalized = value.strip().casefold()
    if not normalized or "update field" in normalized or normalized in {"reference", "error!"}:
        return "PLACEHOLDER"
    return "STALE_OR_UNKNOWN"


def _complex_field_parts(paragraph, instruction_run) -> tuple[str, object | None]:
    if paragraph is None or instruction_run is None:
        return "", None
    elements = list(paragraph.iter())
    stack = []
    begin = None
    for child in elements:
        if child is instruction_run:
            begin = stack[-1] if stack else None
            break
        if child.tag == qn("w:fldChar"):
            kind = child.get(qn("w:fldCharType"))
            if kind == "begin":
                stack.append(child)
            elif kind == "end" and stack:
                stack.pop()
    collecting = False
    values = []
    nested_depth = 0
    seen_instruction = False
    for child in elements:
        if child is instruction_run:
            seen_instruction = True
            continue
        if not seen_instruction:
            continue
        if child.tag == qn("w:fldChar"):
            kind = child.get(qn("w:fldCharType"))
            if kind == "begin":
                nested_depth += 1
            elif kind == "separate" and nested_depth == 0:
                collecting = True
            elif kind == "end":
                if nested_depth == 0:
                    break
                nested_depth -= 1
        elif collecting and child.tag == qn("w:t") and child.text:
            values.append(child.text)
    return "".join(values).strip(), begin


def _complex_cached_result(paragraph, instruction_run) -> str:
    return _complex_field_parts(paragraph, instruction_run)[0]


def insert_field(
    input_path,
    *,
    output_path,
    field_type: str,
    argument: str | None = None,
    switches: str | None = None,
    selector=None,
    display_text="0",
    locked=False,
    dry_run=False,
):
    document = Document(input_path)
    if selector is None:
        paragraph = document.add_paragraph()
    else:
        paragraph = resolve_single_paragraph(document, selector)
    instruction = normalize_instruction(field_type, argument, switches)
    append_field(paragraph, instruction, display_text, locked)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success(
        "field.insert",
        "ooxml",
        Fidelity.HIGH,
        {"type": field_type.upper(), "instruction": instruction},
        output_path=None if dry_run else output,
    )


def lock_fields(input_path, *, output_path, locked=True, field_type=None, dry_run=False):
    document = Document(input_path)
    target = field_type.upper() if field_type else None
    changed = 0
    for _, root in _all_story_roots(document):
        for field in root.xpath(".//w:fldSimple"):
            instruction = field.get(qn("w:instr"), "").strip()
            if target and (not instruction or instruction.split()[0].upper() != target):
                continue
            if locked:
                field.set(qn("w:fldLock"), "true")
            else:
                field.attrib.pop(qn("w:fldLock"), None)
            changed += 1
        for paragraph in root.xpath(".//w:p"):
            instructions = " ".join(paragraph.xpath(".//w:instrText/text()")).strip()
            if target and (not instructions or instructions.split()[0].upper() != target):
                continue
            for begin in paragraph.xpath(".//w:fldChar[@w:fldCharType='begin']"):
                if locked:
                    begin.set(qn("w:fldLock"), "true")
                else:
                    begin.attrib.pop(qn("w:fldLock"), None)
                changed += 1
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("field.lock" if locked else "field.unlock", "ooxml", Fidelity.HIGH, {"fields": changed}, output_path=None if dry_run else output)


def remove_fields(input_path, *, output_path, field_type=None, dry_run=False):
    document = Document(input_path)
    target = field_type.upper() if field_type else None
    removed = 0
    for _, root in _all_story_roots(document):
        for field in list(root.xpath(".//w:fldSimple")):
            instruction = field.get(qn("w:instr"), "").strip()
            if target and (not instruction or instruction.split()[0].upper() != target):
                continue
            field.getparent().remove(field)
            removed += 1
        for paragraph in root.xpath(".//w:p"):
            children = list(paragraph)
            index = 0
            while index < len(children):
                child = children[index]
                begins = [
                    element
                    for element in child.iter(qn("w:fldChar"))
                    if element.get(qn("w:fldCharType")) == "begin"
                ]
                if not begins:
                    index += 1
                    continue
                end = index
                instruction_parts: list[str] = []
                while end < len(children):
                    instruction_parts.extend(
                        element.text or ""
                        for element in children[end].iter(qn("w:instrText"))
                    )
                    if any(
                        element.get(qn("w:fldCharType")) == "end"
                        for element in children[end].iter(qn("w:fldChar"))
                    ):
                        break
                    end += 1
                if end >= len(children):
                    # Malformed complex field: fail closed rather than deleting
                    # an unbounded remainder of the paragraph.
                    index += 1
                    continue
                instruction = " ".join(instruction_parts).strip()
                field_name = instruction.split()[0].upper() if instruction else ""
                if not target or field_name == target:
                    for element in children[index : end + 1]:
                        paragraph.remove(element)
                    removed += 1
                    children = list(paragraph)
                    continue
                index = end + 1
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("field.remove", "ooxml", Fidelity.HIGH, {"removed": removed}, output_path=None if dry_run else output)


def insert_toc(
    input_path,
    *,
    output_path,
    kind="contents",
    heading_depth=3,
    caption_label=None,
    hyperlinks=True,
    page_numbers=True,
    selector=None,
    dry_run=False,
):
    document = Document(input_path)
    if kind == "contents":
        if heading_depth < 1 or heading_depth > 9:
            raise ValueError("heading_depth must be between 1 and 9")
        switches = f'\\o "1-{heading_depth}"'
    elif kind in {"figures", "tables"}:
        label = caption_label or ("Figure" if kind == "figures" else "Table")
        switches = f'\\c "{label}"'
    else:
        raise ValueError("kind must be contents, figures, or tables")
    if hyperlinks:
        switches += " \\h \\z"
    if not page_numbers:
        switches += " \\n"
    instruction = f"TOC {switches}"
    paragraph = None
    existing = False
    if selector is not None:
        anchor = resolve_single_paragraph(document, selector)
        sibling = anchor._p.getnext()
        if sibling is not None and sibling.tag == qn("w:p"):
            simple_fields = sibling.xpath(".//w:fldSimple")
            expected_label = (
                None
                if kind == "contents"
                else (caption_label or ("Figure" if kind == "figures" else "Table"))
            )
            for field in simple_fields:
                current = field.get(qn("w:instr"), "").strip()
                normalized = " ".join(current.split())
                if not normalized.upper().startswith("TOC "):
                    continue
                label_match = re.search(r'\\c\s+"([^"]+)"', normalized, re.IGNORECASE)
                same_kind = (
                    label_match is None
                    if expected_label is None
                    else label_match is not None
                    and label_match.group(1).casefold() == expected_label.casefold()
                )
                if same_kind:
                    field.set(qn("w:instr"), instruction)
                    paragraph = next(
                        item for item in document.paragraphs if item._p is sibling
                    )
                    existing = True
                    break
    if paragraph is None:
        paragraph = add_paragraph_at_boundary(document, selector)
        append_field(paragraph, instruction, "Update field in Word")
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("toc.insert", "ooxml", Fidelity.HIGH, {"kind": kind, "instruction": instruction, "updated_existing": existing}, output_path=None if dry_run else output, manual_review_required=True)


def append_caption_content(document, paragraph, label, title, sequence, chapter_number=None):
    """Use literal verified chapters with Word's heading-based SEQ reset.

    Explicit prefixes are deliberate: STYLEREF numeric switches require real
    automatic numbering, which the report's literal CHƯƠNG headings lack.
    """
    instruction = f"SEQ {sequence} \\* ARABIC"
    prefix = ""
    if chapter_number is not None:
        if type(chapter_number) is not int or chapter_number < 1:
            raise ValueError("chapter_number must be a positive integer")
        heading = None
        found = False
        for candidate in document.paragraphs:
            if candidate._p is paragraph._p:
                found = True
                break
            if candidate.style.style_id == "Heading1":
                heading = candidate
        if not found or heading is None:
            raise ValueError("chapter caption requires a preceding built-in Heading 1 in the document body")
        match = re.match(r"^CHƯƠNG\s+([1-9][0-9]*)(?=\s|[.:]|$)", heading.text.strip(), re.IGNORECASE)
        if match is None or int(match.group(1)) != chapter_number:
            raise ValueError("chapter_number does not match the nearest Heading 1 literal CHƯƠNG number")
        # Numbering can be inherited; do not confuse literal text with an
        # automatically numbered chapter or permit an incompatible outline.
        property_sets = [heading._p.pPr]
        style = heading.style
        seen_styles = set()
        while style is not None:
            if style.style_id in seen_styles:
                raise ValueError("chapter caption cannot use cyclic heading style inheritance")
            seen_styles.add(style.style_id)
            property_sets.append(style.element.pPr)
            style = style.base_style
        for properties in property_sets:
            if properties is None:
                continue
            if properties.find(qn("w:numPr")) is not None:
                raise ValueError("chapter caption literal mode does not support automatic heading numbering")
            outline = properties.find(qn("w:outlineLvl"))
            if outline is not None and outline.get(qn("w:val")) != "0":
                raise ValueError("chapter caption requires Heading 1 outline level")
        prefix = f"{chapter_number}."
        instruction += " \\s 1"
    paragraph.add_run(f"{label.strip()} {prefix}")
    append_field(paragraph, instruction, "1")
    if title:
        paragraph.add_run(f". {title}")
    return instruction


def insert_caption(
    input_path, *, output_path, label, title, selector=None, sequence_id=None,
    chapter_number: int | None = None, dry_run=False
):
    if not isinstance(label, str) or not label.strip() or len(label) > 80:
        raise ValueError("Caption display label must be a non-empty string of at most 80 characters")
    sequence = sequence_id or label
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{0,39}", sequence):
        raise ValueError("Caption sequence_id must be a Word-safe identifier")
    document = Document(input_path)
    paragraph = add_paragraph_at_boundary(document, selector, style="Caption")
    instruction = append_caption_content(document, paragraph, label, title, sequence, chapter_number)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("caption.insert", "ooxml", Fidelity.HIGH, {"label": label.strip(), "sequence_id": sequence, "chapter_number": chapter_number, "instruction": instruction}, output_path=None if dry_run else output, manual_review_required=chapter_number is not None)
