from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .common import atomic_save
from .models import Fidelity, OperationResult


_TYPES = {"plain_text": "text", "rich_text": None, "checkbox": "checkBox", "dropdown": "dropDownList", "combo_box": "comboBox", "date_picker": "date"}


def _controls(document):
    return list(document.element.body.xpath(".//w:sdt"))


def _property_value(control, name):
    properties = control.find(qn("w:sdtPr"))
    element = properties.find(qn(f"w:{name}")) if properties is not None else None
    return element.get(qn("w:val")) if element is not None else None


def list_content_controls(input_path):
    document = Document(input_path)
    result = []
    for control in _controls(document):
        properties = control.find(qn("w:sdtPr"))
        control_type = "rich_text"
        for public_name, xml_name in _TYPES.items():
            if xml_name and properties.find(qn(f"w:{xml_name}")) is not None:
                control_type = public_name
                break
        result.append(
            {
                "type": control_type,
                "title": _property_value(control, "alias"),
                "tag": _property_value(control, "tag"),
                "value": "".join(value.text or "" for value in control.iter(qn("w:t"))),
                "locked": properties.find(qn("w:lock")) is not None,
            }
        )
    return result


def create_content_control(input_path, *, output_path, control_type, title, tag, value="", locked=False, dry_run=False):
    if control_type not in _TYPES:
        raise ValueError(f"Unsupported content control type: {control_type}")
    if control_type != "plain_text":
        raise ValueError(
            "Only plain_text content controls are verified; advanced Word "
            "content-control types are unsupported"
        )
    document = Document(input_path)
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), title)
    tag_element = OxmlElement("w:tag")
    tag_element.set(qn("w:val"), tag)
    properties.extend([alias, tag_element])
    xml_type = _TYPES[control_type]
    if xml_type:
        properties.append(OxmlElement(f"w:{xml_type}"))
    if locked:
        lock = OxmlElement("w:lock")
        lock.set(qn("w:val"), "sdtLocked")
        properties.append(lock)
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = str(value)
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    sdt.extend([properties, content])
    document.element.body.insert(len(document.element.body) - 1, sdt)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("content-control.create", "ooxml", Fidelity.STRUCTURAL, {"tag": tag, "type": control_type}, output_path=None if dry_run else output, manual_review_required=control_type != "plain_text")


def set_content_control_value(input_path, *, output_path, tag, value, dry_run=False):
    document = Document(input_path)
    matches = [control for control in _controls(document) if _property_value(control, "tag") == tag]
    if len(matches) != 1:
        raise ValueError(f"Expected one content control with tag {tag!r}, found {len(matches)}")
    properties = matches[0].find(qn("w:sdtPr"))
    if properties is None or properties.find(qn("w:text")) is None:
        raise ValueError(
            "Only verified plain_text content controls can be updated"
        )
    content = matches[0].find(qn("w:sdtContent"))
    texts = list(content.iter(qn("w:t"))) if content is not None else []
    if not texts:
        raise ValueError("Content control has no editable text node")
    texts[0].text = str(value)
    for text in texts[1:]:
        text.text = ""
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("content-control.set", "ooxml", Fidelity.STRUCTURAL, {"tag": tag}, output_path=None if dry_run else output)


def delete_content_control(input_path, *, output_path, tag, keep_content=True, dry_run=False):
    document = Document(input_path)
    matches = [control for control in _controls(document) if _property_value(control, "tag") == tag]
    if len(matches) != 1:
        raise ValueError(f"Expected one content control with tag {tag!r}, found {len(matches)}")
    control = matches[0]
    parent = control.getparent()
    index = parent.index(control)
    if keep_content:
        content = control.find(qn("w:sdtContent"))
        for child in list(content):
            content.remove(child)
            parent.insert(index, child)
            index += 1
    parent.remove(control)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("content-control.delete", "ooxml", Fidelity.STRUCTURAL, {"tag": tag}, output_path=None if dry_run else output)
