from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from lxml import etree

from ..selectors import Selector, resolve_paragraphs
from .common import atomic_save
from .models import Fidelity, OperationResult


_CONFIG = {
    "footnote": {
        "partname": "/word/footnotes.xml",
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        "reltype": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
        "root": "footnotes",
        "item": "footnote",
        "reference": "footnoteReference",
    },
    "endnote": {
        "partname": "/word/endnotes.xml",
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
        "reltype": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes",
        "root": "endnotes",
        "item": "endnote",
        "reference": "endnoteReference",
    },
}


def _config(kind):
    if kind not in _CONFIG:
        raise ValueError("kind must be footnote or endnote")
    return _CONFIG[kind]


def _initial_xml(config):
    return (
        f'<w:{config["root"]} {nsdecls("w")}>'
        f'<w:{config["item"]} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{config["item"]}>'
        f'<w:{config["item"]} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:{config["item"]}>'
        f'</w:{config["root"]}>'
    ).encode("utf-8")


def _part(document, kind, create=False):
    config = _config(kind)
    for relationship in document.part.rels.values():
        if relationship.reltype == config["reltype"]:
            return relationship.target_part
    if not create:
        return None
    part = Part(
        PackURI(config["partname"]),
        config["content_type"],
        _initial_xml(config),
        document.part.package,
    )
    document.part.relate_to(part, config["reltype"])
    return part


def _root(part):
    return parse_xml(part.blob)


def _save_root(part, root):
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def list_notes(input_path, *, kind="footnote"):
    document = Document(input_path)
    part = _part(document, kind)
    if part is None:
        return []
    config = _config(kind)
    result = []
    for element in _root(part).findall(qn(f'w:{config["item"]}')):
        note_id = int(element.get(qn("w:id")))
        if note_id <= 0:
            continue
        result.append(
            {
                "id": note_id,
                "kind": kind,
                "text": "".join(value.text or "" for value in element.iter(qn("w:t"))),
            }
        )
    return result


def insert_note(input_path, *, output_path, kind, text, selector=None, dry_run=False):
    document = Document(input_path)
    config = _config(kind)
    part = _part(document, kind, create=True)
    root = _root(part)
    ids = [int(element.get(qn("w:id"))) for element in root.findall(qn(f'w:{config["item"]}'))]
    note_id = max(ids, default=0) + 1
    note = OxmlElement(f'w:{config["item"]}')
    note.set(qn("w:id"), str(note_id))
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    paragraph.append(run)
    note.append(paragraph)
    root.append(note)
    _save_root(part, root)
    target = document.add_paragraph() if selector is None else resolve_paragraphs(document, Selector.from_value(selector))[0]
    reference_run = target.add_run()._r
    reference = OxmlElement(f'w:{config["reference"]}')
    reference.set(qn("w:id"), str(note_id))
    reference_run.append(reference)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success(f"{kind}.insert", "ooxml", Fidelity.HIGH, {"id": note_id}, output_path=None if dry_run else output)


def edit_note(input_path, *, output_path, kind, note_id, text, dry_run=False):
    document = Document(input_path)
    config = _config(kind)
    part = _part(document, kind)
    if part is None:
        raise ValueError(f"Document has no {kind}s")
    root = _root(part)
    matches = [element for element in root.findall(qn(f'w:{config["item"]}')) if element.get(qn("w:id")) == str(note_id)]
    if len(matches) != 1:
        raise ValueError(f"{kind} id not found: {note_id}")
    paragraphs = matches[0].findall(qn("w:p"))
    for paragraph in paragraphs:
        matches[0].remove(paragraph)
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    paragraph.append(run)
    matches[0].append(paragraph)
    _save_root(part, root)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success(f"{kind}.edit", "ooxml", Fidelity.HIGH, {"id": note_id}, output_path=None if dry_run else output)


def delete_note(input_path, *, output_path, kind, note_id, dry_run=False):
    document = Document(input_path)
    config = _config(kind)
    part = _part(document, kind)
    if part is None:
        raise ValueError(f"Document has no {kind}s")
    root = _root(part)
    removed = False
    for element in list(root.findall(qn(f'w:{config["item"]}'))):
        if element.get(qn("w:id")) == str(note_id):
            root.remove(element)
            removed = True
    if not removed:
        raise ValueError(f"{kind} id not found: {note_id}")
    _save_root(part, root)
    for reference in list(document.element.body.xpath(f'.//w:{config["reference"]}')):
        if reference.get(qn("w:id")) == str(note_id):
            run = reference.getparent()
            run.getparent().remove(run)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success(f"{kind}.delete", "ooxml", Fidelity.HIGH, {"id": note_id}, output_path=None if dry_run else output)
