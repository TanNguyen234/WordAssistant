from __future__ import annotations

from pathlib import Path
import uuid
from typing import Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from ..selectors import resolve_single_paragraph
from .common import atomic_save
from .models import Fidelity, OperationResult


def _table(document, index: int):
    try:
        return document.tables[index]
    except IndexError as exc:
        raise ValueError(f"Table index out of range: {index}") from exc


def _set_width(element, width: int):
    element.set(qn("w:w"), str(width))
    element.set(qn("w:type"), "dxa")


def _upsert(parent, tag: str):
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    return element


def _section_for_table(document, table):
    section_index = 0
    for child in document.element.body.iterchildren():
        if child is table._tbl:
            return document.sections[section_index]
        if child.tag == qn("w:p"):
            properties = child.find(qn("w:pPr"))
            if properties is not None and properties.find(qn("w:sectPr")) is not None:
                section_index += 1
    raise ValueError("Table is not a top-level document table")


def _default_table_width(document, table) -> int:
    section = _section_for_table(document, table)
    return int(
        section.page_width.twips
        - section.left_margin.twips
        - section.right_margin.twips
    )


def normalize_table_geometry(document, table, width_twips: int | None = None):
    tbl_pr = table._tbl.tblPr
    current = tbl_pr.find(qn("w:tblW"))
    current_width = None
    if current is not None and current.get(qn("w:type")) == "dxa":
        value = current.get(qn("w:w"))
        if value and value.isdigit() and int(value) > 0:
            current_width = int(value)
    total = int(width_twips or current_width or _default_table_width(document, table))
    column_count = len(table._tbl.tblGrid.gridCol_lst) or len(table.columns)
    base, remainder = divmod(total, column_count)
    grid_widths = [base + (1 if index < remainder else 0) for index in range(column_count)]

    _set_width(_upsert(tbl_pr, "w:tblW"), total)
    # A table consuming the printable width must not also carry a positive
    # indent, otherwise its right edge exceeds the text area in Word.
    _set_width(_upsert(tbl_pr, "w:tblInd"), 0)
    layout = _upsert(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    margins = _upsert(tbl_pr, "w:tblCellMar")
    for side in ("top", "start", "bottom", "end"):
        _set_width(
            _upsert(margins, f"w:{side}"),
            120 if side in {"start", "end"} else 80,
        )

    grid = table._tbl.tblGrid
    while len(grid.gridCol_lst) < column_count:
        grid.append(OxmlElement("w:gridCol"))
    for index, grid_column in enumerate(grid.gridCol_lst):
        grid_column.set(qn("w:w"), str(grid_widths[index]))

    for row in table._tbl.tr_lst:
        cursor = 0
        for cell in row.tc_lst:
            grid_span = cell.tcPr.find(qn("w:gridSpan"))
            span = int(grid_span.get(qn("w:val"))) if grid_span is not None else 1
            cell_width = sum(grid_widths[cursor : cursor + span])
            _set_width(_upsert(cell.tcPr, "w:tcW"), cell_width)
            cell_margins = _upsert(cell.tcPr, "w:tcMar")
            for side in ("top", "start", "bottom", "end"):
                _set_width(
                    _upsert(cell_margins, f"w:{side}"),
                    120 if side in {"start", "end"} else 80,
                )
            cursor += span
    table.autofit = False


def _save_result(
    document,
    output_path,
    operation,
    data,
    dry_run=False,
    *,
    normalize_geometry=True,
):
    if normalize_geometry:
        for table in document.tables:
            normalize_table_geometry(document, table)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult(
        "DRY_RUN" if dry_run else "SUCCESS",
        operation,
        "python-docx",
        Fidelity.HIGH,
        None if dry_run else output,
        data=data,
    )


def _set_table_metadata(table, *, caption=None, internal_id=None):
    properties = table._tbl.tblPr
    if caption is not None:
        element = properties.find(qn("w:tblCaption"))
        if element is None:
            element = OxmlElement("w:tblCaption")
            properties.append(element)
        element.set(qn("w:val"), caption)
    identifier = internal_id or f"wa-{uuid.uuid4().hex}"
    description = properties.find(qn("w:tblDescription"))
    if description is None:
        description = OxmlElement("w:tblDescription")
        properties.append(description)
    description.set(qn("w:val"), identifier)
    return identifier


def create_table(
    input_path: str | Path,
    *,
    output_path: str | Path,
    rows: int,
    columns: int,
    data: Sequence[Sequence[str]] | None = None,
    style: str | None = None,
    caption: str | None = None,
    nested: dict | None = None,
    selector=None,
    dry_run: bool = False,
):
    if rows < 1 or columns < 1:
        raise ValueError("rows and columns must be positive")
    if nested is not None:
        raise ValueError(
            "Nested table creation is unsupported until recursive geometry verification is available"
        )
    document = Document(input_path)
    table = document.add_table(rows=rows, cols=columns)
    if selector is not None:
        resolve_single_paragraph(document, selector)._p.addnext(table._tbl)
    if style:
        table.style = style
    if data:
        if len(data) > rows or any(len(row) > columns for row in data):
            raise ValueError("data exceeds requested table dimensions")
        for row_index, values in enumerate(data):
            for column_index, value in enumerate(values):
                table.cell(row_index, column_index).text = str(value)
    identifier = _set_table_metadata(table, caption=caption)
    table_index = next(
        index
        for index, candidate in enumerate(document.tables)
        if candidate._tbl is table._tbl
    )
    return _save_result(document, output_path, "table.create", {"table_index": table_index, "internal_id": identifier}, dry_run)


def delete_table(input_path, *, output_path, table_index: int, dry_run=False):
    document = Document(input_path)
    table = _table(document, table_index)
    table._tbl.getparent().remove(table._tbl)
    return _save_result(document, output_path, "table.delete", {"table_index": table_index}, dry_run)


def insert_row(input_path, *, output_path, table_index: int, index: int, values=None, dry_run=False):
    document = Document(input_path)
    table = _table(document, table_index)
    if index < 0 or index > len(table.rows):
        raise ValueError(f"Row index out of range: {index}")
    row = table.add_row()
    table._tbl.remove(row._tr)
    existing_rows = list(table._tbl.tr_lst)
    if index == len(existing_rows):
        table._tbl.append(row._tr)
    else:
        existing_rows[index].addprevious(row._tr)
    if values is not None:
        if len(values) > len(row.cells):
            raise ValueError("values exceed table columns")
        for cell, value in zip(row.cells, values):
            cell.text = str(value)
    return _save_result(document, output_path, "table.insert_row", {"row": index}, dry_run)


def delete_row(input_path, *, output_path, table_index: int, index: int, dry_run=False):
    document = Document(input_path)
    table = _table(document, table_index)
    if len(table.rows) <= 1:
        raise ValueError("Cannot delete the only table row")
    try:
        row = table.rows[index]
    except IndexError as exc:
        raise ValueError(f"Row index out of range: {index}") from exc
    table._tbl.remove(row._tr)
    return _save_result(document, output_path, "table.delete_row", {"row": index}, dry_run)


def _new_cell():
    cell = OxmlElement("w:tc")
    cell.append(OxmlElement("w:tcPr"))
    cell.append(OxmlElement("w:p"))
    return cell


def _require_simple_grid(table):
    if table._tbl.xpath(".//w:gridSpan | .//w:vMerge | .//w:tbl"):
        raise ValueError(
            "Column mutation on merged or nested tables is unsupported until "
            "a logical-grid engine is available"
        )


def insert_column(input_path, *, output_path, table_index: int, index: int, values=None, dry_run=False):
    document = Document(input_path)
    table = _table(document, table_index)
    _require_simple_grid(table)
    column_count = len(table.columns)
    if index < 0 or index > column_count:
        raise ValueError(f"Column index out of range: {index}")
    grid = table._tbl.tblGrid
    grid_column = OxmlElement("w:gridCol")
    grid.insert(index, grid_column)
    for row_index, row in enumerate(table.rows):
        row._tr.insert(index, _new_cell())
        if values is not None and row_index < len(values):
            row.cells[index].text = str(values[row_index])
    return _save_result(document, output_path, "table.insert_column", {"column": index}, dry_run)


def delete_column(input_path, *, output_path, table_index: int, index: int, dry_run=False):
    document = Document(input_path)
    table = _table(document, table_index)
    _require_simple_grid(table)
    if len(table.columns) <= 1:
        raise ValueError("Cannot delete the only table column")
    if index < 0 or index >= len(table.columns):
        raise ValueError(f"Column index out of range: {index}")
    for row in table.rows:
        row._tr.remove(row._tr.tc_lst[index])
    grid = table._tbl.tblGrid
    grid.remove(grid.gridCol_lst[index])
    return _save_result(document, output_path, "table.delete_column", {"column": index}, dry_run)


def set_cell_text(input_path, *, output_path, table_index: int, row: int, column: int, text: str, dry_run=False):
    document = Document(input_path)
    table = _table(document, table_index)
    table.cell(row, column).text = text
    return _save_result(document, output_path, "table.set_cell", {"row": row, "column": column}, dry_run)


def set_row_text(input_path, *, output_path, table_index: int, row: int, values, dry_run=False):
    document = Document(input_path)
    cells = _table(document, table_index).rows[row].cells
    if len(values) != len(cells):
        raise ValueError("values must match row width")
    for cell, value in zip(cells, values):
        cell.text = str(value)
    return _save_result(document, output_path, "table.set_row", {"row": row}, dry_run)


def set_column_text(input_path, *, output_path, table_index: int, column: int, values, dry_run=False):
    document = Document(input_path)
    table = _table(document, table_index)
    if len(values) != len(table.rows):
        raise ValueError("values must match column height")
    for row, value in zip(table.rows, values):
        row.cells[column].text = str(value)
    return _save_result(document, output_path, "table.set_column", {"column": column}, dry_run)


def merge_cells(input_path, *, output_path, table_index: int, start, end, dry_run=False):
    document = Document(input_path)
    table = _table(document, table_index)
    table.cell(*start).merge(table.cell(*end))
    return _save_result(document, output_path, "table.merge", {"start": list(start), "end": list(end)}, dry_run)


def split_cell(input_path, *, output_path, table_index: int, row: int, column: int, rows: int, columns: int, dry_run=False):
    document = Document(input_path)
    cell = _table(document, table_index).cell(row, column)
    split = getattr(cell, "split", None)
    if not callable(split):
        raise ValueError("Cell splitting is unavailable through the current python-docx backend")
    split(rows, columns)
    return _save_result(document, output_path, "table.split", {"rows": rows, "columns": columns}, dry_run)


def set_table_layout(
    input_path,
    *,
    output_path,
    table_index: int,
    width_cm=None,
    row_height_pt=None,
    autofit=None,
    alignment=None,
    vertical_alignment=None,
    repeat_header=None,
    keep_rows_together=None,
    style=None,
    dry_run=False,
):
    document = Document(input_path)
    table = _table(document, table_index)
    if width_cm is not None:
        normalize_table_geometry(document, table, int(Cm(float(width_cm)).twips))
    elif autofit is False:
        normalize_table_geometry(document, table)
    if row_height_pt is not None:
        for row in table.rows:
            row.height = Pt(float(row_height_pt))
    if autofit is not None:
        table.autofit = bool(autofit)
    if alignment is not None:
        alignments = {"left": WD_TABLE_ALIGNMENT.LEFT, "center": WD_TABLE_ALIGNMENT.CENTER, "right": WD_TABLE_ALIGNMENT.RIGHT}
        if alignment not in alignments:
            raise ValueError(f"Unsupported table alignment: {alignment}")
        table.alignment = alignments[alignment]
    if vertical_alignment is not None:
        vertical = {"top": WD_CELL_VERTICAL_ALIGNMENT.TOP, "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER, "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}
        if vertical_alignment not in vertical:
            raise ValueError(f"Unsupported vertical alignment: {vertical_alignment}")
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = vertical[vertical_alignment]
    if repeat_header is not None:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        existing = tr_pr.find(qn("w:tblHeader"))
        if repeat_header and existing is None:
            tr_pr.append(OxmlElement("w:tblHeader"))
        elif not repeat_header and existing is not None:
            tr_pr.remove(existing)
    if keep_rows_together is not None:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            existing = tr_pr.find(qn("w:cantSplit"))
            if keep_rows_together and existing is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
            elif not keep_rows_together and existing is not None:
                tr_pr.remove(existing)
    if style is not None:
        table.style = style
    return _save_result(
        document,
        output_path,
        "table.layout",
        {"table_index": table_index},
        dry_run,
        normalize_geometry=False,
    )
