from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .common import atomic_save
from .models import Fidelity, OperationResult


def set_header_footer_text(
    input_path,
    *,
    output_path,
    section_index: int,
    story: str,
    text: str,
    variant: str = "primary",
    alignment: str = "left",
    link_previous: bool | None = None,
    dry_run=False,
):
    document = Document(input_path)
    section = document.sections[section_index]
    if story not in {"header", "footer"} or variant not in {"primary", "first", "even"}:
        raise ValueError("Invalid header/footer story or variant")
    attribute = story if variant == "primary" else f"{variant}_page_{story}"
    container = getattr(section, attribute)
    if link_previous is not None:
        container.is_linked_to_previous = link_previous
    paragraph = container.paragraphs[0]
    paragraph.clear()
    paragraph.add_run(text)
    alignments = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
    paragraph.alignment = alignments[alignment]
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success(f"{story}.set", "python-docx", Fidelity.HIGH, {"section_index": section_index, "variant": variant}, output_path=None if dry_run else output)

