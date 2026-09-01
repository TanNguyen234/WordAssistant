from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

from ..selectors import Selector, resolve_paragraphs
from .common import atomic_save
from .fields import append_field
from .models import Fidelity, OperationResult


_BOOKMARK_NAME = re.compile(r"[A-Za-z_][A-Za-z0-9_]{0,39}$")


def _story_roots(document):
    yield document.element
    seen = {id(document.element)}
    for section in document.sections:
        for name in (
            "header",
            "footer",
            "first_page_header",
            "first_page_footer",
            "even_page_header",
            "even_page_footer",
        ):
            root = getattr(section, name)._element
            if id(root) not in seen:
                seen.add(id(root))
                yield root


def _validate_bookmark_name(name):
    if not _BOOKMARK_NAME.fullmatch(name):
        raise ValueError("Bookmark name must start with a letter/underscore and contain at most 40 letters, digits, or underscores")


def inspect_bookmarks(input_path):
    document = Document(input_path)
    result = []
    for element in document.element.body.xpath(".//w:bookmarkStart"):
        name = element.get(qn("w:name"))
        if name and name != "_GoBack":
            result.append({"name": name, "id": int(element.get(qn("w:id")))})
    return result


def _next_bookmark_id(document):
    values = [int(value) for value in document.element.body.xpath(".//w:bookmarkStart/@w:id") if str(value).isdigit()]
    return max(values, default=-1) + 1


def create_bookmark(input_path, *, output_path, name, selector, dry_run=False):
    _validate_bookmark_name(name)
    document = Document(input_path)
    if any(item["name"] == name for item in inspect_bookmarks(input_path)):
        raise ValueError(f"Bookmark already exists: {name}")
    paragraph = resolve_paragraphs(document, Selector.from_value(selector))[0]
    bookmark_id = _next_bookmark_id(document)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("bookmark.create", "ooxml", Fidelity.HIGH, {"name": name, "id": bookmark_id}, output_path=None if dry_run else output)


def rename_bookmark(input_path, *, output_path, old_name, new_name, dry_run=False):
    _validate_bookmark_name(new_name)
    document = Document(input_path)
    starts = [element for element in document.element.body.xpath(".//w:bookmarkStart") if element.get(qn("w:name")) == old_name]
    if len(starts) != 1:
        raise ValueError(f"Expected one bookmark named {old_name!r}, found {len(starts)}")
    if any(
        element.get(qn("w:name")) == new_name
        for element in document.element.body.xpath(".//w:bookmarkStart")
    ):
        raise ValueError(f"Bookmark already exists: {new_name}")
    starts[0].set(qn("w:name"), new_name)
    reference_pattern = re.compile(
        rf"\b(REF|PAGEREF|NOTEREF)(\s+){re.escape(old_name)}\b",
        re.IGNORECASE,
    )
    for root in _story_roots(document):
        for field in root.xpath(".//w:fldSimple"):
            instruction = field.get(qn("w:instr"), "")
            field.set(
                qn("w:instr"),
                reference_pattern.sub(
                    lambda match: f"{match.group(1)}{match.group(2)}{new_name}",
                    instruction,
                ),
            )
        for instruction in root.xpath(".//w:instrText"):
            instruction.text = reference_pattern.sub(
                lambda match: f"{match.group(1)}{match.group(2)}{new_name}",
                instruction.text or "",
            )
        for hyperlink in root.xpath(".//w:hyperlink"):
            if hyperlink.get(qn("w:anchor")) == old_name:
                hyperlink.set(qn("w:anchor"), new_name)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("bookmark.rename", "ooxml", Fidelity.HIGH, {"old_name": old_name, "new_name": new_name}, output_path=None if dry_run else output)


def delete_bookmark(input_path, *, output_path, name, dry_run=False):
    document = Document(input_path)
    starts = [element for element in document.element.body.xpath(".//w:bookmarkStart") if element.get(qn("w:name")) == name]
    if len(starts) != 1:
        raise ValueError(f"Expected one bookmark named {name!r}, found {len(starts)}")
    reference_pattern = re.compile(
        rf"\b(?:REF|PAGEREF|NOTEREF)\s+{re.escape(name)}\b", re.IGNORECASE
    )
    instructions = []
    anchored = False
    for root in _story_roots(document):
        instructions.extend(
            field.get(qn("w:instr"), "")
            for field in root.xpath(".//w:fldSimple")
        )
        instructions.extend(
            field.text or "" for field in root.xpath(".//w:instrText")
        )
        anchored = anchored or any(
            hyperlink.get(qn("w:anchor")) == name
            for hyperlink in root.xpath(".//w:hyperlink")
        )
    if anchored or any(reference_pattern.search(instruction) for instruction in instructions):
        raise ValueError(f"Bookmark {name!r} is referenced and cannot be deleted")
    bookmark_id = starts[0].get(qn("w:id"))
    starts[0].getparent().remove(starts[0])
    for end in list(document.element.body.xpath(".//w:bookmarkEnd")):
        if end.get(qn("w:id")) == bookmark_id:
            end.getparent().remove(end)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("bookmark.delete", "ooxml", Fidelity.HIGH, {"name": name}, output_path=None if dry_run else output)


def insert_hyperlink(input_path, *, output_path, text, url=None, email=None, bookmark=None, selector=None, dry_run=False):
    targets = sum(value is not None for value in (url, email, bookmark))
    if targets != 1:
        raise ValueError("Provide exactly one of url, email, or bookmark")
    document = Document(input_path)
    paragraph = document.add_paragraph() if selector is None else resolve_paragraphs(document, Selector.from_value(selector))[0]
    hyperlink = OxmlElement("w:hyperlink")
    if bookmark:
        _validate_bookmark_name(bookmark)
        names = {
            element.get(qn("w:name"))
            for element in document.element.body.xpath(".//w:bookmarkStart")
        }
        if bookmark not in names:
            raise ValueError(f"Bookmark does not exist: {bookmark}")
        hyperlink.set(qn("w:anchor"), bookmark)
    else:
        target = f"mailto:{email}" if email else url
        relationship_id = document.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    value = OxmlElement("w:t")
    value.text = text
    run.extend([properties, value])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("hyperlink.insert", "ooxml", Fidelity.HIGH, {"text": text, "target": bookmark or email or url}, output_path=None if dry_run else output)


def insert_cross_reference(input_path, *, output_path, bookmark, include_page=False, selector=None, dry_run=False):
    _validate_bookmark_name(bookmark)
    document = Document(input_path)
    names = {
        element.get(qn("w:name"))
        for element in document.element.body.xpath(".//w:bookmarkStart")
    }
    if bookmark not in names:
        raise ValueError(f"Bookmark does not exist: {bookmark}")
    paragraph = document.add_paragraph() if selector is None else resolve_paragraphs(document, Selector.from_value(selector))[0]
    append_field(paragraph, f"REF {bookmark} \\h", "reference")
    if include_page:
        paragraph.add_run(" (page ")
        append_field(paragraph, f"PAGEREF {bookmark} \\h", "1")
        paragraph.add_run(")")
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("reference.insert", "ooxml", Fidelity.HIGH, {"bookmark": bookmark, "include_page": include_page}, output_path=None if dry_run else output)
