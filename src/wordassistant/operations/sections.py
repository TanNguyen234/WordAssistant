from __future__ import annotations

from copy import deepcopy
import math
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import Section
from docx.shared import Cm

from ..selectors import resolve_single_paragraph
from .common import atomic_save
from .models import Fidelity, OperationResult


_BREAK_TYPES = {
    "next_page": WD_SECTION_START.NEW_PAGE,
    "continuous": WD_SECTION_START.CONTINUOUS,
    "odd_page": WD_SECTION_START.ODD_PAGE,
    "even_page": WD_SECTION_START.EVEN_PAGE,
}

_PAGE_SIZES = {
    "A4": (21.0, 29.7),
    "Letter": (21.59, 27.94),
}


def _result(document, output_path, operation, data, dry_run=False):
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult(
        "DRY_RUN" if dry_run else "SUCCESS",
        operation,
        "python-docx",
        Fidelity.HIGH,
        None if dry_run else output,
        data=data,
    )


def _containing_section_properties(document, paragraph):
    children = list(document.element.body.iterchildren())
    try:
        start = next(
            index for index, child in enumerate(children) if child is paragraph._p
        )
    except StopIteration as exc:
        raise ValueError("Anchor paragraph is not in the document body") from exc
    for child in children[start:]:
        if child.tag == qn("w:p"):
            properties = child.pPr.sectPr if child.pPr is not None else None
            if properties is not None:
                return properties
        elif child.tag == qn("w:sectPr"):
            return child
    raise ValueError("Document has no section properties after the anchor")


def add_section_break(
    input_path,
    *,
    output_path,
    break_type="next_page",
    selector=None,
    dry_run=False,
):
    if break_type not in _BREAK_TYPES:
        raise ValueError(f"Unsupported section break: {break_type}")
    document = Document(input_path)
    if selector is None:
        document.add_section(_BREAK_TYPES[break_type])
        break_properties = document.paragraphs[-1]._p.pPr.sectPr
    else:
        target = resolve_single_paragraph(document, selector)
        paragraph_properties = target._p.get_or_add_pPr()
        if paragraph_properties.sectPr is not None:
            raise ValueError("Anchor paragraph already contains a section boundary")
        current_properties = _containing_section_properties(document, target)
        break_properties = deepcopy(current_properties)
        paragraph_properties.append(break_properties)
        Section(current_properties, document.part).start_type = _BREAK_TYPES[
            break_type
        ]
    section_properties = document.element.body.xpath(
        "./w:p/w:pPr/w:sectPr | ./w:sectPr"
    )
    section_index = next(
        index + 1
        for index, properties in enumerate(section_properties)
        if properties is break_properties
    )
    return _result(
        document,
        output_path,
        "section.add",
        {
            "section_index": section_index,
            "break_type": break_type,
            "anchored": selector is not None,
        },
        dry_run,
    )


def set_section_layout(
    input_path,
    *,
    output_path,
    section_index: int,
    page_size: str | None = None,
    width_cm: float | None = None,
    height_cm: float | None = None,
    orientation: str | None = None,
    margins: dict | None = None,
    gutter_cm: float | None = None,
    mirrored_margins: bool | None = None,
    columns: int | None = None,
    different_first_page: bool | None = None,
    different_odd_even: bool | None = None,
    page_border_space_pt: float | None = None,
    dry_run=False,
):
    document = Document(input_path)
    try:
        section = document.sections[section_index]
    except IndexError as exc:
        raise ValueError(f"Section index out of range: {section_index}") from exc
    if page_size is not None:
        if page_size not in _PAGE_SIZES:
            raise ValueError(f"Unsupported page size: {page_size}")
        width_cm, height_cm = _PAGE_SIZES[page_size]
    base_width = float(width_cm) if width_cm is not None else section.page_width.cm
    base_height = float(height_cm) if height_cm is not None else section.page_height.cm
    if orientation is not None:
        if orientation not in {"portrait", "landscape"}:
            raise ValueError("orientation must be portrait or landscape")
        section.orientation = WD_ORIENT.LANDSCAPE if orientation == "landscape" else WD_ORIENT.PORTRAIT
        short, long = sorted((base_width, base_height))
        section.page_width = Cm(long if orientation == "landscape" else short)
        section.page_height = Cm(short if orientation == "landscape" else long)
    else:
        section.page_width = Cm(base_width)
        section.page_height = Cm(base_height)
    if margins:
        for key in ("top", "bottom", "left", "right"):
            if key in margins:
                setattr(section, f"{key}_margin", Cm(float(margins[key])))
    if gutter_cm is not None:
        section.gutter = Cm(float(gutter_cm))
    if columns is not None:
        if columns < 1:
            raise ValueError("columns must be positive")
        cols = section._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            section._sectPr.append(cols)
        cols.set(qn("w:num"), str(columns))
    if different_first_page is not None:
        section.different_first_page_header_footer = bool(different_first_page)
    if page_border_space_pt is not None:
        if (
            isinstance(page_border_space_pt, bool)
            or not isinstance(page_border_space_pt, (int, float))
            or not math.isfinite(float(page_border_space_pt))
            or not 1 <= float(page_border_space_pt) <= 144
        ):
            raise ValueError("page_border_space_pt must be between 1 and 144 points")
        borders = section._sectPr.find(qn("w:pgBorders"))
        if borders is None:
            raise ValueError("Cannot set page border inset because this section has no page border")
        if borders.get(qn("w:offsetFrom"), "page") != "page":
            raise ValueError("page_border_space_pt requires a page-relative border")
        edge_names = ("top", "left", "bottom", "right")
        edges = [borders.find(qn(f"w:{name}")) for name in edge_names]
        if any(edge is None for edge in edges):
            raise ValueError("Cannot set page border inset because one or more border edges are missing")
        space = str(int(round(float(page_border_space_pt))))
        for edge in edges:
            edge.set(qn("w:space"), space)
    settings = document.settings.element
    even = settings.find(qn("w:evenAndOddHeaders"))
    if different_odd_even is not None:
        if different_odd_even and even is None:
            settings.append(OxmlElement("w:evenAndOddHeaders"))
        elif not different_odd_even and even is not None:
            active_even_sections = []
            for candidate_index, candidate in enumerate(document.sections):
                has_even_fields = any(
                    container._element.xpath(".//w:instrText | .//w:fldSimple")
                    for container in (
                        candidate.even_page_header,
                        candidate.even_page_footer,
                    )
                )
                if has_even_fields:
                    active_even_sections.append(candidate_index)
            if len(document.sections) > 1 and active_even_sections:
                raise ValueError(
                    "different_odd_even is a document-level setting and cannot be "
                    "disabled through a section-scoped layout operation while even-page "
                    f"fields exist in sections: {active_even_sections}"
                )
            settings.remove(even)
    mirror = settings.find(qn("w:mirrorMargins"))
    if mirrored_margins is not None:
        if mirrored_margins and mirror is None:
            settings.append(OxmlElement("w:mirrorMargins"))
        elif not mirrored_margins and mirror is not None:
            settings.remove(mirror)
    return _result(
        document,
        output_path,
        "section.layout",
        {"section_index": section_index, "page_border_space_pt": page_border_space_pt},
        dry_run,
    )


def link_header_footer(input_path, *, output_path, section_index: int, linked: bool, story="both", dry_run=False):
    document = Document(input_path)
    section = document.sections[section_index]
    stories = []
    if story in {"header", "both"}:
        stories.extend([section.header, section.first_page_header, section.even_page_header])
    if story in {"footer", "both"}:
        stories.extend([section.footer, section.first_page_footer, section.even_page_footer])
    for value in stories:
        value.is_linked_to_previous = linked
    return _result(document, output_path, "section.link", {"section_index": section_index, "linked": linked}, dry_run)
