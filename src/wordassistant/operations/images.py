from __future__ import annotations

import math
from numbers import Real
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.text.paragraph import Paragraph

from ..selectors import insert_paragraph_after, resolve_single_paragraph
from .common import atomic_save
from .fields import append_caption_content
from .models import Fidelity, OperationResult


EMU_PER_CM = 360000
_PLACEMENTS = {"before", "after", "replace"}
_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def _drawings(document):
    return list(document.element.body.xpath(".//w:drawing"))


def _doc_properties(drawing):
    values = drawing.xpath(".//wp:docPr")
    if not values:
        raise ValueError("Drawing has no wp:docPr metadata")
    return values[0]


def _validate_dimension(name: str, value: float | None) -> None:
    if value is None:
        return
    if isinstance(value, bool) or not isinstance(value, Real):
        raise TypeError(f"{name} must be a positive number")
    if not math.isfinite(float(value)) or float(value) <= 0:
        raise ValueError(f"{name} must be greater than zero")


def _validate_dimensions(
    width_cm: float | None,
    height_cm: float | None,
    preserve_aspect_ratio: bool,
) -> None:
    if type(preserve_aspect_ratio) is not bool:
        raise TypeError("preserve_aspect_ratio must be a boolean")
    _validate_dimension("width_cm", width_cm)
    _validate_dimension("height_cm", height_cm)
    if preserve_aspect_ratio and width_cm is not None and height_cm is not None:
        raise ValueError(
            "width_cm and height_cm conflict when preserve_aspect_ratio is true"
        )


def _insert_paragraph_before(paragraph) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addprevious(element)
    return Paragraph(element, paragraph._parent)


def _image_paragraph(document, selector, placement: str) -> tuple[Paragraph, object | None]:
    if placement not in _PLACEMENTS:
        raise ValueError("placement must be before, after, or replace")
    if selector is None:
        if placement != "after":
            raise ValueError(f"placement={placement!r} requires a selector")
        return document.add_paragraph(), None

    anchor = resolve_single_paragraph(document, selector)
    if placement == "before" or placement == "replace":
        paragraph = _insert_paragraph_before(anchor)
    else:
        paragraph = insert_paragraph_after(anchor)
    if placement == "replace":
        properties = anchor._p.pPr
        if properties is not None and properties.find(qn("w:sectPr")) is not None:
            paragraph._p.getparent().remove(paragraph._p)
            raise ValueError("Cannot replace a paragraph that carries a section break")
    for attribute in ("left_indent", "right_indent"):
        indent = _effective_indent_emu(anchor, attribute)
        if indent:
            setattr(paragraph.paragraph_format, attribute, Emu(indent))
    return paragraph, anchor


def _section_index_for_paragraph(document, paragraph: Paragraph) -> int:
    section_index = 0
    for child in document.element.body.iterchildren():
        if child is paragraph._p:
            return min(section_index, len(document.sections) - 1)
        if child.tag == qn("w:p"):
            properties = child.find(qn("w:pPr"))
            if properties is not None and properties.find(qn("w:sectPr")) is not None:
                section_index += 1
    raise ValueError("Image paragraph is not in the main document body")


def _effective_indent_emu(paragraph: Paragraph, attribute: str) -> int:
    direct = getattr(paragraph.paragraph_format, attribute)
    if direct is not None:
        return int(direct)
    style = paragraph.style
    while style is not None:
        inherited = getattr(style.paragraph_format, attribute)
        if inherited is not None:
            return int(inherited)
        style = style.base_style
    return 0


def _available_width_emu(document, paragraph: Paragraph) -> tuple[int, int]:
    section_index = _section_index_for_paragraph(document, paragraph)
    section = document.sections[section_index]
    printable = int(section.page_width) - int(section.left_margin) - int(section.right_margin)
    left_indent = max(0, _effective_indent_emu(paragraph, "left_indent"))
    right_indent = max(0, _effective_indent_emu(paragraph, "right_indent"))
    available = printable - left_indent - right_indent
    if available <= 0:
        raise ValueError("Effective section width and paragraph indents leave no room for an image")
    return section_index, available


def _requested_extent(
    image: Path,
    *,
    width_cm: float | None,
    height_cm: float | None,
    preserve_aspect_ratio: bool,
) -> tuple[int, int]:
    source = DocxImage.from_file(str(image))
    native_width, _ = source.scaled_dimensions()
    source_ratio = source.px_width / source.px_height
    native_height = max(1, int(round(int(native_width) / source_ratio)))
    width = int(round(float(width_cm) * EMU_PER_CM)) if width_cm is not None else None
    height = int(round(float(height_cm) * EMU_PER_CM)) if height_cm is not None else None
    if width is None and height is None:
        return int(native_width), int(native_height)
    if preserve_aspect_ratio:
        if width is not None:
            return width, max(1, int(round(width / source_ratio)))
        assert height is not None
        return max(1, int(round(height * source_ratio))), height
    if width is None:
        assert height is not None
        width = max(1, int(round(height * source_ratio)))
    if height is None:
        height = max(1, int(round(width / source_ratio)))
    return width, height


def _append_semantic_caption(
    document, image_paragraph: Paragraph, label: str, title: str,
    sequence_id: str | None = None, chapter_number: int | None = None
) -> str:
    sequence = sequence_id or label
    paragraph = insert_paragraph_after(image_paragraph, style="Caption")
    return append_caption_content(document, paragraph, label, title, sequence, chapter_number)


def insert_image(
    input_path: str | Path,
    *,
    output_path: str | Path,
    image_path: str | Path,
    selector=None,
    width_cm: float | None = None,
    height_cm: float | None = None,
    alt_text: str | None = None,
    title: str | None = None,
    caption: str | None = None,
    caption_label: str = "Figure",
    caption_sequence: str | None = None,
    caption_chapter_number: int | None = None,
    placement: str = "after",
    alignment: str = "center",
    preserve_aspect_ratio: bool = True,
    auto_shrink: bool = True,
    inline: bool = True,
    dry_run: bool = False,
) -> OperationResult:
    _validate_dimensions(width_cm, height_cm, preserve_aspect_ratio)
    if caption_chapter_number is not None and not caption:
        raise ValueError("caption_chapter_number requires caption text")
    if type(auto_shrink) is not bool:
        raise TypeError("auto_shrink must be a boolean")
    if type(inline) is not bool:
        raise TypeError("inline must be a boolean")
    if not inline:
        raise ValueError("Floating image insertion requires a proven OOXML/native mode")
    if alignment not in _ALIGNMENTS:
        raise ValueError("alignment must be left, center, or right")
    if caption and (not isinstance(caption_label, str) or not caption_label.strip() or len(caption_label) > 80):
        raise ValueError("Caption label must be a non-empty display string of at most 80 characters")
    sequence = caption_sequence or caption_label
    if caption and not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{0,39}", sequence):
        raise ValueError("Caption sequence must be a Word-safe identifier")
    image = Path(image_path)
    if not image.is_file():
        raise FileNotFoundError(image)
    requested_width, requested_height = _requested_extent(
        image,
        width_cm=width_cm,
        height_cm=height_cm,
        preserve_aspect_ratio=preserve_aspect_ratio,
    )
    document = Document(input_path)
    paragraph, replaced_anchor = _image_paragraph(document, selector, placement)
    paragraph.alignment = _ALIGNMENTS[alignment]
    section_index, available_width = _available_width_emu(document, paragraph)
    width = requested_width
    height = requested_height
    shrunk = False
    if auto_shrink and width > available_width:
        scale = available_width / width
        width = available_width
        height = max(1, int(round(height * scale)))
        shrunk = True
    run = paragraph.add_run()
    shape = run.add_picture(
        str(image),
        width=width,
        height=height,
    )
    properties = shape._inline.docPr
    if alt_text is not None:
        properties.set("descr", alt_text)
    if title is not None:
        properties.set("title", title)
    caption_field = None
    if caption:
        caption_field = _append_semantic_caption(document, paragraph, caption_label.strip(), caption, sequence, caption_chapter_number)
    if replaced_anchor is not None and placement == "replace":
        replaced_anchor._p.getparent().remove(replaced_anchor._p)
    drawing = shape._inline.getparent()
    image_index = next(
        index
        for index, candidate in enumerate(_drawings(document))
        if candidate is drawing
    )
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult(
        "DRY_RUN" if dry_run else "SUCCESS",
        "image.insert",
        "python-docx",
        Fidelity.HIGH,
        None if dry_run else output,
        manual_review_required=caption_chapter_number is not None,
        data={
            "image_index": image_index,
            "inline": True,
            "placement": placement,
            "alignment": alignment,
            "section_index": section_index,
            "width_emu": width,
            "height_emu": height,
            "width_cm": round(width / EMU_PER_CM, 6),
            "height_cm": round(height / EMU_PER_CM, 6),
            "aspect_ratio": round(width / height, 8),
            "available_width_emu": available_width,
            "auto_shrunk": shrunk,
            "caption_field": caption_field,
            "caption_chapter_number": caption_chapter_number,
            "caption_label": caption_label.strip() if caption else None,
        },
    )


def resize_image(
    input_path,
    *,
    output_path,
    image_index: int,
    width_cm: float | None = None,
    height_cm: float | None = None,
    preserve_aspect_ratio: bool = True,
    dry_run: bool = False,
):
    if width_cm is None and height_cm is None:
        raise ValueError("width_cm or height_cm is required")
    _validate_dimensions(width_cm, height_cm, preserve_aspect_ratio)
    document = Document(input_path)
    try:
        drawing = _drawings(document)[image_index]
    except IndexError as exc:
        raise ValueError(f"Image index out of range: {image_index}") from exc
    extents = drawing.xpath(".//wp:extent")
    transform_extents = drawing.xpath(".//a:xfrm/a:ext")
    if not extents:
        raise ValueError("Image has no extent")
    current_width = int(extents[0].get("cx"))
    current_height = int(extents[0].get("cy"))
    width = int(width_cm * EMU_PER_CM) if width_cm is not None else None
    height = int(height_cm * EMU_PER_CM) if height_cm is not None else None
    if preserve_aspect_ratio:
        if width is not None and height is None:
            height = round(current_height * width / current_width)
        elif height is not None and width is None:
            width = round(current_width * height / current_height)
    width = width or current_width
    height = height or current_height
    for extent in [*extents, *transform_extents]:
        extent.set("cx", str(width))
        extent.set("cy", str(height))
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success(
        "image.resize",
        "ooxml",
        Fidelity.HIGH,
        {"image_index": image_index, "width_emu": width, "height_emu": height},
        output_path=None if dry_run else output,
    )


def remove_image(input_path, *, output_path, image_index: int, dry_run=False):
    document = Document(input_path)
    try:
        drawing = _drawings(document)[image_index]
    except IndexError as exc:
        raise ValueError(f"Image index out of range: {image_index}") from exc
    embeds = drawing.xpath(".//a:blip/@r:embed")
    run = drawing.getparent()
    while run is not None and run.tag != qn("w:r"):
        run = run.getparent()
    if run is None:
        raise ValueError("Image drawing is not contained in a run")
    run.getparent().remove(run)
    for relationship_id in embeds:
        if not document.element.body.xpath(f'.//a:blip[@r:embed="{relationship_id}"]'):
            document.part.drop_rel(relationship_id)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success(
        "image.remove", "ooxml", Fidelity.HIGH, {"image_index": image_index}, output_path=None if dry_run else output
    )


def replace_image(input_path, *, output_path, image_index: int, image_path: str | Path, dry_run=False):
    document = Document(input_path)
    try:
        drawing = _drawings(document)[image_index]
    except IndexError as exc:
        raise ValueError(f"Image index out of range: {image_index}") from exc
    embeds = drawing.xpath(".//a:blip/@r:embed")
    if len(embeds) != 1:
        raise ValueError("Selected drawing does not have exactly one image relationship")
    part = document.part.related_parts[embeds[0]]
    replacement = Path(image_path)
    if not replacement.is_file():
        raise FileNotFoundError(replacement)
    part._blob = replacement.read_bytes()
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("image.replace", "ooxml", Fidelity.STRUCTURAL, {"image_index": image_index}, output_path=None if dry_run else output, manual_review_required=True)


def set_image_metadata(input_path, *, output_path, image_index: int, alt_text=None, title=None, dry_run=False):
    document = Document(input_path)
    try:
        properties = _doc_properties(_drawings(document)[image_index])
    except IndexError as exc:
        raise ValueError(f"Image index out of range: {image_index}") from exc
    if alt_text is not None:
        properties.set("descr", alt_text)
    if title is not None:
        properties.set("title", title)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("image.metadata", "ooxml", Fidelity.HIGH, {"image_index": image_index}, output_path=None if dry_run else output)
