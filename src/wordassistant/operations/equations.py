from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement

from .common import atomic_save
from .fields import append_field
from .models import Fidelity, OperationResult


def _math_run(text):
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    return run


def _append_expression(math, latex):
    fraction = re.fullmatch(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", latex)
    if fraction:
        element = OxmlElement("m:f")
        numerator = OxmlElement("m:num")
        numerator.append(_math_run(fraction.group(1)))
        denominator = OxmlElement("m:den")
        denominator.append(_math_run(fraction.group(2)))
        element.extend([numerator, denominator])
        math.append(element)
        return
    superscript = re.fullmatch(r"(.+)\^\{?([^{}]+)\}?", latex)
    if superscript:
        element = OxmlElement("m:sSup")
        base = OxmlElement("m:e")
        base.append(_math_run(superscript.group(1)))
        sup = OxmlElement("m:sup")
        sup.append(_math_run(superscript.group(2)))
        element.extend([base, sup])
        math.append(element)
        return
    math.append(_math_run(latex))


def insert_equation(input_path, *, output_path, latex, number=False, dry_run=False):
    if not latex or not isinstance(latex, str):
        raise ValueError("latex must be a non-empty string")
    document = Document(input_path)
    paragraph = document.add_paragraph()
    math_paragraph = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    _append_expression(math, latex)
    math_paragraph.append(math)
    paragraph._p.append(math_paragraph)
    if number:
        paragraph.add_run(" (")
        append_field(paragraph, "SEQ Equation \\* ARABIC", "1")
        paragraph.add_run(")")
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("equation.insert", "ooxml", Fidelity.STRUCTURAL, {"latex": latex, "numbered": number}, output_path=None if dry_run else output, manual_review_required=True)

