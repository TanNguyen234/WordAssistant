from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .common import atomic_save
from .models import Fidelity, OperationResult


_MODES = {
    "read_only": "readOnly",
    "forms_only": "forms",
    "comments_only": "comments",
    "tracked_changes_only": "trackedChanges",
}


def protect_document(input_path, *, output_path, mode, password=None, dry_run=False):
    if password:
        raise ValueError("Secure password protection requires the Word-native backend")
    if mode not in _MODES:
        raise ValueError(f"Unsupported protection mode: {mode}")
    document = Document(input_path)
    settings = document.settings.element
    current = settings.find(qn("w:documentProtection"))
    if current is not None:
        settings.remove(current)
    protection = OxmlElement("w:documentProtection")
    protection.set(qn("w:edit"), _MODES[mode])
    protection.set(qn("w:enforcement"), "1")
    settings.append(protection)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("protect.set", "ooxml", Fidelity.STRUCTURAL, {"mode": mode, "secure_encryption": False}, output_path=None if dry_run else output, manual_review_required=True)


def inspect_protection(input_path):
    document = Document(input_path)
    element = document.settings.element.find(qn("w:documentProtection"))
    if element is None:
        return {"enforced": False, "mode": None, "secure_encryption": False}
    reverse = {value: key for key, value in _MODES.items()}
    return {
        "enforced": element.get(qn("w:enforcement")) in {"1", "true", "on"},
        "mode": reverse.get(element.get(qn("w:edit")), element.get(qn("w:edit"))),
        "secure_encryption": any(element.get(qn(key)) for key in ("w:hash", "w:cryptProviderType", "w:salt")),
    }


def unprotect_document(input_path, *, output_path, dry_run=False):
    document = Document(input_path)
    settings = document.settings.element
    element = settings.find(qn("w:documentProtection"))
    if element is not None:
        if any(element.get(qn(key)) for key in ("w:hash", "w:cryptProviderType", "w:salt")):
            raise ValueError("Password-protected document requires the Word-native backend")
        settings.remove(element)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("protect.remove", "ooxml", Fidelity.STRUCTURAL, {}, output_path=None if dry_run else output)

