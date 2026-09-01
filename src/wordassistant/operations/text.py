from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from docx import Document

from ..document.edit import _collect_all_paragraphs, _detect_complex_elements
from ..selectors import Selector, resolve_paragraphs
from .common import atomic_save
from .models import Fidelity, OperationResult


def _replace_spans(paragraph, spans: list[tuple[int, int, str]]) -> int:
    runs = paragraph.runs
    if not spans:
        return 0
    if not runs:
        paragraph.add_run(paragraph.text)
        runs = paragraph.runs

    offsets: list[tuple[int, int]] = []
    offset = 0
    for run in runs:
        offsets.append((offset, offset + len(run.text)))
        offset += len(run.text)

    for start, end, replacement in reversed(spans):
        start_index = next(
            (i for i, (left, right) in enumerate(offsets) if left <= start < right),
            None,
        )
        end_index = next(
            (i for i, (left, right) in enumerate(offsets) if left < end <= right),
            None,
        )
        if start_index is None and start == end == offset and runs:
            runs[-1].text += replacement
            continue
        if start_index is None or end_index is None:
            raise ValueError("Replacement span does not map to paragraph runs")
        start_run = runs[start_index]
        end_run = runs[end_index]
        prefix = start_run.text[: start - offsets[start_index][0]]
        suffix = end_run.text[end - offsets[end_index][0] :]
        if start_index == end_index:
            start_run.text = prefix + replacement + suffix
        else:
            start_run.text = prefix + replacement
            for index in range(start_index + 1, end_index):
                runs[index].text = ""
            end_run.text = suffix
    return len(spans)


def _target_paragraphs(document, selector):
    if selector is None:
        return _collect_all_paragraphs(document)
    return resolve_paragraphs(document, Selector.from_value(selector))


def edit_text(
    input_path: str | Path,
    *,
    output_path: str | Path,
    action: str,
    selector: Selector | str | dict[str, Any] | None = None,
    text: str | None = None,
    pattern: str | None = None,
    replacement: str = "",
    regex: bool = False,
    case_sensitive: bool = True,
    whole_word: bool = False,
    max_count: int | None = None,
    dry_run: bool = False,
) -> OperationResult:
    document = Document(input_path)
    paragraphs = _target_paragraphs(document, selector)
    changes = 0

    if action in {"append", "prepend", "insert"}:
        if text is None:
            raise ValueError(f"{action} requires text")
        for paragraph in paragraphs:
            if _detect_complex_elements(paragraph._p):
                raise ValueError("Text insertion target contains complex Word structures")
            if action == "append":
                paragraph.add_run(text)
            elif action == "prepend":
                if paragraph.runs:
                    paragraph.runs[0].text = text + paragraph.runs[0].text
                else:
                    paragraph.add_run(text)
            else:
                paragraph.add_run(text)
            changes += 1
    elif action == "replace":
        if pattern is None or pattern == "":
            raise ValueError("replace requires a non-empty pattern")
        expression = pattern if regex else re.escape(pattern)
        if whole_word:
            expression = rf"(?<!\w)(?:{expression})(?!\w)"
        flags = 0 if case_sensitive else re.IGNORECASE
        compiled = re.compile(expression, flags)
        for paragraph in paragraphs:
            matches = list(compiled.finditer(paragraph.text))
            if max_count is not None:
                if max_count < 0:
                    raise ValueError("max_count must be non-negative")
                matches = matches[: max(0, max_count - changes)]
            if not matches:
                continue
            complex_elements = _detect_complex_elements(paragraph._p)
            if complex_elements:
                raise ValueError(
                    "Text replacement target contains complex Word structures: "
                    + ", ".join(sorted(complex_elements))
                )
            spans = [
                (match.start(), match.end(), match.expand(replacement))
                for match in matches
            ]
            changes += _replace_spans(paragraph, spans)
            if max_count is not None and changes >= max_count:
                break
    else:
        raise ValueError(f"Unsupported text action: {action}")

    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult(
        status="DRY_RUN" if dry_run else "SUCCESS",
        operation=f"text.{action}",
        backend_used="python-docx",
        fidelity=Fidelity.HIGH,
        output_path=None if dry_run else output,
        data={"replacements" if action == "replace" else "changes": changes},
    )
