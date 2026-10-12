from __future__ import annotations

from pathlib import Path
import re
import json
import unicodedata

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..evidence.registry import EvidenceRegistry
from ..selectors import resolve_single_paragraph, _heading_level as _semantic_heading_level
from .common import atomic_save
from .models import Fidelity, OperationResult


_TAG_PREFIX = "wordassistant:citation:"


def _citation_elements(document):
    result = []
    for sdt in document.element.body.xpath(".//w:sdt"):
        properties = sdt.find(qn("w:sdtPr"))
        tag = properties.find(qn("w:tag")) if properties is not None else None
        if tag is not None and tag.get(qn("w:val"), "").startswith(_TAG_PREFIX):
            result.append(sdt)
    return result


def _citation_tag(element):
    properties = element.find(qn("w:sdtPr"))
    tag = properties.find(qn("w:tag")) if properties is not None else None
    return tag.get(qn("w:val"), "") if tag is not None else ""


def _citation_numbers(document, *, displayed=False):
    numbers = {}
    owners = {}
    for element in _citation_elements(document):
        evidence_id = _citation_tag(element)[len(_TAG_PREFIX) :]
        text = ''.join(node.text or '' for node in element.iter(qn('w:t'))).strip()
        match = re.fullmatch(r'\[(?:SYNTHETIC-)?(\d+)\]', text)
        if match:
            number = int(match[1])
            if number < 1 or (number in owners and owners[number] != evidence_id) or (evidence_id in numbers and displayed and numbers[evidence_id] != number):
                raise ValueError('Conflicting citation numeric labels')
            owners[number] = evidence_id
        if evidence_id not in numbers:
            numbers[evidence_id] = int(match[1]) if displayed and match else len(numbers) + 1
    return numbers


_POLICY_NAME = 'wordassistant.citation-policy'
_AUTHOR_ORDER = 'vietnamese_given_name_then_english_family_name'


def _citation_policy(document):
    values = document.settings.element.xpath('./w:docVars/w:docVar')
    matches = [node for node in values if node.get(qn('w:name')) == _POLICY_NAME]
    if not matches:
        return {}
    if len(matches) != 1:
        raise ValueError('Conflicting citation policies')
    policy = json.loads(matches[0].get(qn('w:val'), ''))
    if not isinstance(policy, dict) or policy.get('order') not in ('citation_order', _AUTHOR_ORDER) or not isinstance(policy.get('style'), str) or not policy['style'].strip() or type(policy.get('require_reviewed')) is not bool or (policy.get('heading') is not None and (not isinstance(policy['heading'], str) or not policy['heading'].strip())):
        raise ValueError('Invalid citation policy')
    return policy


def _set_citation_policy(document, policy):
    settings = document.settings.element
    container = settings.find(qn('w:docVars'))
    if container is None:
        container = OxmlElement('w:docVars')
        settings.insert_element_before(container, 'w:rsids', 'w:mathPr', 'w:themeFontLang', 'w:clrSchemeMapping', 'w:doNotIncludeSubdocsInStats')
    for node in list(container):
        if node.get(qn('w:name')) == _POLICY_NAME:
            container.remove(node)
    node = OxmlElement('w:docVar')
    node.set(qn('w:name'), _POLICY_NAME)
    node.set(qn('w:val'), json.dumps(policy, ensure_ascii=False))
    container.append(node)


def _bibliography_sort_key(record):
    first_author = record.get('authors', [''])[0]
    language = str(record.get('language', '')).casefold()
    tokens = str(first_author).replace(',', ' ').split()
    key = record.get('author_sort_key') or (tokens[-1] if language.startswith('vi') and tokens else tokens[0] if tokens else first_author)
    def fold(value):
        return ''.join(c for c in unicodedata.normalize('NFD', str(value).casefold()) if not unicodedata.combining(c))
    return (0 if language.startswith('vi') else 1 if language.startswith('en') else 2, fold(key), fold(first_author), fold(record.get('title', '')))


def _ordered_citation_numbers(document, registry, policy, *, repair=False):
    ids = list(dict.fromkeys(_citation_tag(element)[len(_TAG_PREFIX):] for element in _citation_elements(document))) if repair else list(_citation_numbers(document, displayed=True))
    for evidence_id in ids:
        registry.assert_citable(registry.get(evidence_id), require_reviewed=policy.get('require_reviewed', False))
    if policy.get('order') == _AUTHOR_ORDER:
        ids.sort(key=lambda evidence_id: _bibliography_sort_key(registry.get(evidence_id)))
    return {evidence_id: number for number, evidence_id in enumerate(ids, 1)}


def _assert_legacy_numbering_policy(document, policy):
    if policy:
        return
    try:
        numbers = _citation_numbers(document, displayed=True)
    except ValueError as exc:
        raise ValueError('Conflicting citation labels: explicitly regenerate the bibliography with style and heading to adopt a numbering policy') from exc
    if list(numbers.values()) != list(range(1, len(numbers) + 1)):
        raise ValueError('Legacy citation numbering is ambiguous: compose with explicit style/order or regenerate the bibliography with explicit style and heading to adopt a numbering policy')


def _renumber_citations(document, numbers, registry=None, style=None):
    for element in _citation_elements(document):
        evidence_id = _citation_tag(element)[len(_TAG_PREFIX):]
        nodes = list(element.iter(qn('w:t')))
        text = ''.join(node.text or '' for node in nodes).strip()
        if not nodes:
            continue
        if registry is not None:
            replacement = _format_in_text(registry.get(evidence_id), style, numbers[evidence_id])
        elif re.fullmatch(r'\[(?:SYNTHETIC-)?\d+\]', text):
            replacement = f"[{'SYNTHETIC-' if 'SYNTHETIC-' in text else ''}{numbers[evidence_id]}]"
        else:
            continue
        nodes[0].text = replacement
        for node in nodes[1:]:
            node.text = ''


def _bibliography_entries(registry, numbers, style, *, require_reviewed=False):
    entries = []
    last_group = None
    for evidence_id, number in numbers.items():
        record = registry.get(evidence_id)
        registry.assert_citable(record, require_reviewed=require_reviewed)
        if style.casefold() == 'hcmunre_numeric':
            language = str(record.get('language', '')).casefold()
            group = 'Tiếng Việt:' if language.startswith('vi') else 'Tiếng Anh:' if language.startswith('en') else 'Ngôn ngữ khác:'
            if group != last_group:
                entries.append({'group_heading': group})
                last_group = group
        entries.append(_format_bibliography_runs(record, style, number))
    return entries


def _format_in_text(record, style, number):
    if record.get("synthetic") is True:
        return f"[SYNTHETIC-{number}]"
    normalized = style.lower()
    if normalized in {"numeric", "ieee", "numeric_brackets", "hcmunre_numeric"}:
        return f"[{number}]"
    author = record["authors"][0].split()[-1]
    if normalized in {"apa", "chicago"}:
        return f"({author}, {record['year']})"
    if normalized == "mla":
        return f"({author})"
    raise ValueError(f"Unsupported citation style: {style}")


def _format_bibliography(record, style, number):
    if style.lower() == "hcmunre_numeric":
        return "".join(
            part["text"]
            for part in _format_bibliography_runs(record, style, number)
        )
    authors = ", ".join(record["authors"])
    marker = "[SYNTHETIC TEST DATA] " if record.get("synthetic") is True else ""
    normalized = style.lower()
    if normalized in {"numeric", "ieee", "numeric_brackets", "hcmunre_numeric"}:
        return f'[{number}] {marker}{authors}, "{record["title"]}," {record["year"]}. {record["url"]}'
    if normalized == "apa":
        return f"{marker}{authors}. ({record['year']}). {record['title']}. {record['url']}"
    if normalized == "mla":
        return f'{marker}{authors}. "{record["title"]}." {record["year"]}, {record["url"]}.'
    if normalized == "chicago":
        return f"{marker}{authors}. {record['title']}. {record['year']}. {record['url']}."
    raise ValueError(f"Unsupported bibliography style: {style}")


def _format_bibliography_runs(record, style, number):
    if style.lower() != "hcmunre_numeric":
        return [{"text": _format_bibliography(record, style, number)}]

    authors = record["authors"]
    if len(authors) == 1:
        author_text = authors[0]
    elif len(authors) == 2:
        conjunction = " và " if str(record.get("language", "")).casefold().startswith("vi") else " and "
        author_text = conjunction.join(authors)
    else:
        conjunction = " và " if str(record.get("language", "")).casefold().startswith("vi") else " and "
        author_text = ", ".join(authors[:-1]) + conjunction + authors[-1]

    parts = [
        {"text": f"[{number}] "},
        *([{"text": "[SYNTHETIC TEST DATA] "}] if record.get("synthetic") is True else []),
        {"text": author_text, "bold": True},
        {"text": ". "},
        {"text": record["title"], "italic": True},
    ]
    details = []
    if record.get("container_title"):
        details.append({
            "text": f"In: {record['container_title']}",
            "italic": True,
        })
    if record.get("volume"):
        details.append({"text": f"Vol. {record['volume']}"})
    if record.get("issue"):
        details.append({"text": f"No. {record['issue']}"})
    if record.get("pages"):
        details.append({"text": f"pp. {record['pages']}"})
    if record.get("publisher"):
        details.append({"text": str(record["publisher"])})
    if record.get("publication_place"):
        details.append({"text": str(record["publication_place"])})
    details.append({"text": str(record["year"])})
    for detail in details:
        parts.append({"text": ". "})
        parts.append(detail)

    doi_url = f"https://doi.org/{record['doi']}" if record.get("doi") else record.get("url")
    if doi_url:
        parts.extend(({"text": ". "}, {"text": str(doi_url)}))
    parts.append({"text": "."})
    return parts


def insert_citation(
    input_path,
    *,
    output_path,
    evidence_registry,
    evidence_id,
    style=None,
    selector=None,
    dry_run=False,
):
    registry = EvidenceRegistry(evidence_registry)
    record = registry.get(evidence_id)
    document = Document(input_path)
    policy = _citation_policy(document)
    _assert_legacy_numbering_policy(document, policy)
    style = style or policy.get('style', 'numeric')
    if policy.get('require_reviewed') and style.casefold() != policy['style'].casefold():
        raise ValueError('citation style cannot override the source-backed document policy')
    policy = {**policy, 'style': style, 'order': policy.get('order', 'citation_order'), 'require_reviewed': policy.get('require_reviewed', False)}
    registry.assert_citable(record, require_reviewed=policy['require_reviewed'])
    numbers = _citation_numbers(document, displayed=True)
    number = numbers.get(evidence_id, max(numbers.values(), default=0) + 1)
    text = _format_in_text(record, style, number)
    paragraph = (
        document.add_paragraph()
        if selector is None
        else resolve_single_paragraph(document, selector)
    )
    if selector is not None and policy.get('heading'):
        if _bibliography_contains_paragraph(document, policy['heading'], paragraph):
            raise ValueError('Citation anchor must be outside the managed bibliography')
    if selector is None and policy.get('heading'):
        headings = [p for p in document.paragraphs if p.text.strip().casefold() == policy['heading'].strip().casefold() and _heading_level(p) is not None]
        if len(headings) != 1:
            raise ValueError('Bibliography heading must be unique for citation insertion')
        headings[0]._p.addprevious(paragraph._p)
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), f"{_TAG_PREFIX}{evidence_id}")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), f"Citation {evidence_id}")
    properties.extend([alias, tag])
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    content.append(run)
    sdt.extend([properties, content])
    if paragraph.text:
        paragraph.add_run(" ")
    paragraph._p.append(sdt)
    numbers = _ordered_citation_numbers(document, registry, policy)
    _renumber_citations(document, numbers, registry, style)
    number = numbers[evidence_id]
    if policy.get('heading'):
        _replace_bibliography_content(document, policy['heading'], _bibliography_entries(registry, numbers, style, require_reviewed=policy['require_reviewed']))
    _set_citation_policy(document, policy)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success(
        "reference.cite",
        "ooxml",
        Fidelity.HIGH,
        {
            "evidence_id": evidence_id,
            "number": number,
            "style": style,
            "synthetic": record.get("synthetic") is True,
            "anchored": selector is not None,
        },
        output_path=None if dry_run else output,
    )


def list_citations(input_path):
    document = Document(input_path)
    result = []
    for index, element in enumerate(_citation_elements(document)):
        tag = _citation_tag(element)
        result.append(
            {
                "index": index,
                "evidence_id": tag[len(_TAG_PREFIX) :],
                "text": "".join(value.text or "" for value in element.iter(qn("w:t"))),
            }
        )
    return result


def remove_citation(input_path, *, output_path, citation_index: int, dry_run=False):
    document = Document(input_path)
    policy = _citation_policy(document)
    _assert_legacy_numbering_policy(document, policy)
    old_numbers = _citation_numbers(document, displayed=True)
    if citation_index < 0:
        raise ValueError(f"Citation index out of range: {citation_index}")
    try:
        element = _citation_elements(document)[citation_index]
    except IndexError as exc:
        raise ValueError(f"Citation index out of range: {citation_index}") from exc
    element.getparent().remove(element)
    ids = list(_citation_numbers(document, displayed=True))
    if policy.get('order') == _AUTHOR_ORDER:
        ids.sort(key=old_numbers.__getitem__)
    numbers = {key: number for number, key in enumerate(ids, 1)}
    _renumber_citations(document, numbers)
    if policy.get('heading'):
        _relabel_existing_bibliography(document, policy['heading'], old_numbers, numbers)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success("reference.remove", "ooxml", Fidelity.HIGH, {"citation_index": citation_index}, output_path=None if dry_run else output)


def _heading_level(paragraph) -> int | None:
    level = _semantic_heading_level(paragraph)
    if level is not None:
        return level
    name = paragraph.style.name if paragraph.style is not None else ""
    return 1 if name == "Title" else None


def _bibliography_contains_paragraph(document, heading, paragraph, *, allow_missing=False):
    paragraphs = document.paragraphs
    matches = [p for p in paragraphs
               if p.text.strip().casefold() == heading.strip().casefold()
               and _heading_level(p) is not None]
    if not matches and allow_missing:
        return False
    if len(matches) != 1:
        raise ValueError('Managed bibliography heading must be unique')
    target = matches[0]
    level = _heading_level(target)
    by_element = {p._p: p for p in paragraphs}
    node = target._p
    while node is not None:
        current = by_element.get(node)
        if current is not None:
            current_level = _heading_level(current)
            if node is not target._p and current_level is not None and current_level <= level:
                break
            if node is paragraph._p:
                return True
        node = node.getnext()
    return False


def _replace_bibliography_content(document, heading: str, lines: list):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().casefold() == heading.strip().casefold()
        and _heading_level(paragraph) is not None
    ]
    if len(matches) > 1:
        raise ValueError(
            f"Bibliography heading must be unique; found {len(matches)} exact matches for {heading!r}"
        )
    if matches:
        target = matches[0]
    else:
        target = document.add_heading(heading, level=1)

    target_level = _heading_level(target) or 1
    node = target._p.getnext()
    while node is not None:
        next_node = node.getnext()
        paragraph = next(
            (item for item in document.paragraphs if item._p is node), None
        )
        if paragraph is not None:
            level = _heading_level(paragraph)
            if level is not None and level <= target_level:
                break
            node.getparent().remove(node)
        node = next_node

    cursor = target._p
    bibliography_style = (
        "WA References"
        if "WA References" in document.styles
        else "Bibliography"
        if "Bibliography" in document.styles
        else "Normal"
    )
    for line in lines:
        if isinstance(line, dict) and "group_heading" in line:
            group_style = (
                "WA Reference Language Heading"
                if "WA Reference Language Heading" in document.styles
                else bibliography_style
            )
            paragraph = document.add_paragraph(style=group_style)
            paragraph.add_run(line["group_heading"])
        else:
            paragraph = document.add_paragraph(style=bibliography_style)
        if isinstance(line, str):
            paragraph.add_run(line)
        elif isinstance(line, list):
            for run_spec in line:
                run = paragraph.add_run(run_spec["text"])
                if "bold" in run_spec:
                    run.bold = bool(run_spec["bold"])
                if "italic" in run_spec:
                    run.italic = bool(run_spec["italic"])
        cursor.addnext(paragraph._p)
        cursor = paragraph._p
    return target


def _relabel_existing_bibliography(document, heading, old_numbers, numbers):
    matches = [p for p in document.paragraphs if p.text.strip().casefold() == heading.strip().casefold() and _heading_level(p) is not None]
    if len(matches) != 1:
        raise ValueError('Bibliography heading must be unique for citation removal')
    target = matches[0]
    level = _heading_level(target) or 1
    paragraphs = []
    cursor = target._p.getnext()
    while cursor is not None:
        paragraph = next((p for p in document.paragraphs if p._p is cursor), None)
        if paragraph is not None:
            current_level = _heading_level(paragraph)
            if current_level is not None and current_level <= level:
                break
            paragraphs.append(paragraph)
        cursor = cursor.getnext()
    owners = {number: key for key, number in old_numbers.items()}
    rows = {}
    group = None
    group_labels = {'Tiếng Việt:', 'Tiếng Anh:', 'Ngôn ngữ khác:'}
    for paragraph in paragraphs:
        if paragraph.text in group_labels:
            group = paragraph
            continue
        match = re.match(r'^\[(\d+)\]', paragraph.text)
        if not match:
            # Author-date entries have no numeric label. Their original order
            # is the order captured when the managed bibliography was generated.
            number = len(rows) + 1
        else:
            number = int(match[1])
        if number not in owners or owners[number] in rows:
            raise ValueError('Conflicting citation bibliography labels')
        rows[owners[number]] = (paragraph, group)
    if set(rows) != set(old_numbers):
        raise ValueError('Conflicting citation bibliography entries')
    for paragraph in paragraphs:
        paragraph._p.getparent().remove(paragraph._p)
    cursor = target._p
    last_group = None
    for key, number in numbers.items():
        paragraph, group = rows[key]
        if group is not None and group is not last_group:
            cursor.addnext(group._p)
            cursor = group._p
            last_group = group
        nodes = list(paragraph._p.iter(qn('w:t')))
        prefix = re.match(r'^\[\d+\]', paragraph.text)
        if prefix:
            remaining = len(prefix[0])
            for node in nodes:
                value = node.text or ''
                take = min(remaining, len(value))
                node.text = value[take:]
                remaining -= take
                if remaining == 0:
                    break
            nodes[0].text = f'[{number}]' + (nodes[0].text or '')
        cursor.addnext(paragraph._p)
        cursor = paragraph._p


def generate_bibliography(input_path, *, output_path, evidence_registry, style=None, heading=None, dry_run=False):
    document = Document(input_path)
    policy = _citation_policy(document)
    repair = not policy and style is not None
    if not repair:
        _assert_legacy_numbering_policy(document, policy)
    style = style or policy.get('style', 'IEEE')
    heading = heading or policy.get('heading') or 'REFERENCES'
    if policy.get('require_reviewed') and (style.casefold() != policy['style'].casefold() or (policy.get('heading') and heading.casefold() != policy['heading'].casefold())):
        raise ValueError('bibliography cannot override the source-backed document policy')
    policy = {**policy, 'style': style, 'heading': heading,
              'order': policy.get('order', 'citation_order'),
              'require_reviewed': policy.get('require_reviewed', False)}
    registry = EvidenceRegistry(evidence_registry)
    numbers = _ordered_citation_numbers(document, registry, policy, repair=repair)
    _renumber_citations(document, numbers, registry, style)
    lines = _bibliography_entries(registry, numbers, style, require_reviewed=policy['require_reviewed'])
    _replace_bibliography_content(document, heading, lines)
    _set_citation_policy(document, policy)
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return OperationResult.success('reference.bibliography', 'python-docx', Fidelity.HIGH, {'entries': len(numbers), 'style': style, 'heading': heading}, output_path=None if dry_run else output)
