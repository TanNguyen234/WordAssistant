from __future__ import annotations

from dataclasses import dataclass
import json
import re
from typing import Any

from docx.document import Document as Doc
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SELECTOR_SCHEMA = {
    "allowed_keys": ["type", "text", "index", "style", "case_sensitive"],
    "types": {
        "document": {
            "required": ["type"],
            "example": {"type": "document"},
            "meaning": "all top-level paragraphs",
        },
        "heading": {
            "required": ["type", "text"],
            "example": {"type": "heading", "text": "Introduction"},
            "meaning": "the uniquely named real Word Heading paragraph",
        },
        "paragraph": {
            "required": ["type", "index"],
            "example": {"type": "paragraph", "index": 0},
            "meaning": "zero-based top-level paragraph index; negative indexes are allowed",
        },
        "text": {
            "required": ["type", "text"],
            "example": {"type": "text", "text": "Exact anchor text"},
            "meaning": "paragraph whose full text matches exactly",
        },
        "style": {
            "required": ["type", "style"],
            "example": {"type": "style", "style": "Caption"},
            "meaning": "paragraphs with the named Word style",
        },
    },
    "defaults": {"case_sensitive": False},
    "single_target_policy": "fail closed when zero or multiple paragraphs match",
}

SELECTOR_HELP = (
    "Selector JSON; keys: type, text, index, style, case_sensitive. "
    'Examples: {"type":"heading","text":"Introduction"}; '
    '{"type":"text","text":"Exact anchor text"}; '
    '{"type":"paragraph","index":0}'
)


@dataclass(frozen=True)
class Selector:
    type: str
    text: str | None = None
    index: int | None = None
    style: str | None = None
    case_sensitive: bool = False

    @classmethod
    def from_value(cls, value: "Selector | str | dict[str, Any]") -> "Selector":
        if isinstance(value, cls):
            return value
        if isinstance(value, str):
            try:
                value = json.loads(value)
            except json.JSONDecodeError as exc:
                raise ValueError("Selector string must be valid JSON") from exc
        if not isinstance(value, dict):
            raise TypeError("Selector must be a Selector, JSON object, or mapping")
        allowed = set(SELECTOR_SCHEMA["allowed_keys"])
        unknown = set(value) - allowed
        if unknown:
            raise ValueError(f"Unknown selector keys: {sorted(unknown)}")
        selector_type = value.get("type")
        if not isinstance(selector_type, str):
            raise TypeError("Selector type must be a string")
        if selector_type not in SELECTOR_SCHEMA["types"]:
            raise ValueError(f"Unsupported selector type: {selector_type}")
        required = set(SELECTOR_SCHEMA["types"][selector_type]["required"])
        missing = required - set(value)
        if missing:
            raise ValueError(
                f"{selector_type.title()} selector requires: {sorted(missing)}"
            )
        allowed_for_type = required | (
            {"case_sensitive"} if selector_type in {"heading", "text", "style"} else set()
        )
        irrelevant = set(value) - allowed_for_type
        if irrelevant:
            raise ValueError(
                f"Selector type {selector_type!r} does not accept keys: {sorted(irrelevant)}"
            )
        if "case_sensitive" in value and type(value["case_sensitive"]) is not bool:
            raise TypeError("Selector case_sensitive must be a boolean")
        if selector_type == "paragraph" and (
            type(value["index"]) is not int
        ):
            raise TypeError("Paragraph selector index must be an integer")
        for key in ("text", "style"):
            if key in required and (
                not isinstance(value[key], str) or not value[key].strip()
            ):
                raise ValueError(f"Selector {key} must be a non-empty string")
        selector = cls(**value)
        return selector


class SelectorNotFound(LookupError):
    def __init__(self, selector: Selector):
        super().__init__(f"Selector did not match: {selector}")
        self.selector = selector


class SelectorAmbiguous(LookupError):
    def __init__(self, selector: Selector, match_count: int):
        super().__init__(f"Selector matched {match_count} objects: {selector}")
        self.selector = selector
        self.match_count = match_count


def _outline_level(paragraph) -> int | None:
    """Resolve an explicit outline level through paragraph and style inheritance."""
    properties = paragraph._p.pPr
    outline = properties.find(qn('w:outlineLvl')) if properties is not None else None
    if outline is not None:
        return int(outline.get(qn('w:val')))
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        properties = style.element.pPr
        outline = properties.find(qn('w:outlineLvl')) if properties is not None else None
        if outline is not None:
            return int(outline.get(qn('w:val')))
        style = style.base_style
    return None


def _heading_level(paragraph) -> int | None:
    """Resolve real headings; explicit body outline level 9 suppresses inheritance."""
    outline_level = _outline_level(paragraph)
    if outline_level is not None:
        return outline_level + 1 if 0 <= outline_level <= 8 else None
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        match = re.fullmatch(r"Heading\s*([1-9])", style.name, flags=re.IGNORECASE)
        if match:
            return int(match.group(1))
        style = style.base_style
    return None


def _equal(left: str, right: str, case_sensitive: bool) -> bool:
    return left == right if case_sensitive else left.casefold() == right.casefold()


def resolve_paragraphs(
    document: Doc,
    selector: Selector | str | dict[str, Any],
    *,
    require_single: bool = False,
) -> list:
    selector = Selector.from_value(selector)
    paragraphs = list(document.paragraphs)

    if selector.type == "document":
        matches = paragraphs
    elif selector.type == "paragraph":
        if selector.index is None:
            raise ValueError("Paragraph selector requires index")
        matches = (
            [paragraphs[selector.index]]
            if -len(paragraphs) <= selector.index < len(paragraphs)
            else []
        )
    elif selector.type == "style":
        expected = selector.style or selector.text
        if not expected:
            raise ValueError("Style selector requires style or text")
        matches = [
            paragraph
            for paragraph in paragraphs
            if paragraph.style is not None
            and _equal(paragraph.style.name, expected, selector.case_sensitive)
        ]
    elif selector.type == "text":
        if selector.text is None:
            raise ValueError("Text selector requires text")
        matches = [
            paragraph
            for paragraph in paragraphs
            if _equal(paragraph.text, selector.text, selector.case_sensitive)
        ]
    elif selector.type == "heading":
        if selector.text is None:
            raise ValueError("Heading selector requires text")
        heading_matches = [
            index
            for index, paragraph in enumerate(paragraphs)
            if _heading_level(paragraph) is not None
            and _equal(paragraph.text, selector.text, selector.case_sensitive)
        ]
        if len(heading_matches) > 1:
            raise SelectorAmbiguous(selector, len(heading_matches))
        start = heading_matches[0] if heading_matches else None
        if start is None:
            matches = []
        else:
            start_level = _heading_level(paragraphs[start]) or 9
            end = len(paragraphs)
            for index in range(start + 1, len(paragraphs)):
                level = _heading_level(paragraphs[index])
                if level is not None and level <= start_level:
                    end = index
                    break
            matches = paragraphs[start:end]
    else:
        raise ValueError(
            f"Selector type {selector.type!r} does not resolve paragraphs"
        )

    if not matches:
        raise SelectorNotFound(selector)
    if require_single and len(matches) != 1:
        raise SelectorAmbiguous(selector, len(matches))
    return matches


def resolve_single_paragraph(
    document: Doc,
    selector: Selector | str | dict[str, Any],
):
    """Resolve one paragraph, treating a heading as the heading itself.

    ``resolve_paragraphs`` intentionally expands a heading to its bounded section.
    Operations that attach to one paragraph must not inherit that multi-paragraph
    behavior.
    """
    parsed = Selector.from_value(selector)
    if parsed.type != "heading":
        return resolve_paragraphs(document, parsed, require_single=True)[0]
    if parsed.text is None:
        raise ValueError("Heading selector requires text")
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if _heading_level(paragraph) is not None
        and _equal(paragraph.text, parsed.text, parsed.case_sensitive)
    ]
    if not matches:
        raise SelectorNotFound(parsed)
    if len(matches) != 1:
        raise SelectorAmbiguous(parsed, len(matches))
    return matches[0]


def insert_paragraph_after(paragraph, *, style: str | None = None):
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    if style is not None:
        inserted.style = style
    return inserted


def add_paragraph_at_boundary(
    document: Doc,
    selector: Selector | str | dict[str, Any] | None = None,
    *,
    style: str | None = None,
):
    if selector is None:
        return document.add_paragraph(style=style)
    return insert_paragraph_after(
        resolve_single_paragraph(document, selector), style=style
    )
