from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import http.client
import ipaddress
import json
import os
from pathlib import Path
import re
import socket
import ssl
import tempfile
from typing import Any, Callable, Mapping
from urllib.parse import quote, urlencode, urljoin, urlsplit

from .registry import EvidenceError, EvidenceRegistry


class SourceResearchError(ValueError):
    """A source discovery or verification step failed closed."""


Transport = Callable[[str, Mapping[str, str], int], tuple[str, Mapping[str, str], bytes]]
_MAX_METADATA_BYTES = 4 * 1024 * 1024
_MAX_PDF_BYTES = 40 * 1024 * 1024
_MAX_REDIRECTS = 5
_CROSSREF_AGENT = "WordAssistant academic source workflow/1.0"


def _validate_https_url(url: str, *, resolve: bool) -> tuple[str, int]:
    parsed = urlsplit(url)
    if parsed.scheme.casefold() != "https" or not parsed.hostname:
        raise SourceResearchError("Source URL must use HTTPS and include a hostname")
    if parsed.username or parsed.password or parsed.port not in (None, 443):
        raise SourceResearchError("Credentials and nonstandard HTTPS ports are not allowed")
    host = parsed.hostname.rstrip(".").casefold()
    if host in {"localhost", "localhost.localdomain"} or host.endswith((".localhost", ".local", ".internal")):
        raise SourceResearchError("Local and internal network hosts are not allowed")
    try:
        address = ipaddress.ip_address(host)
    except ValueError:
        address = None
    if address is not None and not address.is_global:
        raise SourceResearchError("Private, loopback, link-local, and reserved IP addresses are not allowed")
    if resolve:
        if address is None:
            try:
                addresses = {
                    result[4][0]
                    for result in socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)
                }
            except OSError as exc:
                raise SourceResearchError("Source host could not be resolved") from exc
        else:
            addresses = {str(address)}
        if not addresses or any(not ipaddress.ip_address(item).is_global for item in addresses):
            raise SourceResearchError("Source host resolves to a non-public network address")
    return host, parsed.port or 443


class _PinnedHTTPSConnection(http.client.HTTPSConnection):
    """HTTPS connection pinned to a previously checked public IP address."""

    def __init__(self, host: str, address: str, timeout: float):
        super().__init__(host, port=443, timeout=timeout, context=ssl.create_default_context())
        self._pinned_address = address

    def connect(self) -> None:
        raw_socket = socket.create_connection((self._pinned_address, self.port), self.timeout)
        self.sock = self._context.wrap_socket(raw_socket, server_hostname=self.host)


def _safe_https_get(
    url: str,
    headers: Mapping[str, str],
    max_bytes: int,
    *,
    timeout: float = 20.0,
    allowed_redirect_hosts: set[str] | None = None,
) -> tuple[str, Mapping[str, str], bytes]:
    current_url = url
    for redirect_count in range(_MAX_REDIRECTS + 1):
        host, _ = _validate_https_url(current_url, resolve=True)
        if allowed_redirect_hosts is not None and host not in allowed_redirect_hosts:
            raise SourceResearchError("Metadata provider redirected outside its approved host")
        parsed = urlsplit(current_url)
        if _is_ip_address(host):
            address = host
        else:
            results = socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)
            addresses = list(dict.fromkeys(result[4][0] for result in results))
            if not addresses or any(not ipaddress.ip_address(item).is_global for item in addresses):
                raise SourceResearchError("Source host resolved to a non-public network address")
            address = addresses[0]
        path = parsed.path or "/"
        if parsed.query:
            path += f"?{parsed.query}"
        connection = _PinnedHTTPSConnection(host, address, timeout)
        try:
            request_headers = {"User-Agent": _CROSSREF_AGENT, **dict(headers)}
            connection.request("GET", path, headers=request_headers)
            response = connection.getresponse()
            response_headers = {key.casefold(): value for key, value in response.getheaders()}
            if response.status in {301, 302, 303, 307, 308}:
                location = response_headers.get("location")
                if not location:
                    raise SourceResearchError("Provider returned a redirect without a Location")
                if redirect_count >= _MAX_REDIRECTS:
                    raise SourceResearchError("Provider exceeded the redirect limit")
                current_url = urljoin(current_url, location)
                continue
            if response.status < 200 or response.status >= 300:
                raise SourceResearchError(f"Provider request failed with HTTP status {response.status}")
            content_length = response_headers.get("content-length")
            if content_length:
                try:
                    if int(content_length) > max_bytes:
                        raise SourceResearchError("Provider response exceeds the configured size limit")
                except ValueError as exc:
                    raise SourceResearchError("Provider returned an invalid Content-Length") from exc
            body = bytearray()
            while True:
                chunk = response.read(min(64 * 1024, max_bytes + 1 - len(body)))
                if not chunk:
                    break
                body.extend(chunk)
                if len(body) > max_bytes:
                    raise SourceResearchError("Provider response exceeds the configured size limit")
            return current_url, response_headers, bytes(body)
        except (OSError, http.client.HTTPException, ssl.SSLError) as exc:
            raise SourceResearchError("HTTPS provider request failed") from exc
        finally:
            connection.close()
    raise SourceResearchError("Provider exceeded the redirect limit")


def _is_ip_address(host: str) -> bool:
    try:
        ipaddress.ip_address(host)
        return True
    except ValueError:
        return False


def _get(
    url: str,
    *,
    accept: str,
    max_bytes: int = _MAX_METADATA_BYTES,
    transport: Transport | None = None,
    allowed_redirect_hosts: set[str] | None = None,
) -> tuple[str, Mapping[str, str], bytes]:
    headers = {"Accept": accept}
    if transport is None:
        return _safe_https_get(
            url, headers, max_bytes, allowed_redirect_hosts=allowed_redirect_hosts
        )
    final_url, response_headers, body = transport(url, headers, max_bytes)
    _validate_https_url(final_url, resolve=False)
    if allowed_redirect_hosts is not None and urlsplit(final_url).hostname.casefold() not in allowed_redirect_hosts:
        raise SourceResearchError("Metadata provider redirected outside its approved host")
    normalized_headers = {key.casefold(): value for key, value in response_headers.items()}
    if len(body) > max_bytes:
        raise SourceResearchError("Provider response exceeds the configured size limit")
    return final_url, normalized_headers, body


def _json_response(url: str, *, transport: Transport | None = None) -> dict[str, Any]:
    _, headers, body = _get(
        url,
        accept="application/json",
        transport=transport,
        allowed_redirect_hosts={urlsplit(url).hostname.casefold()},
    )
    content_type = headers.get("content-type", "").split(";", 1)[0].strip().casefold()
    if content_type not in {"application/json", "application/vnd.api+json", "text/json"}:
        raise SourceResearchError("Metadata provider did not return JSON")
    try:
        data = json.loads(body.decode("utf-8"))
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise SourceResearchError("Metadata provider returned invalid JSON") from exc
    if not isinstance(data, dict):
        raise SourceResearchError("Metadata provider JSON root must be an object")
    return data


def _normalize_doi(value: str) -> str:
    doi = value.strip()
    for prefix in ("https://doi.org/", "http://doi.org/", "doi:"):
        if doi.casefold().startswith(prefix):
            doi = doi[len(prefix):].strip()
            break
    if not re.fullmatch(r"10\.\d{4,9}/\S+", doi, re.IGNORECASE) or any(char.isspace() for char in doi):
        raise SourceResearchError("DOI must use the form 10.xxxx/suffix")
    return doi


def _year(item: Mapping[str, Any]) -> int:
    for key in ("published", "published-print", "published-online", "issued"):
        date_parts = item.get(key, {}).get("date-parts") if isinstance(item.get(key), dict) else None
        try:
            year = int(date_parts[0][0]) if date_parts and date_parts[0] else None
        except (IndexError, TypeError, ValueError, OverflowError):
            continue
        if year is not None and 1800 <= year <= 2100:
            return year
    raise SourceResearchError("Crossref metadata has no usable publication year")


def _author_names(item: Mapping[str, Any], *, language: str) -> list[str]:
    authors = item.get("author")
    if not isinstance(authors, list) or not authors:
        raise SourceResearchError("Crossref metadata has no authors")
    names = []
    for author in authors:
        if not isinstance(author, dict):
            continue
        given = str(author.get("given", "")).strip()
        family = str(author.get("family", "")).strip()
        literal = str(author.get("name", "")).strip()
        if language.startswith("vi"):
            name = " ".join(part for part in (family, given) if part) or literal
        else:
            initials = "".join(
                f"{token[0].upper()}."
                for token in re.findall(r"[^\W_]+", given, re.UNICODE)
            )
            name = " ".join(part for part in (family, initials) if part) or literal
        if name:
            names.append(name)
    if not names:
        raise SourceResearchError("Crossref metadata has no usable author names")
    return names


def _source_record(item: Mapping[str, Any]) -> dict[str, Any]:
    doi = _normalize_doi(str(item.get("DOI", "")))
    title_values = item.get("title")
    if not isinstance(title_values, list) or not title_values or not str(title_values[0]).strip():
        raise SourceResearchError("Crossref metadata has no title")
    language = str(item.get("language") or "").casefold().strip()
    authors = _author_names(item, language=language)
    first_author = item.get("author", [{}])[0]
    sort_field = "given" if language.startswith("vi") else "family"
    sort_key = str(first_author.get(sort_field, "")).strip() if isinstance(first_author, dict) else ""
    if not sort_key:
        sort_key = authors[0]
    raw_author_metadata = item.get("author")
    author_metadata = [
        {
            key: str(author.get(key, "")).strip()
            for key in ("given", "family", "name")
            if str(author.get(key, "")).strip()
        }
        for author in raw_author_metadata
        if isinstance(author, dict)
    ] if isinstance(raw_author_metadata, list) else []
    record_id = "doi-" + re.sub(r"[^a-z0-9]+", "-", doi.casefold()).strip("-")
    return {
        "id": record_id,
        "title": str(title_values[0]).strip(),
        "authors": authors,
        "year": _year(item),
        "url": f"https://doi.org/{doi}",
        "doi": doi,
        "language": language or "unknown",
        "author_metadata": author_metadata,
        "container_title": _first_metadata_value(item.get("container-title")),
        "volume": _optional_metadata_text(item.get("volume")),
        "issue": _optional_metadata_text(item.get("issue")),
        "pages": _optional_metadata_text(item.get("page")),
        "publisher": _optional_metadata_text(item.get("publisher")),
        "publication_place": _optional_metadata_text(item.get("publisher-location")),
        "verified_by": "Crossref metadata only; source content not reviewed",
        "synthetic": False,
        "review_status": "candidate",
        "source_version": str(item.get("type", "")),
        "author_sort_key": sort_key or None,
    }


def _optional_metadata_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _first_metadata_value(value: Any) -> str | None:
    if isinstance(value, list):
        return _optional_metadata_text(value[0]) if value else None
    return _optional_metadata_text(value)


def search_crossref(query: str, *, limit: int = 10, transport: Transport | None = None) -> list[dict[str, Any]]:
    if not isinstance(query, str) or not query.strip():
        raise SourceResearchError("Search query must not be empty")
    if type(limit) is not int or not 1 <= limit <= 50:
        raise SourceResearchError("Search limit must be an integer from 1 to 50")
    params = urlencode({"query.bibliographic": query.strip(), "rows": limit})
    data = _json_response(f"https://api.crossref.org/works?{params}", transport=transport)
    message = data.get("message")
    if not isinstance(message, dict) or not isinstance(message.get("items"), list):
        raise SourceResearchError("Crossref response is missing message.items")
    result = []
    for item in message["items"][:limit]:
        if not isinstance(item, dict):
            continue
        try:
            result.append(_source_record(item))
        except SourceResearchError:
            continue
    return result


def import_crossref_doi(
    doi_value: str,
    *,
    evidence_registry: str | Path,
    transport: Transport | None = None,
) -> dict[str, Any]:
    doi = _normalize_doi(doi_value)
    encoded = quote(doi, safe="")
    data = _json_response(f"https://api.crossref.org/works/{encoded}", transport=transport)
    message = data.get("message")
    if not isinstance(message, dict):
        raise SourceResearchError("Crossref DOI response is missing message metadata")
    record = _source_record(message)
    registry = EvidenceRegistry(evidence_registry)
    registry.add(record)
    return record


def find_open_access(
    evidence_id: str,
    *,
    evidence_registry: str | Path,
    email: str,
    transport: Transport | None = None,
) -> dict[str, Any]:
    if not isinstance(email, str) or "@" not in email or len(email) > 254:
        raise SourceResearchError("Unpaywall requires a valid contact email supplied at runtime")
    registry = EvidenceRegistry(evidence_registry)
    record = registry.get(evidence_id)
    doi = _normalize_doi(str(record.get("doi", "")))
    url = f"https://api.unpaywall.org/v2/{quote(doi, safe='')}?{urlencode({'email': email})}"
    data = _json_response(url, transport=transport)
    locations = data.get("oa_locations")
    best = data.get("best_oa_location")
    candidates = ([best] if isinstance(best, dict) else []) + (locations if isinstance(locations, list) else [])
    selected = None
    for location in candidates:
        if not isinstance(location, dict):
            continue
        pdf_url = location.get("url_for_pdf")
        license_name = location.get("license")
        if isinstance(pdf_url, str) and isinstance(license_name, str) and _has_reuse_license(license_name):
            selected = (pdf_url, license_name)
            break
    if selected is None:
        raise SourceResearchError("Unpaywall found no PDF with an explicit supported reuse license")
    oa_url, license_name = selected
    _validate_https_url(oa_url, resolve=False)
    record.update({
        "oa_url": oa_url,
        "oa_license": license_name,
        "oa_host": urlsplit(oa_url).hostname,
        "verified_by": "Crossref metadata and Unpaywall OA location; source content not reviewed",
        "review_status": "candidate",
    })
    registry.add(record, replace=True)
    return {"id": evidence_id, "doi": doi, "oa_host": record["oa_host"], "license": license_name}


def _has_reuse_license(value: str) -> bool:
    normalized = value.casefold().strip()
    if normalized in {
        "cc0", "public domain", "public-domain", "cc-by", "cc-by-sa",
        "cc-by-nc", "cc-by-nc-sa", "cc-by-nd", "cc-by-nc-nd",
    }:
        return True
    if re.fullmatch(r"cc-(?:by|by-sa|by-nc|by-nc-sa|by-nd|by-nc-nd)-[0-9.]+", normalized):
        return True
    return bool(re.search(r"creativecommons\.org/licenses/(?:by|by-sa|by-nc|by-nc-sa|by-nd|by-nc-nd|zero)/[0-9.]+", normalized))


def download_open_access_pdf(
    evidence_id: str,
    *,
    evidence_registry: str | Path,
    transport: Transport | None = None,
) -> dict[str, Any]:
    registry = EvidenceRegistry(evidence_registry)
    record = registry.get(evidence_id)
    if record.get("review_status") != "candidate":
        raise SourceResearchError("Only a discovery candidate can enter the source download workflow")
    url = record.get("oa_url")
    license_name = record.get("oa_license")
    if not isinstance(url, str) or not isinstance(license_name, str) or not _has_reuse_license(license_name):
        raise SourceResearchError("Download refused: an explicit supported reuse license is required")
    final_url, headers, body = _get(
        url,
        accept="application/pdf,application/octet-stream;q=0.9",
        max_bytes=_MAX_PDF_BYTES,
        transport=transport,
    )
    media_type = headers.get("content-type", "").split(";", 1)[0].casefold()
    if media_type not in {"application/pdf", "application/octet-stream", "binary/octet-stream"}:
        raise SourceResearchError("OA endpoint did not return a PDF content type")
    if not body.startswith(b"%PDF-"):
        raise SourceResearchError("OA response does not have a PDF signature")
    source_root = Path(evidence_registry).expanduser().resolve().parent / "sources"
    source_root.mkdir(parents=True, exist_ok=True)
    filename = f"{record['id']}.pdf"
    destination = source_root / filename
    if destination.exists():
        raise SourceResearchError("A source file with this evidence ID already exists; review or move it before downloading again")
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{record['id']}.", suffix=".tmp", dir=source_root)
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        temporary_path.write_bytes(body)
        digest = hashlib.sha256(body).hexdigest()
        temporary_path.replace(destination)
        record.update({
            "source_file": filename,
            "source_sha256": digest,
            "source_license": license_name,
            "source_url": final_url,
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
        })
        registry.add(record, replace=True)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise
    return {"id": evidence_id, "path": str(destination), "sha256": digest, "license": license_name, "bytes": len(body)}


def _extract_pdf_pages(path: Path) -> list[tuple[str, str]]:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise SourceResearchError("PDF review needs the configured pypdfium2 dependency") from exc
    try:
        document = pdfium.PdfDocument(str(path))
        if len(document) > 500:
            raise SourceResearchError("Source has more than 500 pages; inspect it in bounded page ranges")
        pages: list[tuple[str, str]] = []
        for index in range(len(document)):
            page = document[index]
            try:
                text_page = page.get_textpage()
                try:
                    pages.append((f"Page {index + 1}", text_page.get_text_range()))
                finally:
                    text_page.close()
            finally:
                page.close()
        return pages
    except SourceResearchError:
        raise
    except Exception as exc:
        raise SourceResearchError("Downloaded PDF could not be read") from exc
    finally:
        if "document" in locals():
            document.close()


def _source_path(record: Mapping[str, Any], evidence_registry: str | Path) -> Path:
    if not record.get("source_file") or not record.get("source_sha256"):
        raise SourceResearchError("Download an explicitly licensed source before reviewing it")
    source_root = Path(evidence_registry).expanduser().resolve().parent / "sources"
    source_path = (source_root / record["source_file"]).resolve()
    try:
        source_path.relative_to(source_root.resolve())
    except ValueError as exc:
        raise SourceResearchError("Source file escapes the workspace source store") from exc
    if not source_path.is_file():
        raise SourceResearchError("Downloaded source file is missing")
    digest = hashlib.sha256(source_path.read_bytes()).hexdigest()
    if digest.casefold() != record["source_sha256"].casefold():
        raise SourceResearchError("Downloaded source hash changed; review it again before use")
    return source_path


def _extract_source_pages(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.casefold()
    if suffix == ".pdf":
        return _extract_pdf_pages(path)
    if suffix == ".docx":
        try:
            from docx import Document
            document = Document(path)
            return [(f"Paragraph {index + 1}", paragraph.text) for index, paragraph in enumerate(document.paragraphs)]
        except Exception as exc:
            raise SourceResearchError("Downloaded DOCX could not be read") from exc
    if suffix in {".txt", ".md"}:
        return [("Text file", path.read_text(encoding="utf-8"))]
    raise SourceResearchError("Source review supports PDF, DOCX, TXT, and Markdown files")


def _normalize_excerpt(value: str) -> str:
    return " ".join(value.split()).casefold()


def search_source_excerpts(
    evidence_id: str,
    *,
    evidence_registry: str | Path,
    query: str,
    limit: int = 5,
    context_chars: int = 300,
) -> dict[str, Any]:
    if not isinstance(query, str) or not query.strip():
        raise SourceResearchError("Excerpt query must not be empty")
    if type(limit) is not int or not 1 <= limit <= 20:
        raise SourceResearchError("Excerpt limit must be an integer from 1 to 20")
    if type(context_chars) is not int or not 50 <= context_chars <= 1200:
        raise SourceResearchError("context_chars must be an integer from 50 to 1200")
    registry = EvidenceRegistry(evidence_registry)
    record = registry.get(evidence_id)
    source_path = _source_path(record, evidence_registry)
    needle = _normalize_excerpt(query)
    matches = []
    for location, raw_text in _extract_source_pages(source_path):
        normalized = " ".join(raw_text.split())
        folded = normalized.casefold()
        position = folded.find(needle)
        if position < 0:
            continue
        start = max(0, position - context_chars // 2)
        end = min(len(normalized), position + len(needle) + context_chars // 2)
        excerpt = normalized[start:end]
        matches.append({"location": location, "excerpt": excerpt})
        if len(matches) >= limit:
            break
    return {
        "id": evidence_id,
        "source_sha256": record["source_sha256"],
        "query": query.strip(),
        "matches": matches,
    }


def review_downloaded_source(
    evidence_id: str,
    *,
    evidence_registry: str | Path,
    claim: str,
    support_excerpt: str,
    support_location: str,
    reviewed_by: str,
    confirm_support: bool = False,
    quote: bool = False,
    language: str | None = None,
) -> dict[str, Any]:
    if confirm_support is not True:
        raise SourceResearchError("Review requires explicit confirmation that the source passage supports the stated claim")
    required = {"claim": claim, "support excerpt": support_excerpt, "support location": support_location, "reviewed by": reviewed_by}
    missing = [name for name, value in required.items() if not isinstance(value, str) or not value.strip()]
    if missing:
        raise SourceResearchError(f"Review requires non-empty fields: {', '.join(missing)}")
    registry = EvidenceRegistry(evidence_registry)
    record = registry.get(evidence_id)
    if record.get("review_status") != "candidate":
        raise SourceResearchError("Only a discovery candidate can be reviewed")
    if not record.get("source_file") or not record.get("source_sha256"):
        raise SourceResearchError("Download an explicitly licensed source before marking it reviewed")
    source_path = _source_path(record, evidence_registry)
    if language is not None:
        language = language.casefold().strip()
        if language not in {"vi", "en"}:
            raise SourceResearchError("language must be 'vi' or 'en'")
        record["language"] = language
        author_metadata = record.get("author_metadata", [])
        if author_metadata:
            record["authors"] = _author_names(
                {"author": author_metadata}, language=language
            )
            first_author = author_metadata[0]
            sort_field = "given" if language == "vi" else "family"
            record["author_sort_key"] = (
                str(first_author.get(sort_field, "")).strip() or record["authors"][0]
            )
    if not str(record.get("language", "")).casefold().startswith(("vi", "en")):
        raise SourceResearchError("Set the checked source language with --language vi or --language en before review")
    if not str(record.get("author_sort_key", "")).strip():
        raise SourceResearchError("Set a language-appropriate author_sort_key before review")
    digest = hashlib.sha256(source_path.read_bytes()).hexdigest()
    source_pages = _extract_source_pages(source_path)
    page_match = re.search(r"\bpage\s+(\d+)\b", support_location, re.IGNORECASE)
    eligible_pages = [
        text for location, text in source_pages
        if page_match is None or location.casefold() == f"page {page_match.group(1)}".casefold()
    ]
    if not eligible_pages or not any(
        _normalize_excerpt(support_excerpt) in _normalize_excerpt(text)
        for text in eligible_pages
    ):
        raise SourceResearchError("Support excerpt was not found in the downloaded source")
    now = datetime.now(timezone.utc).isoformat()
    record.update({
        "review_status": "reviewed",
        "reviewed_by": reviewed_by.strip(),
        "reviewed_at": now,
        "claim": claim.strip(),
        "support_location": support_location.strip(),
        "support_excerpt": support_excerpt.strip(),
        "verified_by": reviewed_by.strip(),
    })
    if quote:
        record["quote"] = support_excerpt.strip()
    registry.add(record, replace=True)
    return {"id": evidence_id, "review_status": "reviewed", "source_sha256": digest, "support_location": support_location.strip()}
