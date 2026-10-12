from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..evidence.registry import EvidenceRegistry
from ..operations.citations import (
    _TAG_PREFIX,
    _citation_numbers,
    _citation_policy,
    _set_citation_policy,
    _ordered_citation_numbers,
    _renumber_citations,
    _bibliography_entries,
    _bibliography_contains_paragraph,
    _format_in_text,
    _replace_bibliography_content,
)
from ..operations.common import atomic_save
from ..profiles.loader import DocumentProfile, resolve_profile


_ROOT_KEYS = {"version", "sections", "bibliography_heading", "citation_style", "bibliography_order"}
_SECTION_KEYS = {"heading", "paragraphs", "subsections"}
_SUBSECTION_KEYS = {"title", "level", "paragraphs"}
_PARAGRAPH_KEYS = {"text", "citations"}
_PLACEHOLDERS = {"[section content...]", "[front matter content]", "[document content]"}


def _unknown_keys(value: Mapping[str, Any], allowed: set[str], context: str) -> None:
    unknown = sorted(set(value) - allowed)
    if unknown:
        raise ValueError(f"Unsupported {context} keys: {', '.join(unknown)}")


def _paragraph_spec(value: Any, context: str) -> tuple[str, list[str]]:
    if isinstance(value, str):
        text = value
        citations: list[str] = []
    elif isinstance(value, Mapping):
        _unknown_keys(value, _PARAGRAPH_KEYS, context)
        text = value.get("text")
        citations = value.get("citations", [])
        if not isinstance(citations, list) or any(
            not isinstance(item, str) or not item.strip() for item in citations
        ):
            raise TypeError(f"{context}.citations must be an array of non-empty strings")
        citations = [item.strip() for item in citations]
    else:
        raise TypeError(f"{context} must be a string or object")
    if not isinstance(text, str) or not text.strip():
        raise ValueError(f"{context}.text must be a non-empty string")
    if len(text) > 20_000:
        raise ValueError(f"{context}.text exceeds 20000 characters")
    return text.strip(), citations


def _append_citation(paragraph, evidence_id: str, text: str) -> None:
    if len(paragraph._p):
        paragraph.add_run(" ")
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), f"{_TAG_PREFIX}{evidence_id}")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), f"Citation {evidence_id}")
    properties.extend([alias, tag])
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    content.append(run)
    sdt.extend([properties, content])
    paragraph._p.append(sdt)


def _insert_after(document, cursor, text: str, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    cursor.addnext(paragraph._p)
    return paragraph


def _remove_placeholder_after(paragraph) -> bool:
    node = paragraph._p.getnext()
    if node is None or node.tag != qn("w:p"):
        return False
    text = "".join(item.text or "" for item in node.iter(qn("w:t"))).strip().casefold()
    if text not in _PLACEHOLDERS:
        return False
    node.getparent().remove(node)
    return True


def compose_document(
    input_path: str | Path,
    *,
    output_path: str | Path,
    content: Mapping[str, Any],
    evidence_registry: str | Path | None = None,
    profile: str | DocumentProfile | None = None,
    variant_name: str | None = None,
    min_visible_characters: int | None = None,
    max_visible_characters: int | None = None,
    dry_run: bool = False,
) -> dict[str, Any]:
    """Populate an existing scaffold from a bounded, fail-closed JSON contract."""
    if not isinstance(content, Mapping):
        raise TypeError("Composition content must be a JSON object")
    _unknown_keys(content, _ROOT_KEYS, "composition")
    if content.get("version") != 1 or isinstance(content.get("version"), bool):
        raise ValueError("Composition content version must be 1")
    sections = content.get("sections")
    if not isinstance(sections, list) or not sections:
        raise ValueError("Composition sections must be a non-empty array")
    if len(sections) > 100:
        raise ValueError("Composition cannot contain more than 100 sections")
    profile_variant = None
    if profile is not None:
        loaded_profile = (
            profile
            if isinstance(profile, DocumentProfile)
            else resolve_profile(str(profile), start_path=input_path)
        )
        if variant_name is None:
            variant_name = "default" if "default" in loaded_profile.variants else next(iter(loaded_profile.variants))
        profile_variant = loaded_profile.variant(variant_name)
    for name, value in (
        ("min_visible_characters", min_visible_characters),
        ("max_visible_characters", max_visible_characters),
    ):
        if value is not None and (type(value) is not int or value < 0):
            raise ValueError(f"{name} must be a non-negative integer")
    if (
        min_visible_characters is not None
        and max_visible_characters is not None
        and min_visible_characters > max_visible_characters
    ):
        raise ValueError("min_visible_characters cannot exceed max_visible_characters")

    document = Document(input_path)
    existing_policy = _citation_policy(document)
    _citation_numbers(document, displayed=True)
    citation_style = content.get(
        "citation_style",
        profile_variant.citations.get("style", "numeric") if profile_variant else existing_policy.get("style", "numeric"),
    )
    if not isinstance(citation_style, str) or not citation_style.strip():
        raise ValueError("citation_style must be a non-empty string")
    citation_style = citation_style.strip()
    required_citation_style = (
        profile_variant.citations.get("style") if profile_variant else existing_policy.get("style")
    )
    require_reviewed_sources = bool(
        (profile_variant and profile_variant.citations.get("require_reviewed_sources", False))
        or existing_policy.get("require_reviewed", False)
    )
    if (
        require_reviewed_sources
        and "citation_style" in content
        and required_citation_style
        and citation_style.casefold() != str(required_citation_style).casefold()
    ):
        raise ValueError("citation_style cannot override the selected source-backed profile")
    bibliography_order = content.get(
        "bibliography_order",
        profile_variant.citations.get("bibliography_order", "citation_order")
        if profile_variant
        else existing_policy.get("order", "citation_order"),
    )
    allowed_orders = {
        "citation_order",
        "vietnamese_given_name_then_english_family_name",
    }
    if bibliography_order not in allowed_orders:
        raise ValueError(f"bibliography_order must be one of {sorted(allowed_orders)}")
    required_bibliography_order = (
        profile_variant.citations.get("bibliography_order")
        if profile_variant
        else existing_policy.get("order")
    )
    if (
        require_reviewed_sources
        and "bibliography_order" in content
        and required_bibliography_order
        and bibliography_order != required_bibliography_order
    ):
        raise ValueError("bibliography_order cannot override the selected source-backed profile")
    registry = EvidenceRegistry(evidence_registry) if evidence_registry else None
    citation_numbers: dict[str, int] = _citation_numbers(document, displayed=True)
    paragraph_count = 0
    subsection_count = 0
    removed_placeholders = 0
    used_targets: set[str] = set()
    managed_heading = existing_policy.get('heading')
    if profile_variant and evidence_registry:
        reference_headings = [
            str(item['title']) for item in profile_variant.section_blueprint
            if 'tham khảo' in str(item.get('title', '')).casefold()
            or 'references' in str(item.get('title', '')).casefold()
        ]
        if reference_headings:
            managed_heading = reference_headings[-1]
    managed_heading = content.get('bibliography_heading', managed_heading)

    for section_index, section in enumerate(sections):
        context = f"sections[{section_index}]"
        if not isinstance(section, Mapping):
            raise TypeError(f"{context} must be an object")
        _unknown_keys(section, _SECTION_KEYS, context)
        heading = section.get("heading")
        if not isinstance(heading, str) or not heading.strip():
            raise ValueError(f"{context}.heading must be a non-empty string")
        key = heading.strip().casefold()
        if key in used_targets:
            raise ValueError(f"Duplicate composition target heading: {heading!r}")
        used_targets.add(key)
        matches = [
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.strip().casefold() == key
        ]
        if len(matches) != 1:
            raise ValueError(
                f"{context}.heading must match exactly one paragraph; found {len(matches)}"
            )
        target = matches[0]
        if managed_heading and isinstance(managed_heading, str):
            existing_heading = existing_policy.get('heading')
            allow_missing = not existing_heading or managed_heading.casefold() != existing_heading.casefold()
            if _bibliography_contains_paragraph(document, managed_heading, target, allow_missing=allow_missing):
                raise ValueError('Composition target must be outside the managed bibliography')
        removed_placeholders += int(_remove_placeholder_after(target))
        cursor = target._p

        raw_paragraphs = section.get("paragraphs", [])
        if not isinstance(raw_paragraphs, list):
            raise TypeError(f"{context}.paragraphs must be an array")
        for paragraph_index, raw in enumerate(raw_paragraphs):
            text, citations = _paragraph_spec(
                raw, f"{context}.paragraphs[{paragraph_index}]"
            )
            paragraph = _insert_after(document, cursor, text)
            for evidence_id in citations:
                if registry is None:
                    raise ValueError("evidence_registry is required when citations are used")
                record = registry.get(evidence_id)
                registry.assert_citable(record, require_reviewed=require_reviewed_sources)
                number = citation_numbers.setdefault(evidence_id, max(citation_numbers.values(), default=0) + 1)
                _append_citation(
                    paragraph,
                    evidence_id,
                    _format_in_text(record, citation_style, number),
                )
            cursor = paragraph._p
            paragraph_count += 1

        raw_subsections = section.get("subsections", [])
        if not isinstance(raw_subsections, list):
            raise TypeError(f"{context}.subsections must be an array")
        for subsection_index, subsection in enumerate(raw_subsections):
            sub_context = f"{context}.subsections[{subsection_index}]"
            if not isinstance(subsection, Mapping):
                raise TypeError(f"{sub_context} must be an object")
            _unknown_keys(subsection, _SUBSECTION_KEYS, sub_context)
            title = subsection.get("title")
            level = subsection.get("level", 2)
            if not isinstance(title, str) or not title.strip():
                raise ValueError(f"{sub_context}.title must be a non-empty string")
            if type(level) is not int or not 2 <= level <= 9:
                raise ValueError(f"{sub_context}.level must be an integer from 2 to 9")
            heading_paragraph = _insert_after(
                document, cursor, title.strip(), f"Heading {level}"
            )
            cursor = heading_paragraph._p
            subsection_count += 1
            raw_sub_paragraphs = subsection.get("paragraphs", [])
            if not isinstance(raw_sub_paragraphs, list):
                raise TypeError(f"{sub_context}.paragraphs must be an array")
            for paragraph_index, raw in enumerate(raw_sub_paragraphs):
                text, citations = _paragraph_spec(
                    raw, f"{sub_context}.paragraphs[{paragraph_index}]"
                )
                paragraph = _insert_after(document, cursor, text)
                for evidence_id in citations:
                    if registry is None:
                        raise ValueError("evidence_registry is required when citations are used")
                    record = registry.get(evidence_id)
                    registry.assert_citable(record, require_reviewed=require_reviewed_sources)
                    number = citation_numbers.setdefault(
                        evidence_id, max(citation_numbers.values(), default=0) + 1
                    )
                    _append_citation(
                        paragraph,
                        evidence_id,
                        _format_in_text(record, citation_style, number),
                    )
                cursor = paragraph._p
                paragraph_count += 1

    default_bibliography_heading = existing_policy.get("heading")
    if profile_variant and evidence_registry:
        references = [
            item["title"]
            for item in profile_variant.section_blueprint
            if "tham khảo" in str(item.get("title", "")).casefold()
            or "references" in str(item.get("title", "")).casefold()
        ]
        default_bibliography_heading = references[-1] if references else None
    bibliography_heading = content.get("bibliography_heading", default_bibliography_heading or existing_policy.get("heading"))
    if existing_policy.get('require_reviewed') and (
        citation_style.casefold() != existing_policy['style'].casefold()
        or bibliography_order != existing_policy['order']
        or bibliography_heading != existing_policy.get('heading')
    ):
        raise ValueError('composition cannot override the source-backed document policy')
    if bibliography_heading is not None:
        if not isinstance(bibliography_heading, str) or not bibliography_heading.strip():
            raise ValueError("bibliography_heading must be a non-empty string")
        if (
            require_reviewed_sources
            and "bibliography_heading" in content
            and default_bibliography_heading
            and bibliography_heading.strip().casefold()
            != default_bibliography_heading.strip().casefold()
        ):
            raise ValueError("bibliography_heading cannot override the selected source-backed profile")
        if registry is None and (citation_numbers or 'bibliography_heading' in content):
            raise ValueError("evidence_registry is required for bibliography generation")
    policy = {"style": citation_style, "order": bibliography_order,
              "heading": bibliography_heading, "require_reviewed": require_reviewed_sources}
    if citation_numbers:
        if registry is None:
            raise ValueError("evidence_registry is required for existing citations")
        citation_numbers = _ordered_citation_numbers(document, registry, policy)
        _renumber_citations(document, citation_numbers, registry, citation_style)
    if bibliography_heading and registry is not None:
        entries = _bibliography_entries(registry, citation_numbers, citation_style,
                                       require_reviewed=require_reviewed_sources)
        _replace_bibliography_content(document, bibliography_heading, entries)
    _set_citation_policy(document, policy)

    if paragraph_count > 5000:
        raise ValueError("Composition cannot contain more than 5000 content paragraphs")
    visible_characters = sum(
        len("".join(node.text or "" for node in paragraph._p.iter(qn("w:t"))).strip())
        for paragraph in document.paragraphs
    )
    if (
        min_visible_characters is not None
        and visible_characters < min_visible_characters
    ):
        raise ValueError(
            "Composition visible character count is below the required minimum: "
            f"{visible_characters} < {min_visible_characters}"
        )
    if (
        max_visible_characters is not None
        and visible_characters > max_visible_characters
    ):
        raise ValueError(
            "Composition visible character count exceeds the allowed maximum: "
            f"{visible_characters} > {max_visible_characters}"
        )
    output = Path(output_path).resolve()
    if not dry_run:
        atomic_save(document, output)
    return {
        "status": "DRY_RUN" if dry_run else "SUCCESS",
        "output": None if dry_run else str(output),
        "sections": len(sections),
        "subsections": subsection_count,
        "paragraphs": paragraph_count,
        "visible_characters": visible_characters,
        "citations": len(citation_numbers),
        "bibliography_entries": len(citation_numbers) if bibliography_heading else 0,
        "removed_placeholders": removed_placeholders,
    }
