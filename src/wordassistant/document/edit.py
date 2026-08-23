from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
import os
import tempfile
from docx import Document

@dataclass(frozen=True)
class ReplaceResult:
    output: Path
    replacements: int
    dry_run: bool = False
    complex_structures_detected: tuple[str, ...] = ()
    advanced_backend_recommended: bool = False

def _detect_complex_elements(element) -> list[str]:
    detected = []
    xml = element.xml if hasattr(element, 'xml') else ''
    if '<w:hyperlink' in xml:
        detected.append('hyperlink')
    if '<w:fldSimple' in xml or '<w:fldChar' in xml or '<w:instrText' in xml:
        detected.append('field')
    if '<w:sdt' in xml:
        detected.append('content_control')
    if '<w:ins' in xml or '<w:del' in xml:
        detected.append('tracked_change')
    if '<w:footnoteReference' in xml:
        detected.append('footnote')
    if '<w:endnoteReference' in xml:
        detected.append('endnote')
    if '<w:commentReference' in xml:
        detected.append('comment')
    if '<w:drawing' in xml:
        detected.append('drawing')
    return detected

def _replace_paragraph_runs_safe(paragraph, old: str, new: str) -> int:
    # 1. Fast path: check if any run contains the entire match
    single_run_count = 0
    for run in paragraph.runs:
        c = run.text.count(old)
        if c:
            run.text = run.text.replace(old, new)
            single_run_count += c
    if single_run_count > 0:
        return single_run_count
        
    # 2. Cross-run match: find occurrences across runs
    full_text = paragraph.text
    if old not in full_text:
        return 0
        
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(full_text.replace(old, new))
        return full_text.count(old)
        
    # Build run character offset intervals
    run_spans = []
    curr_offset = 0
    for r in runs:
        r_len = len(r.text)
        run_spans.append((curr_offset, curr_offset + r_len, r))
        curr_offset += r_len
        
    # Find all occurrences of old in full_text
    match_indices = []
    start = 0
    while True:
        idx = full_text.find(old, start)
        if idx == -1:
            break
        match_indices.append((idx, idx + len(old)))
        start = idx + len(old)
        
    if not match_indices:
        return 0
        
    # For cross-run matches, safely update the affected run span
    # We process from last match to first to avoid offset shifting issues if multiple
    for m_start, m_end in reversed(match_indices):
        start_run_idx = None
        end_run_idx = None
        
        for i, (r_start, r_end, r) in enumerate(run_spans):
            if r_start <= m_start < r_end:
                start_run_idx = i
            if r_start < m_end <= r_end:
                end_run_idx = i
                
        if start_run_idx is not None and end_run_idx is not None:
            r_start_obj = runs[start_run_idx]
            r_end_obj = runs[end_run_idx]
            
            offset_in_start = m_start - run_spans[start_run_idx][0]
            offset_in_end = m_end - run_spans[end_run_idx][0]
            
            prefix = r_start_obj.text[:offset_in_start]
            suffix = r_end_obj.text[offset_in_end:]
            
            if start_run_idx == end_run_idx:
                r_start_obj.text = prefix + new + suffix
            else:
                r_start_obj.text = prefix + new
                for mid_idx in range(start_run_idx + 1, end_run_idx):
                    runs[mid_idx].text = ""
                r_end_obj.text = suffix
        else:
            # Fallback across boundary
            first = runs[0]
            first.text = full_text.replace(old, new)
            for r in runs[1:]:
                r.text = ""
            break
            
    return len(match_indices)

def replace_text(
    input_path: str | Path,
    old: str,
    new: str,
    output_path: str | Path | None = None,
    dry_run: bool = False
) -> ReplaceResult:
    if not old:
        raise ValueError("old text must be non-empty")
        
    src = Path(input_path)
    out = Path(output_path) if output_path else src.with_name(f"{src.stem}_edited.docx")
    
    d = Document(src)
    total_replacements = 0
    complex_detected: set[str] = set()
    
    for p in d.paragraphs:
        if old in p.text:
            complex_in_p = _detect_complex_elements(p._p)
            complex_detected.update(complex_in_p)
            total_replacements += _replace_paragraph_runs_safe(p, old, new)
            
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if old in p.text:
                        complex_in_p = _detect_complex_elements(p._p)
                        complex_detected.update(complex_in_p)
                        total_replacements += _replace_paragraph_runs_safe(p, old, new)
                        
    adv_recommended = bool(complex_detected.intersection({'field', 'content_control', 'tracked_change', 'footnote', 'endnote'}))
    
    if not dry_run and total_replacements > 0:
        out.parent.mkdir(parents=True, exist_ok=True)
        # Atomic file write using a temporary file in the destination folder
        tmp_fd, tmp_path = tempfile.mkstemp(dir=out.parent, prefix=".tmp_wa_", suffix=".docx")
        os.close(tmp_fd)
        try:
            d.save(tmp_path)
            os.replace(tmp_path, out)
        except Exception:
            if os.path.exists(tmp_path):
                try:
                    os.remove(tmp_path)
                except Exception:
                    pass
            raise
            
    return ReplaceResult(
        output=out,
        replacements=total_replacements,
        dry_run=dry_run,
        complex_structures_detected=tuple(sorted(complex_detected)),
        advanced_backend_recommended=adv_recommended
    )
