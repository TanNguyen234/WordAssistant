from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile, ZipFile
import hashlib
import hmac
import json
import shutil
import subprocess
import tempfile

from docx import Document
from docx.oxml.ns import qn

from .render import RenderResult, render_pdf_pages
from .structure import inspect_document
from .validate import validate_document
from ..operations.fields import inspect_fields, inspect_field_structure


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _visual_review_challenge(
    source: Path, pdf_path: Path, page_images: tuple[Path, ...]
) -> dict:
    pages = [
        {
            "page": index,
            "path": str(page.resolve()),
            "size_bytes": page.stat().st_size,
            "sha256": _sha256_file(page),
        }
        for index, page in enumerate(page_images, start=1)
    ]
    # Word's PDF exporter may rewrite non-visual metadata (for example creation
    # timestamps) on every export. Bind the attestation to the source document
    # and the ordered rendered-page evidence that was actually inspected, while
    # still reporting the raw PDF hash for provenance.
    token_payload = {
        "document_sha256": _sha256_file(source),
        "pages": [
            {
                "page": page["page"],
                "size_bytes": page["size_bytes"],
                "sha256": page["sha256"],
            }
            for page in pages
        ],
    }
    token = hashlib.sha256(
        json.dumps(token_payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return {
        **token_payload,
        "pdf_sha256": _sha256_file(pdf_path),
        "pdf_path": str(pdf_path.resolve()),
        "page_count": len(pages),
        "page_images": pages,
        "token": token,
        "token_scope": "document SHA-256 plus ordered page-image SHA-256 and size; PDF SHA-256 is provenance only",
    }


def _reuse_manifest(source: Path, pdf_path: Path, pages: tuple[Path, ...], engine: str) -> dict:
    return {
        "version": 1,
        "document_sha256": _sha256_file(source),
        "pdf_sha256": _sha256_file(pdf_path),
        "renderer": {"engine": engine, "page_dpi": 120},
        "page_count": len(pages),
        "pages": [
            {"page": i, "path": str(p.resolve()), "size_bytes": p.stat().st_size, "sha256": _sha256_file(p)}
            for i, p in enumerate(pages, 1)
        ],
    }


def _reuse_token(manifest: dict) -> str:
    payload = json.dumps(manifest, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return "wa-visual-reuse-v1:" + hashlib.sha256(payload).hexdigest()


def _pdf_page_count(pdf_path: Path) -> int:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise RuntimeError("PDF page-count verification requires pypdfium2.") from exc
    document = pdfium.PdfDocument(pdf_path)
    try:
        count = len(document)
    finally:
        document.close()
    if count <= 0:
        raise RuntimeError("Rendered PDF contains no pages.")
    return count


def audit_pdf_page_content(
    pdf_path: str | Path,
    *,
    page_count: int,
    minimum_alphanumeric: int = 20,
) -> dict:
    """Detect rendered pages containing only pagination or similarly sparse text."""
    executable = shutil.which("pdftotext")
    pages = []
    near_blank_pages = []
    pdfium_document = None

    def ensure_pdfium():
        nonlocal pdfium_document
        if pdfium_document is not None:
            return pdfium_document
        try:
            import pypdfium2 as pdfium

            pdfium_document = pdfium.PdfDocument(Path(pdf_path).resolve())
            if len(pdfium_document) != page_count:
                raise RuntimeError(
                    f"PDF contains {len(pdfium_document)} page(s), expected {page_count}."
                )
        except Exception as exc:
            raise RuntimeError(f"PDFium text extraction failed: {exc}") from exc
        return pdfium_document

    poppler_texts = None
    poppler_error = None
    if executable:
        try:
            completed = subprocess.run(
                [executable, "-layout", str(Path(pdf_path).resolve()), "-"],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=max(30, min(300, page_count * 2)),
            )
            if completed.returncode != 0:
                raise RuntimeError(completed.stderr.strip() or "pdftotext failed")
            pieces = completed.stdout.split("\f")
            # pdftotext terminates output with a form feed. Remove only that
            # delimiter's trailing empty segment; retain real blank pages.
            if pieces and pieces[-1] == "":
                pieces.pop()
            if len(pieces) != page_count:
                raise RuntimeError(
                    f"pdftotext returned {len(pieces)} page(s), expected {page_count}."
                )
            poppler_texts = pieces
        except Exception as exc:
            poppler_error = str(exc)
    if poppler_texts is None:
        try:
            ensure_pdfium()
        except Exception as exc:
            reason = f"{poppler_error}; {exc}" if poppler_error else str(exc)
            return {"status": "UNAVAILABLE", "pages": [], "near_blank_pages": [], "reason": reason}
    try:
        for page_number in range(1, page_count + 1):
            if poppler_texts is not None:
                text = poppler_texts[page_number - 1]
            else:
                page = pdfium_document[page_number - 1]
                try:
                    text_page = page.get_textpage()
                    try:
                        text = text_page.get_text_bounded()
                    finally:
                        text_page.close()
                finally:
                    page.close()
            count = sum(character.isalnum() for character in text)
            near_blank = count < minimum_alphanumeric
            if near_blank:
                near_blank_pages.append(page_number)
            pages.append(
                {
                    "page": page_number,
                    "alphanumeric_characters": count,
                    "near_blank": near_blank,
                }
            )
    finally:
        if pdfium_document is not None:
            pdfium_document.close()
    return {
        "status": "FAIL" if near_blank_pages else "PASS",
        "pages": pages,
        "near_blank_pages": near_blank_pages,
        "minimum_alphanumeric": minimum_alphanumeric,
    }


def visual_preflight(path: str | Path) -> dict:
    document = Document(path)
    consecutive = False
    previous_break_only = False
    for paragraph in document.paragraphs:
        breaks = paragraph._p.xpath(".//w:br[@w:type='page']")
        text = paragraph.text.strip()
        break_only = bool(breaks) and not text
        if break_only and previous_break_only:
            consecutive = True
            break
        previous_break_only = break_only
    warnings = ["consecutive_page_breaks"] if consecutive else []
    return {
        "status": "WARNING" if warnings else "PASS",
        "warnings": warnings,
        "manual_review_required": bool(warnings),
    }


def _render_via_execution_service(
    path: str | Path,
    output_dir: str | Path | None = None,
    *,
    engine: str = "auto",
    timeout_seconds: int = 180,
) -> RenderResult:
    from ..execution import ExecutionService
    from ..operations.models import OperationRequest

    source = Path(path).resolve()
    destination_dir = Path(output_dir).resolve() if output_dir else source.parent
    destination = destination_dir / f"{source.stem}.pdf"
    result = ExecutionService.default().execute(
        OperationRequest(
            "render.pdf",
            input_path=source,
            output_path=destination,
            options={"engine": engine, "timeout_seconds": timeout_seconds},
        )
    )
    return RenderResult(
        pdf_path=destination,
        file_size_bytes=destination.stat().st_size,
        engine="word" if result.backend_used == "word-native" else "libreoffice",
        page_count=result.data.get("page_count"),
        duration_seconds=float(result.data.get("duration_seconds", 0.0)),
    )


def _structural_check(path: Path) -> tuple[str, list[dict]]:
    try:
        with ZipFile(path) as package:
            names = set(package.namelist())
            missing = {"[Content_Types].xml", "word/document.xml"} - names
            if missing:
                raise ValueError(f"Missing required DOCX parts: {sorted(missing)}")
        Document(path)
        return "PASS", []
    except (OSError, BadZipFile, ValueError, KeyError) as exc:
        return "FAIL", [
            {"severity": "ERROR", "code": "invalid_docx", "message": str(exc)}
        ]


def _combined_native_fields_render(
    source: Path, render_output_dir: str | Path | None
) -> tuple[dict, RenderResult | None]:
    """Run field verification and PDF export in one private Word session."""
    from ..backends.word_native.backend import (
        VERIFIED_NATIVE_OPERATIONS,
        _run_native_worker,
    )
    from .render import _atomic_copy_replace

    required = {"render.pdf", "field.update_all"}
    if not required.issubset(VERIFIED_NATIVE_OPERATIONS):
        raise RuntimeError(
            "Combined native verification requires verified operations: "
            + ", ".join(sorted(required))
        )
    output_dir = Path(render_output_dir).resolve() if render_output_dir else source.parent
    output_dir.mkdir(parents=True, exist_ok=True)
    published_pdf = output_dir / f"{source.stem}.pdf"
    with tempfile.TemporaryDirectory(prefix="wordassistant-verify-combined-") as temporary_directory:
        private_dir = Path(temporary_directory)
        private_docx = private_dir / source.name
        private_pdf = private_dir / f"{source.stem}.pdf"
        shutil.copy2(source, private_docx)
        completed = _run_native_worker(
            {
                "action": "verify_fields_and_render",
                "input_path": str(private_docx),
                "output_path": str(private_pdf),
            },
            300,
        )
        lines = [line for line in completed.stdout.splitlines() if line.strip()]
        payload = json.loads(lines[-1]) if lines else {}
        if completed.returncode != 0 or payload.get("status") != "SUCCESS":
            raise RuntimeError(
                payload.get("message")
                or completed.stderr.strip()
                or "Combined Microsoft Word verification failed"
            )
        data = payload.get("data")
        if not isinstance(data, dict) or set(data) != {"render", "fields"}:
            raise ValueError("Combined Word worker returned malformed stage results")
        render_stage = data["render"]
        fields_stage = data["fields"]
        if not isinstance(render_stage, dict) or render_stage.get("status") not in {"SUCCESS", "ERROR"}:
            render_stage = {"status": "ERROR", "message": "Combined Word worker returned malformed render result"}
        if not isinstance(fields_stage, dict) or fields_stage.get("status") not in {"SUCCESS", "ERROR"}:
            fields_stage = {"status": "ERROR", "message": "Combined Word worker returned malformed field result"}
        render_result = None
        if render_stage["status"] == "SUCCESS":
            render_data = render_stage.get("data")
            if not isinstance(render_data, dict):
                render_stage = {"status": "ERROR", "message": "Combined Word worker returned malformed PDF metadata"}
            else:
                page_count = render_data.get("page_count")
                if isinstance(page_count, bool) or not isinstance(page_count, int) or page_count <= 0:
                    render_stage = {"status": "ERROR", "message": "Combined Word worker returned an invalid PDF page count"}
                elif not private_pdf.is_file() or private_pdf.stat().st_size <= 0:
                    render_stage = {"status": "ERROR", "message": "Combined Word worker did not create a nonempty PDF"}
                else:
                    try:
                        _atomic_copy_replace(private_pdf, published_pdf)
                        render_result = RenderResult(
                            pdf_path=published_pdf,
                            file_size_bytes=published_pdf.stat().st_size,
                            engine="word",
                            page_count=page_count,
                        )
                    except Exception as exc:
                        render_stage = {"status": "ERROR", "message": str(exc)}
        if fields_stage["status"] == "SUCCESS" and not isinstance(fields_stage.get("data"), dict):
            fields_stage = {"status": "ERROR", "message": "Combined Word worker returned malformed field data"}
        return {"render": render_stage, "fields": fields_stage}, render_result


def _record_native_field_result(report: dict, data: dict) -> None:
    native_results = data.get("fields")
    if not isinstance(native_results, list):
        raise ValueError("Microsoft Word returned malformed field results")
    bad_results = [
        field
        for field in native_results
        if not isinstance(field, dict)
        or field.get("locked")
        or not str(field.get("result", "")).strip()
        or str(field.get("result", "")).strip().casefold().startswith("error!")
    ]
    if bad_results:
        raise ValueError("Microsoft Word returned locked, empty, or error fields")
    persisted_fresh = bool(data.get("source_results_were_current"))
    if data.get("field_freshness_comparison"):
        report["field_freshness_comparison"] = data["field_freshness_comparison"]
    report["field_evaluation_status"] = "PASS"
    report["native_word_status"] = "PASS"
    report["native_field_results"] = native_results
    report["field_result_basis"] = "EVALUATED_IN_ISOLATED_MICROSOFT_WORD_SESSION"
    if persisted_fresh:
        report["field_freshness_status"] = "PASS"
    else:
        report["field_freshness_status"] = "STALE_OR_UNKNOWN"
        report["manual_review_required"] = True
        report["issues"].append(
            {
                "severity": "WARNING",
                "code": "native_evaluation_not_persisted",
                "message": (
                    "Microsoft Word evaluated a temporary copy successfully, but "
                    "its results differed from the cached results in the source. "
                    "Run field update-all to a distinct output before delivery."
                ),
            }
        )


def verify_document(
    path: str | Path,
    *,
    profile=None,
    variant_name: str = "default",
    render: bool = False,
    render_output_dir: str | Path | None = None,
    renderer=None,
    native_fields: bool = False,
    visual_reviewed: bool = False,
    visual_review_token: str | None = None,
    reuse_render_report: str | Path | None = None,
) -> dict:
    if reuse_render_report is not None and render:
        raise ValueError("render and reuse_render_report are mutually exclusive")
    source = Path(path).resolve()
    structural_status, issues = _structural_check(source)
    report = {
        "overall_status": "FAIL" if structural_status == "FAIL" else "PASS",
        "structural_status": structural_status,
        "format_status": "NOT_REQUESTED",
        "numbering_status": "NOT_APPLICABLE",
        "table_geometry_status": "NOT_APPLICABLE",
        "field_structure_status": "NOT_APPLICABLE",
        "field_freshness_status": "NOT_APPLICABLE",
        "field_evaluation_status": "NOT_REQUESTED",
        "field_status": "NOT_APPLICABLE",
        "render_status": "NOT_REQUESTED",
        "native_word_status": "NOT_REQUESTED",
        "visual_status": "NOT_REQUESTED",
        "manual_review_required": False,
        "issues": issues,
    }
    if structural_status == "FAIL":
        report["delivery_status"] = "FAIL"
        report["delivery_gates"] = {"overall": "FAIL"}
        return report

    inventory = inspect_document(source)
    if profile is not None:
        validation = validate_document(
            source, profile, variant_name, verify_mode=True
        )
        report["format_status"] = "PASS" if validation.valid else "FAIL"
        report["issues"].extend(
            {
                "severity": issue.severity.upper(),
                "code": issue.code,
                "message": issue.message,
            }
            for issue in validation.issues
        )

    heading_literals = [
        paragraph
        for paragraph in inventory["paragraphs_detail"]
        if paragraph["heading_level"]
        and __import__("re").match(r"^\s*\d+(?:\.\d+)*[.)]?\s+", paragraph["text"])
        and paragraph["numbering"] is None
    ]
    if heading_literals:
        report["numbering_status"] = "FAIL"
        report["issues"].append(
            {
                "severity": "ERROR",
                "code": "literal_heading_numbering",
                "message": "Heading numbers are literal text, not automatic numbering.",
            }
        )
    elif any(p["heading_level"] and p["numbering"] for p in inventory["paragraphs_detail"]):
        report["numbering_status"] = "PASS"

    if inventory["tables_detail"]:
        report["table_geometry_status"] = (
            "PASS"
            if all(t["geometry_status"] == "PASS" for t in inventory["tables_detail"])
            else "FAIL"
        )
    field_structure = inspect_field_structure(source)
    if field_structure["status"] == "FAIL":
        report["field_structure_status"] = field_structure["status"]
        report["issues"].extend(
            {
                "severity": "ERROR",
                "code": issue["code"],
                "message": str(issue),
            }
            for issue in field_structure["issues"]
        )
    combined_render_result = None
    combined_render_stage = None
    combined_word_mode = bool(
        render
        and native_fields
        and inventory["fields_detail"]
        and renderer is None
        and reuse_render_report is None
    )
    if inventory["fields_detail"]:
        if report["field_structure_status"] == "NOT_APPLICABLE":
            report["field_structure_status"] = "PASS"
        freshness = {field["freshness"] for field in inventory["fields_detail"]}
        evidence_path = source.with_suffix(source.suffix + ".wa-evidence.json")
        native_evidence = None
        if evidence_path.is_file():
            try:
                candidate = json.loads(evidence_path.read_text(encoding="utf-8"))
                if (
                    candidate.get("version") == 1
                    and candidate.get("operation") == "field.update_all"
                    and candidate.get("backend") == "word-native"
                    and candidate.get("sha256")
                    == hashlib.sha256(source.read_bytes()).hexdigest()
                ):
                    native_evidence = candidate
            except (OSError, ValueError, TypeError):
                native_evidence = None
        if "PLACEHOLDER" in freshness:
            report["field_freshness_status"] = "FAIL"
        elif native_evidence is not None:
            report["field_freshness_status"] = "STALE_OR_UNKNOWN"
            report["field_evidence"] = native_evidence
            report["field_evidence_status"] = "INFORMATIONAL_UNTRUSTED"
        else:
            report["field_freshness_status"] = "STALE_OR_UNKNOWN"
        if native_fields and combined_word_mode:
            try:
                stages, combined_render_result = _combined_native_fields_render(
                    source, render_output_dir
                )
                combined_render_stage = stages["render"]
                field_stage = stages["fields"]
                if field_stage["status"] != "SUCCESS":
                    raise RuntimeError(
                        field_stage.get("message") or "Microsoft Word field update failed"
                    )
                field_data = field_stage.get("data")
                if not isinstance(field_data, dict):
                    raise ValueError("Combined Word worker returned malformed field data")
                _record_native_field_result(report, field_data)
            except Exception as exc:
                if combined_render_stage is None:
                    combined_render_stage = {"status": "ERROR", "message": str(exc)}
                report["native_word_status"] = "FAIL"
                report["field_evaluation_status"] = "FAIL"
                report["field_freshness_status"] = "FAIL"
                report["issues"].append(
                    {"severity": "ERROR", "code": "native_field_verification_failed", "message": str(exc)}
                )
        elif native_fields:
            try:
                from ..backends.word_native.backend import _run_native_worker

                with tempfile.TemporaryDirectory(
                    prefix="wordassistant-verify-fields-"
                ) as temporary_directory:
                    verification_copy = (
                        Path(temporary_directory) / source.name
                    )
                    shutil.copy2(source, verification_copy)
                    completed = _run_native_worker(
                        {
                            "action": "verify_fields",
                            "input_path": str(verification_copy),
                        },
                        90,
                    )
                lines = [line for line in completed.stdout.splitlines() if line.strip()]
                payload = json.loads(lines[-1]) if lines else {}
                if completed.returncode != 0 or payload.get("status") != "SUCCESS":
                    raise RuntimeError(
                        payload.get("message")
                        or completed.stderr.strip()
                        or "Microsoft Word field verification failed"
                    )
                native_results = payload["data"].get("fields", [])
                bad_results = [
                    field
                    for field in native_results
                    if field.get("locked")
                    or not str(field.get("result", "")).strip()
                    or str(field.get("result", "")).strip().casefold().startswith("error!")
                ]
                if bad_results:
                    raise ValueError("Microsoft Word returned locked, empty, or error fields")
                persisted_fresh = bool(
                    payload["data"].get("source_results_were_current")
                )
                if payload["data"].get("field_freshness_comparison"):
                    report["field_freshness_comparison"] = payload["data"][
                        "field_freshness_comparison"
                    ]
                report["field_evaluation_status"] = "PASS"
                report["native_word_status"] = "PASS"
                report["native_field_results"] = native_results
                report["field_result_basis"] = "EVALUATED_IN_ISOLATED_MICROSOFT_WORD_SESSION"
                if persisted_fresh:
                    report["field_freshness_status"] = "PASS"
                else:
                    report["field_freshness_status"] = "STALE_OR_UNKNOWN"
                    report["manual_review_required"] = True
                    report["issues"].append(
                        {
                            "severity": "WARNING",
                            "code": "native_evaluation_not_persisted",
                            "message": (
                                "Microsoft Word evaluated a temporary copy successfully, but "
                                "its results differed from the cached results in the source. "
                                "Run field update-all to a distinct output before delivery."
                            ),
                        }
                    )
            except Exception as exc:
                report["native_word_status"] = "FAIL"
                report["field_evaluation_status"] = "FAIL"
                report["field_freshness_status"] = "FAIL"
                report["issues"].append(
                    {
                        "severity": "ERROR",
                        "code": "native_field_verification_failed",
                        "message": str(exc),
                    }
                )
        report["field_status"] = report["field_structure_status"]
        report["manual_review_required"] = report["field_freshness_status"] != "PASS"

    preflight = visual_preflight(source)
    if preflight["warnings"]:
        report["issues"].extend(
            {"severity": "WARNING", "code": warning, "message": warning}
            for warning in preflight["warnings"]
        )
        report["manual_review_required"] = True

    if reuse_render_report is not None:
        try:
            prior = json.loads(Path(reuse_render_report).read_text(encoding="utf-8"))
            manifest = prior.get("render_reuse_manifest")
            if not isinstance(manifest, dict) or manifest.get("version") != 1:
                raise ValueError("Reusable render report has no supported manifest.")
            if manifest != _reuse_manifest(
                source,
                Path(prior["render"]["pdf_path"]),
                tuple(Path(item["path"]) for item in manifest["pages"]),
                str(prior["render"]["engine"]),
            ):
                raise ValueError("Reusable render artifacts or manifest have changed.")
            if len(manifest["pages"]) != manifest["page_count"] or [p["page"] for p in manifest["pages"]] != list(range(1, manifest["page_count"] + 1)):
                raise ValueError("Reusable page manifest is not complete and ordered.")
            token = _reuse_token(manifest)
            if not visual_reviewed or not visual_review_token or not hmac.compare_digest(token, visual_review_token):
                raise ValueError("Versioned visual-reuse token is missing or invalid.")
            pdf_path = Path(prior["render"]["pdf_path"])
            pages = tuple(Path(item["path"]) for item in manifest["pages"])
            count = _pdf_page_count(pdf_path)
            if count != manifest["page_count"]:
                raise ValueError("Reusable PDF page count changed.")
            reported_count = prior["render"].get("page_count")
            if reported_count is not None and reported_count != count:
                raise ValueError(
                    f"Reusable renderer reported {reported_count} page(s), but the PDF contains {count}."
                )
            audit = audit_pdf_page_content(pdf_path, page_count=count)
            report["render_status"] = "PASS"
            report["render"] = {**prior["render"], "verified_pdf_page_count": count, "page_images": [str(p) for p in pages], "page_content_audit": audit}
            report["render_reuse_manifest"] = manifest
            report["visual_review_reuse"] = {"version": 1, "token": token, "basis": "CURRENT_ARTIFACT_HASHES_REVALIDATED_WITHOUT_RERENDER"}
            if audit["status"] == "PASS":
                report["visual_status"] = "PASS"
            elif audit["status"] == "UNAVAILABLE":
                report["visual_status"] = "REVIEW_REQUIRED"
                report["manual_review_required"] = True
                report["issues"].append(
                    {
                        "severity": "WARNING",
                        "code": "page_content_audit_unavailable",
                        "message": audit.get("reason", "Rendered-page text audit is unavailable."),
                    }
                )
            else:
                report["visual_status"] = audit["status"]
            if audit["status"] == "PASS":
                report["visual_review_attestation"] = {"reviewed": True, "page_count": count, "challenge_token": token, "basis": "VERSIONED_RENDER_REPORT_REUSE"}
            elif audit["status"] == "FAIL":
                report["issues"].append({"severity": "ERROR", "code": "near_blank_rendered_pages", "message": "Reused PDF has near-blank page(s)."})
        except Exception as exc:
            report["render_status"] = "FAIL"
            report["visual_status"] = "FAIL"
            report["issues"].append({"severity": "ERROR", "code": "render_reuse_invalid", "message": str(exc)})

    if render:
        try:
            render_function = renderer or _render_via_execution_service
            if combined_word_mode and combined_render_result is None:
                failure = (combined_render_stage or {}).get("message") or (
                    "The combined Word worker did not produce a usable PDF."
                )
                report["issues"].append(
                    {
                        "severity": "WARNING",
                        "code": "native_render_fallback",
                        "message": (
                            "Combined Microsoft Word rendering failed; the configured "
                            f"fallback renderer was attempted. Reason: {failure}"
                        ),
                    }
                )
            result = combined_render_result or render_function(
                source, output_dir=render_output_dir
            )
            report["render_status"] = "PASS"
            report["render"] = {
                "engine": result.engine,
                "pdf_path": str(result.pdf_path),
                "file_size_bytes": result.file_size_bytes,
                "page_count": result.page_count,
            }
            page_images = render_pdf_pages(result.pdf_path)
            report["render"]["page_images"] = [str(path) for path in page_images]
            actual_pdf_page_count = _pdf_page_count(Path(result.pdf_path))
            report["render"]["verified_pdf_page_count"] = actual_pdf_page_count
            if result.engine == "word" and report["native_word_status"] != "FAIL":
                report["native_word_status"] = "PASS"
            if page_images:
                pages_tuple = tuple(page_images)
                manifest = _reuse_manifest(source, Path(result.pdf_path), pages_tuple, result.engine)
                report["render_reuse_manifest"] = manifest
                report["visual_review_reuse"] = {"version": 1, "token": _reuse_token(manifest)}
                challenge = _visual_review_challenge(
                    source, Path(result.pdf_path), page_images
                )
                report["visual_review_challenge"] = challenge
                reported_mismatch = (
                    result.page_count is not None
                    and result.page_count != actual_pdf_page_count
                )
                image_mismatch = actual_pdf_page_count != len(page_images)
                if reported_mismatch or image_mismatch:
                    report["visual_status"] = "FAIL"
                    report["issues"].append(
                        {
                            "severity": "ERROR",
                            "code": "rendered_page_count_mismatch",
                            "message": (
                                f"Renderer reported {result.page_count!r} page(s), PDFium found "
                                f"{actual_pdf_page_count}, and {len(page_images)} page image(s) were produced."
                            ),
                        }
                    )
                else:
                    content_audit = audit_pdf_page_content(
                        result.pdf_path, page_count=len(page_images)
                    )
                    report["render"]["page_content_audit"] = content_audit
                    if content_audit["status"] == "FAIL":
                        report["visual_status"] = "FAIL"
                        report["issues"].append(
                            {
                                "severity": "ERROR",
                                "code": "near_blank_rendered_pages",
                                "message": (
                                    "Rendered page(s) contain too little visible text: "
                                    + ", ".join(
                                        str(value)
                                        for value in content_audit["near_blank_pages"]
                                    )
                                ),
                            }
                        )
                    elif content_audit["status"] == "PASS" and visual_reviewed:
                        if visual_review_token and hmac.compare_digest(
                            visual_review_token, challenge["token"]
                        ):
                            report["visual_status"] = "PASS"
                            report["visual_review_attestation"] = {
                                "reviewed": True,
                                "page_count": len(page_images),
                                "challenge_token": challenge["token"],
                                "basis": "CALLER_ATTESTATION_BOUND_TO_CURRENT_PAGE_IMAGE_HASHES",
                            }
                        else:
                            report["visual_status"] = "FAIL"
                            report["issues"].append(
                                {
                                    "severity": "ERROR",
                                    "code": "visual_review_attestation_invalid",
                                    "message": (
                                        "--visual-reviewed requires the token from the current "
                                        "visual_review_challenge after every listed page image is inspected."
                                    ),
                                }
                            )
                    else:
                        report["visual_status"] = "REVIEW_REQUIRED"
                        if content_audit["status"] == "UNAVAILABLE":
                            report["issues"].append(
                                {
                                    "severity": "WARNING",
                                    "code": "page_content_audit_unavailable",
                                    "message": content_audit.get(
                                        "reason", "Rendered-page text audit is unavailable."
                                    ),
                                }
                            )
            else:
                report["visual_status"] = "UNAVAILABLE"
                report["issues"].append(
                    {
                        "severity": "WARNING",
                        "code": "page_images_unavailable",
                        "message": "PDF rendered, but page-image rendering is unavailable.",
                    }
                )
            if report["visual_status"] != "PASS":
                report["manual_review_required"] = True
        except Exception as exc:
            report["render_status"] = "FAIL"
            report["manual_review_required"] = True
            report["issues"].append(
                {"severity": "ERROR", "code": "render_failed", "message": str(exc)}
            )

    failing = {
        report["structural_status"],
        report["format_status"],
        report["numbering_status"],
        report["table_geometry_status"],
        report["field_structure_status"],
        report["field_freshness_status"],
        report["render_status"],
        report["visual_status"],
    }
    if "FAIL" in failing:
        report["overall_status"] = "FAIL"
    elif report["manual_review_required"] or any(
        value in {"STALE_OR_UNKNOWN", "REVIEW_REQUIRED"}
        for value in failing | {report["visual_status"], report["native_word_status"]}
    ):
        report["overall_status"] = "INCOMPLETE"
    else:
        report["overall_status"] = "PASS"
    delivery_required = {
        "overall": report["overall_status"],
        "render": report["render_status"],
        "visual": report["visual_status"],
        "native_word": report["native_word_status"],
    }
    if report["overall_status"] == "FAIL":
        report["delivery_status"] = "FAIL"
    elif (
        report["overall_status"] == "PASS"
        and report["render_status"] == "PASS"
        and report["visual_status"] == "PASS"
        and report["native_word_status"] == "PASS"
    ):
        report["delivery_status"] = "PASS"
    else:
        report["delivery_status"] = "INCOMPLETE"
    report["delivery_gates"] = delivery_required
    return report
