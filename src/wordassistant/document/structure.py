from __future__ import annotations
from collections import defaultdict
import hashlib
from pathlib import Path
from typing import Any, Iterator
import re
import unicodedata
from docx import Document
from docx.document import Document as Doc
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from .model import DocumentSection
from wordassistant.operations.fields import inspect_fields
from wordassistant.selectors import _heading_level, _outline_level

_MAJOR = re.compile(
    r'^(?:'
    r'chương\s+\d+|chapter\s+\d+|'
    r'phần\s*:\s*[A-ZÀ-Ỹ]|'
    r'mở đầu(?:\s*[:.]?.*)?$|lời mở đầu(?:\s*[:.]?.*)?$|'
    r'kết luận(?:\s+và\s+hướng\s+phát\s+triển)?\s*$|'
    r'phụ lục(?:\s*\([^)]*\))?\s*$|'
    r'danh mục tài liệu tham khảo(?:\s+\d+)?\s*$|'
    r'tài liệu tham khảo\.?\s*$|'
    r'quy định về trích dẫn(?:\s*[:.]?.*)?$|'
    r'hướng dẫn\s*$|trình bày đồ án.*|'
    r'nội dung\s*$|nội dung của các chương mục.*|'
    r'dự kiến kế hoạch thực hiện.*|các cơ quan, đơn vị cần liên hệ.*|'
    r'i{1,3}v?\.|v\.'
    r')', re.I)
_TOP_GUIDANCE = re.compile(r'^(?:trang\s+|cách thể hiện báo cáo|nội dung của các chương mục|phụ lục:|tài liệu tham khảo\.?$|quy định về trích dẫn)', re.I)
_MINOR = re.compile(r'^(?:\d+\.\s+|\d+\.\d+\.?\s+|\d+\.\d+\.\d+\.?\s+|\d+\.\s*[A-ZÀ-Ỹ])')
_EMU_PER_CM = 360000

def index_docx(path: str | Path) -> list[DocumentSection]:
    d = Document(path)
    out = []
    counts = defaultdict(int)
    stack = {}
    title = 'Document preamble'
    sid = '__preamble__'
    level = 0
    parent = None
    buf = []
    start_idx = 0
    sample_mode = False
    
    def flush():
        nonlocal buf
        text = '\n\n'.join(x for x in buf if x.strip()).strip()
        if text or sid != '__preamble__':
            out.append(DocumentSection(sid, title, text, level, parent, start_idx))
        buf = []
        
    for idx, block in enumerate(iter_blocks(d)):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            is_sample_marker = bool(text) and text.casefold().startswith('(mẫu')
            if is_sample_marker:
                h = 1
                sample_mode = True
            elif sample_mode:
                h = None
            else:
                h = detect_heading(block)
            if text and h is not None:
                flush()
                title = text
                level = h
                base = slug(text)
                counts[base] += 1
                sid = base if counts[base] == 1 else f"{base}-{counts[base]}"
                parent = next((stack[x] for x in sorted(stack, reverse=True) if x < level), None)
                stack[level] = sid
                for k in list(stack):
                    if k > level:
                        del stack[k]
                start_idx = idx
                continue
            if text:
                buf.append(text)
        else:
            t = table_text(block)
            if t:
                buf.append(t)
    flush()
    return _coalesce(out)

def detect_heading(p: Paragraph) -> int | None:
    text = p.text.strip()
    if not text:
        return None
    semantic_level = _heading_level(p)
    if semantic_level is not None:
        return semantic_level
    ppr = p._p.pPr
    if _outline_level(p) == 9:
        # Direct Body overrides and overridden real headings suppress heading
        # semantics. Legacy body/list styles still need conservative heuristics.
        if ppr is not None and ppr.find(qn('w:outlineLvl')) is not None:
            return None
        style = p.style
        seen = set()
        while style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            if re.fullmatch(r'Heading [1-9]', style.name or ''):
                return None
            outline = style.element.find(f"{qn('w:pPr')}/{qn('w:outlineLvl')}")
            if style.name != 'Normal' and outline is not None and outline.get(qn('w:val')) in {str(i) for i in range(9)}:
                return None
            style = style.base_style
    num_ilvl = None
    if ppr is not None and ppr.numPr is not None and ppr.numPr.ilvl is not None:
        num_ilvl = int(ppr.numPr.ilvl.val)
    meaningful = [r for r in p.runs if r.text.strip()]
    bold = bool(meaningful) and sum(1 for r in meaningful if r.bold is True) >= max(1, len(meaningful) // 2)
    if num_ilvl == 0 and bold and len(text) <= 180 and (_MAJOR.match(text) or _TOP_GUIDANCE.match(text)):
        return 1
    if num_ilvl is not None and num_ilvl > 0 and bold and len(text) <= 180:
        return min(6, num_ilvl + 1)
    if _MAJOR.match(text) and (bold or text.upper() == text or len(text) <= 90):
        return 1
    if _MINOR.match(text) and bold and len(text) <= 160:
        dots = text.split()[0].count('.')
        return min(6, max(2, dots + 1))
    return None

def iter_blocks(d: Doc) -> Iterator[Paragraph | Table]:
    for child in d.element.body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, d)
        elif child.tag == qn('w:tbl'):
            yield Table(child, d)

def table_text(t: Table) -> str:
    rows = []
    for row in t.rows:
        cells = [' '.join(p.text.strip() for p in c.paragraphs if p.text.strip()) for c in row.cells]
        if any(cells):
            rows.append(' | '.join(cells))
    return '\n'.join(rows)

def slug(s: str) -> str:
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower()
    s = re.sub(r'[^a-z0-9]+', '-', s).strip('-')[:64]
    return s or 'section'

def _coalesce(sections: list[DocumentSection]) -> list[DocumentSection]:
    if len(sections) < 2:
        return sections
    result = []
    i = 0
    while i < len(sections):
        s = sections[i]
        if s.section_id != '__preamble__' and len(s.text) < 20 and len(s.title) < 35 and i + 1 < len(sections) and s.level >= 2:
            n = sections[i + 1]
            merged = DocumentSection(n.section_id, n.title, (s.title + '\n' + s.text + '\n' + n.text).strip(), n.level, n.parent_id, s.source_index, n.metadata)
            sections = sections[:i] + [merged] + sections[i + 2:]
            continue
        result.append(s)
        i += 1
    return result


def _paragraph_style_name(document, paragraph_element) -> str | None:
    style_values = paragraph_element.xpath("./w:pPr/w:pStyle/@w:val")
    if not style_values:
        return None
    style_id = style_values[0]
    try:
        return document.styles[style_id].name
    except (KeyError, TypeError):
        return style_id


def _field_evidence(paragraph_element) -> list[dict[str, str]]:
    evidence = []
    for field in paragraph_element.xpath(".//w:fldSimple"):
        instruction = (field.get(qn("w:instr")) or "").strip()
        if instruction:
            evidence.append({"kind": "simple", "instruction": instruction})
    complex_instruction = "".join(
        element.text or ""
        for element in paragraph_element.xpath(".//w:instrText")
    ).strip()
    if complex_instruction:
        evidence.append({"kind": "complex", "instruction": complex_instruction})
    return evidence


def _body_paragraph_context(document) -> tuple[list, dict]:
    paragraphs = list(document.element.body.xpath(".//w:p"))
    tree = document.element.getroottree()
    paragraph_indexes = {
        tree.getpath(paragraph): index for index, paragraph in enumerate(paragraphs)
    }
    contexts = {}
    section_index = 0
    for child in document.element.body.iterchildren():
        child_paragraphs = (
            [child, *child.xpath(".//w:p")]
            if child.tag == qn("w:p")
            else child.xpath(".//w:p")
        )
        for paragraph in child_paragraphs:
            paragraph_path = tree.getpath(paragraph)
            contexts[paragraph_path] = {
                "paragraph_index": paragraph_indexes[paragraph_path],
                "section_index": min(section_index, len(document.sections) - 1),
            }
        if child.tag == qn("w:p"):
            properties = child.find(qn("w:pPr"))
            if properties is not None and properties.find(qn("w:sectPr")) is not None:
                section_index += 1
    return paragraphs, contexts


def _inspect_images(document) -> list[dict[str, Any]]:
    paragraphs, contexts = _body_paragraph_context(document)
    tree = document.element.getroottree()
    images = []
    for paragraph in paragraphs:
        context = contexts[tree.getpath(paragraph)]
        alignment_values = paragraph.xpath("./w:pPr/w:jc/@w:val")
        alignment = alignment_values[0] if alignment_values else "left"
        next_paragraph = paragraph.getnext()
        caption = None
        caption_fields = []
        if next_paragraph is not None and next_paragraph.tag == qn("w:p"):
            next_fields = _field_evidence(next_paragraph)
            seq_fields = [
                field
                for field in next_fields
                if field["instruction"].lstrip().upper().startswith("SEQ ")
            ]
            caption_style = _paragraph_style_name(document, next_paragraph)
            if caption_style == "Caption" or seq_fields:
                caption_fields = seq_fields
                caption = {
                    "paragraph_index": contexts.get(tree.getpath(next_paragraph), {}).get("paragraph_index"),
                    "style": caption_style,
                    "text": "".join(
                        element.text or ""
                        for element in next_paragraph.iter(qn("w:t"))
                    ),
                    "seq_field": bool(seq_fields),
                    "field_evidence": seq_fields,
                }
        for drawing in paragraph.xpath(".//w:drawing"):
            inline_nodes = drawing.xpath("./wp:inline")
            anchor_nodes = drawing.xpath("./wp:anchor")
            container = inline_nodes[0] if inline_nodes else (anchor_nodes[0] if anchor_nodes else None)
            if container is None:
                continue
            extents = container.xpath("./wp:extent")
            properties = container.xpath("./wp:docPr")
            blips = container.xpath(".//a:blip/@r:embed")
            width = int(extents[0].get("cx")) if extents else 0
            height = int(extents[0].get("cy")) if extents else 0
            relationship_id = blips[0] if blips else None
            media_part = None
            media_hash = None
            if relationship_id is not None:
                part = document.part.related_parts.get(relationship_id)
                if part is not None:
                    media_part = str(part.partname)
                    media_hash = hashlib.sha256(part.blob).hexdigest()
            doc_properties = properties[0] if properties else None
            images.append(
                {
                    "index": len(images),
                    **context,
                    "paragraph_text": "".join(
                        element.text or "" for element in paragraph.iter(qn("w:t"))
                    ),
                    "paragraph_style": _paragraph_style_name(document, paragraph),
                    "width_emu": width,
                    "height_emu": height,
                    "width_cm": round(width / _EMU_PER_CM, 6),
                    "height_cm": round(height / _EMU_PER_CM, 6),
                    "aspect_ratio": round(width / height, 8) if height else None,
                    "alignment": alignment,
                    "inline": bool(inline_nodes),
                    "alt_text": doc_properties.get("descr") if doc_properties is not None else None,
                    "title": doc_properties.get("title") if doc_properties is not None else None,
                    "relationship_id": relationship_id,
                    "media_part": media_part,
                    "media_sha256": media_hash,
                    "associated_caption": caption,
                    "caption_field_evidence": caption_fields,
                }
            )
    return images

def inspect_document(path: str | Path) -> dict[str, Any]:
    src = Path(path).resolve()
    d = Document(src)
    doc_xml_parts = [d._element.xml]
    for s in d.sections:
        try:
            doc_xml_parts.append(s.header._element.xml)
            doc_xml_parts.append(s.footer._element.xml)
        except Exception:
            pass
            
    has_protection = False
    extra_comments = 0
    extra_footnotes = 0
    extra_endnotes = 0
    try:
        import zipfile
        with zipfile.ZipFile(src) as zf:
            for name in zf.namelist():
                if name.startswith('word/'):
                    try:
                        content = zf.read(name).decode('utf-8', errors='ignore')
                        if name == 'word/settings.xml':
                            if '<w:documentProtection' in content:
                                has_protection = True
                        elif name == 'word/comments.xml':
                            extra_comments = content.count('<w:comment ')
                        elif name == 'word/footnotes.xml':
                            extra_footnotes = content.count('<w:footnote ')
                        elif name == 'word/endnotes.xml':
                            extra_endnotes = content.count('<w:endnote ')
                        if name not in ('word/document.xml', 'word/header1.xml', 'word/header2.xml', 'word/header3.xml',
                                        'word/footer1.xml', 'word/footer2.xml', 'word/footer3.xml'):
                            doc_xml_parts.append(content)
                    except Exception:
                        pass
    except Exception:
        pass
        
    full_xml = "\n".join(doc_xml_parts)
    sections = index_docx(src)
    
    # Inventory counts
    p_count = len(d.paragraphs)
    tbl_count = len(d.tables)
    sec_count = len(d.sections)
    
    img_count = len(re.findall(r'<w:drawing\b', full_xml)) + len(re.findall(r'<w:pict\b', full_xml))
    hyperlink_count = len(re.findall(r'<w:hyperlink\b', full_xml))
    field_tag_count = len(re.findall(r'<w:fldSimple\b', full_xml)) + len(re.findall(r'<w:fldChar\b', full_xml))
    bookmark_count = len(re.findall(r'<w:bookmarkStart\b', full_xml))
    comment_count = max(extra_comments, len(re.findall(r'<w:commentReference\b', full_xml)))
    footnote_count = max(extra_footnotes, len(re.findall(r'<w:footnoteReference\b', full_xml)))
    endnote_count = max(extra_endnotes, len(re.findall(r'<w:endnoteReference\b', full_xml)))
    sdt_count = len(re.findall(r'<w:sdt\b', full_xml))
    ins_count = len(re.findall(r'<w:ins\b', full_xml))
    del_count = len(re.findall(r'<w:del\b', full_xml))
    
    # Structural field parsing (avoiding false positives from visible paragraph text)
    fld_simple_instrs = re.findall(r'<w:fldSimple\b[^>]*w:instr=["\']([^"\']+)["\']', full_xml, re.I)
    instr_text_instrs = re.findall(r'<w:instrText\b[^>]*>([^<]+)</w:instrText>', full_xml, re.I)
    all_instrs = [x.strip() for x in (fld_simple_instrs + instr_text_instrs) if x.strip()]
    
    page_fields_count = 0
    toc_fields_count = 0
    other_fields_count = 0
    
    for instr in all_instrs:
        up = f" {instr.upper()} "
        if " PAGE " in up or up.strip().startswith("PAGE"):
            page_fields_count += 1
        elif " TOC " in up or up.strip().startswith("TOC"):
            toc_fields_count += 1
        else:
            other_fields_count += 1
            
    has_page = page_fields_count > 0
    has_toc = toc_fields_count > 0
    
    adv_rec = (ins_count + del_count + comment_count + sdt_count > 0 or has_protection)

    paragraphs_detail = []
    lists_detail = []
    for index, paragraph in enumerate(d.paragraphs):
        style_name = paragraph.style.name if paragraph.style else None
        heading_level = detect_heading(paragraph)
        p_pr = paragraph._p.pPr
        num_pr = p_pr.numPr if p_pr is not None else None
        if num_pr is None and paragraph.style is not None:
            style_p_pr = paragraph.style.element.pPr
            num_pr = style_p_pr.numPr if style_p_pr is not None else None
        numbering = None
        if num_pr is not None and num_pr.numId is not None:
            numbering = {
                "num_id": int(num_pr.numId.val),
                "level": int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0,
            }
            lists_detail.append(
                {
                    "paragraph_index": index,
                    "text": paragraph.text,
                    **numbering,
                }
            )
        paragraphs_detail.append(
            {
                "index": index,
                "text": paragraph.text,
                "style": style_name,
                "heading_level": heading_level,
                "numbering": numbering,
            }
        )

    table_sections = {}
    section_index = 0
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            table_sections[id(child)] = section_index
        elif child.tag == qn("w:p"):
            properties = child.find(qn("w:pPr"))
            if properties is not None and properties.find(qn("w:sectPr")) is not None:
                section_index += 1

    all_tables = []
    seen_tables = set()

    def collect_table(table, depth, available_width):
        identity = id(table._tbl)
        if identity in seen_tables:
            return
        seen_tables.add(identity)
        all_tables.append((table, depth, available_width))
        for row in table.rows:
            for cell in row.cells:
                cell_width_element = cell._tc.tcPr.tcW
                cell_width = (
                    int(cell_width_element.get(qn("w:w"), "0"))
                    if cell_width_element is not None
                    else 0
                )
                for nested_table in cell.tables:
                    collect_table(nested_table, depth + 1, cell_width)

    for table in d.tables:
        table_section = d.sections[
            table_sections.get(id(table._tbl), len(d.sections) - 1)
        ]
        printable_width = int(
            table_section.page_width.twips
            - table_section.left_margin.twips
            - table_section.right_margin.twips
        )
        collect_table(table, 0, printable_width)

    tbl_count = len(all_tables)

    tables_detail = []
    for index, (table, depth, available_width) in enumerate(all_tables):
        tbl_pr = table._tbl.tblPr
        caption = tbl_pr.find(qn("w:tblCaption"))
        description = tbl_pr.find(qn("w:tblDescription"))
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        tbl_layout = tbl_pr.find(qn("w:tblLayout"))
        grid = [int(col.get(qn("w:w"), "0")) for col in table._tbl.tblGrid.gridCol_lst]
        width = int(tbl_w.get(qn("w:w"), "0")) if tbl_w is not None else 0
        row_widths = []
        for row in table._tbl.tr_lst:
            row_widths.append(
                sum(
                    int(cell.tcPr.tcW.get(qn("w:w"), "0"))
                    for cell in row.tc_lst
                    if cell.tcPr.tcW is not None
                )
            )
        geometry_ok = (
            tbl_w is not None
            and tbl_w.get(qn("w:type")) == "dxa"
            and width > 0
            and sum(grid) == width
            and all(value == width for value in row_widths)
            and width
            + (int(tbl_ind.get(qn("w:w"), "0")) if tbl_ind is not None else 0)
            <= available_width
        )
        tables_detail.append(
            {
                "index": index,
                "depth": depth,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "caption": caption.get(qn("w:val")) if caption is not None else None,
                "internal_id": description.get(qn("w:val")) if description is not None else None,
                "geometry_status": "PASS" if geometry_ok else "FAIL",
                "width_dxa": width,
                "indent_dxa": int(tbl_ind.get(qn("w:w"))) if tbl_ind is not None else None,
                "layout": tbl_layout.get(qn("w:type")) if tbl_layout is not None else None,
                "grid_widths_dxa": grid,
                "row_widths_dxa": row_widths,
            }
        )

    bookmarks_detail = [
        {
            "name": element.get(qn("w:name")),
            "id": element.get(qn("w:id")),
        }
        for element in d.element.xpath(".//w:bookmarkStart")
        if element.get(qn("w:name")) and not element.get(qn("w:name")).startswith("_")
    ]
    fields_detail = inspect_fields(src)
    images_detail = _inspect_images(d)
    sections_detail = []
    for index, section in enumerate(d.sections):
        sections_detail.append(
            {
                "index": index,
                "orientation": "landscape" if section.page_width > section.page_height else "portrait",
                "width_cm": round(section.page_width.cm, 3),
                "height_cm": round(section.page_height.cm, 3),
                "margins_cm": {
                    "top": round(section.top_margin.cm, 3),
                    "bottom": round(section.bottom_margin.cm, 3),
                    "left": round(section.left_margin.cm, 3),
                    "right": round(section.right_margin.cm, 3),
                },
                "header_linked": bool(section.header.is_linked_to_previous),
                "footer_linked": bool(section.footer.is_linked_to_previous),
            }
        )
    
    return {
        'file': str(src),
        'file_name': src.name,
        'size_bytes': src.stat().st_size if src.exists() else 0,
        'paragraphs': p_count,
        'tables': tbl_count,
        'sections': sec_count,
        'images': img_count,
        'hyperlinks': hyperlink_count,
        'bookmarks': bookmark_count,
        'comments': comment_count,
        'footnotes': footnote_count,
        'endnotes': endnote_count,
        'content_controls': sdt_count,
        'document_protection': has_protection,
        'revisions': {
            'insertions': ins_count,
            'deletions': del_count
        },
        'fields': {
            'total_field_tags': field_tag_count,
            'toc_field': has_toc,
            'page_field': has_page,
            'page_fields': page_fields_count,
            'toc_fields': toc_fields_count,
            'other_fields': other_fields_count
        },
        'advanced_backend_recommended': adv_rec,
        'paragraphs_detail': paragraphs_detail,
        'lists_detail': lists_detail,
        'tables_detail': tables_detail,
        'images_detail': images_detail,
        'bookmarks_detail': bookmarks_detail,
        'fields_detail': fields_detail,
        'sections_detail': sections_detail,
        'outline': [
            {'id': s.section_id, 'title': s.title, 'level': s.level}
            for s in sections
        ]
    }
