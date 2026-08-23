from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
import re
import unicodedata
from docx import Document
from wordassistant.profiles.loader import DocumentProfile, Variant

@dataclass(frozen=True)
class Issue:
    severity: str  # 'error', 'warning', 'info'
    code: str
    message: str

@dataclass(frozen=True)
class ValidationResult:
    valid: bool
    issues: tuple[Issue, ...]

GENERIC_DRAFT_PLACEHOLDERS = (
    "[front matter content]",
    "[section content...]",
    "[list of figures]",
    "[list of tables]",
    "[document content]",
    "untitled document",
)

def _normalize_text(s: str) -> str:
    if not s:
        return ""
    norm = unicodedata.normalize('NFC', s)
    return " ".join(norm.split()).casefold()

def _find_draft_placeholders(d: Document, v: Variant) -> list[str]:
    found = []
    check_set = set(GENERIC_DRAFT_PLACEHOLDERS)
    
    # Add defaults from profile metadata fields if they have draft brackets or angle brackets
    if 'metadata_fields' in v.cover:
        for mf in v.cover['metadata_fields']:
            d_val = mf.get('default')
            if d_val and (d_val.startswith('[') or d_val.startswith('<')):
                check_set.add(_normalize_text(d_val))
                
    # Check paragraphs
    for p in d.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        txt_norm = _normalize_text(txt)
        for pl in check_set:
            if pl in txt_norm:
                found.append(txt[:80])
                break
                
    # Check tables
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                c_txt = cell.text.strip()
                if not c_txt:
                    continue
                c_norm = _normalize_text(c_txt)
                for pl in check_set:
                    if pl in c_norm:
                        found.append(c_txt[:80])
                        break
    return found

def _extract_cover_metadata(d: Document, v: Variant) -> dict[str, str]:
    meta: dict[str, str] = {}
    cover_lines: list[str] = []
    
    front_titles = [t.casefold() for t in v.front_matter_titles.values() if len(t) > 3]
    blueprint_titles = [sec['title'].casefold() for sec in v.section_blueprint[:2]]
    
    for p in d.paragraphs[:50]:
        txt = p.text.strip()
        if not txt:
            continue
        txt_lower = txt.lower()
        if any(ft in txt_lower for ft in front_titles):
            break
        if any(bt in txt_lower for bt in blueprint_titles):
            break
        for line in txt.split('\n'):
            clean_line = line.strip()
            if clean_line:
                cover_lines.append(clean_line)
                
    # Match metadata fields from profile cover declarations
    for line in cover_lines:
        line_norm = _normalize_text(line)
        for mf in v.cover.get('metadata_fields', []):
            label = mf.get('label', '').strip()
            key = mf.get('value_key', '')
            if label and key:
                label_norm = _normalize_text(label)
                if label_norm in line_norm:
                    idx = line.lower().find(label.lower())
                    if idx != -1:
                        val = line[idx + len(label):].strip().lstrip(':').strip()
                        if val and not (val.startswith('[') and val.endswith(']')) and not (val.startswith('<') and val.endswith('>')):
                            meta[key] = val
                            
    # Extract topic from document_type
    doc_type = v.cover.get('document_type', '').strip()
    if doc_type:
        doc_type_norm = _normalize_text(doc_type)
        dt_idx = None
        for i, line in enumerate(cover_lines):
            if doc_type_norm in _normalize_text(line):
                dt_idx = i
                break
        if dt_idx is not None:
            title_lines = []
            for next_line in cover_lines[dt_idx + 1:]:
                nl_norm = _normalize_text(next_line)
                if any(_normalize_text(mf.get('label', '')) in nl_norm for mf in v.cover.get('metadata_fields', []) if mf.get('label')):
                    break
                if v.cover.get('location_year') and _normalize_text(v.cover['location_year']) in nl_norm:
                    break
                if any(_normalize_text(inst) in nl_norm for inst in v.cover.get('institution_lines', [])):
                    break
                title_lines.append(next_line)
            if title_lines:
                topic_str = " ".join(title_lines).strip()
                t_prefix = v.cover.get('title_prefix', '').strip()
                if t_prefix and topic_str.lower().startswith(t_prefix.lower()):
                    topic_str = topic_str[len(t_prefix):].strip().lstrip(':').strip()
                meta['topic'] = topic_str
                
    return meta

def _resolve_run_font(run, p, doc_default_font: str | None) -> tuple[str | None, bool]:
    if hasattr(run, 'font') and run.font and run.font.name:
        return run.font.name, True
    if p.style and hasattr(p.style, 'font') and p.style.font and p.style.font.name:
        return p.style.font.name, True
    if doc_default_font:
        return doc_default_font, True
    return None, False

def _resolve_run_size(run, p, doc_default_size: float | None) -> tuple[float | None, bool]:
    if hasattr(run, 'font') and run.font and run.font.size:
        return run.font.size.pt, True
    if p.style and hasattr(p.style, 'font') and p.style.font and p.style.font.size:
        return p.style.font.size.pt, True
    if doc_default_size:
        return doc_default_size, True
    return None, False

def _resolve_line_spacing(p, doc_default_ls: float | None) -> tuple[float | None, bool]:
    if p.paragraph_format.line_spacing is not None:
        ls = p.paragraph_format.line_spacing
        if isinstance(ls, (int, float)):
            return float(ls), True
    if p.style and hasattr(p.style, 'paragraph_format') and p.style.paragraph_format.line_spacing is not None:
        ls = p.style.paragraph_format.line_spacing
        if isinstance(ls, (int, float)):
            return float(ls), True
    if doc_default_ls:
        return doc_default_ls, True
    return None, False

def validate_document(
    path: str | Path,
    profile: DocumentProfile,
    variant_name: str,
    tolerance_cm: float = 0.08,
    verify_mode: bool = False
) -> ValidationResult:
    v = profile.variant(variant_name)
    try:
        d = Document(path)
    except Exception as exc:
        return ValidationResult(
            valid=False,
            issues=(Issue('error', 'unreadable_document', f"Cannot open document: {exc}"),)
        )
        
    issues: list[Issue] = []
    is_template_clone = bool(v.template_asset)
    
    # 0. Draft Placeholder Detection
    placeholders = _find_draft_placeholders(d, v)
    if placeholders:
        sev = 'error' if verify_mode else 'warning'
        issues.append(
            Issue(
                sev,
                'draft_placeholder_detected',
                f"Unresolved draft placeholder(s) detected in document ({len(placeholders)} item(s)): {', '.join(placeholders[:3])}."
            )
        )
    
    # 1. Section & Page Geometry Validation
    if not d.sections:
        issues.append(Issue('error', 'no_sections', "Document has no sections."))
    else:
        for i, s in enumerate(d.sections):
            checks = [
                ('width', s.page_width.cm, v.page.width_cm),
                ('height', s.page_height.cm, v.page.height_cm),
                ('top_margin', s.top_margin.cm, v.page.margins.top_cm),
                ('bottom_margin', s.bottom_margin.cm, v.page.margins.bottom_cm),
                ('left_margin', s.left_margin.cm, v.page.margins.left_cm),
                ('right_margin', s.right_margin.cm, v.page.margins.right_cm),
            ]
            for label, actual, expected in checks:
                if abs(actual - expected) > tolerance_cm:
                    issues.append(
                        Issue(
                            'error',
                            'page_geometry',
                            f"Section {i+1} {label}: measured {actual:.2f} cm, expected {expected:.2f} cm."
                        )
                    )
                    
    # 2. Typography & Line Spacing Validation
    try:
        normal_style = d.styles['Normal']
        doc_default_font = normal_style.font.name if normal_style and normal_style.font else None
        doc_default_size = normal_style.font.size.pt if normal_style and normal_style.font and normal_style.font.size else None
        doc_default_ls = normal_style.paragraph_format.line_spacing if normal_style and normal_style.paragraph_format else None
        if not isinstance(doc_default_ls, (int, float)):
            doc_default_ls = None
    except Exception:
        doc_default_font = None
        doc_default_size = None
        doc_default_ls = None

    expected_font = v.typography.body_font.lower()
    expected_size = v.typography.body_size_pt
    expected_ls = v.typography.line_spacing
    
    body_paragraphs = []
    in_body = False or len(d.sections) == 1 or not any(x in v.front_matter for x in ('cover', 'inner_cover'))
    
    for p in d.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if any(sec['title'].lower() in txt.lower() for sec in v.section_blueprint[:2]):
            in_body = True
        if not in_body:
            continue
        if p.style.name.startswith('Heading') or p.style.name in ('Title', 'Subtitle', 'Header', 'Footer'):
            continue
        if txt.startswith('[') and txt.endswith(']'):
            continue
        body_paragraphs.append(p)
        
    font_errors = 0
    size_errors = 0
    ls_errors = 0
    unresolved_warnings = 0
    
    if not is_template_clone:
        for p in body_paragraphs:
            if expected_ls is not None:
                actual_ls, ls_ok = _resolve_line_spacing(p, doc_default_ls)
                if not ls_ok:
                    unresolved_warnings += 1
                elif actual_ls is not None and abs(actual_ls - expected_ls) > 0.15:
                    ls_errors += 1
                    
            runs_to_check = [r for r in p.runs if r.text.strip()]
            if not runs_to_check and p.text.strip():
                font_name, font_ok = _resolve_run_font(p, p, doc_default_font)
                if not font_ok:
                    unresolved_warnings += 1
                elif font_name and font_name.lower() != expected_font:
                    font_errors += 1
                size_val, size_ok = _resolve_run_size(p, p, doc_default_size)
                if not size_ok:
                    unresolved_warnings += 1
                elif size_val is not None and abs(size_val - expected_size) > 0.6:
                    size_errors += 1
            else:
                for r in runs_to_check:
                    font_name, font_ok = _resolve_run_font(r, p, doc_default_font)
                    if not font_ok:
                        unresolved_warnings += 1
                    elif font_name and font_name.lower() != expected_font:
                        font_errors += 1
                        
                    size_val, size_ok = _resolve_run_size(r, p, doc_default_size)
                    if not size_ok:
                        unresolved_warnings += 1
                    elif size_val is not None and abs(size_val - expected_size) > 0.6:
                        size_errors += 1

        if font_errors > 0:
            issues.append(
                Issue(
                    'error',
                    'typography_font',
                    f"Body font mismatch detected in {font_errors} text run(s); expected '{v.typography.body_font}'."
                )
            )
        if size_errors > 0:
            issues.append(
                Issue(
                    'error',
                    'typography_size',
                    f"Body font size mismatch detected in {size_errors} text run(s); expected {expected_size} pt."
                )
            )
        if ls_errors > 0:
            issues.append(
                Issue(
                    'error',
                    'line_spacing',
                    f"Line spacing mismatch detected in {ls_errors} paragraph(s); expected {expected_ls}."
                )
            )
        if unresolved_warnings > 0:
            issues.append(
                Issue(
                    'warning',
                    'unresolvable_formatting',
                    f"Unresolvable formatting properties in {unresolved_warnings} run(s)."
                )
            )

    # 3. Front Matter Presence & Sequential Order
    headings_in_doc: list[tuple[int, str]] = []
    for idx, p in enumerate(d.paragraphs):
        if p.style.name.startswith('Heading') or p.style.name == 'Title':
            headings_in_doc.append((idx, p.text.strip()))
            
    # Check Cover Signals
    front_titles = [t.casefold() for t in v.front_matter_titles.values() if len(t) > 3]
    blueprint_titles = [sec['title'].casefold() for sec in v.section_blueprint[:2]]
    
    if 'cover' in v.front_matter and not is_template_clone:
        cover_paras = []
        for p in d.paragraphs[:50]:
            t = p.text.strip()
            if not t:
                continue
            if any(ft in t.casefold() for ft in front_titles):
                break
            if any(bt in t.casefold() for bt in blueprint_titles):
                break
            cover_paras.append(t)
            
        cover_text_full = "\n".join(cover_paras)
        inst_lines = v.cover.get('institution_lines', [])
        if inst_lines:
            missing_inst = False
            for line in inst_lines:
                if line.casefold() not in cover_text_full.casefold():
                    missing_inst = True
                    break
            if missing_inst:
                issues.append(
                    Issue(
                        'error' if verify_mode else 'warning',
                        'missing_cover_institution',
                        "Required cover institution header lines are missing or altered."
                    )
                )
                
        doc_type = v.cover.get('document_type', '')
        if doc_type and doc_type.casefold() not in cover_text_full.casefold():
            issues.append(
                Issue(
                    'error' if verify_mode else 'warning',
                    'missing_cover_document_type',
                    f"Required cover document type line '{doc_type}' is missing."
                )
            )
            
        # Inner Cover Verification
        if 'inner_cover' in v.front_matter:
            cover_count = 0
            if doc_type:
                cover_count = cover_text_full.casefold().count(doc_type.casefold())
            elif inst_lines:
                cover_count = cover_text_full.casefold().count(inst_lines[0].casefold())
                
            if cover_count < 2:
                issues.append(
                    Issue(
                        'error' if verify_mode else 'warning',
                        'missing_inner_cover',
                        "Profile specifies both cover and inner_cover, but inner cover page/block was not found."
                    )
                )
            
    non_cover_front = [x for x in v.front_matter if x not in ('cover', 'inner_cover', 'proposal_cover')]
    if non_cover_front and not is_template_clone:
        last_idx = -1
        for item in non_cover_front:
            item_title = v.front_matter_titles.get(item, item.replace('_', ' ')).casefold()
            match_idx = None
            for p_idx, h_text in headings_in_doc:
                if item_title in h_text.casefold() or h_text.casefold() in item_title:
                    match_idx = p_idx
                    break
            if match_idx is None and verify_mode:
                issues.append(
                    Issue(
                        'error',
                        'missing_front_matter',
                        f"Missing required front-matter section: '{item}' ('{v.front_matter_titles.get(item, item)}')."
                    )
                )
            elif match_idx is not None:
                if match_idx < last_idx:
                    issues.append(
                        Issue(
                            'error',
                            'front_matter_order',
                            f"Front matter section '{item}' appears out of order in document."
                        )
                    )
                last_idx = match_idx

    # 4. Section Blueprint Presence, Level & Sequential Order
    if v.section_blueprint:
        missing_blueprint = []
        blueprint_indices: list[int] = []
        
        for sec in v.section_blueprint:
            target_title = str(sec['title']).casefold()
            expected_level = sec.get('level', 1)
            is_required = sec.get('required', True)
            matched_p = None
            match_idx = None
            
            for p_idx, p in enumerate(d.paragraphs):
                h_text = p.text.strip()
                if (p.style.name.startswith('Heading') or p.style.name == 'Title') and (target_title in h_text.casefold() or h_text.casefold() in target_title):
                    match_idx = p_idx
                    matched_p = p
                    break
                    
            if match_idx is None:
                doc_full_text = "\n".join(p.text for p in d.paragraphs)
                if target_title not in doc_full_text.casefold():
                    if is_required:
                        missing_blueprint.append(str(sec['title']))
            else:
                blueprint_indices.append(match_idx)
                if matched_p is not None and not is_template_clone:
                    style_name = matched_p.style.name
                    actual_level = None
                    if style_name.startswith('Heading '):
                        try:
                            actual_level = int(style_name.replace('Heading ', '').strip())
                        except ValueError:
                            actual_level = None
                    elif style_name == 'Title':
                        actual_level = 1
                        
                    if actual_level is not None:
                        if actual_level != expected_level:
                            issues.append(
                                Issue(
                                    'error' if verify_mode else 'warning',
                                    'blueprint_heading_level_mismatch',
                                    f"Section '{sec['title']}' has heading level {actual_level} ('{style_name}'), expected level {expected_level}."
                                )
                            )
                    else:
                        issues.append(
                            Issue(
                                'warning',
                                'blueprint_heading_level_unresolvable',
                                f"Section '{sec['title']}' uses custom style '{style_name}'; heading level could not be deterministically verified."
                            )
                        )
                
        if missing_blueprint:
            sev = 'warning' if (is_template_clone and not verify_mode) else 'error'
            issues.append(
                Issue(
                    sev,
                    'missing_blueprint_sections',
                    f"Missing required section(s) from profile blueprint: {', '.join(missing_blueprint[:5])}."
                )
            )
            
        if len(blueprint_indices) >= 2 and not is_template_clone:
            if blueprint_indices != sorted(blueprint_indices):
                issues.append(
                    Issue(
                        'error',
                        'blueprint_order',
                        "Blueprint sections appear out of sequence in document."
                    )
                )

    # 5. Header & Footer Requirements
    hf = v.header_footer
    if hf:
        req_page = v.field_requirements.get('page_number')
        if (hf.get('footer_right') == 'page_number' or req_page is True) and req_page is not False:
            has_page_field = False
            for s in d.sections:
                f_xml = s.footer._element.xml
                if 'w:fldSimple' in f_xml and ('PAGE' in f_xml or 'w:instr="PAGE"' in f_xml):
                    has_page_field = True
                    break
                if '<w:instrText' in f_xml and 'PAGE' in f_xml:
                    has_page_field = True
                    break
            if not has_page_field:
                issues.append(
                    Issue(
                        'error' if verify_mode else 'warning',
                        'missing_page_number_field',
                        "Profile specifies footer page numbering, but no PAGE field was detected in footers."
                    )
                )
                
        # Check page_number_start restart
        p_start = hf.get('page_number_start')
        if p_start is not None and verify_mode and not is_template_clone:
            has_pg_restart = False
            for s in d.sections:
                s_xml = s._sectPr.xml
                if 'w:pgNumType' in s_xml and ('w:start=' in s_xml or 'w:start' in s_xml):
                    has_pg_restart = True
                    break
            if not has_pg_restart:
                issues.append(
                    Issue(
                        'error',
                        'missing_page_number_restart',
                        f"Profile specifies page_number_start={p_start}, but no section with pgNumType restart was found."
                    )
                )
                
        # In verify_mode, check semantic roles
        if verify_mode and not is_template_clone:
            body_section = d.sections[-1]
            header_paras = [p for p in body_section.header.paragraphs if p.text.strip()]
            footer_paras = [p for p in body_section.footer.paragraphs if p.text.strip()]
            header_text = " ".join(p.text.strip() for p in header_paras)
            footer_text = " ".join(p.text.strip() for p in footer_paras)
            
            cover_meta = _extract_cover_metadata(d, v)
            expected_topic = cover_meta.get('topic')
            expected_sup = cover_meta.get('supervisor')
            expected_student = cover_meta.get('student')
            
            header_norm = _normalize_text(header_text)
            footer_norm = _normalize_text(footer_text)
            
            if hf.get('header_left'):
                role = hf['header_left']
                if role == 'topic':
                    if not header_text:
                        issues.append(Issue('error', 'missing_header_topic', "Profile requires header topic, but header is empty."))
                    elif expected_topic:
                        exp_topic_norm = _normalize_text(expected_topic)
                        if exp_topic_norm not in header_norm:
                            issues.append(Issue('error', 'header_topic_mismatch', f"Header text does not match document topic: '{header_text}' vs expected '{expected_topic}'."))
                    elif "untitled" in header_norm or any(_normalize_text(p) in header_norm for p in GENERIC_DRAFT_PLACEHOLDERS):
                        issues.append(Issue('error', 'header_topic_placeholder', f"Header topic contains draft placeholder: '{header_text}'."))

            if hf.get('header_right') == 'supervisor':
                if not header_text:
                    issues.append(Issue('error', 'missing_header_supervisor', "Profile requires supervisor in header_right, but header is empty."))
                elif expected_sup:
                    exp_sup_norm = _normalize_text(expected_sup)
                    if '\t' in header_text:
                        right_part = header_text.split('\t')[-1].strip()
                        right_norm = _normalize_text(right_part)
                        if right_norm != exp_sup_norm and exp_sup_norm not in header_norm:
                            issues.append(Issue('error', 'header_supervisor_mismatch', f"Header supervisor '{right_part}' does not match cover supervisor '{expected_sup}'."))
                    elif exp_sup_norm not in header_norm:
                        issues.append(Issue('error', 'header_supervisor_mismatch', f"Header text does not contain cover supervisor '{expected_sup}'."))
                else:
                    if '\t' not in header_text or len(header_text.split('\t')[-1].strip()) < 3:
                        issues.append(Issue('error', 'missing_header_supervisor', "Profile requires supervisor in header_right, but supervisor was missing."))

            if hf.get('footer_left') == 'student':
                if not footer_text:
                    issues.append(Issue('error', 'missing_footer_student', "Profile requires footer student info, but footer is empty."))
                elif expected_student:
                    exp_stu_norm = _normalize_text(expected_student)
                    if '\t' in footer_text:
                        left_part = footer_text.split('\t')[0].strip()
                        left_norm = _normalize_text(left_part)
                        if left_norm != exp_stu_norm and exp_stu_norm not in footer_norm:
                            issues.append(Issue('error', 'footer_student_mismatch', f"Footer student text '{left_part}' does not match cover student '{expected_student}'."))
                    elif exp_stu_norm not in footer_norm:
                        issues.append(Issue('error', 'footer_student_mismatch', f"Footer text does not contain cover student '{expected_student}'."))
                else:
                    if any(_normalize_text(p) in footer_norm for p in GENERIC_DRAFT_PLACEHOLDERS):
                        issues.append(Issue('error', 'footer_student_placeholder', f"Footer student contains placeholder: '{footer_text}'."))

    # 6. TOC Field Validation
    req_toc = v.field_requirements.get('table_of_contents')
    if (('table_of_contents' in v.front_matter or req_toc is True) and req_toc is not False) and not is_template_clone:
        doc_xml = d._element.xml
        has_toc = ('w:fldSimple' in doc_xml and 'TOC' in doc_xml) or ('<w:instrText' in doc_xml and 'TOC' in doc_xml)
        if not has_toc:
            issues.append(
                Issue(
                    'error' if verify_mode else 'warning',
                    'missing_toc_field',
                    "Profile includes table_of_contents front-matter, but no TOC field code was found."
                )
            )

    # 7. Bibliography / Reference Count (Scoped strictly under Bibliography section)
    minrefs = v.citations.get('minimum_references')
    if isinstance(minrefs, int) and minrefs > 0 and not is_template_clone:
        ref_pattern = re.compile(r'^\[\d+\]')
        
        bib_start_idx = None
        bib_level = 1
        for idx, p in enumerate(d.paragraphs):
            txt = p.text.strip()
            if not txt:
                continue
            txt_norm = _normalize_text(txt)
            if ('tài liệu tham khảo' in txt_norm or 'tai lieu tham khao' in txt_norm or 
                'references' in txt_norm or 'bibliography' in txt_norm):
                if p.style.name.startswith('Heading') or p.style.name == 'Title' or any(_normalize_text(sec['title']) in txt_norm for sec in v.section_blueprint):
                    bib_start_idx = idx
                    if p.style.name.startswith('Heading '):
                        try:
                            bib_level = int(p.style.name.replace('Heading ', '').strip())
                        except ValueError:
                            bib_level = 1
                    break
                    
        count = 0
        if bib_start_idx is not None:
            for p in d.paragraphs[bib_start_idx + 1:]:
                txt = p.text.strip()
                if not txt:
                    continue
                if p.style.name.startswith('Heading '):
                    try:
                        lvl = int(p.style.name.replace('Heading ', '').strip())
                        if lvl <= bib_level:
                            break
                    except ValueError:
                        pass
                if ref_pattern.match(txt):
                    count += 1
                elif txt and not (txt.startswith('[') and txt.endswith(']')) and len(txt) > 10:
                    count += 1
                    
        if count == 0:
            issues.append(
                Issue(
                    'error' if verify_mode else 'warning',
                    'missing_references',
                    f"Found 0 reference items in bibliography section; profile requires at least {minrefs}."
                )
            )
        elif count < minrefs:
            issues.append(
                Issue(
                    'error' if verify_mode else 'warning',
                    'insufficient_references',
                    f"Found {count} reference item(s) in bibliography section; profile requires at least {minrefs}."
                )
            )

    valid = not any(i.severity == 'error' for i in issues)
    return ValidationResult(valid=valid, issues=tuple(issues))
