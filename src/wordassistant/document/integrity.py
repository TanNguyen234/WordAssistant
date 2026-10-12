from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from typing import Any, Iterable, Mapping
import re

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from wordassistant.evidence.registry import EvidenceError, EvidenceRegistry
from wordassistant.profiles.loader import DocumentProfile, Variant, resolve_profile
from wordassistant.selectors import _heading_level


INTEGRITY_SCHEMA_VERSION = "1.0"

_INVISIBLE = {
    "\u00ad": "SOFT_HYPHEN",
    "\u200b": "ZERO_WIDTH_SPACE",
    "\u200c": "ZERO_WIDTH_NON_JOINER",
    "\u200d": "ZERO_WIDTH_JOINER",
    "\u2060": "WORD_JOINER",
    "\ufeff": "BYTE_ORDER_MARK",
}
_PUNCTUATION = {
    "em_dash": "\u2014",
    "en_dash": "\u2013",
    "left_curly_double_quote": "\u201c",
    "right_curly_double_quote": "\u201d",
    "left_curly_single_quote": "\u2018",
    "right_curly_single_quote": "\u2019",
    "ellipsis": "\u2026",
    "non_breaking_space": "\u00a0",
}
_CITATION = re.compile(r"\[(\d+(?:\s*(?:-|,|;)\s*\d+)*)\]")
_AUTHOR_YEAR_CITATION = re.compile(
    r"\([^()]{1,120},\s*(?:18|19|20)\d{2}[a-z]?(?:\s*[,;][^()]*)?\)",
    re.IGNORECASE,
)
_BIBLIOGRAPHY_ENTRY = re.compile(r"^\s*\[(\d+)\]\s*\S")
_BIBLIOGRAPHY_HEADING = re.compile(
    r"^(?:references|bibliography|tài liệu tham khảo|danh mục tài liệu tham khảo)\s*[:.]?$",
    re.IGNORECASE,
)
_BIBLIOGRAPHY_LANGUAGE_HEADING = re.compile(
    r"^(?:tiếng việt|tiếng anh|vietnamese|english)\s*[:.]?$",
    re.IGNORECASE,
)
_QUOTE = re.compile(r'"([^"\n]+)"|“([^”\n]+)”')
_NUMBER_OR_PERCENT = re.compile(r"(?<!\w)\d+(?:[.,]\d+)?(?:\s*%|\b)")
_PLACEHOLDER = re.compile(
    r"(?:\b(?:TBD|TODO|FIXME|LOREM\s+IPSUM)\b|"
    r"\[(?:INSERT|ADD|PLACEHOLDER|CITATION\s+NEEDED)[^\]]*\]|"
    r"\{\{[^{}]+\}\}|<\s*(?:insert|placeholder)[^>]*>)",
    re.IGNORECASE,
)
_DISCLOSURE = re.compile(
    r"(?:AI[- ]use disclosure|declaration of (?:generative )?AI use|"
    r"tuyên bố (?:về )?việc sử dụng (?:trí tuệ nhân tạo|AI)|"
    r"công bố (?:việc )?sử dụng (?:trí tuệ nhân tạo|AI)|"
    r"generative AI (?:was|was not|has been|has not been) used)",
    re.IGNORECASE,
)
_WORD = re.compile(r"[^\W_]+(?:[-'][^\W_]+)*", re.UNICODE)


def _hash_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _safe_style_name(paragraph: Paragraph) -> str:
    try:
        return paragraph.style.name if paragraph.style is not None else ""
    except Exception:
        return ""


def _table_paragraphs(table: Table) -> Iterable[Paragraph]:
    seen_cells: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            identity = id(cell._tc)
            if identity in seen_cells:
                continue
            seen_cells.add(identity)
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _table_paragraphs(nested)


def _body_paragraphs(document: DocxDocument) -> Iterable[tuple[str, Paragraph]]:
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield "body", Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for paragraph in _table_paragraphs(table):
                yield "table", paragraph


def _collect_paragraphs(document: DocxDocument) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    story_indexes: Counter[str] = Counter()

    def append(story: str, paragraph: Paragraph) -> None:
        index = story_indexes[story]
        story_indexes[story] += 1
        # python-docx Paragraph.text omits text nested in structured document
        # tags, including WordAssistant's evidence-linked citation controls.
        # Read the OOXML text nodes so the audit evaluates what Word displays.
        visible_text = "".join(
            node.text or "" for node in paragraph._p.iter(qn("w:t"))
        )
        records.append(
            {
                "story": story,
                "paragraph_index": index,
                "text": visible_text,
                "style": _safe_style_name(paragraph),
                "heading_level": _heading_level(paragraph),
            }
        )

    for story, paragraph in _body_paragraphs(document):
        append(story, paragraph)

    seen_parts: set[str] = set()
    for section_index, section in enumerate(document.sections):
        for kind, suffix, container in (
            ("header", "default", section.header),
            ("header", "first", section.first_page_header),
            ("header", "even", section.even_page_header),
            ("footer", "default", section.footer),
            ("footer", "first", section.first_page_footer),
            ("footer", "even", section.even_page_footer),
        ):
            part_name = str(container.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            story = f"{kind}[{section_index}].{suffix}"
            for paragraph in container.paragraphs:
                append(story, paragraph)
            for table in container.tables:
                for paragraph in _table_paragraphs(table):
                    append(f"{story}.table", paragraph)
    return records


def _first_variant(profile: DocumentProfile) -> Variant:
    if "default" in profile.variants:
        return profile.variants["default"]
    return next(iter(profile.variants.values()))


def _profile_policy(
    profile: Any,
    document_path: Path,
    variant_name: str | None = None,
) -> dict[str, Any]:
    policy: dict[str, Any] = {
        "profile_id": None,
        "profile_version": None,
        "minimum_references": 0,
        "prefer_primary_source": False,
        "require_reviewed_sources": False,
        "long_quote_words": 40,
        "quote_citation_words": 8,
        "repeated_opening_words": 3,
        "repeated_opening_minimum": 3,
        "ngram_size": 5,
        "repeated_ngram_minimum": 3,
        "review_punctuation": list(_PUNCTUATION),
    }
    if profile is None:
        return policy

    integrity: Mapping[str, Any] = {}
    citations: Mapping[str, Any] = {}
    if isinstance(profile, (str, Path)):
        loaded = resolve_profile(str(profile), start_path=document_path)
        variant = loaded.variant(variant_name) if variant_name else _first_variant(loaded)
        policy["profile_id"] = loaded.profile_id
        policy["profile_version"] = loaded.version
        citations = variant.citations
        integrity = variant.integrity
    elif isinstance(profile, DocumentProfile):
        variant = profile.variant(variant_name) if variant_name else _first_variant(profile)
        policy["profile_id"] = profile.profile_id
        policy["profile_version"] = profile.version
        citations = variant.citations
        integrity = variant.integrity
    elif isinstance(profile, Variant):
        citations = profile.citations
        integrity = profile.integrity
    elif isinstance(profile, Mapping):
        policy["profile_id"] = profile.get("id") or profile.get("profile_id")
        policy["profile_version"] = profile.get("version")
        raw_variant: Mapping[str, Any] = profile
        variants = profile.get("variants")
        if isinstance(variants, Mapping) and variants:
            if variant_name is not None and variant_name not in variants:
                raise ValueError(f"Unknown integrity profile variant: {variant_name}")
            selected = variants.get(variant_name) if variant_name else variants.get("default")
            selected = selected or next(iter(variants.values()))
            if isinstance(selected, Mapping):
                raw_variant = selected
        raw_citations = raw_variant.get("citations", profile.get("citations", {}))
        if isinstance(raw_citations, Mapping):
            citations = raw_citations
        raw_integrity = raw_variant.get("integrity", profile.get("integrity", {}))
        if isinstance(raw_integrity, Mapping):
            integrity = raw_integrity
    else:
        raise TypeError("profile must be a profile name/path, mapping, DocumentProfile, Variant, or None")

    minimum = citations.get("minimum_references", policy["minimum_references"])
    if type(minimum) is not int or minimum < 0:
        raise ValueError("minimum_references must be an integer >= 0")
    policy["minimum_references"] = minimum
    policy["prefer_primary_source"] = citations.get("prefer_primary_source", False) is True
    policy["require_reviewed_sources"] = citations.get("require_reviewed_sources", False) is True

    integer_options = {
        "long_quote_words": 1,
        "quote_citation_words": 1,
        "repeated_opening_words": 1,
        "repeated_opening_minimum": 2,
        "ngram_size": 2,
        "repeated_ngram_minimum": 2,
    }
    for key, minimum_value in integer_options.items():
        value = integrity.get(key, policy[key])
        if type(value) is not int or value < minimum_value:
            raise ValueError(f"{key} must be an integer >= {minimum_value}")
        policy[key] = value

    punctuation = integrity.get("review_punctuation", policy["review_punctuation"])
    if not isinstance(punctuation, (list, tuple, set)):
        raise ValueError("review_punctuation must be an array of punctuation names")
    unknown = set(punctuation) - set(_PUNCTUATION)
    if unknown:
        raise ValueError(f"Unknown review_punctuation values: {sorted(unknown)}")
    policy["review_punctuation"] = list(punctuation)
    return policy


def _gate(status: str, findings: list[dict[str, Any]], **details: Any) -> dict[str, Any]:
    return {"status": status, "finding_count": len(findings), "findings": findings, **details}


def _citation_numbers(text: str) -> set[int]:
    numbers: set[int] = set()
    for match in _CITATION.finditer(text):
        value = match.group(1)
        for part in re.split(r"\s*[,;]\s*", value):
            if "-" in part:
                left, right = (int(item.strip()) for item in part.split("-", 1))
                if 0 < left <= right <= 10_000 and right - left <= 100:
                    numbers.update(range(left, right + 1))
            else:
                number = int(part)
                if number > 0:
                    numbers.add(number)
    return numbers


def _has_citation_marker(text: str) -> bool:
    return bool(_citation_numbers(text) or _AUTHOR_YEAR_CITATION.search(text))


def _bibliography_indexes(records: list[dict[str, Any]]) -> tuple[set[int], set[int], int]:
    bibliography_record_indexes: set[int] = set()
    entry_numbers: set[int] = set()
    entry_count = 0
    inside = False
    for index, record in enumerate(records):
        if record["story"] != "body":
            continue
        text = record["text"].strip()
        if _BIBLIOGRAPHY_HEADING.match(text):
            inside = True
            bibliography_record_indexes.add(index)
            continue
        if inside and _BIBLIOGRAPHY_LANGUAGE_HEADING.match(text):
            bibliography_record_indexes.add(index)
            continue
        if inside and record["heading_level"] is not None and text:
            inside = False
        if not inside or not text:
            continue
        bibliography_record_indexes.add(index)
        entry_count += 1
        match = _BIBLIOGRAPHY_ENTRY.match(text)
        if match:
            entry_numbers.add(int(match.group(1)))

    if not bibliography_record_indexes:
        for index, record in enumerate(records):
            match = _BIBLIOGRAPHY_ENTRY.match(record["text"])
            if match:
                bibliography_record_indexes.add(index)
                entry_numbers.add(int(match.group(1)))
        entry_count = len(entry_numbers)
    return bibliography_record_indexes, entry_numbers, entry_count


def _location(record: Mapping[str, Any], **extra: Any) -> dict[str, Any]:
    return {
        "story": record["story"],
        "paragraph_index": record["paragraph_index"],
        **extra,
    }


def _registry_report(path: str | Path | None) -> tuple[list[dict[str, Any]], dict[str, Any] | None, str | None]:
    if path is None:
        return [], None, None
    registry_path = Path(path).expanduser().resolve()
    if not registry_path.is_file():
        return [], {"path": str(registry_path), "exists": False}, "Evidence registry does not exist"
    try:
        records = EvidenceRegistry(registry_path).list()
    except EvidenceError as exc:
        return [], {"path": str(registry_path), "exists": registry_path.exists()}, str(exc)
    return records, {
        "path": str(registry_path),
        "exists": registry_path.exists(),
        "sha256": _hash_file(registry_path) if registry_path.exists() else None,
        "record_count": len(records),
        "record_ids": sorted(str(item["id"]) for item in records),
    }, None


def _linked_evidence_ids(document: DocxDocument) -> set[str]:
    linked_ids: set[str] = set()
    for element in document.element.body.xpath(".//w:sdt"):
        properties = element.find(qn("w:sdtPr"))
        tag_node = properties.find(qn("w:tag")) if properties is not None else None
        tag = tag_node.get(qn("w:val"), "") if tag_node is not None else ""
        citation_prefix = "wordassistant:citation:"
        if tag.startswith(citation_prefix):
            evidence_id = tag[len(citation_prefix):].strip()
            if evidence_id:
                linked_ids.add(evidence_id)
    return linked_ids


def audit_integrity(
    docx_path: str | Path,
    evidence_registry_path: str | Path | None = None,
    profile: Any = None,
    variant_name: str | None = None,
) -> dict[str, Any]:
    """Audit a DOCX for academic-integrity risks without AI detection.

    The function is read-only, deterministic for fixed inputs, and returns a
    JSON-serializable report. Heuristic prose/style findings are WARN-only and
    must be reviewed by a person; they are never treated as evidence of AI use.
    """

    source = Path(docx_path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"DOCX does not exist: {source}")
    if source.suffix.lower() != ".docx":
        raise ValueError(f"Expected a .docx file: {source}")

    document = Document(source)
    records = _collect_paragraphs(document)
    policy = _profile_policy(profile, source, variant_name)
    bibliography_indexes, bibliography_numbers, bibliography_count = _bibliography_indexes(records)
    gates: dict[str, dict[str, Any]] = {}

    invisible_findings: list[dict[str, Any]] = []
    for record in records:
        for offset, character in enumerate(record["text"]):
            if character in _INVISIBLE:
                invisible_findings.append(
                    _location(
                        record,
                        offset=offset,
                        character_name=_INVISIBLE[character],
                        codepoint=f"U+{ord(character):04X}",
                    )
                )
    gates["invisible_unicode"] = _gate(
        "FAIL" if invisible_findings else "PASS",
        invisible_findings,
        classification="DETERMINISTIC_INTEGRITY_CHECK",
    )

    punctuation_findings: list[dict[str, Any]] = []
    for name in policy["review_punctuation"]:
        character = _PUNCTUATION[name]
        occurrences = sum(record["text"].count(character) for record in records)
        if occurrences:
            punctuation_findings.append(
                {"punctuation": name, "codepoint": f"U+{ord(character):04X}", "occurrences": occurrences}
            )
    gates["typography_punctuation_review"] = _gate(
        "WARN" if punctuation_findings else "PASS",
        punctuation_findings,
        classification="STYLE_REVIEW_NOT_AI_EVIDENCE",
        configured=list(policy["review_punctuation"]),
    )

    long_quote_findings: list[dict[str, Any]] = []
    uncited_quote_findings: list[dict[str, Any]] = []
    by_story: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        by_story.setdefault(record["story"], []).append(record)
    for story_records in by_story.values():
        for index, record in enumerate(story_records):
            quote_texts = [next(group for group in match.groups() if group is not None) for match in _QUOTE.finditer(record["text"])]
            if "quote" in record["style"].lower() and record["text"].strip():
                quote_texts.append(record["text"].strip())
            for quote in quote_texts:
                word_count = len(_WORD.findall(quote))
                if word_count >= policy["long_quote_words"]:
                    long_quote_findings.append(
                        _location(record, word_count=word_count, excerpt=quote[:180])
                    )
                adjacent = story_records[max(0, index - 1) : index + 2]
                if word_count >= policy["quote_citation_words"] and not any(
                    _has_citation_marker(item["text"]) for item in adjacent
                ):
                    uncited_quote_findings.append(
                        _location(record, word_count=word_count, excerpt=quote[:180])
                    )
    gates["long_verbatim_quotes"] = _gate(
        "WARN" if long_quote_findings else "PASS",
        long_quote_findings,
        classification="HUMAN_REVIEW_REQUIRED",
        word_threshold=policy["long_quote_words"],
    )
    gates["quotes_lacking_adjacent_citation"] = _gate(
        "WARN" if uncited_quote_findings else "PASS",
        uncited_quote_findings,
        classification="HUMAN_REVIEW_REQUIRED",
        word_threshold=policy["quote_citation_words"],
        adjacency="same, previous, or next paragraph in the same story",
    )

    cited_numbers: set[int] = set()
    for index, record in enumerate(records):
        if index not in bibliography_indexes:
            cited_numbers.update(_citation_numbers(record["text"]))
    missing_bibliography = sorted(cited_numbers - bibliography_numbers)
    uncited_bibliography = sorted(bibliography_numbers - cited_numbers)
    citation_findings = [
        {"kind": "citation_missing_bibliography_entry", "number": number}
        for number in missing_bibliography
    ] + [
        {"kind": "bibliography_entry_not_cited", "number": number}
        for number in uncited_bibliography
    ]
    citation_status = "FAIL" if missing_bibliography else ("WARN" if uncited_bibliography else "PASS")
    gates["numeric_citation_bibliography_consistency"] = _gate(
        citation_status,
        citation_findings,
        classification="DETERMINISTIC_CROSS_CHECK",
        citation_numbers=sorted(cited_numbers),
        bibliography_numbers=sorted(bibliography_numbers),
    )

    evidence_records, registry_meta, registry_error = _registry_report(evidence_registry_path)
    evidence_by_id = {str(item.get("id")): item for item in evidence_records}
    linked_evidence_ids = _linked_evidence_ids(document)
    linked_source_findings: list[dict[str, Any]] = []
    if linked_evidence_ids and evidence_registry_path is not None and not registry_error:
        registry = EvidenceRegistry(evidence_registry_path)
        for evidence_id in sorted(linked_evidence_ids):
            record = evidence_by_id.get(evidence_id)
            if record is None:
                linked_source_findings.append({
                    "kind": "cited_evidence_id_missing",
                    "evidence_id": evidence_id,
                })
                continue
            if record.get("review_status") == "candidate":
                linked_source_findings.append({
                    "kind": "cited_source_is_candidate",
                    "evidence_id": evidence_id,
                })
                continue
            if policy["require_reviewed_sources"] and record.get("review_status") != "reviewed":
                linked_source_findings.append({
                    "kind": "cited_source_not_reviewed",
                    "evidence_id": evidence_id,
                })
                continue
            try:
                registry.assert_citable(
                    record, require_reviewed=policy["require_reviewed_sources"]
                )
            except EvidenceError as exc:
                error_text = str(exc).casefold()
                if "hash changed" in error_text:
                    kind = "reviewed_source_hash_changed"
                elif "missing" in error_text:
                    kind = "reviewed_source_file_missing"
                else:
                    kind = "reviewed_source_file_invalid"
                linked_source_findings.append({
                    "kind": kind,
                    "evidence_id": evidence_id,
                })
    primary_records = [
        item
        for item in evidence_records
        if any(str(tag).strip().lower() in {"primary", "primary-source", "primary_source"} for tag in item.get("tags", []))
    ]
    synthetic_records = [item for item in evidence_records if item.get("synthetic") is True]
    source_findings: list[dict[str, Any]] = []
    if bibliography_count < policy["minimum_references"]:
        source_findings.append(
            {
                "kind": "minimum_references_not_met",
                "required": policy["minimum_references"],
                "actual": bibliography_count,
            }
        )
    if registry_error:
        registry_kind = "missing_evidence_registry" if registry_meta and not registry_meta.get("exists") else "invalid_evidence_registry"
        source_findings.append({"kind": registry_kind, "message": registry_error})
    elif evidence_registry_path is None and (cited_numbers or bibliography_count):
        source_findings.append({"kind": "evidence_registry_not_provided"})
    elif evidence_registry_path is not None and cited_numbers and len(evidence_records) < max(cited_numbers):
        source_findings.append(
            {
                "kind": "evidence_registry_has_fewer_records_than_numeric_citations",
                "registry_records": len(evidence_records),
                "highest_citation": max(cited_numbers),
            }
        )
    if synthetic_records:
        source_findings.append(
            {"kind": "synthetic_evidence_present", "record_ids": sorted(item["id"] for item in synthetic_records)}
        )
    source_findings.extend(linked_source_findings)
    if policy["prefer_primary_source"] and evidence_records and not primary_records:
        source_findings.append({"kind": "no_explicitly_tagged_primary_source"})

    hard_source_failure = any(
        finding["kind"]
        in {
            "minimum_references_not_met",
            "missing_evidence_registry",
            "invalid_evidence_registry",
            "evidence_registry_has_fewer_records_than_numeric_citations",
            "synthetic_evidence_present",
            "cited_evidence_id_missing",
            "cited_source_is_candidate",
            "cited_source_not_reviewed",
            "reviewed_source_file_invalid",
            "reviewed_source_hash_changed",
            "reviewed_source_file_missing",
        }
        for finding in source_findings
    )
    source_status = "FAIL" if hard_source_failure else ("WARN" if source_findings else "PASS")
    gates["source_evidence"] = _gate(
        source_status,
        source_findings,
        classification="REGISTRY_AND_DOCUMENT_CROSS_CHECK",
        minimum_references=policy["minimum_references"],
        document_reference_count=bibliography_count,
        registry=registry_meta,
        verified_record_count=len(evidence_records),
        primary_record_count=len(primary_records),
        synthetic_record_count=len(synthetic_records),
        linked_evidence_id_count=len(linked_evidence_ids),
        require_reviewed_sources=policy["require_reviewed_sources"],
        primary_status="EXPLICIT_TAG_ONLY",
    )

    quantitative_findings: list[dict[str, Any]] = []
    for index, record in enumerate(records):
        text = record["text"].strip()
        if (
            index in bibliography_indexes
            or not text
            or record["style"].lower().startswith("heading")
            or not _NUMBER_OR_PERCENT.search(text)
            or _has_citation_marker(text)
        ):
            continue
        quantitative_findings.append(_location(record, excerpt=text[:220]))
    gates["quantitative_claims_lacking_citation"] = _gate(
        "WARN" if quantitative_findings else "PASS",
        quantitative_findings,
        classification="HUMAN_REVIEW_HEURISTIC_NOT_PROOF",
    )

    prose_records = [
        record
        for index, record in enumerate(records)
        if index not in bibliography_indexes
        and record["story"] in {"body", "table"}
        and not record["style"].lower().startswith("heading")
        and not record["style"].lower().startswith(("wa cover", "wa front matter"))
    ]
    opening_locations: dict[str, list[dict[str, Any]]] = {}
    ngram_locations: dict[str, list[dict[str, Any]]] = {}
    for record in prose_records:
        tokens = [token.lower() for token in _WORD.findall(record["text"])]
        opening_size = policy["repeated_opening_words"]
        sentences = [item.strip() for item in re.split(r"(?<=[.!?])\s+", record["text"]) if item.strip()]
        for sentence_index, sentence in enumerate(sentences):
            sentence_tokens = [token.lower() for token in _WORD.findall(sentence)]
            if len(sentence_tokens) >= opening_size:
                opening = " ".join(sentence_tokens[:opening_size])
                opening_locations.setdefault(opening, []).append(
                    _location(record, sentence_index=sentence_index)
                )
        ngram_size = policy["ngram_size"]
        for offset in range(max(0, len(tokens) - ngram_size + 1)):
            ngram = " ".join(tokens[offset : offset + ngram_size])
            ngram_locations.setdefault(ngram, []).append(_location(record, token_offset=offset))

    repetition_findings: list[dict[str, Any]] = []
    for opening, locations in sorted(opening_locations.items()):
        if len(locations) >= policy["repeated_opening_minimum"]:
            repetition_findings.append(
                {"kind": "repeated_sentence_opening", "text": opening, "occurrences": len(locations), "locations": locations[:10]}
            )
    for ngram, locations in sorted(ngram_locations.items()):
        if len(locations) >= policy["repeated_ngram_minimum"]:
            repetition_findings.append(
                {"kind": "repeated_ngram", "text": ngram, "occurrences": len(locations), "locations": locations[:10]}
            )
    gates["repetitive_prose_review"] = _gate(
        "WARN" if repetition_findings else "PASS",
        repetition_findings[:100],
        classification="HUMAN_REVIEW_HEURISTIC_NOT_AI_DETECTION",
        finding_count_total=len(repetition_findings),
        ngram_size=policy["ngram_size"],
    )

    placeholder_findings: list[dict[str, Any]] = []
    for record in records:
        for match in _PLACEHOLDER.finditer(record["text"]):
            placeholder_findings.append(_location(record, text=match.group(0), offset=match.start()))
    gates["unresolved_placeholders"] = _gate(
        "FAIL" if placeholder_findings else "PASS",
        placeholder_findings,
        classification="DETERMINISTIC_COMPLETENESS_CHECK",
    )

    disclosure_findings: list[dict[str, Any]] = []
    for record in records:
        if _DISCLOSURE.search(record["text"]):
            disclosure_findings.append(_location(record, excerpt=record["text"][:300]))
    gates["ai_use_disclosure_provenance"] = _gate(
        "PASS" if disclosure_findings else "WARN",
        disclosure_findings if disclosure_findings else [{"kind": "ai_use_disclosure_not_found"}],
        classification="DISCLOSURE_CHECK_NOT_AI_DETECTION",
        declared=bool(disclosure_findings),
    )

    counts = Counter(gate["status"] for gate in gates.values())
    overall = "FAIL" if counts["FAIL"] else ("WARN" if counts["WARN"] else "PASS")
    story_counts = Counter(record["story"].split("[")[0].split(".")[0] for record in records if record["text"].strip())
    return {
        "schema_version": INTEGRITY_SCHEMA_VERSION,
        "audit_kind": "ETHICAL_ACADEMIC_INTEGRITY",
        "ai_detection_performed": False,
        "detector_evasion_assistance": False,
        "document": {
            "path": str(source),
            "sha256": _hash_file(source),
            "size_bytes": source.stat().st_size,
        },
        "profile": policy,
        "coverage": {
            "paragraph_records": len(records),
            "nonempty_story_counts": dict(sorted(story_counts.items())),
            "stories": ["body", "table", "header", "footer"],
        },
        "status": overall,
        "summary": {name: counts[name] for name in ("PASS", "WARN", "FAIL")},
        "gates": gates,
        "limitations": [
            "Similarity to external publications is not measured because no corpus or external provider is used.",
            "Prose repetition and citation-adjacency findings are review heuristics, not proof of misconduct or AI authorship.",
            "Primary-source status is accepted only from an explicit primary tag in the evidence registry.",
        ],
    }


def build_provenance_manifest(
    docx_path: str | Path,
    evidence_registry_path: str | Path | None = None,
    profile: Any = None,
    disclosure: str | Mapping[str, Any] | None = None,
    audit_report: Mapping[str, Any] | None = None,
    source_registry: Mapping[str, Any] | None = None,
    variant_name: str | None = None,
) -> dict[str, Any]:
    """Build a JSON-ready provenance manifest without writing files."""

    source = Path(docx_path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"DOCX does not exist: {source}")
    supplied_audit = audit_report is not None
    report = dict(audit_report) if supplied_audit else audit_integrity(
        source,
        evidence_registry_path=evidence_registry_path,
        profile=profile,
        variant_name=variant_name,
    )
    if report.get("document", {}).get("sha256") != _hash_file(source):
        raise ValueError("audit_report does not match the current DOCX SHA-256")

    evidence_records, registry_meta, registry_error = _registry_report(evidence_registry_path)
    if registry_error:
        raise EvidenceError(registry_error)
    if supplied_audit:
        audited_registry = report.get("gates", {}).get("source_evidence", {}).get("registry")
        audited_hash = audited_registry.get("sha256") if isinstance(audited_registry, Mapping) else None
        current_hash = registry_meta.get("sha256") if isinstance(registry_meta, Mapping) else None
        if audited_hash != current_hash:
            raise ValueError("audit_report does not match the current evidence registry SHA-256")
        # Registry and DOCX hashes do not bind the external source-store bytes.
        # Recheck only linked reviewed sources before reusing the audit summary.
        if evidence_registry_path is not None:
            registry = EvidenceRegistry(evidence_registry_path)
            linked_ids = _linked_evidence_ids(Document(source))
            for record in evidence_records:
                if record['id'] in linked_ids and record.get('review_status') == 'reviewed':
                    registry.assert_citable(record)
    detected_disclosure = report.get("gates", {}).get("ai_use_disclosure_provenance", {})
    if disclosure is None:
        disclosure_value: Any = {
            "declared": bool(detected_disclosure.get("declared")),
            "source": "document" if detected_disclosure.get("declared") else "not_found",
            "statement": None,
        }
    elif isinstance(disclosure, str):
        disclosure_value = {"declared": True, "source": "caller", "statement": disclosure}
    elif isinstance(disclosure, Mapping):
        disclosure_value = dict(disclosure)
    else:
        raise TypeError("disclosure must be a string, mapping, or None")

    try:
        package_version = version("wordassistant")
    except PackageNotFoundError:
        package_version = "unknown"
    stat = source.stat()
    return {
        "schema_version": INTEGRITY_SCHEMA_VERSION,
        "manifest_kind": "WORDASSISTANT_ACADEMIC_PROVENANCE",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "document": {
            "path": str(source),
            "sha256": _hash_file(source),
            "size_bytes": stat.st_size,
            "modified_at_utc": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
        },
        "producer": {
            "name": "wordassistant",
            "version": package_version,
            "integrity_schema_version": INTEGRITY_SCHEMA_VERSION,
        },
        "source_registry": {
            "evidence": registry_meta,
            "record_count": len(evidence_records),
            "metadata": dict(source_registry or {}),
        },
        "profile": report.get("profile"),
        "disclosure": disclosure_value,
        "audit_summary": {
            "status": report.get("status"),
            "gate_counts": report.get("summary"),
            "document_sha256": report.get("document", {}).get("sha256"),
            "ai_detection_performed": False,
            "detector_evasion_assistance": False,
        },
    }
