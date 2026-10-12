from __future__ import annotations
from pathlib import Path
import os
import shutil
import tempfile
from collections import defaultdict
from importlib.resources import files as resource_files
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from wordassistant.document.fonts import set_word_font
from wordassistant.operations.tables import normalize_table_geometry
from wordassistant.operations.citations import _set_citation_policy
from wordassistant.profiles.loader import DocumentProfile, Variant

def _add_page_number_field(paragraph):
    p = paragraph._p
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    p.append(fldSimple)

def _add_toc_field(paragraph):
    p = paragraph._p
    r1 = OxmlElement('w:r')
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    r1.append(fld1)
    p.append(r1)
    
    r2 = OxmlElement('w:r')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    r2.append(instr)
    p.append(r2)
    
    r3 = OxmlElement('w:r')
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    r3.append(fld2)
    p.append(r3)
    
    r4 = OxmlElement('w:r')
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    r4.append(fld3)
    p.append(r4)

def _set_section_page_restart(section, start_num: int = 1):
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:start'), str(start_num))
    sectPr.append(pgNumType)

def _apply_section_page(section, v: Variant):
    if v.page.orientation == 'landscape':
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(max(v.page.width_cm, v.page.height_cm))
        section.page_height = Cm(min(v.page.width_cm, v.page.height_cm))
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(min(v.page.width_cm, v.page.height_cm))
        section.page_height = Cm(max(v.page.width_cm, v.page.height_cm))
    section.top_margin = Cm(v.page.margins.top_cm)
    section.bottom_margin = Cm(v.page.margins.bottom_cm)
    section.left_margin = Cm(v.page.margins.left_cm)
    section.right_margin = Cm(v.page.margins.right_cm)

def _apply_styles(d: Document, v: Variant):
    normal = d.styles['Normal']
    set_word_font(normal.font, v.typography.body_font)
    normal.font.size = Pt(v.typography.body_size_pt)
    if v.typography.line_spacing is not None:
        normal.paragraph_format.line_spacing = v.typography.line_spacing
    if v.typography.paragraph_spacing_before_pt is not None:
        normal.paragraph_format.space_before = Pt(v.typography.paragraph_spacing_before_pt)
    if v.typography.paragraph_spacing_after_pt is not None:
        normal.paragraph_format.space_after = Pt(v.typography.paragraph_spacing_after_pt)
        
    for s_name, s_spec in v.typography.styles.items():
        style_obj = None
        if s_name == 'title':
            style_obj = d.styles['Title']
        elif s_name.startswith('heading') and s_name[7:].isdigit():
            level = int(s_name[7:])
            if 1 <= level <= 9:
                style_obj = d.styles[f'Heading {level}']
                
        if style_obj is None and s_name not in {'title'}:
            try:
                style_obj = d.styles[s_name]
            except KeyError:
                style_obj = d.styles.add_style(s_name, WD_STYLE_TYPE.PARAGRAPH)
                style_obj.base_style = normal
                style_obj.hidden = False
                style_obj.quick_style = True

        if style_obj:
            if s_name == 'WA Front Matter Title':
                # Preserve real heading semantics when this custom style replaces
                # Heading 1, without inheriting its numbering or typography.
                properties = style_obj.element.get_or_add_pPr()
                outline = properties.find(qn('w:outlineLvl'))
                if outline is None:
                    outline = OxmlElement('w:outlineLvl')
                    properties.append(outline)
                outline.set(qn('w:val'), '0')
            if 'font' in s_spec:
                set_word_font(style_obj.font, s_spec['font'])
            else:
                set_word_font(style_obj.font, v.typography.body_font)

            if 'size_pt' in s_spec:
                style_obj.font.size = Pt(s_spec['size_pt'])
            if 'bold' in s_spec:
                style_obj.font.bold = s_spec['bold']
            if 'italic' in s_spec:
                style_obj.font.italic = s_spec['italic']
            if 'alignment' in s_spec:
                align_map = {
                    'left': WD_ALIGN_PARAGRAPH.LEFT,
                    'center': WD_ALIGN_PARAGRAPH.CENTER,
                    'right': WD_ALIGN_PARAGRAPH.RIGHT,
                    'justify': WD_ALIGN_PARAGRAPH.JUSTIFY
                }
                if s_spec['alignment'] in align_map:
                    style_obj.paragraph_format.alignment = align_map[s_spec['alignment']]
            if 'line_spacing' in s_spec:
                style_obj.paragraph_format.line_spacing = s_spec['line_spacing']
            if 'space_before_pt' in s_spec:
                style_obj.paragraph_format.space_before = Pt(s_spec['space_before_pt'])
            if 'space_after_pt' in s_spec:
                style_obj.paragraph_format.space_after = Pt(s_spec['space_after_pt'])


def _apply_cover_border(section, cover_spec):
    border_spec = cover_spec.get('border')
    if not border_spec:
        return
    sect_pr = section._sectPr
    existing = sect_pr.find(qn('w:pgBorders'))
    if existing is not None:
        sect_pr.remove(existing)
    borders = OxmlElement('w:pgBorders')
    borders.set(qn('w:offsetFrom'), border_spec.get('offset_from', 'page'))
    borders.set(qn('w:display'), border_spec['display'])
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_spec = border_spec['edges'][edge]
        node = OxmlElement(f"w:{edge}")
        node.set(qn('w:val'), edge_spec['style'])
        node.set(qn('w:sz'), str(int(edge_spec['size_eighth_points'])))
        node.set(qn('w:space'), str(int(edge_spec['space_points'])))
        borders.append(node)
    sect_pr.append(borders)


def _remove_page_border(section):
    borders = section._sectPr.find(qn('w:pgBorders'))
    if borders is not None:
        section._sectPr.remove(borders)


def _profile_asset_path(relative_path: str) -> Path:
    package_root = Path(__file__).resolve().parents[1]
    direct = package_root / 'data' / relative_path
    if direct.is_file():
        return direct
    return Path(str(resource_files('wordassistant').joinpath('data', relative_path)))

_DEFAULT_FRONT_TITLES = {
    'acknowledgements': 'ACKNOWLEDGEMENTS',
    'declaration': 'DECLARATION',
    'supervisor_review': 'SUPERVISOR REVIEW',
    'reviewer_review': 'REVIEWER REVIEW',
    'work_schedule': 'WORK SCHEDULE',
    'abstract': 'ABSTRACT',
    'table_of_contents': 'TABLE OF CONTENTS',
    'list_of_figures': 'LIST OF FIGURES',
    'list_of_tables': 'LIST OF TABLES',
    'abbreviations': 'ABBREVIATIONS'
}

def _render_cover_page(
    d: Document,
    v: Variant,
    title_text: str,
    student: str,
    supervisor: str,
    cover_metadata: Mapping[str, str] | None = None,
):
    cover_spec = v.cover
    metadata = dict(cover_metadata or {})
    named_styles = cover_spec.get('styles', {})
    inst_lines = cover_spec.get('institution_lines', [])
    for index, line in enumerate(inst_lines):
        role = 'institution' if index == 0 else 'faculty'
        p_top = d.add_paragraph(style=named_styles.get(role, named_styles.get('institution', 'Normal')))
        p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_top.add_run(line)
        if p_top.style.name == 'Normal':
            r.bold = True
            r.font.size = Pt(13)

    logo_asset = cover_spec.get('logo_asset')
    if logo_asset:
        logo_path = _profile_asset_path(logo_asset)
        if not logo_path.is_file():
            raise FileNotFoundError(f"Profile cover logo is missing: {logo_path}")
        logo_paragraph = d.add_paragraph()
        logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_paragraph.add_run().add_picture(str(logo_path), width=Cm(2.5))
        
    doc_type = cover_spec.get('document_type', '')
    if doc_type:
        p_dt = d.add_paragraph(style=named_styles.get('document_type', 'Normal'))
        p_dt.add_run(f"\n\n{doc_type}")
        p_dt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p_dt.style.name == 'Normal' and p_dt.runs:
            p_dt.runs[0].bold = True
            p_dt.runs[0].font.size = Pt(16)
            
    prefix = cover_spec.get('title_prefix', '')
    p_title = d.add_paragraph("\n\n", style=named_styles.get('title', 'Normal'))
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if prefix:
        r_pre = p_title.add_run(f"{prefix}\n")
        if p_title.style.name == 'Normal':
            r_pre.bold = True
            r_pre.font.size = Pt(14)
    r_t = p_title.add_run(title_text)
    if p_title.style.name == 'Normal':
        r_t.bold = True
        r_t.font.size = Pt(18)
    
    meta_fields = cover_spec.get('metadata_fields', [])
    if meta_fields:
        p_meta = d.add_paragraph(style=named_styles.get('metadata', 'Normal'))
        p_meta.paragraph_format.left_indent = Cm(3.0)
        for mf in meta_fields:
            lbl = mf.get('label', '')
            val_key = mf.get('value_key')
            val = mf.get('value')
            default_val = mf.get('default', '')
            if val_key == 'student':
                res_val = student or default_val
            elif val_key == 'supervisor':
                res_val = supervisor or default_val
            else:
                res_val = metadata.get(val_key, val or default_val)
            p_meta.add_run(f"{lbl}:\t{res_val or ''}\n")
    elif student or supervisor:
        p_meta = d.add_paragraph()
        p_meta.paragraph_format.left_indent = Cm(3.0)
        if student:
            p_meta.add_run(f"Author:\t{student}\n")
        if supervisor:
            p_meta.add_run(f"Supervisor:\t{supervisor}\n")
            
    loc_year = cover_spec.get('location_year', '')
    if loc_year:
        year_key = cover_spec.get('location_year_key')
        template = cover_spec.get('location_year_template')
        if template:
            location = template.format_map(defaultdict(str, {key: str(value) for key, value in metadata.items()}))
        else:
            location = metadata.get(year_key, loc_year) if year_key else loc_year
        p_bottom = d.add_paragraph(f"\n\n\n{location}", style=named_styles.get('location_year', 'Normal'))
        p_bottom.alignment = WD_ALIGN_PARAGRAPH.CENTER

def _build_scaffold_document(
    output: str | Path,
    profile: DocumentProfile,
    variant_name: str,
    *,
    title: str | None = None,
    student: str | None = None,
    supervisor: str | None = None,
    cover_metadata: Mapping[str, str] | None = None,
    dry_run: bool = False
) -> Path:
    v = profile.variant(variant_name)
    out = Path(output)
    if not dry_run:
        out.parent.mkdir(parents=True, exist_ok=True)
        
    eff_student = student or ''
    eff_supervisor = supervisor or ''
    
    # Template asset clone path (e.g. for official proposal template)
    if v.template_asset:
        root = Path(__file__).resolve().parents[3]
        asset = root / v.template_asset
        if not asset.exists():
            rel = v.template_asset.split('use_cases/', 1)[-1]
            asset = Path(str(resource_files('wordassistant').joinpath('data', 'use_cases', rel)))
        if asset.exists():
            if not dry_run:
                shutil.copy2(asset, out)
                if title is not None or eff_student or eff_supervisor:
                    _fill_template_metadata(out, title=title or '', student=eff_student, supervisor=eff_supervisor)
            return out
            
    effective_title = title if title is not None else 'UNTITLED DOCUMENT'
    d = Document()
    _apply_section_page(d.sections[0], v)
    _apply_styles(d, v)
    _apply_cover_border(d.sections[0], v.cover)
    settings = d.settings.element
    update_fields = settings.find(qn('w:updateFields'))
    if update_fields is None:
        update_fields = OxmlElement('w:updateFields')
        settings.append(update_fields)
    update_fields.set(qn('w:val'), 'true')
    mirror = settings.find(qn('w:mirrorMargins'))
    if not v.page.one_sided and mirror is None:
        settings.append(OxmlElement('w:mirrorMargins'))
    elif v.page.one_sided and mirror is not None:
        settings.remove(mirror)
    
    # 1. Front matter & Covers
    has_covers = any(x in v.front_matter for x in ('cover', 'inner_cover', 'proposal_cover'))
    non_cover_front = [x for x in v.front_matter if x not in ('cover', 'inner_cover', 'proposal_cover')]
    
    if has_covers:
        _render_cover_page(d, v, effective_title, eff_student, eff_supervisor, cover_metadata)
        if 'inner_cover' in v.front_matter:
            d.add_page_break()
            _render_cover_page(d, v, effective_title, eff_student, eff_supervisor, cover_metadata)
            
    if non_cover_front:
        front_sec = d.add_section(WD_SECTION_START.NEW_PAGE)
        _apply_section_page(front_sec, v)
        _remove_page_border(front_sec)
        front_sec.header.is_linked_to_previous = False
        front_sec.footer.is_linked_to_previous = False
        from wordassistant.operations.page_numbers import apply_page_numbers_to_section
        front_number_format = v.header_footer.get('front_matter_page_number_format', 'lowerRoman')
        if front_number_format != 'none':
            apply_page_numbers_to_section(
                d,
                len(d.sections) - 1,
                position='footer',
                alignment='center',
                number_format=front_number_format,
                start=1,
                field='PAGE',
                unlink_previous=True,
            )
        
        for idx, item in enumerate(non_cover_front):
            title_text = v.front_matter_titles.get(item, _DEFAULT_FRONT_TITLES.get(item, item.replace('_', ' ').upper()))
            h = d.add_heading(title_text, level=1)
            if 'WA Front Matter Title' in d.styles:
                h.style = d.styles['WA Front Matter Title']
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if idx > 0:
                # Attach pagination to the destination heading. A standalone
                # page-break paragraph after a dynamic TOC/LOF/LOT can spill
                # onto the following page when Word expands the field, then
                # break again and leave a numbered blank page behind.
                h.paragraph_format.page_break_before = True
            
            if item == 'table_of_contents':
                p_toc = d.add_paragraph()
                _add_toc_field(p_toc)
            elif item == 'list_of_figures':
                from wordassistant.operations.fields import append_field
                p_lof = d.add_paragraph()
                labels = v.captions.get('figure_labels') or [
                    {'display': 'Figure', 'sequence': 'Figure'}
                ]
                for label_index, caption_label in enumerate(labels):
                    if label_index:
                        p_lof.add_run('\n')
                    append_field(
                        p_lof,
                        f'TOC \\h \\z \\c "{caption_label["sequence"]}"',
                        'Update field in Word',
                    )
            elif item == 'list_of_tables':
                from wordassistant.operations.fields import append_field
                p_lot = d.add_paragraph()
                caption_label = v.captions.get('table_label') or {
                    'display': 'Table', 'sequence': 'Table'
                }
                append_field(
                    p_lot,
                    f'TOC \\h \\z \\c "{caption_label["sequence"]}"',
                    'Update field in Word',
                )
            elif item == 'work_schedule':
                table = d.add_table(rows=7, cols=4)
                table.style = 'Table Grid'
                for idx, header in enumerate(('STT', 'Nội dung công việc', 'Thời gian dự kiến', 'Xác nhận')):
                    table.rows[0].cells[idx].text = header
                    for run in table.rows[0].cells[idx].paragraphs[0].runs:
                        run.bold = True
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                set_word_font(run.font, v.typography.body_font)
                                run.font.size = Pt(v.typography.body_size_pt)
            elif item == 'abbreviations':
                table = d.add_table(rows=7, cols=3)
                table.style = 'Table Grid'
                for idx, header in enumerate(('STT', 'Ký hiệu', 'Giải thích')):
                    table.rows[0].cells[idx].text = header
                    for run in table.rows[0].cells[idx].paragraphs[0].runs:
                        run.bold = True
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                set_word_font(run.font, v.typography.body_font)
                                run.font.size = Pt(v.typography.body_size_pt)
            else:
                if item in {'supervisor_review', 'reviewer_review'}:
                    for _ in range(5):
                        d.add_paragraph()
                else:
                    d.add_paragraph("[Front matter content]")
                
    # 2. Main Body Section
    if v.section_blueprint:
        body_sec = d.add_section(WD_SECTION_START.NEW_PAGE)
        _apply_section_page(body_sec, v)
        _remove_page_border(body_sec)
        body_sec.header.is_linked_to_previous = False
        body_sec.footer.is_linked_to_previous = False
        
        usable_width_cm = max(5.0, v.page.width_cm - v.page.margins.left_cm - v.page.margins.right_cm)
        
        hf = v.header_footer
        if hf:
            hp = body_sec.header.paragraphs[0]
            header_l = effective_title if hf.get('header_left') == 'topic' else (eff_supervisor if hf.get('header_left') == 'supervisor' else '')
            header_r = f"GVHD: {eff_supervisor}" if (hf.get('header_right') == 'supervisor' and eff_supervisor) else (effective_title if hf.get('header_right') == 'topic' else '')
            if header_l or header_r:
                hp.text = f"{header_l}\t{header_r}".strip()
                hp.paragraph_format.tab_stops.add_tab_stop(Cm(usable_width_cm), WD_TAB_ALIGNMENT.RIGHT)
                
            fp = body_sec.footer.paragraphs[0]
            foot_l = eff_student if hf.get('footer_left') == 'student' else ''
            fp.text = f"{foot_l}\t" if foot_l else "\t"
            fp.paragraph_format.tab_stops.add_tab_stop(Cm(usable_width_cm), WD_TAB_ALIGNMENT.RIGHT)
            if hf.get('footer_right') == 'page_number':
                _add_page_number_field(fp)
                
            if hf.get('page_number_start') in ('introduction', '1', 1):
                from wordassistant.operations.page_numbers import _set_page_number_type
                _set_page_number_type(body_sec, 'decimal', 1)
                
        for sec in v.section_blueprint:
            sec_title = str(sec['title'])
            sec_level = max(1, min(9, int(sec.get('level', 1))))
            heading = d.add_heading(sec_title, level=sec_level)
            if sec.get('page_break_before', False):
                heading.paragraph_format.page_break_before = True
            d.add_paragraph("[Section content...]")
            
    elif not has_covers:
        p = d.add_paragraph()
        p.style = 'Title'
        p.add_run(effective_title)
        d.add_paragraph("[Document content]")

    for table in d.tables:
        normalize_table_geometry(d, table)

    if v.citations.get('require_reviewed_sources', False):
        reference_headings = [
            str(item['title']) for item in v.section_blueprint
            if 'tham khảo' in str(item.get('title', '')).casefold()
            or 'references' in str(item.get('title', '')).casefold()
        ]
        _set_citation_policy(d, {
            'style': v.citations.get('style', 'numeric'),
            'order': v.citations.get('bibliography_order', 'citation_order'),
            'heading': reference_headings[-1] if reference_headings else None,
            'require_reviewed': True,
        })
        
    if not dry_run:
        d.save(out)
    return out


def scaffold_document(
    output: str | Path,
    profile: DocumentProfile,
    variant_name: str,
    *,
    title: str | None = None,
    student: str | None = None,
    supervisor: str | None = None,
    cover_metadata: Mapping[str, str] | None = None,
    dry_run: bool = False,
) -> Path:
    """Build and validate a DOCX before atomically publishing it."""
    destination = Path(output).resolve()
    if destination.suffix.casefold() != ".docx":
        raise ValueError("Scaffold output must use the .docx extension")
    if dry_run:
        return _build_scaffold_document(
            destination,
            profile,
            variant_name,
            title=title,
            student=student,
            supervisor=supervisor,
            cover_metadata=cover_metadata,
            dry_run=True,
        )
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}.", suffix=".docx", dir=destination.parent
    )
    os.close(descriptor)
    staging = Path(temporary_name)
    staging.unlink(missing_ok=True)
    try:
        _build_scaffold_document(
            staging,
            profile,
            variant_name,
            title=title,
            student=student,
            supervisor=supervisor,
            cover_metadata=cover_metadata,
            dry_run=False,
        )
        Document(staging)
        staging.replace(destination)
        return destination
    finally:
        staging.unlink(missing_ok=True)

def _fill_template_metadata(doc_path: Path, title: str, student: str, supervisor: str):
    from wordassistant.document.edit import replace_text
    unresolved = []
    
    if title:
        count = 0
        for old_pat in ('<TÊN ĐỀ TÀI>', '[TÊN ĐỀ TÀI]', 'TÊN ĐỀ TÀI', '<ĐỀ TÀI>', '[ĐỀ TÀI]'):
            res = replace_text(doc_path, old=old_pat, new=title, output_path=doc_path)
            count += res.replacements
        if count == 0:
            unresolved.append('title')
            
    if student:
        count = 0
        for old_pat in ('<TÊN SINH VIÊN>', '<SINH VIÊN>', '[Họ và tên sinh viên]', '[SINH VIÊN]', 'Tên sinh viên'):
            res = replace_text(doc_path, old=old_pat, new=student, output_path=doc_path)
            count += res.replacements
        if count == 0:
            unresolved.append('student')
            
    if supervisor:
        count = 0
        for old_pat in ('Học vị. Tên giảng viên hướng dẫn', 'Tên giảng viên hướng dẫn', '<CÁN BỘ HƯỚNG DẪN>', '<GIẢNG VIÊN HƯỚNG DẪN>', '<GIẢNG VIÊN>', '[Họ và tên GVHD]', '[CÁN BỘ HƯỚNG DẪN]'):
            res = replace_text(doc_path, old=old_pat, new=supervisor, output_path=doc_path)
            count += res.replacements
        if count == 0:
            unresolved.append('supervisor')
            
    if unresolved:
        raise ValueError(f"Could not resolve template placeholder(s) for explicitly provided metadata: {', '.join(unresolved)}")
