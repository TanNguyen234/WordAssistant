from __future__ import annotations
from pathlib import Path
import pytest
from docx import Document
from wordassistant.document.scaffold import scaffold_document
from wordassistant.document.validate import validate_document
from wordassistant.profiles.loader import resolve_profile


def test_verify_summary_writes_full_report_and_reports_hash(tmp_path, monkeypatch, capsys):
    import hashlib
    import json
    from wordassistant import cli

    doc_path = tmp_path / "report.docx"
    Document().save(doc_path)
    report_path = tmp_path / "full.json"
    monkeypatch.setattr(cli, "load_project_config", lambda _: {})
    monkeypatch.setattr(cli, "resolve_profile", lambda *args, **kwargs: type("Profile", (), {"profile_id": "generic"})())
    monkeypatch.setattr(cli, "inspect_document", lambda _: {"paragraphs": 1, "tables": 0, "sections": 1, "images": 0, "fields": 0})
    monkeypatch.setattr(cli, "verify_document", lambda *args, **kwargs: {
        "overall_status": "PASS", "delivery_status": "INCOMPLETE", "delivery_gates": {},
        "format_status": "PASS", "issues": [], "render_status": "NOT_REQUESTED",
        "native_field_results": [{"result": "private-detail"}],
    })
    result = cli.run(["verify", str(doc_path), "--summary", "--report-output", str(report_path), "--json"])
    assert result == 0
    compact = json.loads(capsys.readouterr().out)
    assert compact["report_path"] == str(report_path.resolve())
    assert compact["report_sha256"] == hashlib.sha256(report_path.read_bytes()).hexdigest()
    assert "native_field_results" not in compact
    assert "native_field_results" in json.loads(report_path.read_text(encoding="utf-8"))


@pytest.mark.parametrize("collision", ["source", "reuse", "pdf", "image"])
def test_verify_report_output_rejects_protected_path_collisions(
    tmp_path, monkeypatch, capsys, collision
):
    import json
    from wordassistant import cli

    source = tmp_path / "report.docx"
    Document().save(source)
    reuse_report = tmp_path / "reuse.json"
    reuse_report.write_text("{}", encoding="utf-8")
    pdf = tmp_path / "rendered.pdf"
    pdf.write_bytes(b"%PDF")
    image = tmp_path / "page-1.png"
    image.write_bytes(b"png")
    output = {
        "source": source,
        "reuse": reuse_report,
        "pdf": pdf,
        "image": image,
    }[collision]
    before = output.read_bytes()
    monkeypatch.setattr(cli, "load_project_config", lambda _: {})
    monkeypatch.setattr(cli, "resolve_profile", lambda *args, **kwargs: type("Profile", (), {"profile_id": "generic"})())
    monkeypatch.setattr(cli, "inspect_document", lambda _: {"paragraphs": 1, "tables": 0, "sections": 1, "images": 0, "fields": 0})
    monkeypatch.setattr(cli, "verify_document", lambda *args, **kwargs: {
        "overall_status": "PASS", "delivery_status": "INCOMPLETE", "delivery_gates": {},
        "format_status": "PASS", "issues": [], "render_status": "PASS",
        "render": {"pdf_path": str(pdf), "page_images": [str(image)]},
    })
    result = cli.run([
        "verify", str(source), "--report-output", str(output),
        "--reuse-render-report", str(reuse_report), "--json",
    ])
    assert result == 1
    assert "must not overwrite" in json.loads(capsys.readouterr().out)["error"]["message"]
    assert output.read_bytes() == before

def test_verify_fails_when_required_front_matter_missing(tmp_path: Path):
    doc_path = tmp_path / "report.docx"
    prof = resolve_profile("hcmunre_dacn_2026_2027")
    scaffold_document(doc_path, prof, "report", title="TEST TITLE", student="STUDENT A", supervisor="SUPERVISOR B")
    
    # Add enough references so reference check passes
    d = Document(doc_path)
    for i in range(1, 12):
        d.add_paragraph(f"[{i}] Author {i}, Title {i}, 2024.")
        
    # Delete the heading "LỜI CẢM ƠN"
    for p in d.paragraphs:
        if "cảm ơn" in p.text.lower() or "cam on" in p.text.lower():
            p.text = "THAY THE"
            p.style = d.styles["Normal"]
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues]
    assert "missing_front_matter" in codes

def test_verify_fails_when_header_topic_missing(tmp_path: Path):
    doc_path = tmp_path / "report.docx"
    prof = resolve_profile("hcmunre_dacn_2026_2027")
    scaffold_document(doc_path, prof, "report", title="TEST TITLE", student="STUDENT A", supervisor="SUPERVISOR B")
    
    d = Document(doc_path)
    for i in range(1, 12):
        d.add_paragraph(f"[{i}] Author {i}, Title {i}, 2024.")
    # Clear header text in all sections
    for s in d.sections:
        for p in s.header.paragraphs:
            p.text = ""
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues]
    assert any("header" in c for c in codes)
