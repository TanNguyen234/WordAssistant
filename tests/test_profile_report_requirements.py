from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from wordassistant.document.validate import validate_document
from wordassistant.operations.lists import apply_heading_numbering
from wordassistant.profiles.loader import validate_profile_dict


def _profile(*, typography=None, front_matter=None, cover=None, section_blueprint=None):
    return validate_profile_dict(
        {
            "id": "course-report",
            "name": "Course report",
            "version": 1,
            "variants": {
                "default": {
                    "page": {
                        "size": "A4",
                        "width_cm": 21,
                        "height_cm": 29.7,
                        "margins_cm": {"top": 2.5, "bottom": 2.5, "left": 2.5, "right": 2.5},
                        "one_sided": True,
                    },
                    "typography": typography
                    or {"body_font": "Times New Roman", "body_size_pt": 13},
                    "front_matter": front_matter or [],
                    "cover": cover or {},
                    "section_blueprint": section_blueprint or [],
                }
            },
        }
    )


def test_profile_checks_heading_and_caption_color_and_size(tmp_path: Path):
    profile = _profile(
        typography={
            "body_font": "Times New Roman",
            "body_size_pt": 13,
            "styles": {
                "heading2": {"font_color": "000000", "size_pt": 14},
                "Caption": {"font_color": "000000", "size_pt": 12},
            },
        }
    )
    path = tmp_path / "blue-heading.docx"
    document = Document()
    heading = document.add_heading("Nested heading", level=2)
    heading.runs[0].font.color.rgb = RGBColor(0, 0, 255)
    heading.runs[0].font.size = Pt(12)
    caption = document.add_paragraph("Hình 1.1: Giao diện", style="Caption")
    caption.runs[0].font.color.rgb = RGBColor(0, 0, 255)
    caption.runs[0].font.size = Pt(10)
    document.save(path)

    result = validate_document(path, profile, "default", verify_mode=True)

    codes = {issue.code for issue in result.issues}
    assert "style_color_mismatch" in codes
    assert "style_size_mismatch" in codes


def test_profile_checks_page_break_and_true_multilevel_heading_numbering(tmp_path: Path):
    profile = _profile(
        section_blueprint=[
            {"title": "Main", "level": 1, "numbered": True, "page_break_before": True},
            {"title": "Nested", "level": 2, "numbered": True},
        ]
    )
    source = tmp_path / "unnumbered.docx"
    document = Document()
    main = document.add_heading("Main", level=1)
    nested = document.add_heading("Nested", level=2)
    nested.paragraph_format.page_break_before = True
    document.save(source)

    initial = validate_document(source, profile, "default", verify_mode=True)
    initial_codes = {issue.code for issue in initial.issues}
    assert "blueprint_page_break_missing" in initial_codes
    assert "blueprint_heading_numbering_missing" in initial_codes

    numbered = tmp_path / "numbered.docx"
    apply_heading_numbering(source, output_path=numbered, levels=2)
    result = validate_document(numbered, profile, "default", verify_mode=True)

    assert "blueprint_heading_numbering_missing" not in {
        issue.code for issue in result.issues
    }


def _add_page_border(document, space_points: int):
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    borders.set(qn("w:display"), "firstPage")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:space"), str(space_points))
        borders.append(node)
    document.sections[0]._sectPr.append(borders)


def test_profile_rejects_cover_frame_too_close_to_page_edge(tmp_path: Path):
    border = {
        "display": "firstPage",
        "offset_from": "page",
        "edges": {
            edge: {"style": "single", "size_eighth_points": 8, "space_points": 36}
            for edge in ("top", "left", "bottom", "right")
        },
    }
    profile = _profile(
        front_matter=["cover"],
        cover={"document_type": "TIỂU LUẬN", "border": border},
    )
    path = tmp_path / "tight-frame.docx"
    document = Document()
    _add_page_border(document, 6)
    document.save(path)

    result = validate_document(path, profile, "default", verify_mode=True)

    assert "cover_border_inset" in {issue.code for issue in result.issues}


def test_profile_requires_inner_cover_on_a_separate_page(tmp_path: Path):
    profile = _profile(
        front_matter=["cover", "inner_cover"],
        cover={
            "document_type": "TIỂU LUẬN",
            "require_distinct_inner_cover": True,
        },
    )
    path = tmp_path / "same-page-covers.docx"
    document = Document()
    document.add_paragraph("TRƯỜNG ĐẠI HỌC")
    document.add_paragraph("TIỂU LUẬN")
    document.add_paragraph("TRƯỜNG ĐẠI HỌC")
    document.add_paragraph("TIỂU LUẬN")
    document.save(path)

    result = validate_document(path, profile, "default", verify_mode=True)

    assert "inner_cover_not_distinct" in {issue.code for issue in result.issues}


def test_profile_checks_inner_cover_break_at_cover_block_start(tmp_path: Path):
    profile = _profile(
        front_matter=["cover", "inner_cover"],
        cover={
            "institution_lines": ["TRƯỜNG ĐẠI HỌC"],
            "document_type": "TIỂU LUẬN",
            "require_distinct_inner_cover": True,
        },
    )
    path = tmp_path / "separate-cover-blocks.docx"
    document = Document()
    document.add_paragraph("TRƯỜNG ĐẠI HỌC")
    document.add_paragraph("TIỂU LUẬN")
    inner_cover_start = document.add_paragraph("TRƯỜNG ĐẠI HỌC")
    inner_cover_start.paragraph_format.page_break_before = True
    document.add_paragraph("TIỂU LUẬN")
    document.save(path)

    result = validate_document(path, profile, "default", verify_mode=True)

    assert "inner_cover_not_distinct" not in {issue.code for issue in result.issues}
