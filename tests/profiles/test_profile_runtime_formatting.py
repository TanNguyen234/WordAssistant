import json
from zipfile import ZipFile

from docx import Document
from docx.shared import RGBColor

from wordassistant.document.scaffold import scaffold_document
from wordassistant.profiles.loader import load_profile


def test_profile_spacing_properties_and_heading_nine_are_written(tmp_path):
    profile_path = tmp_path / "profile.json"
    profile_path.write_text(
        json.dumps(
            {
                "id": "runtime-format",
                "name": "Runtime format",
                "version": 1,
                "variants": {
                    "default": {
                        "page": {
                            "size": "A4",
                            "width_cm": 21,
                            "height_cm": 29.7,
                            "margins_cm": {"top": 2, "bottom": 2, "left": 2, "right": 2},
                            "one_sided": True,
                        },
                        "typography": {
                            "body_font": "Times New Roman",
                            "body_size_pt": 12,
                            "line_spacing": 1.25,
                            "paragraph_spacing_before_pt": 3,
                            "paragraph_spacing_after_pt": 6,
                            "styles": {
                                "heading9": {
                                    "size_pt": 9,
                                    "line_spacing": 1.1,
                                    "space_before_pt": 8,
                                    "space_after_pt": 4,
                                }
                            },
                        },
                        "section_blueprint": [{"title": "Deep", "level": 9}],
                    }
                },
            }
        ),
        encoding="utf-8",
    )
    profile = load_profile(profile_path)
    output = tmp_path / "output.docx"

    scaffold_document(output, profile, "default", title="Runtime")

    document = Document(output)
    normal = document.styles["Normal"].paragraph_format
    heading = document.styles["Heading 9"].paragraph_format
    assert normal.space_before.pt == 3
    assert normal.space_after.pt == 6
    assert normal.line_spacing == 1.25
    assert heading.space_before.pt == 8
    assert heading.space_after.pt == 4
    assert heading.line_spacing == 1.1
    assert any(p.style.name == "Heading 9" and p.text == "Deep" for p in document.paragraphs)
    with ZipFile(output) as package:
        styles_xml = package.read("word/styles.xml").decode("utf-8")
    normal_xml = styles_xml.split('<w:style w:type="paragraph" w:default="1" w:styleId="Normal">', 1)[1].split("</w:style>", 1)[0]
    assert 'w:ascii="Times New Roman"' in normal_xml
    assert 'w:hAnsi="Times New Roman"' in normal_xml
    assert 'w:eastAsia="Times New Roman"' in normal_xml
    assert 'w:cs="Times New Roman"' in normal_xml
    assert "w:asciiTheme" not in normal_xml


def test_profile_heading_color_and_indentation_are_written(tmp_path):
    profile_path = tmp_path / "heading-style-profile.json"
    profile_path.write_text(
        json.dumps(
            {
                "id": "heading-style-format",
                "name": "Heading style format",
                "version": 1,
                "variants": {
                    "default": {
                        "page": {
                            "size": "A4",
                            "width_cm": 21,
                            "height_cm": 29.7,
                            "margins_cm": {"top": 2, "bottom": 2, "left": 2, "right": 2},
                            "one_sided": True,
                        },
                        "typography": {
                            "body_font": "Times New Roman",
                            "body_size_pt": 13,
                            "styles": {
                                "heading2": {
                                    "font_color": "000000",
                                    "size_pt": 14,
                                    "left_indent_pt": 18,
                                    "first_line_indent_pt": -9,
                                },
                                "Caption": {"font_color": "000000", "size_pt": 12},
                            },
                        },
                        "section_blueprint": [{"title": "Nested", "level": 2}],
                    }
                },
            }
        ),
        encoding="utf-8",
    )
    profile = load_profile(profile_path)
    output = tmp_path / "heading-styles.docx"

    scaffold_document(output, profile, "default", title="Report")

    document = Document(output)
    heading = document.styles["Heading 2"]
    assert heading.font.color.rgb == RGBColor(0, 0, 0)
    assert heading.font.size.pt == 14
    assert heading.paragraph_format.left_indent.pt == 18
    assert heading.paragraph_format.first_line_indent.pt == -9
    caption = document.styles["Caption"]
    assert caption.font.color.rgb == RGBColor(0, 0, 0)
    assert caption.font.size.pt == 12
