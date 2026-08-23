from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest
from wordassistant.document.edit import replace_text, ReplaceResult

def test_replace_blocked_in_complex_container_with_field(tmp_path):
    doc_path = tmp_path / 'field_doc.docx'
    d = Document()
    p = d.add_paragraph("Paragraph with field ")
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p._p.append(fld)
    p.add_run(" and target keyword.")
    d.save(doc_path)
    
    out_path = tmp_path / 'field_doc_out.docx'
    res = replace_text(doc_path, old="keyword", new="replacement", output_path=out_path)
    
    # Must NOT mutate or write file when complex structure is detected
    assert res.status == 'UNSAFE_EDIT_BLOCKED'
    assert res.advanced_backend_required is True
    assert res.replacements == 0
    assert not out_path.exists()
    assert 'field' in res.complex_structures_detected

def test_replace_blocked_in_content_control_or_revision(tmp_path):
    doc_path = tmp_path / 'sdt_doc.docx'
    d = Document()
    p = d.add_paragraph("Paragraph with ")
    sdt = OxmlElement('w:sdt')
    p._p.append(sdt)
    p.add_run(" target keyword.")
    d.save(doc_path)
    
    out_path = tmp_path / 'sdt_doc_out.docx'
    res = replace_text(doc_path, old="keyword", new="replacement", output_path=out_path)
    
    assert res.status == 'UNSAFE_EDIT_BLOCKED'
    assert res.advanced_backend_required is True
    assert not out_path.exists()
    assert 'content_control' in res.complex_structures_detected

def test_replace_works_in_headers_footers_tables_and_body(tmp_path):
    doc_path = tmp_path / 'all_parts.docx'
    d = Document()
    # Body paragraph
    d.add_paragraph("Body TARGET_1")
    # Table
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).paragraphs[0].text = "Cell TARGET_2"
    # Header
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Header TARGET_3"
    # Footer
    sec.footer.paragraphs[0].text = "Footer TARGET_4"
    d.save(doc_path)
    
    out_path = tmp_path / 'all_parts_out.docx'
    res = replace_text(doc_path, old="TARGET", new="REPLACED", output_path=out_path)
    
    assert res.status == 'SUCCESS'
    assert res.replacements == 4
    assert out_path.exists()
    
    d2 = Document(out_path)
    assert "Body REPLACED_1" in d2.paragraphs[0].text
    assert "Cell REPLACED_2" in d2.tables[0].cell(0, 0).paragraphs[0].text
    assert "Header REPLACED_3" in d2.sections[0].header.paragraphs[0].text
    assert "Footer REPLACED_4" in d2.sections[0].footer.paragraphs[0].text
