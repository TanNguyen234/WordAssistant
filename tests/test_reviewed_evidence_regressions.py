from hashlib import sha256
import json

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

from wordassistant.document.integrity import audit_integrity, build_provenance_manifest
from wordassistant.evidence.registry import EvidenceError, EvidenceRegistry, validate_record


def _reviewed_record():
    return {
        "id": "reviewed-1", "title": "Checked source", "authors": ["Researcher"],
        "year": 2026, "url": "https://example.org/source", "verified_by": "human",
        "synthetic": False, "review_status": "reviewed", "reviewed_by": "Researcher",
        "reviewed_at": "2026-10-12T00:00:00+00:00", "support_location": "page 1",
        "support_excerpt": "Checked passage", "claim": "Supported claim",
        "source_file": "source.pdf", "source_sha256": sha256(b"source bytes").hexdigest(),
        "source_license": "cc-by", "source_url": "https://example.org/source.pdf",
        "retrieved_at": "2026-10-12T00:00:00+00:00", "language": "en",
        "author_sort_key": "Researcher",
    }


@pytest.mark.parametrize("field", [
    "reviewed_by", "reviewed_at", "support_location", "support_excerpt", "claim",
    "source_license", "source_url", "retrieved_at", "language", "author_sort_key",
    "source_file", "source_sha256",
])
@pytest.mark.parametrize("value", ["", "   ", None, 12])
def test_reviewed_fields_require_nonempty_strings(field, value):
    record = _reviewed_record()
    record[field] = value
    with pytest.raises(EvidenceError):
        validate_record(record)


@pytest.mark.parametrize("url", ["doi:10.1234/example", "file:///source.pdf", "https://", "https://bad host/source.pdf"])
def test_reviewed_source_url_requires_http_hostname(url):
    record = _reviewed_record()
    record["source_url"] = url
    with pytest.raises(EvidenceError, match="source_url"):
        validate_record(record)


@pytest.mark.parametrize("status", [None, "candidate"])
@pytest.mark.parametrize("synthetic", [False, True])
def test_unreviewed_optional_fields_remain_compatible(status, synthetic):
    record = {key: value for key, value in _reviewed_record().items() if key in {
        "id", "title", "authors", "year", "url", "verified_by", "synthetic",
    }}
    record.update(synthetic=synthetic, claim=None, support_excerpt="", source_url=None)
    if status is not None:
        record["review_status"] = status
    validate_record(record)


def test_registry_get_does_not_require_pdf_bytes(tmp_path):
    registry = EvidenceRegistry(tmp_path / "evidence.json")
    registry.add(_reviewed_record())
    assert registry.get("reviewed-1")["review_status"] == "reviewed"


def test_supplied_audit_does_not_require_unlinked_reviewed_pdf(tmp_path):
    path = tmp_path / "unlinked.docx"
    document = Document()
    document.add_paragraph("AI-use disclosure: no generative AI was used.")
    document.save(path)
    registry_path = tmp_path / "evidence.json"
    EvidenceRegistry(registry_path).add(_reviewed_record())
    audit = audit_integrity(path, registry_path)
    manifest = build_provenance_manifest(path, registry_path, audit_report=audit)
    assert manifest["audit_summary"]["status"] == audit["status"]


@pytest.mark.parametrize("heading_style", [False, True])
@pytest.mark.parametrize("numbered", [False, True])
def test_bibliography_language_headings_do_not_count_as_sources(tmp_path, heading_style, numbered):
    document = Document()
    document.add_heading("Tài liệu tham khảo", level=1)
    for language, numbers in (("Tiếng Việt:", range(1, 5)), ("Tiếng Anh:", range(5, 9))):
        if heading_style:
            document.add_heading(language, level=2)
        else:
            document.add_paragraph(language)
        for number in numbers:
            prefix = f"[{number}] " if numbered else ""
            document.add_paragraph(f"{prefix}Researcher, Source {number} (2026).")
    path = tmp_path / "references.docx"
    document.save(path)
    report = audit_integrity(path, profile={"citations": {"minimum_references": 10}})
    gate = report["gates"]["source_evidence"]
    assert gate["document_reference_count"] == 8
    assert gate["status"] == "FAIL"
    assert any(f["kind"] == "minimum_references_not_met" for f in gate["findings"])


@pytest.mark.parametrize("mutation", ["changed", "missing", "unchanged"])
def test_supplied_audit_rechecks_linked_reviewed_source_bytes(tmp_path, mutation, monkeypatch):
    document = Document()
    paragraph = document.add_paragraph("Supported claim ")
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "wordassistant:citation:reviewed-1")
    properties.append(tag)
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "[1]"
    run.append(text)
    content.append(run)
    sdt.extend([properties, content])
    paragraph._p.append(sdt)
    document.add_paragraph("AI-use disclosure: no generative AI was used.")
    document.add_heading("References", level=1)
    document.add_paragraph("[1] Checked source.")
    path = tmp_path / "reviewed.docx"
    document.save(path)
    registry_path = tmp_path / "evidence.json"
    EvidenceRegistry(registry_path).add(_reviewed_record())
    source = tmp_path / "sources" / "source.pdf"
    source.parent.mkdir()
    source.write_bytes(b"source bytes")
    audit = audit_integrity(path, registry_path)
    assert audit["gates"]["source_evidence"]["status"] == "PASS"
    if mutation == "unchanged":
        assert build_provenance_manifest(path, registry_path, audit_report=audit)["audit_summary"]["status"] == audit["status"]
    else:
        if mutation == "missing":
            source.unlink()
        else:
            source.write_bytes(b"altered source bytes")
        with pytest.raises(EvidenceError, match="missing|hash changed"):
            build_provenance_manifest(path, registry_path, audit_report=audit)

    if mutation == "unchanged":
        # Read failures happen after existence checks, including deletion races.
        original_read = type(source).read_bytes
        for error in (PermissionError("denied"), OSError("unreadable"), FileNotFoundError("deleted during read")):
            def read_bytes(file, failure=error):
                if file == source:
                    raise failure
                return original_read(file)

            monkeypatch.setattr(type(source), "read_bytes", read_bytes)
            report = audit_integrity(path, registry_path)
            assert report["gates"]["source_evidence"]["status"] == "FAIL"
            with pytest.raises(EvidenceError):
                build_provenance_manifest(path, registry_path, audit_report=audit)


@pytest.mark.parametrize("status", [[], {}])
def test_invalid_review_status_is_a_registry_failure(tmp_path, status):
    record = _reviewed_record()
    record["review_status"] = status
    with pytest.raises(EvidenceError, match="review_status"):
        validate_record(record)
    registry_path = tmp_path / "evidence.json"
    registry_path.write_text(json.dumps({record["id"]: record}))
    path = tmp_path / "invalid-record.docx"
    Document().save(path)
    report = audit_integrity(path, registry_path)
    assert report["gates"]["source_evidence"]["status"] == "FAIL"
    assert report["gates"]["source_evidence"]["findings"][0]["kind"] == "invalid_evidence_registry"


@pytest.mark.parametrize("outline_source", ["direct", "style", "inherited"])
def test_semantic_appendix_heading_ends_bibliography(tmp_path, outline_source):
    document = Document()
    document.add_heading("References", level=1)
    document.add_paragraph("[1] Checked source.")
    appendix = document.add_paragraph("Appendix")
    if outline_source == "direct":
        properties = appendix._p.get_or_add_pPr()
    else:
        style = document.styles.add_style("Custom appendix", WD_STYLE_TYPE.PARAGRAPH)
        appendix.style = style
        if outline_source == "inherited":
            base = document.styles.add_style("Appendix base", WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = base
            style = base
        properties = style.element.get_or_add_pPr()
    level = OxmlElement("w:outlineLvl")
    level.set(qn("w:val"), "0")
    properties.append(level)
    document.add_paragraph("An appendix claim refers to [2].")
    path = tmp_path / "appendix.docx"
    document.save(path)
    report = audit_integrity(path)
    assert report["gates"]["source_evidence"]["document_reference_count"] == 1
    consistency = report["gates"]["numeric_citation_bibliography_consistency"]
    assert consistency["citation_numbers"] == [2]
    assert consistency["status"] == "FAIL"
