from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
import pytest
from wordassistant.profiles.loader import load_profile
from wordassistant.document.validate import validate_document
from wordassistant.document.scaffold import scaffold_document

ROOT = Path(__file__).resolve().parents[1]

def test_validator_fails_on_any_body_paragraph_font_violation(tmp_path):
    profile = load_profile(ROOT / 'profiles/hcmunre_dacn_2026_2027.json')
    out = tmp_path / 'bad_run.docx'
    scaffold_document(out, profile, 'report', title='TEST PROJECT')
    
    d = Document(out)
    # Add a paragraph at the very end with wrong font in second run
    p = d.add_paragraph()
    p.add_run("Normal prefix. ")
    bad_r = p.add_run("Bad Arial Run")
    bad_r.font.name = "Arial"
    d.save(out)
    
    res = validate_document(out, profile, 'report')
    assert not res.valid
    assert any('font' in i.code.lower() or 'typography_font' in i.code for i in res.issues)

def test_validator_detects_out_of_order_blueprint_sections(tmp_path):
    profile = load_profile(ROOT / 'profiles/hcmunre_dacn_2026_2027.json')
    out = tmp_path / 'reversed_blueprint.docx'
    d = Document()
    s = d.sections[0]
    s.page_width = Cm(21.0); s.page_height = Cm(29.7)
    s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0); s.left_margin = Cm(3.0); s.right_margin = Cm(2.0)
    
    # Put Chapter 2 BEFORE Chapter 1
    d.add_heading("CHƯƠNG 2: PHƯƠNG PHÁP THỰC HIỆN", level=1)
    d.add_paragraph("Content 2")
    d.add_heading("CHƯƠNG 1: TỔNG QUAN TÀI LIỆU", level=1)
    d.add_paragraph("Content 1")
    d.add_heading("MỞ ĐẦU", level=1)
    d.add_paragraph("Content Intro")
    d.save(out)
    
    res = validate_document(out, profile, 'report')
    assert not res.valid
    assert any('order' in i.code.lower() or 'blueprint_order' in i.code for i in res.issues)

def test_validator_verify_mode_flags_zero_references_as_error(tmp_path):
    profile = load_profile(ROOT / 'profiles/hcmunre_dacn_2026_2027.json')
    out = tmp_path / 'no_refs.docx'
    scaffold_document(out, profile, 'report', title='TEST PROJECT')
    
    # In verify mode (final delivery), zero references must produce an ERROR
    res_verify = validate_document(out, profile, 'report', verify_mode=True)
    assert not res_verify.valid
    assert any('reference' in i.code.lower() for i in res_verify.issues)
    
    # In regular draft validate mode, it produces a WARNING but valid remains True for scaffold
    res_draft = validate_document(out, profile, 'report', verify_mode=False)
    assert any('reference' in i.code.lower() for i in res_draft.issues)
