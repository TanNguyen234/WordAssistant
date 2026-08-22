from docx import Document
from wordassistant.document.edit import replace_text

def test_replace_text_in_body_and_tables(tmp_path):
    p=tmp_path/'x.docx'; d=Document(); d.add_paragraph('Alpha old value'); t=d.add_table(1,1); t.cell(0,0).text='old value in table'; d.save(p)
    out=tmp_path/'y.docx'; result=replace_text(p,'old value','new value',out)
    assert result.replacements==2
    d2=Document(out); text='\n'.join([x.text for x in d2.paragraphs]+[c.text for tb in d2.tables for row in tb.rows for c in row.cells])
    assert 'new value' in text and 'old value' not in text
