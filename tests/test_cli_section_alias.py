from __future__ import annotations
from pathlib import Path
from docx import Document
from wordassistant.cli import run

def test_context_cli_accepts_section_alias(tmp_path: Path):
    doc_path = tmp_path / "test_sec.docx"
    d = Document()
    d.add_heading("Section One Title", 1)
    d.add_paragraph("Content of section 1")
    d.add_heading("Section Two Title", 1)
    d.add_paragraph("Content of section 2")
    d.save(doc_path)
    
    code1 = run(["context", str(doc_path), "--query", "content", "--active-section-id", "section-one-title", "--json"])
    assert code1 == 0
    
    code2 = run(["context", str(doc_path), "--query", "content", "--section", "section-two-title", "--json"])
    assert code2 == 0
