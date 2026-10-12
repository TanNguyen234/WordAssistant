from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wordassistant.document.scaffold import scaffold_document
from wordassistant.document.structure import detect_heading, index_docx
from wordassistant.operations.text import edit_text
from wordassistant.profiles.loader import load_profile
from wordassistant.selectors import SelectorNotFound, resolve_paragraphs, resolve_single_paragraph


ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def front_report(tmp_path):
    profile = load_profile(ROOT / 'profiles/hcm-university-dacn-2026-2027.json')
    path = tmp_path / 'front.docx'
    scaffold_document(path, profile, 'report')
    return path, profile.variant('report')


def test_scaffold_front_titles_have_distinct_level_one_sections(front_report):
    path, variant = front_report
    expected = [variant.front_matter_titles[item] for item in variant.front_matter
                if item not in {'cover', 'inner_cover', 'proposal_cover'}]
    sections = index_docx(path)
    actual = [section.title for section in sections if section.title in expected]
    assert actual == expected
    assert all(section.level == 1 for section in sections if section.title in expected)
    document = Document(path)
    for paragraph in document.paragraphs:
        if paragraph.text in expected:
            assert paragraph.style.name == 'WA Front Matter Title'
            assert paragraph.style.font.bold is True
            assert paragraph.style.font.size.pt == 16


def test_front_heading_selector_replacement_stops_at_next_peer(front_report, tmp_path):
    path, _variant = front_report
    document = Document(path)
    titles = [paragraph for paragraph in document.paragraphs
              if paragraph.style.name == 'WA Front Matter Title']
    first, peer = titles[:2]
    first_index = next(
        index for index, paragraph in enumerate(document.paragraphs) if paragraph._p is first._p
    )
    document.paragraphs[first_index + 1].text = 'SCOPED TARGET'
    peer_index = next(index for index, paragraph in enumerate(document.paragraphs)
                      if paragraph._p is peer._p)
    document.paragraphs[peer_index + 1].text = 'SCOPED TARGET'
    document.save(path)
    output = tmp_path / 'edited.docx'
    result = edit_text(path, output_path=output, action='replace',
                       selector={'type': 'heading', 'text': first.text},
                       pattern='SCOPED TARGET', replacement='CHANGED')
    assert result.data['replacements'] == 1
    reopened = Document(output)
    assert reopened.paragraphs[first_index + 1].text == 'CHANGED'
    assert reopened.paragraphs[peer_index + 1].text == 'SCOPED TARGET'
    assert Document(path).paragraphs[first_index + 1].text == 'SCOPED TARGET'


@pytest.mark.parametrize('origin', ['paragraph', 'style', 'inherited'])
def test_custom_outline_styles_are_semantic_heading_boundaries(origin):
    document = Document()
    style = document.styles.add_style('Custom section', WD_STYLE_TYPE.PARAGRAPH)
    if origin == 'inherited':
        style.base_style = document.styles['Heading 1']
    elif origin == 'style':
        outline = OxmlElement('w:outlineLvl')
        outline.set(qn('w:val'), '0')
        style.element.get_or_add_pPr().append(outline)
    heading = document.add_paragraph('Custom heading', style=style)
    if origin == 'paragraph':
        outline = OxmlElement('w:outlineLvl')
        outline.set(qn('w:val'), '0')
        heading._p.get_or_add_pPr().append(outline)
    document.add_paragraph('Section content')
    document.add_heading('Nested heading', level=2)
    document.add_paragraph('Nested content')
    peer = document.add_paragraph('Peer heading', style=style)
    if origin == 'paragraph':
        outline = OxmlElement('w:outlineLvl')
        outline.set(qn('w:val'), '0')
        peer._p.get_or_add_pPr().append(outline)
    assert detect_heading(heading) == 1
    selected = resolve_paragraphs(document, {'type': 'heading', 'text': 'Custom heading'})
    assert [paragraph.text for paragraph in selected] == [
        'Custom heading', 'Section content', 'Nested heading', 'Nested content'
    ]
    assert resolve_single_paragraph(document, {'type': 'heading', 'text': 'Custom heading'})._p is heading._p


def test_body_outline_override_and_bold_title_are_not_semantic_headings():
    document = Document()
    paragraph = document.add_heading('Body override', level=1)
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), '9')
    paragraph._p.get_or_add_pPr().append(outline)
    document.add_paragraph('Decorative title').runs[0].bold = True
    for text in ('Body override', 'Decorative title'):
        with pytest.raises(SelectorNotFound):
            resolve_paragraphs(document, {'type': 'heading', 'text': text})
    assert detect_heading(paragraph) is None


def test_style_body_outline_override_suppresses_name_heuristics():
    document = Document()
    style = document.styles.add_style('Body override', WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles['Heading 1']
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), '9')
    style.element.get_or_add_pPr().append(outline)
    paragraph = document.add_paragraph('CHAPTER 1', style=style)
    assert detect_heading(paragraph) is None
    with pytest.raises(SelectorNotFound):
        resolve_paragraphs(document, {'type': 'heading', 'text': paragraph.text})


def test_legacy_body_list_style_keeps_bounded_heading_heuristics():
    document = Document()
    style = document.styles.add_style('Legacy Bullet', WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles['Normal']
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), '9')
    style.element.get_or_add_pPr().append(outline)
    paragraph = document.add_paragraph('17.3. Cách trích dẫn nguồn tài liệu:', style=style)
    paragraph.runs[0].bold = True
    assert detect_heading(paragraph) == 3


@pytest.mark.parametrize('origin', ['paragraph', 'style', 'inherited'])
def test_heading_resolution_supports_docx_without_outline_descriptor(tmp_path, monkeypatch, origin):
    from docx.oxml.text.parfmt import CT_PPr
    from wordassistant.cli import run
    # Supported older python-docx versions expose the XML but no descriptor.
    monkeypatch.delattr(CT_PPr, 'outlineLvl', raising=False)
    document = Document()
    style = document.styles.add_style('Portable heading', WD_STYLE_TYPE.PARAGRAPH)
    if origin == 'inherited':
        style.base_style = document.styles['Heading 1']
    paragraph = document.add_paragraph('Portable heading', style=style)
    if origin != 'inherited':
        properties = paragraph._p.get_or_add_pPr() if origin == 'paragraph' else style.element.get_or_add_pPr()
        outline = OxmlElement('w:outlineLvl')
        outline.set(qn('w:val'), '0')
        properties.append(outline)
    document.add_paragraph('Bounded content')
    path = tmp_path / 'portable.docx'
    document.save(path)
    assert detect_heading(paragraph) == 1
    assert run(['inspect', str(path), '--summary', '--json']) == 0
