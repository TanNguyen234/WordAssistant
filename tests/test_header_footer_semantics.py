from __future__ import annotations
from pathlib import Path
from docx import Document
from wordassistant.document.validate import validate_document
from wordassistant.profiles.loader import resolve_profile
from wordassistant.document.scaffold import scaffold_document

def _create_clean_report(path: Path):
    prof = resolve_profile("hcmunre_dacn_2026_2027")
    scaffold_document(
        path,
        prof,
        "report",
        title="XAY DUNG HE THONG RAG CHO VAN BAN PHAP LUAT",
        student="NGUYEN VAN A - 2351010001",
        supervisor="TS. NGUYEN VAN B"
    )
    d = Document(path)
    for p in d.paragraphs:
        if "[Front matter content]" in p.text:
            p.text = "Lời cảm ơn chi tiết."
        elif "[Section content...]" in p.text:
            p.text = "Nội dung chương hoàn chỉnh."
        elif "[List of figures]" in p.text:
            p.text = "Hình 1.1: Sơ đồ."
        elif "[List of tables]" in p.text:
            p.text = "Bảng 1.1: Bảng dữ liệu."
    for i in range(1, 12):
        d.add_paragraph(f"[{i}] Tac gia {i}, Sach {i}, 2024.")
    d.save(path)
    return prof

def test_header_arbitrary_garbage_fails_verify(tmp_path: Path):
    doc_path = tmp_path / "garbage_header.docx"
    prof = _create_clean_report(doc_path)
    
    d = Document(doc_path)
    body_sec = d.sections[-1]
    # Replace header with garbage that does not match topic or supervisor
    body_sec.header.paragraphs[0].text = "RANDOM GARBAGE HEADER TEXT"
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues if i.severity == "error"]
    assert any("header" in c for c in codes)

def test_header_missing_supervisor_fails_verify(tmp_path: Path):
    doc_path = tmp_path / "missing_sup.docx"
    prof = _create_clean_report(doc_path)
    
    d = Document(doc_path)
    body_sec = d.sections[-1]
    # Keep topic on left but omit supervisor on right
    body_sec.header.paragraphs[0].text = "XAY DUNG HE THONG RAG CHO VAN BAN PHAP LUAT\t"
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues if i.severity == "error"]
    assert "missing_header_supervisor" in codes or any("supervisor" in c for c in codes)

def test_footer_student_mismatch_fails_verify(tmp_path: Path):
    doc_path = tmp_path / "mismatched_student.docx"
    prof = _create_clean_report(doc_path)
    
    d = Document(doc_path)
    body_sec = d.sections[-1]
    # Keep page number but change student to arbitrary text
    for p in body_sec.footer.paragraphs:
        if "NGUYEN VAN A" in p.text:
            p.text = p.text.replace("NGUYEN VAN A - 2351010001", "SOME OTHER UNRELATED PERSON")
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues if i.severity == "error"]
    assert any("footer" in c or "student" in c for c in codes)
