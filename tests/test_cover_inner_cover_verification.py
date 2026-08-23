from __future__ import annotations
from pathlib import Path
from docx import Document
from wordassistant.document.validate import validate_document
from wordassistant.profiles.loader import resolve_profile
from wordassistant.document.scaffold import scaffold_document

def _create_clean_report(path: Path):
    prof = resolve_profile("hcmunre_dacn_2026_2027")
    scaffold_document(
        path,
        prof,
        "report",
        title="XAY DUNG HE THONG RAG CHO VAN BAN PHAP LUAT",
        student="NGUYEN VAN A - 2351010001",
        supervisor="TS. NGUYEN VAN B"
    )
    d = Document(path)
    for p in d.paragraphs:
        if "[Front matter content]" in p.text:
            p.text = "Lời cảm ơn chi tiết."
        elif "[Section content...]" in p.text:
            p.text = "Nội dung chương hoàn chỉnh."
        elif "[List of figures]" in p.text:
            p.text = "Hình 1.1: Sơ đồ."
        elif "[List of tables]" in p.text:
            p.text = "Bảng 1.1: Bảng dữ liệu."
    for i in range(1, 12):
        d.add_paragraph(f"[{i}] Tac gia {i}, Sach {i}, 2024.")
    d.save(path)
    return prof

def test_missing_cover_institution_fails_verify(tmp_path: Path):
    doc_path = tmp_path / "no_institution.docx"
    prof = _create_clean_report(doc_path)
    
    d = Document(doc_path)
    # Remove institution lines from cover
    for p in d.paragraphs:
        if "TÀI NGUYÊN VÀ MÔI TRƯỜNG" in p.text.upper() or "KHOA CÔNG NGHỆ THÔNG TIN" in p.text.upper():
            p.text = ""
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues if i.severity == "error"]
    assert "missing_cover_institution" in codes or any("cover" in c for c in codes)

def test_missing_document_type_fails_verify(tmp_path: Path):
    doc_path = tmp_path / "no_doctype.docx"
    prof = _create_clean_report(doc_path)
    
    d = Document(doc_path)
    for p in d.paragraphs:
        if "BÁO CÁO ĐỒ ÁN CHUYÊN NGÀNH" in p.text.upper():
            p.text = ""
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues if i.severity == "error"]
    assert "missing_cover_document_type" in codes or any("cover" in c for c in codes)
