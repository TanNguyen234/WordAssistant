from __future__ import annotations

import hashlib
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from PIL import Image

from wordassistant.document.structure import inspect_document
from wordassistant.operations.images import insert_image
from wordassistant.selectors import SelectorAmbiguous


def _make_image(path, *, size=(1200, 600)):
    Image.new("RGB", size, color=(24, 92, 148)).save(path, dpi=(96, 96))


def _make_anchor_document(path):
    document = Document()
    document.add_paragraph("Before")
    document.add_paragraph("Anchor")
    document.add_paragraph("After")
    document.save(path)


def test_insert_has_dedicated_paragraph_metadata_and_semantic_caption(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "architecture.png"
    output = tmp_path / "illustrated.docx"
    _make_anchor_document(source)
    _make_image(image)

    result = insert_image(
        source,
        output_path=output,
        image_path=image,
        selector={"type": "text", "text": "Anchor"},
        placement="after",
        alignment="right",
        width_cm=6,
        preserve_aspect_ratio=True,
        alt_text="Architecture flow",
        title="System architecture",
        caption="Production processing pipeline",
        caption_label="Figure",
    )

    reopened = Document(result.output_path)
    assert [paragraph.text for paragraph in reopened.paragraphs] == [
        "Before",
        "Anchor",
        "",
        "Figure . Production processing pipeline",
        "After",
    ]
    image_paragraph = reopened.paragraphs[2]
    assert image_paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert len(image_paragraph._p.xpath(".//w:drawing")) == 1
    assert reopened.paragraphs[3].style.name == "Caption"

    with ZipFile(output) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "SEQ Figure \\* ARABIC" in document_xml

    detail = inspect_document(output)["images_detail"]
    assert len(detail) == 1
    image_detail = detail[0]
    assert image_detail["paragraph_index"] == 2
    assert image_detail["section_index"] == 0
    assert image_detail["alignment"] == "right"
    assert image_detail["inline"] is True
    assert image_detail["alt_text"] == "Architecture flow"
    assert image_detail["title"] == "System architecture"
    assert image_detail["width_emu"] == 6 * 360000
    assert image_detail["height_emu"] == 3 * 360000
    assert image_detail["width_cm"] == 6
    assert image_detail["height_cm"] == 3
    assert image_detail["aspect_ratio"] == 2
    assert image_detail["media_sha256"] == hashlib.sha256(image.read_bytes()).hexdigest()
    assert image_detail["associated_caption"]["style"] == "Caption"
    assert image_detail["associated_caption"]["seq_field"] is True
    assert image_detail["caption_field_evidence"] == [
        {"kind": "simple", "instruction": "SEQ Figure \\* ARABIC"}
    ]


@pytest.mark.parametrize(
    ("placement", "alignment", "expected_text"),
    [
        ("before", "left", ["Before", "", "Anchor", "After"]),
        ("after", "center", ["Before", "Anchor", "", "After"]),
        ("replace", "right", ["Before", "", "After"]),
    ],
)
def test_placement_is_structural_and_uses_a_dedicated_paragraph(
    tmp_path, placement, alignment, expected_text
):
    source = tmp_path / f"{placement}-source.docx"
    image = tmp_path / f"{placement}.png"
    output = tmp_path / f"{placement}-output.docx"
    _make_anchor_document(source)
    _make_image(image, size=(100, 50))

    insert_image(
        source,
        output_path=output,
        image_path=image,
        selector={"type": "text", "text": "Anchor"},
        placement=placement,
        alignment=alignment,
    )

    reopened = Document(output)
    assert [paragraph.text for paragraph in reopened.paragraphs] == expected_text
    image_paragraphs = [
        paragraph
        for paragraph in reopened.paragraphs
        if paragraph._p.xpath(".//w:drawing")
    ]
    assert len(image_paragraphs) == 1
    assert image_paragraphs[0].text == ""
    assert inspect_document(output)["images_detail"][0]["alignment"] == alignment


def test_auto_shrink_uses_effective_section_and_inherited_anchor_indents(tmp_path):
    source = tmp_path / "sections.docx"
    image = tmp_path / "wide.png"
    output = tmp_path / "shrunk.docx"
    _make_image(image)
    document = Document()
    document.add_paragraph("First section")
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Cm(30)
    section.page_height = Cm(21)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    anchor = document.add_paragraph("Landscape anchor")
    anchor.paragraph_format.left_indent = Cm(3)
    anchor.paragraph_format.right_indent = Cm(2)
    document.save(source)

    result = insert_image(
        source,
        output_path=output,
        image_path=image,
        selector={"type": "text", "text": "Landscape anchor"},
        placement="after",
        width_cm=28,
        preserve_aspect_ratio=True,
    )

    assert result.data["section_index"] == 1
    assert result.data["auto_shrunk"] is True
    assert result.data["available_width_emu"] == pytest.approx(
        21 * 360000, abs=635
    )
    detail = inspect_document(output)["images_detail"][0]
    assert detail["section_index"] == 1
    assert detail["width_cm"] == pytest.approx(21, abs=0.001)
    assert detail["height_cm"] == pytest.approx(10.5, abs=0.001)
    assert detail["aspect_ratio"] == pytest.approx(2, abs=0.0001)


@pytest.mark.parametrize(
    "dimensions",
    [
        {"width_cm": 0},
        {"height_cm": -1},
        {"width_cm": 4, "height_cm": 2, "preserve_aspect_ratio": True},
    ],
)
def test_invalid_dimensions_fail_before_output_is_written(tmp_path, dimensions):
    source = tmp_path / "source.docx"
    image = tmp_path / "image.png"
    output = tmp_path / "must-not-exist.docx"
    _make_anchor_document(source)
    _make_image(image)

    with pytest.raises((TypeError, ValueError)):
        insert_image(source, output_path=output, image_path=image, **dimensions)

    assert not output.exists()


def test_selector_remains_fail_closed_without_writing(tmp_path):
    source = tmp_path / "ambiguous.docx"
    image = tmp_path / "image.png"
    output = tmp_path / "must-not-exist.docx"
    document = Document()
    document.add_paragraph("Duplicate")
    document.add_paragraph("Duplicate")
    document.save(source)
    _make_image(image)

    with pytest.raises(SelectorAmbiguous):
        insert_image(
            source,
            output_path=output,
            image_path=image,
            selector={"type": "text", "text": "Duplicate"},
            placement="after",
        )

    assert not output.exists()
