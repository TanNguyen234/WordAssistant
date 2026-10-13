import pytest
from docx import Document
from docx.shared import RGBColor

from wordassistant.operations.styles import update_style


def test_style_update_sets_font_color_size_and_heading_indentation(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Second-level heading", style="Heading 2")
    document.save(source)

    result = update_style(
        source,
        output_path=tmp_path / "styled.docx",
        name="Heading 2",
        font_family="Times New Roman",
        font_size_pt=13,
        bold=True,
        font_color="000000",
        left_indent_pt=0,
        first_line_indent_pt=0,
    )

    style = Document(result.output_path).styles["Heading 2"]
    assert style.font.size.pt == 13
    assert style.font.color.rgb == RGBColor(0, 0, 0)
    assert style.font.bold is True
    assert style.paragraph_format.left_indent.pt == 0
    assert style.paragraph_format.first_line_indent.pt == 0


def test_style_update_rejects_invalid_font_color_without_output(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "styled.docx"
    Document().save(source)

    with pytest.raises(ValueError, match="font_color"):
        update_style(source, output_path=output, name="Caption", font_color="black")
    assert not output.exists()
