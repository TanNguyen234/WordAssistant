from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from wordassistant.operations.lists import (
    apply_heading_numbering,
    create_list,
    inspect_heading_numbering,
    remove_heading_numbering,
    restart_numbering,
    set_list_level,
)


def _source(path):
    Document().save(path)
    return path


def _num_id(paragraph):
    return int(paragraph._p.pPr.numPr.numId.val)


def test_multilevel_numbering_writes_levels_and_paragraph_assignments(tmp_path):
    result = create_list(
        _source(tmp_path / "source.docx"),
        output_path=tmp_path / "list.docx",
        items=["One", "Nested", "Two"],
        kind="numbered",
        levels=[0, 1, 0],
        formats=["decimal", "lowerLetter", "lowerRoman"],
    )

    document = Document(result.output_path)
    assert [p.text for p in document.paragraphs] == ["One", "Nested", "Two"]
    assert [int(p._p.pPr.numPr.ilvl.val) for p in document.paragraphs] == [0, 1, 0]
    with ZipFile(result.output_path) as package:
        numbering = package.read("word/numbering.xml").decode("utf-8")
    assert 'w:val="decimal"' in numbering
    assert 'w:val="lowerLetter"' in numbering
    assert numbering.count("<w:lvl ") >= 3


def test_restart_numbering_creates_new_num_instance_with_start_override(tmp_path):
    first = create_list(
        _source(tmp_path / "source.docx"),
        output_path=tmp_path / "list.docx",
        items=["A", "B", "C"],
        kind="numbered",
    )
    before = Document(first.output_path)
    old_id = _num_id(before.paragraphs[0])

    restarted = restart_numbering(
        first.output_path,
        output_path=tmp_path / "restart.docx",
        paragraph_index=2,
        start=1,
    )

    after = Document(restarted.output_path)
    assert _num_id(after.paragraphs[2]) != old_id
    with ZipFile(restarted.output_path) as package:
        numbering = package.read("word/numbering.xml").decode("utf-8")
    assert "<w:startOverride w:val=\"1\"" in numbering


def test_set_list_level_updates_only_selected_paragraph(tmp_path):
    first = create_list(
        _source(tmp_path / "source.docx"),
        output_path=tmp_path / "list.docx",
        items=["A", "B"],
        kind="bullet",
    )

    set_list_level(
        first.output_path,
        output_path=tmp_path / "nested.docx",
        paragraph_index=1,
        level=2,
    )

    paragraphs = Document(tmp_path / "nested.docx").paragraphs
    assert int(paragraphs[0]._p.pPr.numPr.ilvl.val) == 0
    assert int(paragraphs[1]._p.pPr.numPr.ilvl.val) == 2


def test_heading_numbering_links_real_heading_styles_without_literal_prefixes(tmp_path):
    source = tmp_path / "headings.docx"
    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_heading("Purpose", level=2)
    document.add_heading("Scope", level=2)
    document.save(source)

    result = apply_heading_numbering(source, output_path=tmp_path / "numbered.docx")

    reopened = Document(result.output_path)
    assert [p.text for p in reopened.paragraphs] == ["Introduction", "Purpose", "Scope"]
    report = inspect_heading_numbering(result.output_path)
    assert report["linked_levels"][:2] == ["Heading1", "Heading2"]
    assert report["heading_paragraphs"] == 3
    assert report["literal_prefixes"] == []
    with ZipFile(result.output_path) as package:
        numbering = package.read("word/numbering.xml").decode("utf-8")
        styles = package.read("word/styles.xml").decode("utf-8")
    assert '<w:pStyle w:val="Heading1"' in numbering
    assert '<w:pStyle w:val="Heading2"' in numbering
    assert '<w:lvlText w:val="%1.%2"' in numbering
    assert '<w:style w:type="paragraph" w:styleId="Heading1"' in styles
    assert "<w:numPr>" in styles

    removed = remove_heading_numbering(
        result.output_path, output_path=tmp_path / "unnumbered.docx"
    )
    assert inspect_heading_numbering(removed.output_path)["linked_levels"] == []


def test_heading_numbering_rejects_existing_literal_numbers_by_default(tmp_path):
    source = tmp_path / "literal.docx"
    document = Document()
    document.add_heading("1 Introduction", level=1)
    document.save(source)

    try:
        apply_heading_numbering(source, output_path=tmp_path / "out.docx")
    except ValueError as exc:
        assert "literal numbering prefix" in str(exc)
    else:
        raise AssertionError("Expected double-numbering protection")


def test_heading_numbering_can_hide_level_one_label_for_literal_chapter_titles(tmp_path):
    source = tmp_path / "chapters.docx"
    document = Document()
    document.add_heading("CHƯƠNG 1: TỔNG QUAN", level=1)
    document.add_heading("Phạm vi", level=2)
    document.save(source)

    result = apply_heading_numbering(
        source,
        output_path=tmp_path / "numbered.docx",
        suppress_first_level_label=True,
    )

    with ZipFile(result.output_path) as package:
        numbering = package.read("word/numbering.xml").decode("utf-8")
    assert '<w:lvlText w:val=""' in numbering
    reopened = Document(result.output_path)
    assert reopened.paragraphs[0].style.name == "Heading 1"
    assert reopened.paragraphs[1].style.name == "Heading 2"
