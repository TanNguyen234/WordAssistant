from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from wordassistant.operations.fields import insert_caption, insert_toc


def test_caption_and_list_of_figures_use_real_fields(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Figure content")
    doc.save(source)
    captioned = insert_caption(
        source,
        output_path=tmp_path / "captioned.docx",
        label="Figure",
        title="System architecture",
    )
    listed = insert_toc(
        captioned.output_path,
        output_path=tmp_path / "listed.docx",
        kind="figures",
        caption_label="Figure",
    )

    with ZipFile(listed.output_path) as package:
        xml = package.read("word/document.xml").decode("utf-8")
    assert "SEQ Figure" in xml
    assert "TOC" in xml and "Figure" in xml
    assert "System architecture" in xml


def test_table_of_contents_and_list_of_tables_have_distinct_instructions(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    toc = insert_toc(source, output_path=tmp_path / "toc.docx", kind="contents", heading_depth=4)
    lot = insert_toc(toc.output_path, output_path=tmp_path / "lot.docx", kind="tables", caption_label="Table")

    with ZipFile(lot.output_path) as package:
        xml = package.read("word/document.xml").decode("utf-8")
    assert 'TOC \\o &quot;1-4&quot;' in xml or 'TOC \\o "1-4"' in xml
    assert '\\c &quot;Table&quot;' in xml or '\\c "Table"' in xml


def test_toc_and_caption_can_be_placed_after_heading_boundaries(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("Contents", level=1)
    document.add_paragraph("Body follows contents")
    document.add_heading("Figure area", level=1)
    document.add_paragraph("Body follows figure")
    document.save(source)

    toc = insert_toc(
        source,
        output_path=tmp_path / "toc.docx",
        selector={"type": "heading", "text": "Contents"},
    )
    captioned = insert_caption(
        toc.output_path,
        output_path=tmp_path / "caption.docx",
        label="Figure",
        title="Architecture",
        selector={"type": "heading", "text": "Figure area"},
    )

    reopened = Document(captioned.output_path)
    assert reopened.paragraphs[1]._p.xpath(".//w:fldSimple")[0].get(
        qn("w:instr")
    ).startswith("TOC ")
    assert reopened.paragraphs[2].text == "Body follows contents"
    assert reopened.paragraphs[4].style.name == "Caption"
    assert reopened.paragraphs[5].text == "Body follows figure"


def test_toc_insert_updates_matching_scaffold_field_without_duplication(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("MỤC LỤC", level=1)
    field_paragraph = document.add_paragraph()
    from wordassistant.operations.fields import append_field

    append_field(
        field_paragraph,
        'TOC \\o "1-3" \\h \\z \\u',
        "Update field in Word",
    )
    document.add_paragraph("Body")
    document.save(source)

    result = insert_toc(
        source,
        output_path=tmp_path / "updated.docx",
        kind="contents",
        heading_depth=3,
        selector={"type": "heading", "text": "MỤC LỤC"},
    )

    reopened = Document(result.output_path)
    fields = reopened.element.xpath(".//w:fldSimple")
    assert len(fields) == 1
    assert result.data["updated_existing"] is True
