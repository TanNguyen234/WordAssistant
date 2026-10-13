import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from wordassistant.operations.fields import insert_caption, inspect_fields
from wordassistant.operations.images import insert_image
from wordassistant.operations.lists import apply_heading_numbering
from wordassistant.profiles.loader import load_profile, ProfileError


@pytest.mark.parametrize('label,sequence', [('HÌNH', 'Hinh'), ('SƠ ĐỒ', 'SoDo'), ('BẢNG', 'Bang')])
def test_chapter_caption_sequence_resets_and_keeps_list_identifier(tmp_path, label, sequence):
    source = tmp_path / 'source.docx'
    output = tmp_path / 'output.docx'
    doc = Document()
    doc.add_heading('CHƯƠNG 1. METHODS', 1)
    doc.add_paragraph('first')
    doc.add_heading('CHƯƠNG 2. RESULTS', 1)
    doc.add_paragraph('second')
    doc.save(source)
    # Insert in reverse order: reset must be tied to headings, not call order.
    for chapter, anchor in [(2, 'second'), (1, 'first'), (1, 'first')]:
        insert_caption(source, output_path=output, label=label, title='Checked title',
                       sequence_id=sequence, chapter_number=chapter,
                       selector={'type': 'text', 'text': anchor})
        source = output
    fields = inspect_fields(output)
    assert [f['instruction'] for f in fields] == [f'SEQ {sequence} \\* ARABIC \\s 1'] * 3
    captions = [p for p in Document(output).paragraphs if p.style.name == 'Caption']
    visible = [''.join(p._p.xpath('.//w:t/text()')) for p in captions]
    # Cached sequence results are placeholders until Word updates fields.
    assert visible == [f'{label} 1.1. Checked title'] * 2 + [f'{label} 2.1. Checked title']


@pytest.mark.parametrize('heading,style,chapter', [
    ('CHƯƠNG 2. RESULTS', 'Heading 1', 1),
    ('INTRODUCTION', 'Heading 1', 1),
    ('CHƯƠNG 1. METHODS', 'Normal', 1),
    ('CHƯƠNG 1. METHODS', 'Heading 1', True),
    ('CHƯƠNG 1. METHODS', 'Heading 1', 0),
])
def test_chapter_caption_fails_closed(tmp_path, heading, style, chapter):
    source, output = tmp_path / 'source.docx', tmp_path / 'output.docx'
    doc = Document()
    doc.add_paragraph(heading, style)
    doc.save(source)
    with pytest.raises(ValueError, match='chapter|Heading 1'):
        insert_caption(source, output_path=output, label='HÌNH', title='Test',
                       sequence_id='Hinh', chapter_number=chapter)
    assert not output.exists()


def test_automatically_numbered_heading_rejected(tmp_path):
    source, output = tmp_path / 'source.docx', tmp_path / 'output.docx'
    doc = Document()
    heading = doc.add_heading('CHƯƠNG 1. METHODS', 1)
    num = OxmlElement('w:numPr')
    num_id = OxmlElement('w:numId')
    num_id.set(qn('w:val'), '1')
    num.append(num_id)
    heading._p.get_or_add_pPr().append(num)
    doc.save(source)
    with pytest.raises(ValueError, match='automatic'):
        insert_caption(source, output_path=output, label='HÌNH', title='Test',
                       sequence_id='Hinh', chapter_number=1)
    assert not output.exists()


def test_image_uses_same_chapter_contract(tmp_path):
    source, output = tmp_path / 'source.docx', tmp_path / 'output.docx'
    image = tmp_path / 'image.png'
    Image.new('RGB', (10, 10)).save(image)
    doc = Document()
    doc.add_heading('CHƯƠNG 3. RESULTS', 1)
    doc.save(source)
    result = insert_image(source, output_path=output, image_path=image,
                          caption='Measured figure', caption_label='HÌNH',
                          caption_sequence='Hinh', caption_chapter_number=3)
    assert inspect_fields(output)[0]['instruction'] == 'SEQ Hinh \\* ARABIC \\s 1'
    assert result.data['caption_chapter_number'] == 3
    assert any(p.text.startswith('HÌNH 3.') for p in Document(output).paragraphs)


def test_image_caption_uses_numeric_chapter_and_colon(tmp_path):
    source, output, image = tmp_path / 'source.docx', tmp_path / 'output.docx', tmp_path / 'image.png'
    Image.new('RGB', (10, 10)).save(image)
    doc = Document()
    doc.add_heading('5 LUỒNG NGHIỆP VỤ QUẢN TRỊ VÀ KHO', 1)
    doc.save(source)

    insert_image(source, output_path=output, image_path=image, caption='Tồn kho',
                 caption_label='Hình', caption_sequence='Hinh', caption_chapter_number=5,
                 caption_separator=':')

    caption = next(p for p in Document(output).paragraphs if p.style.name == 'Caption')
    assert ''.join(caption._p.xpath('.//w:t/text()')) == 'Hình 5.1: Tồn kho'


def test_numbered_heading_caption_accepts_colon_separator(tmp_path):
    source, output = tmp_path / 'source.docx', tmp_path / 'output.docx'
    doc = Document()
    doc.add_heading('4 LUỒNG NGHIỆP VỤ KHÁCH HÀNG', 1)
    doc.add_paragraph('screen anchor')
    doc.save(source)

    insert_caption(source, output_path=output, label='Hình', title='Giao diện trang đăng nhập',
                   sequence_id='Hinh', chapter_number=4, caption_separator=':',
                   selector={'type': 'text', 'text': 'screen anchor'})

    fields = inspect_fields(output)
    caption = next(p for p in Document(output).paragraphs if p.style.name == 'Caption')
    visible = ''.join(caption._p.xpath('.//w:t/text()'))
    assert fields[0]['instruction'] == 'SEQ Hinh \\* ARABIC \\s 1'
    assert visible == 'Hình 4.1: Giao diện trang đăng nhập'


def test_caption_number_verifies_automatic_multilevel_heading(tmp_path):
    source, numbered, output = (
        tmp_path / 'source.docx', tmp_path / 'numbered.docx', tmp_path / 'captioned.docx'
    )
    doc = Document()
    doc.add_heading('Front matter', 1)
    doc.add_heading('CƠ SỞ LÝ THUYẾT', 1)
    doc.add_heading('KHẢO SÁT TỔNG QUAN', 1)
    doc.add_heading('PHÂN TÍCH MÔ HÌNH', 1)
    doc.add_heading('LUỒNG NGHIỆP VỤ KHÁCH HÀNG', 1)
    doc.add_paragraph('screen anchor')
    doc.add_heading('Appendix', 1)
    doc.save(source)

    apply_heading_numbering(
        source,
        output_path=numbered,
        levels=3,
        start_text='CƠ SỞ LÝ THUYẾT',
        end_text='Appendix',
    )
    insert_caption(
        numbered,
        output_path=output,
        label='Hình',
        title='Giao diện danh mục',
        sequence_id='Hinh',
        chapter_number=4,
        caption_separator=':',
        selector={'type': 'text', 'text': 'screen anchor'},
    )

    caption = next(p for p in Document(output).paragraphs if p.style.name == 'Caption')
    assert ''.join(caption._p.xpath('.//w:t/text()')) == 'Hình 4.1: Giao diện danh mục'
    assert inspect_fields(output)[0]['instruction'] == 'SEQ Hinh \\* ARABIC \\s 1'


@pytest.mark.parametrize('separator', ['-', ';', ''])
def test_caption_rejects_unsupported_separator(tmp_path, separator):
    source, output = tmp_path / 'source.docx', tmp_path / 'output.docx'
    Document().save(source)
    with pytest.raises(ValueError, match='separator'):
        insert_caption(source, output_path=output, label='Hình', title='Test',
                       sequence_id='Hinh',
                       caption_separator=separator)
    assert not output.exists()


def test_generic_caption_remains_global(tmp_path):
    source, output = tmp_path / 'source.docx', tmp_path / 'output.docx'
    Document().save(source)
    insert_caption(source, output_path=output, label='Figure', title='Generic')
    assert inspect_fields(output)[0]['instruction'] == 'SEQ Figure \\* ARABIC'


def test_nearest_heading_blocks_prior_chapter(tmp_path):
    source, output = tmp_path / 'source.docx', tmp_path / 'output.docx'
    doc = Document()
    doc.add_heading('CHƯƠNG 1. METHODS', 1)
    doc.add_heading('TÀI LIỆU THAM KHẢO', 1)
    doc.save(source)
    with pytest.raises(ValueError, match='nearest Heading 1'):
        insert_caption(source, output_path=output, label='BẢNG', title='Test',
                       sequence_id='Bang', chapter_number=1)
    assert not output.exists()


@pytest.mark.parametrize('chapter,caption', [(2, 'Wrong chapter'), (1, None)])
def test_image_chapter_prerequisite_failure_preserves_output(tmp_path, chapter, caption):
    source, output, image = tmp_path / 'source.docx', tmp_path / 'output.docx', tmp_path / 'image.png'
    Image.new('RGB', (10, 10)).save(image)
    doc = Document()
    doc.add_heading('CHƯƠNG 1. METHODS', 1)
    doc.save(source)
    output.write_bytes(b'existing output')
    with pytest.raises(ValueError, match='chapter|caption'):
        insert_image(source, output_path=output, image_path=image, caption=caption,
                     caption_label='HÌNH', caption_sequence='Hinh', caption_chapter_number=chapter)
    assert output.read_bytes() == b'existing output'


def test_chapter_profile_metadata_validates_and_copies_match():
    root = Path(__file__).resolve().parents[2]
    local = root / 'profiles/hcm-university-dacn-2026-2027.json'
    package = root / 'src/wordassistant/data/profiles/hcm-university-dacn-2026-2027.json'
    assert json.loads(local.read_text(encoding='utf-8')) == json.loads(package.read_text(encoding='utf-8'))
    assert load_profile(local).variant('report').captions['numbering']['format'] == 'chapter.sequence'


@pytest.mark.parametrize('key,value', [('heading_level', True), ('heading_level', 2),
    ('mode', 'automatic'), ('separator', '-'), ('unrecognized', 'value'),
    ('chapter_option', 'wrong'), ('format', 1)])
def test_chapter_profile_rejects_unsupported_numbering(tmp_path, key, value):
    root = Path(__file__).resolve().parents[2]
    profile = json.loads((root / 'profiles/hcm-university-dacn-2026-2027.json').read_text(encoding='utf-8'))
    profile['variants']['report']['captions']['numbering'][key] = value
    path = tmp_path / 'invalid.json'
    path.write_text(json.dumps(profile), encoding='utf-8')
    with pytest.raises(ProfileError, match='numbering'):
        load_profile(path)


@pytest.mark.parametrize('two_styles', [False, True])
@pytest.mark.parametrize('image_caption', [False, True])
def test_chapter_caption_rejects_cyclic_heading_styles(tmp_path, two_styles, image_caption):
    source, output = tmp_path / 'source.docx', tmp_path / 'output.docx'
    doc = Document()
    heading_style = doc.styles['Heading 1']
    if two_styles:
        peer = doc.styles['Heading 2']
        heading_style.base_style = peer
        peer.base_style = heading_style
        # Keep outline level compatible so the tested failure is the cycle.
        peer.element.pPr.find(qn('w:outlineLvl')).set(qn('w:val'), '0')
    else:
        heading_style.base_style = heading_style
    doc.add_heading('CHƯƠNG 1. METHODS', 1)
    doc.save(source)
    output.write_bytes(b'existing output')
    with pytest.raises(ValueError, match='cyclic.*style'):
        if image_caption:
            image = tmp_path / 'image.png'
            Image.new('RGB', (10, 10)).save(image)
            insert_image(source, output_path=output, image_path=image, caption='Measured figure',
                         caption_label='HÌNH', caption_sequence='Hinh', caption_chapter_number=1)
        else:
            insert_caption(source, output_path=output, label='BẢNG', title='Measured table',
                           sequence_id='Bang', chapter_number=1)
    assert output.read_bytes() == b'existing output'
