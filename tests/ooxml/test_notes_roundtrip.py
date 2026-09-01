from zipfile import ZipFile

from docx import Document

from wordassistant.operations.notes import delete_note, edit_note, insert_note, list_notes


def test_footnote_insert_edit_delete_roundtrip(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Claim")
    doc.save(source)
    inserted = insert_note(
        source,
        output_path=tmp_path / "inserted.docx",
        kind="footnote",
        text="Original note",
        selector={"type": "paragraph", "index": 0},
    )
    assert list_notes(inserted.output_path, kind="footnote")[0]["text"] == "Original note"
    edited = edit_note(inserted.output_path, output_path=tmp_path / "edited.docx", kind="footnote", note_id=1, text="Edited note")
    assert list_notes(edited.output_path, kind="footnote")[0]["text"] == "Edited note"
    with ZipFile(edited.output_path) as package:
        assert "word/footnotes.xml" in package.namelist()
        assert b"footnoteReference" in package.read("word/document.xml")

    deleted = delete_note(edited.output_path, output_path=tmp_path / "deleted.docx", kind="footnote", note_id=1)
    assert list_notes(deleted.output_path, kind="footnote") == []


def test_endnote_has_distinct_part_and_reference(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    inserted = insert_note(source, output_path=tmp_path / "endnote.docx", kind="endnote", text="End note")
    with ZipFile(inserted.output_path) as package:
        assert "word/endnotes.xml" in package.namelist()
        assert b"endnoteReference" in package.read("word/document.xml")

