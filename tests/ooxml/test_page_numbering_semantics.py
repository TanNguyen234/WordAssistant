import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START

from wordassistant.operations.page_numbers import configure_page_numbers, inspect_page_numbers


def _two_section_document(path):
    doc = Document()
    doc.add_heading("FRONT MATTER", level=1)
    doc.add_paragraph("Abstract")
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading("MỞ ĐẦU", level=1)
    doc.add_paragraph("Body")
    doc.save(path)
    return path


def test_front_roman_body_arabic_page_numbering_and_page_x_of_y(tmp_path):
    source = _two_section_document(tmp_path / "source.docx")
    front = configure_page_numbers(
        source,
        output_path=tmp_path / "front.docx",
        section_index=0,
        position="footer",
        alignment="center",
        number_format="lowerRoman",
        start=1,
        field="PAGE",
    )
    body = configure_page_numbers(
        front.output_path,
        output_path=tmp_path / "body.docx",
        section_index=1,
        position="footer",
        alignment="right",
        number_format="decimal",
        start=1,
        page_x_of_y=True,
        unlink_previous=True,
    )

    sections = inspect_page_numbers(body.output_path)
    assert sections[0]["format"] == "lowerRoman"
    assert sections[0]["start"] == 1
    assert sections[1]["format"] == "decimal"
    assert sections[1]["start"] == 1
    assert sections[1]["first_heading"] == "MỞ ĐẦU"
    assert sections[1]["fields"] == ["PAGE", "NUMPAGES"]
    assert sections[1]["position"] == "footer"
    assert sections[1]["alignment"] == "right"


def test_odd_even_page_numbering_populates_both_stories_and_can_disable(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    enabled = configure_page_numbers(
        source,
        output_path=tmp_path / "enabled.docx",
        section_index=0,
        different_odd_even=True,
    )
    document = Document(enabled.output_path)
    assert "PAGE" in document.sections[0].footer._element.xml
    assert "PAGE" in document.sections[0].even_page_footer._element.xml
    report = inspect_page_numbers(enabled.output_path)[0]
    assert report["variants"]["primary"]["fields"] == ["PAGE"]
    assert report["variants"]["even"]["fields"] == ["PAGE"]

    disabled = configure_page_numbers(
        enabled.output_path,
        output_path=tmp_path / "disabled.docx",
        section_index=0,
        different_odd_even=False,
    )
    assert "evenAndOddHeaders" not in Document(disabled.output_path).settings.element.xml


def test_configuring_another_section_does_not_implicitly_disable_odd_even(tmp_path):
    source = _two_section_document(tmp_path / "source.docx")
    enabled = configure_page_numbers(
        source,
        output_path=tmp_path / "enabled.docx",
        section_index=0,
        different_odd_even=True,
    )
    configured = configure_page_numbers(
        enabled.output_path,
        output_path=tmp_path / "configured.docx",
        section_index=1,
        number_format="decimal",
    )

    report = inspect_page_numbers(configured.output_path)
    assert all(section["odd_even_enabled"] for section in report)

    with pytest.raises(ValueError, match="document-level"):
        configure_page_numbers(
            configured.output_path,
            output_path=tmp_path / "unsafe-disable.docx",
            section_index=1,
            different_odd_even=False,
        )


def test_page_x_of_y_complex_fields_keep_independent_cached_results(tmp_path):
    source = _two_section_document(tmp_path / "source.docx")
    numbered = configure_page_numbers(
        source,
        output_path=tmp_path / "numbered.docx",
        section_index=0,
        page_x_of_y=True,
    )
    document = Document(numbered.output_path)
    result_nodes = [
        node
        for node in document.sections[0].footer._element.xpath(".//w:t")
        if (node.text or "") == "1"
    ]
    result_nodes[0].text = "2"
    result_nodes[1].text = "35"
    document.save(numbered.output_path)

    from wordassistant.operations.fields import inspect_fields

    footer = [
        field
        for field in inspect_fields(numbered.output_path)
        if field["story"].endswith("footer")
    ]
    assert [(field["type"], field["cached_result"]) for field in footer] == [
        ("PAGE", "2"),
        ("NUMPAGES", "35"),
    ]


def test_page_number_update_preserves_existing_footer_identity(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    footer = document.sections[0].footer.paragraphs[0]
    footer.add_run("NGUYEN VAN A - 2351010001\t")
    document.save(source)

    result = configure_page_numbers(
        source,
        output_path=tmp_path / "numbered.docx",
        section_index=0,
        alignment="right",
        number_format="decimal",
    )

    reopened = Document(result.output_path)
    paragraph = reopened.sections[0].footer.paragraphs[0]
    assert "NGUYEN VAN A - 2351010001" in paragraph.text
    assert "PAGE" in reopened.sections[0].footer._element.xml


def test_page_number_can_continue_without_a_restart_marker(tmp_path):
    source = _two_section_document(tmp_path / "source.docx")
    result = configure_page_numbers(
        source,
        output_path=tmp_path / "continued.docx",
        section_index=1,
        continue_from_previous=True,
    )

    assert inspect_page_numbers(result.output_path)[1]["start"] is None
