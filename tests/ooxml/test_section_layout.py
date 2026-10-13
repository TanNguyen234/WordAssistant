import pytest
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

from wordassistant.operations.sections import add_section_break, set_section_layout
from wordassistant.operations.page_numbers import configure_page_numbers


def _add_first_page_border(section, space_points=6):
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    borders.set(qn("w:display"), "firstPage")
    for edge_name in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "8")
        edge.set(qn("w:space"), str(space_points))
        borders.append(edge)
    section._sectPr.append(borders)


def test_section_layout_can_inset_existing_first_page_frame(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    _add_first_page_border(document.sections[0])
    document.add_paragraph("Cover")
    document.save(source)

    result = set_section_layout(
        source,
        output_path=tmp_path / "inset.docx",
        section_index=0,
        page_border_space_pt=34,
    )

    borders = Document(result.output_path).sections[0]._sectPr.find(qn("w:pgBorders"))
    assert borders.get(qn("w:offsetFrom")) == "page"
    assert borders.get(qn("w:display")) == "firstPage"
    assert {borders.find(qn(f"w:{edge}")).get(qn("w:space")) for edge in ("top", "left", "bottom", "right")} == {"34"}


def test_section_layout_requires_existing_frame_to_set_inset(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)

    with pytest.raises(ValueError, match="page border"):
        set_section_layout(
            source,
            output_path=tmp_path / "inset.docx",
            section_index=0,
            page_border_space_pt=34,
        )
    assert not (tmp_path / "inset.docx").exists()


def test_portrait_landscape_portrait_roundtrip_with_columns_and_mirroring(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Portrait")
    doc.save(source)

    first = add_section_break(source, output_path=tmp_path / "a.docx", break_type="next_page")
    second = add_section_break(first.output_path, output_path=tmp_path / "b.docx", break_type="continuous")
    set_section_layout(second.output_path, output_path=tmp_path / "c.docx", section_index=0, orientation="portrait", page_size="A4")
    set_section_layout(tmp_path / "c.docx", output_path=tmp_path / "d.docx", section_index=1, orientation="landscape", page_size="A4", columns=2, gutter_cm=0.5)
    result = set_section_layout(tmp_path / "d.docx", output_path=tmp_path / "final.docx", section_index=2, orientation="portrait", page_size="Letter", mirrored_margins=True)

    reopened = Document(result.output_path)
    assert [section.orientation for section in reopened.sections] == [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT]
    assert reopened.sections[1].start_type == WD_SECTION_START.NEW_PAGE
    assert 'w:num="2"' in reopened.sections[1]._sectPr.xml
    assert round(reopened.sections[1].gutter.cm, 1) == 0.5
    assert "w:mirrorMargins" in reopened.settings.element.xml


def test_section_layout_cannot_silently_disable_document_global_odd_even(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("First")
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph("Second")
    document.save(source)
    enabled = configure_page_numbers(
        source,
        output_path=tmp_path / "enabled.docx",
        section_index=0,
        different_odd_even=True,
        page_x_of_y=True,
    )

    with pytest.raises(ValueError, match="document-level"):
        set_section_layout(
            enabled.output_path,
            output_path=tmp_path / "unsafe.docx",
            section_index=1,
            different_odd_even=False,
        )

    assert "evenAndOddHeaders" in Document(enabled.output_path).settings.element.xml


def test_section_break_can_be_anchored_after_existing_content(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Portrait content")
    document.add_paragraph("Landscape content starts here")
    document.save(source)

    result = add_section_break(
        source,
        output_path=tmp_path / "anchored.docx",
        break_type="next_page",
        selector={"type": "text", "text": "Portrait content"},
    )

    reopened = Document(result.output_path)
    assert len(reopened.sections) == 2
    assert reopened.paragraphs[0].text == "Portrait content"
    assert "w:sectPr" in reopened.paragraphs[0]._p.xml
    assert reopened.paragraphs[1].text == "Landscape content starts here"


def test_anchored_section_break_returns_actual_inserted_section_index(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("First section", level=1)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph("Second section")
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph("Third section")
    document.save(source)

    inserted = add_section_break(
        source,
        output_path=tmp_path / "inserted.docx",
        selector={"type": "heading", "text": "First section"},
    )

    assert inserted.data["section_index"] == 1
    laid_out = set_section_layout(
        inserted.output_path,
        output_path=tmp_path / "landscape.docx",
        section_index=inserted.data["section_index"],
        orientation="landscape",
        page_size="A4",
    )
    reopened = Document(laid_out.output_path)
    assert len(reopened.sections) == 4
    assert reopened.sections[1].orientation == WD_ORIENT.LANDSCAPE


def test_anchored_break_preserves_existing_layouts_and_clones_anchor_section(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    first = document.sections[0]
    first.orientation = WD_ORIENT.PORTRAIT
    first.page_width, first.page_height = Cm(21), Cm(29.7)
    first.top_margin, first.bottom_margin = Cm(1.1), Cm(1.2)
    first.left_margin, first.right_margin = Cm(1.3), Cm(1.4)
    document.add_heading("Split here", level=1)
    document.add_paragraph("First remainder")

    second = document.add_section(WD_SECTION_START.NEW_PAGE)
    second.orientation = WD_ORIENT.LANDSCAPE
    second.page_width, second.page_height = Cm(29.7), Cm(21)
    second.top_margin, second.bottom_margin = Cm(2.1), Cm(2.2)
    second.left_margin, second.right_margin = Cm(2.3), Cm(2.4)
    document.add_paragraph("Second section")

    third = document.add_section(WD_SECTION_START.CONTINUOUS)
    third.orientation = WD_ORIENT.PORTRAIT
    third.page_width, third.page_height = Cm(21.59), Cm(27.94)
    third.top_margin, third.bottom_margin = Cm(3.1), Cm(3.2)
    third.left_margin, third.right_margin = Cm(3.3), Cm(3.4)
    document.add_paragraph("Third section")
    document.save(source)

    def layout(section):
        return (
            section.orientation,
            section.page_width,
            section.page_height,
            section.top_margin,
            section.bottom_margin,
            section.left_margin,
            section.right_margin,
        )

    before = [layout(section) for section in Document(source).sections]
    result = add_section_break(
        source,
        output_path=tmp_path / "inserted.docx",
        break_type="odd_page",
        selector={"type": "heading", "text": "Split here"},
    )
    reopened = Document(result.output_path)
    after = [layout(section) for section in reopened.sections]

    assert result.data["section_index"] == 1
    assert after == [before[0], before[0], before[1], before[2]]
    assert reopened.sections[0].start_type == WD_SECTION_START.NEW_PAGE
    assert reopened.sections[1].start_type == WD_SECTION_START.ODD_PAGE
    assert reopened.sections[2].start_type == WD_SECTION_START.NEW_PAGE
    assert reopened.sections[3].start_type == WD_SECTION_START.CONTINUOUS
