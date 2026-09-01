from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wordassistant.operations.fields import (
    inspect_fields,
    insert_field,
    lock_fields,
    remove_fields,
)


def _source(path):
    doc = Document()
    doc.add_paragraph("Field target")
    doc.save(path)
    return path


def test_insert_inspect_lock_unlock_and_remove_fields(tmp_path):
    first = insert_field(
        _source(tmp_path / "source.docx"),
        output_path=tmp_path / "page.docx",
        field_type="PAGE",
        selector={"type": "paragraph", "index": 0},
    )
    second = insert_field(
        first.output_path,
        output_path=tmp_path / "property.docx",
        field_type="DOCPROPERTY",
        argument="ProjectCode",
        selector={"type": "paragraph", "index": 0},
    )

    assert [field["type"] for field in inspect_fields(second.output_path)] == ["PAGE", "DOCPROPERTY"]

    locked = lock_fields(second.output_path, output_path=tmp_path / "locked.docx", locked=True)
    assert all(field["locked"] for field in inspect_fields(locked.output_path))
    unlocked = lock_fields(locked.output_path, output_path=tmp_path / "unlocked.docx", locked=False)
    assert not any(field["locked"] for field in inspect_fields(unlocked.output_path))
    removed = remove_fields(unlocked.output_path, output_path=tmp_path / "removed.docx", field_type="PAGE")
    assert [field["type"] for field in inspect_fields(removed.output_path)] == ["DOCPROPERTY"]


def _field_char(kind):
    element = OxmlElement("w:fldChar")
    element.set(qn("w:fldCharType"), kind)
    return element


def test_complex_field_instruction_split_across_runs_is_one_field(tmp_path):
    source = tmp_path / "split.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run()._r.append(_field_char("begin"))
    first = OxmlElement("w:instrText")
    first.text = " REF "
    paragraph.add_run()._r.append(first)
    second = OxmlElement("w:instrText")
    second.text = "target_bookmark \\h "
    paragraph.add_run()._r.append(second)
    paragraph.add_run()._r.append(_field_char("separate"))
    paragraph.add_run("resolved text")
    paragraph.add_run()._r.append(_field_char("end"))
    document.save(source)

    fields = inspect_fields(source)

    assert len(fields) == 1
    assert fields[0]["type"] == "REF"
    assert fields[0]["instruction"] == "REF target_bookmark \\h"
    assert fields[0]["cached_result"] == "resolved text"


def test_complex_field_structure_checks_order_not_only_global_counts(tmp_path):
    from wordassistant.operations.fields import inspect_field_structure

    source = tmp_path / "misordered.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run()._r.append(_field_char("begin"))
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGE "
    paragraph.add_run()._r.append(instruction)
    paragraph.add_run()._r.append(_field_char("end"))
    paragraph.add_run()._r.append(_field_char("separate"))
    document.save(source)

    result = inspect_field_structure(source)

    assert result["status"] == "FAIL"
    assert {issue["code"] for issue in result["issues"]} & {
        "complex_field_missing_separator",
        "orphan_complex_field_separator",
    }
