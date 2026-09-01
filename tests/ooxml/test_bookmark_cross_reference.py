import pytest
from docx import Document

from wordassistant.operations.fields import inspect_fields
from wordassistant.operations.references import (
    create_bookmark,
    delete_bookmark,
    insert_cross_reference,
    insert_hyperlink,
    inspect_bookmarks,
    rename_bookmark,
)


def _source(path):
    doc = Document()
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Measured value")
    doc.save(path)
    return path


def test_bookmark_ref_pageref_and_internal_hyperlink_roundtrip(tmp_path):
    bookmarked = create_bookmark(
        _source(tmp_path / "source.docx"),
        output_path=tmp_path / "bookmark.docx",
        name="results",
        selector={"type": "heading", "text": "Results"},
    )
    renamed = rename_bookmark(bookmarked.output_path, output_path=tmp_path / "renamed.docx", old_name="results", new_name="analysis_results")
    referenced = insert_cross_reference(
        renamed.output_path,
        output_path=tmp_path / "reference.docx",
        bookmark="analysis_results",
        include_page=True,
    )
    linked = insert_hyperlink(
        referenced.output_path,
        output_path=tmp_path / "linked.docx",
        text="Jump to results",
        bookmark="analysis_results",
    )

    assert inspect_bookmarks(linked.output_path)[0]["name"] == "analysis_results"
    assert {field["type"] for field in inspect_fields(linked.output_path)} >= {"REF", "PAGEREF"}
    assert "w:anchor=\"analysis_results\"" in Document(linked.output_path).element.xml

    with pytest.raises(ValueError, match="referenced"):
        delete_bookmark(
            linked.output_path,
            output_path=tmp_path / "deleted.docx",
            name="analysis_results",
        )


def test_external_hyperlink_creates_external_relationship(tmp_path):
    linked = insert_hyperlink(
        _source(tmp_path / "source.docx"),
        output_path=tmp_path / "linked.docx",
        text="OpenAI",
        url="https://openai.com",
    )
    relationships = Document(linked.output_path).part.rels.values()
    assert any(rel.target_ref == "https://openai.com" and rel.is_external for rel in relationships)
