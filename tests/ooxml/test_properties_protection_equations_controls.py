import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wordassistant.operations.content_controls import create_content_control, list_content_controls, set_content_control_value
from wordassistant.operations.equations import insert_equation
from wordassistant.operations.properties import get_properties, set_properties
from wordassistant.operations.protection import inspect_protection, protect_document, unprotect_document


def _source(path):
    doc = Document()
    doc.add_paragraph("Controlled value")
    doc.save(path)
    return path


def test_core_custom_properties_and_doc_protection_are_truthful(tmp_path):
    source = _source(tmp_path / "source.docx")
    properties = set_properties(
        source,
        output_path=tmp_path / "properties.docx",
        title="Research Report",
        author="Student",
        custom={"Company": "HCMUNRE", "RevisionCode": 7},
    )
    values = get_properties(properties.output_path)
    assert values["title"] == "Research Report"
    assert values["custom"]["Company"] == "HCMUNRE"

    protected = protect_document(properties.output_path, output_path=tmp_path / "protected.docx", mode="read_only")
    protection = inspect_protection(protected.output_path)
    assert protection["enforced"] is True
    assert protection["mode"] == "read_only"
    assert protection["secure_encryption"] is False
    unprotected = unprotect_document(protected.output_path, output_path=tmp_path / "unprotected.docx")
    assert inspect_protection(unprotected.output_path)["enforced"] is False


def test_equation_and_content_control_roundtrip(tmp_path):
    equation = insert_equation(
        _source(tmp_path / "source.docx"),
        output_path=tmp_path / "equation.docx",
        latex=r"E=mc^2",
        number=True,
    )
    created = create_content_control(
        equation.output_path,
        output_path=tmp_path / "control.docx",
        control_type="plain_text",
        title="Project code",
        tag="project_code",
        value="WA-001",
    )
    updated = set_content_control_value(
        created.output_path,
        output_path=tmp_path / "updated.docx",
        tag="project_code",
        value="WA-002",
    )
    controls = list_content_controls(updated.output_path)
    assert controls == [{"type": "plain_text", "title": "Project code", "tag": "project_code", "value": "WA-002", "locked": False}]
    xml = Document(updated.output_path).element.xml
    assert "m:oMath" in xml and "SEQ Equation" in xml


def test_set_value_rejects_unverified_advanced_control(tmp_path):
    source = _source(tmp_path / "source.docx")
    document = Document(source)
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "advanced")
    properties.extend([tag, OxmlElement("w:checkBox")])
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "false"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    sdt.extend([properties, content])
    document.element.body.insert(0, sdt)
    document.save(source)

    with pytest.raises(ValueError, match="plain_text"):
        set_content_control_value(
            source,
            output_path=tmp_path / "out.docx",
            tag="advanced",
            value="true",
        )
