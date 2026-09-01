from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

from wordassistant.document.validate import validate_document
from wordassistant.profiles.loader import resolve_profile


def _document_with_landscape_and_caption(path):
    document = Document()
    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    document.add_paragraph("Substantive body paragraph.")
    caption = document.add_paragraph("Figure 1. Architecture", style="Caption")
    caption.runs[0].font.size = Pt(9)
    section = document.add_section()
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    document.add_paragraph("Landscape appendix content.")
    document.save(path)


def test_generic_profile_accepts_explicit_mixed_orientation_and_caption_style(tmp_path):
    path = tmp_path / "mixed.docx"
    _document_with_landscape_and_caption(path)

    result = validate_document(path, resolve_profile("generic"), "default", verify_mode=True)
    codes = {issue.code for issue in result.issues if issue.severity == "error"}

    assert "page_geometry" not in codes
    assert "typography_size" not in codes


def test_generic_blueprint_is_guidance_not_required_titles():
    variant = resolve_profile("generic").variant("default")
    assert variant.section_blueprint
    assert all(section.get("required") is False for section in variant.section_blueprint)
