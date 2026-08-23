from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from wordassistant.profiles.loader import load_profile
from wordassistant.document.validate import validate_document

ROOT = Path(__file__).resolve().parents[1]

def test_validator_fails_on_wrong_font_and_size(tmp_path):
    profile = load_profile(ROOT / 'profiles/hcmunre_dacn_2026_2027.json')
    bad_doc = tmp_path / 'bad_font.docx'
    d = Document()
    s = d.sections[0]
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(2.0)
    s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(3.0)
    s.right_margin = Cm(2.0)
    
    # Intentionally set Arial 8pt instead of Times New Roman 13pt
    p = d.add_paragraph()
    r = p.add_run("This text is Arial 8pt.")
    r.font.name = "Arial"
    r.font.size = Pt(8.0)
    d.save(bad_doc)
    
    res = validate_document(bad_doc, profile, 'report')
    assert not res.valid
    assert any('font' in i.message.lower() or 'size' in i.message.lower() for i in res.issues)

def test_validator_fails_on_wrong_margins(tmp_path):
    profile = load_profile(ROOT / 'profiles/hcmunre_dacn_2026_2027.json')
    bad_doc = tmp_path / 'bad_margin.docx'
    d = Document()
    s = d.sections[0]
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.left_margin = Cm(1.0) # Should be 3.0cm
    d.save(bad_doc)
    
    res = validate_document(bad_doc, profile, 'report')
    assert not res.valid
    assert any('left_margin' in i.message.lower() or 'margin' in i.message.lower() for i in res.issues)

def test_validator_detects_missing_blueprint_sections(tmp_path):
    profile = load_profile(ROOT / 'profiles/hcmunre_dacn_2026_2027.json')
    doc_path = tmp_path / 'missing_sections.docx'
    d = Document()
    s = d.sections[0]
    s.page_width = Cm(21.0); s.page_height = Cm(29.7)
    s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0); s.left_margin = Cm(3.0); s.right_margin = Cm(2.0)
    d.add_heading("Only Chapter 1", level=1)
    d.save(doc_path)
    
    res = validate_document(doc_path, profile, 'report')
    assert not res.valid
    assert any('MỞ ĐẦU' in i.message or 'blueprint' in i.message.lower() for i in res.issues)
