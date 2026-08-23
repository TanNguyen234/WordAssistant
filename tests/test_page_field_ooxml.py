from pathlib import Path
from docx import Document
from wordassistant.profiles.loader import load_profile
from wordassistant.document.scaffold import scaffold_document

ROOT = Path(__file__).resolve().parents[1]

def test_page_field_ooxml_valid_structure_and_right_tab_stop(tmp_path):
    profile = load_profile(ROOT / 'profiles/hcmunre_dacn_2026_2027.json')
    out = tmp_path / 'report.docx'
    scaffold_document(
        out,
        profile,
        'report',
        title='HỆ THỐNG KIỂM TRA ĐỊNH DẠNG',
        student='Nguyễn Văn Test',
        supervisor='TS. Trần Hướng Dẫn'
    )
    
    d = Document(out)
    body_sec = d.sections[-1]
    
    # 1. Inspect footer OOXML structure
    footer_p = body_sec.footer.paragraphs[0]
    p_xml = footer_p._p.xml
    
    # In WordprocessingML, w:fldSimple must be a child of w:p, not nested inside w:r
    assert '<w:fldSimple' in p_xml
    assert 'w:instr="PAGE"' in p_xml
    # Make sure w:fldSimple is not inside <w:r>
    assert '<w:r><w:fldSimple' not in p_xml
    
    # 2. Inspect Tab Stop on header and footer
    # Usable width: 21.0 - 3.0 (left) - 2.0 (right) = 16.0 cm
    header_p = body_sec.header.paragraphs[0]
    tabs = header_p.paragraph_format.tab_stops
    assert len(tabs) >= 1
    tab = tabs[0]
    assert round(tab.position.cm, 1) == 16.0
    
    footer_tabs = footer_p.paragraph_format.tab_stops
    assert len(footer_tabs) >= 1
    assert round(footer_tabs[0].position.cm, 1) == 16.0
