from __future__ import annotations
from pathlib import Path
from docx import Document
from wordassistant.document.validate import validate_document
from wordassistant.profiles.loader import resolve_profile
from wordassistant.document.scaffold import scaffold_document

def _create_clean_report(path: Path):
    prof = resolve_profile("hcmunre_dacn_2026_2027")
    scaffold_document(
        path,
        prof,
        "report",
        title="XAY DUNG HE THONG RAG CHO VAN BAN PHAP LUAT",
        student="NGUYEN VAN A - 2351010001",
        supervisor="TS. NGUYEN VAN B"
    )
    d = Document(path)
    for p in d.paragraphs:
        if "[Front matter content]" in p.text:
            p.text = "Lời cảm ơn chi tiết."
        elif "[Section content...]" in p.text:
            p.text = "Nội dung chương hoàn chỉnh."
        elif "[List of figures]" in p.text:
            p.text = "Hình 1.1: Sơ đồ."
        elif "[List of tables]" in p.text:
            p.text = "Bảng 1.1: Bảng dữ liệu."
    for i in range(1, 12):
        d.add_paragraph(f"[{i}] Tac gia {i}, Sach {i}, 2024.")
    d.save(path)
    return prof

def test_missing_inner_cover_fails_verify(tmp_path: Path):
    doc_path = tmp_path / "missing_inner.docx"
    prof = _create_clean_report(doc_path)
    
    d = Document(doc_path)
    
    # Locate where the second cover page starts and delete its paragraphs
    # Cover block has institution lines. Find the second occurrence and clear it
    found_first_inst = False
    for p in d.paragraphs:
        txt_u = p.text.upper()
        if "TÀI NGUYÊN VÀ MÔI TRƯỜNG" in txt_u or "CÔNG NGHỆ THÔNG TIN" in txt_u or "BÁO CÁO ĐỒ ÁN" in txt_u:
            if not found_first_inst:
                found_first_inst = True
            else:
                p.text = ""
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues if i.severity == "error"]
    assert "missing_inner_cover" in codes
