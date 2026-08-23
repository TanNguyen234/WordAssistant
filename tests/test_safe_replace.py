from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest
from wordassistant.document.edit import replace_text, ReplaceResult

def test_replace_preserves_single_run_formatting(tmp_path):
    doc_path = tmp_path / 'styled.docx'
    d = Document()
    p = d.add_paragraph()
    r1 = p.add_run("Normal prefix. ")
    r2 = p.add_run("BOLD AND RED TARGET")
    r2.bold = True
    r2.font.color.rgb = RGBColor(255, 0, 0)
    r3 = p.add_run(" Normal suffix.")
    d.save(doc_path)
    
    out_path = tmp_path / 'styled_out.docx'
    res = replace_text(doc_path, old="BOLD AND RED TARGET", new="NEW REPLACEMENT", output_path=out_path)
    assert res.replacements == 1
    assert res.status == 'SUCCESS'
    
    d2 = Document(out_path)
    p2 = d2.paragraphs[0]
    assert len(p2.runs) == 3
    assert p2.runs[1].text == "NEW REPLACEMENT"
    assert p2.runs[1].bold is True
    assert p2.runs[1].font.color.rgb == RGBColor(255, 0, 0)
    assert p2.runs[0].text == "Normal prefix. "
    assert p2.runs[2].text == " Normal suffix."

def test_replace_dry_run_does_not_modify_disk(tmp_path):
    doc_path = tmp_path / 'doc.docx'
    d = Document()
    d.add_paragraph("Hello world")
    d.save(doc_path)
    
    out_path = tmp_path / 'doc_edited.docx'
    res = replace_text(doc_path, old="world", new="universe", output_path=out_path, dry_run=True)
    assert res.replacements == 1
    assert res.dry_run is True
    assert res.status == 'DRY_RUN'
    assert not out_path.exists()

def test_replace_detects_complex_structures_and_blocks_mutation(tmp_path):
    doc_path = tmp_path / 'complex.docx'
    d = Document()
    p = d.add_paragraph("Section title with field ")
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p._p.append(fld)
    p.add_run(" and more text target_term.")
    d.save(doc_path)
    
    out_path = tmp_path / 'complex_out.docx'
    res = replace_text(doc_path, old="target_term", new="replaced_term", output_path=out_path)
    assert res.status == 'UNSAFE_EDIT_BLOCKED'
    assert res.advanced_backend_required is True
    assert res.replacements == 0
    assert not out_path.exists()
    assert 'field' in res.complex_structures_detected

def test_atomic_file_write_on_replace(tmp_path):
    doc_path = tmp_path / 'atomic.docx'
    d = Document()
    d.add_paragraph("Original text")
    d.save(doc_path)
    
    out_path = tmp_path / 'atomic_out.docx'
    res = replace_text(doc_path, old="Original", new="Atomic", output_path=out_path)
    assert res.status == 'SUCCESS'
    assert res.replacements == 1
    assert out_path.exists()
    d_out = Document(out_path)
    assert "Atomic text" in d_out.paragraphs[0].text
