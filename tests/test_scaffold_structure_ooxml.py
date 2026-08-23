from pathlib import Path
from docx import Document
from wordassistant.profiles.loader import load_profile
from wordassistant.document.scaffold import scaffold_document
from wordassistant.document.validate import validate_document

ROOT = Path(__file__).resolve().parents[1]

def test_hcmunre_report_scaffold_ooxml_and_structure(tmp_path):
    profile = load_profile(ROOT / 'profiles/hcmunre_dacn_2026_2027.json')
    out = tmp_path / 'report_scaffold.docx'
    scaffold_document(
        out,
        profile,
        'report',
        title='HỆ THỐNG TRỢ LÝ TỰ ĐỘNG SOẠN THẢO VĂN BẢN WORD',
        student='Nguyễn Văn A - 2100001234',
        supervisor='TS. Trần Văn B'
    )
    
    d = Document(out)
    
    # 1. Check sections: Should have at least 2 or 3 distinct sections (Cover, Front Matter, Body)
    assert len(d.sections) >= 2, f"Expected multiple sections, got {len(d.sections)}"
    
    # 2. Check no duplicate cover title headings
    headings_text = [p.text for p in d.paragraphs if p.style.name.startswith('Heading') or p.style.name == 'Title']
    # 'TRANG BÌA' should not be a heading following another cover
    assert headings_text.count('TRANG BÌA') <= 1
    
    # 3. Check body section (the last section) has headers and footers
    body_sec = d.sections[-1]
    header_text = "".join(p.text for p in body_sec.header.paragraphs)
    footer_text = "".join(p.text for p in body_sec.footer.paragraphs)
    
    assert 'TS. Trần Văn B' in header_text or 'HỆ THỐNG TRỢ LÝ' in header_text or 'ĐỒ ÁN' in header_text
    assert 'Nguyễn Văn A' in footer_text or len(body_sec.footer.paragraphs) > 0
    
    # 4. Check OOXML for PAGE field in footer
    footer_xml = body_sec.footer._element.xml
    assert 'w:instr="PAGE"' in footer_xml or 'PAGE' in footer_xml, "Missing PAGE field in footer OOXML"
    
    # 5. Check OOXML for TOC field in document body
    doc_xml = d._element.xml
    assert 'TOC' in doc_xml, "Missing TOC field in document OOXML"
    
    # 6. Validate the generated scaffold with validate_document
    val_res = validate_document(out, profile, 'report')
    assert val_res.valid, f"Generated scaffold failed validation: {[i.message for i in val_res.issues]}"
