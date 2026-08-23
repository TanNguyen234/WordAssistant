from __future__ import annotations
import json
from pathlib import Path
from docx import Document
from wordassistant.document.validate import validate_document
from wordassistant.profiles.loader import load_profile
from wordassistant.document.scaffold import scaffold_document

def test_optional_section_blueprint_missing_does_not_fail(tmp_path: Path):
    custom_prof = {
        "id": "optional_sec_prof",
        "name": "Optional Section Profile",
        "version": 1,
        "variants": {
            "default": {
                "page": {
                    "size": "A4",
                    "width_cm": 21.0,
                    "height_cm": 29.7,
                    "margins_cm": {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0},
                    "one_sided": True
                },
                "typography": {
                    "body_font": "Arial",
                    "body_size_pt": 11.0
                },
                "section_blueprint": [
                    {"title": "Required Chapter", "level": 1, "required": True},
                    {"title": "Optional Appendix", "level": 1, "required": False}
                ]
            }
        }
    }
    prof_file = tmp_path / "custom.json"
    prof_file.write_text(json.dumps(custom_prof), encoding="utf-8")
    prof = load_profile(prof_file)
    
    doc_path = tmp_path / "opt_doc.docx"
    scaffold_document(doc_path, prof, "default", title="Test Doc")
    d = Document(doc_path)
    # Remove all placeholders
    for p in d.paragraphs:
        if "[Document content]" in p.text or "[Section content...]" in p.text:
            p.text = "Valid body content."
    # Remove optional appendix heading if present
    for p in d.paragraphs:
        if "Optional Appendix" in p.text:
            p.text = ""
    d.save(doc_path)
    
    val_res = validate_document(doc_path, prof, "default", verify_mode=True)
    # Missing optional section must NOT produce missing_blueprint_sections error
    errors = [i for i in val_res.issues if i.severity == "error"]
    assert len(errors) == 0

def test_field_requirements_boolean_false_does_not_require_toc(tmp_path: Path):
    custom_prof = {
        "id": "no_toc_prof",
        "name": "No TOC Profile",
        "version": 1,
        "variants": {
            "default": {
                "page": {
                    "size": "A4",
                    "width_cm": 21.0,
                    "height_cm": 29.7,
                    "margins_cm": {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0},
                    "one_sided": True
                },
                "typography": {
                    "body_font": "Arial",
                    "body_size_pt": 11.0
                },
                "field_requirements": {
                    "table_of_contents": False
                }
            }
        }
    }
    prof_file = tmp_path / "custom_notoc.json"
    prof_file.write_text(json.dumps(custom_prof), encoding="utf-8")
    prof = load_profile(prof_file)
    
    doc_path = tmp_path / "notoc_doc.docx"
    scaffold_document(doc_path, prof, "default", title="Test Doc")
    d = Document(doc_path)
    for p in d.paragraphs:
        if "[Document content]" in p.text or "[Section content...]" in p.text:
            p.text = "Valid body content."
    d.save(doc_path)
    
    val_res = validate_document(doc_path, prof, "default", verify_mode=True)
    issues = [i for i in val_res.issues if i.code == "missing_toc_field"]
    assert len(issues) == 0

def test_page_number_start_restart_enforcement(tmp_path: Path):
    custom_prof = {
        "id": "page_start_prof",
        "name": "Page Start Profile",
        "version": 1,
        "variants": {
            "default": {
                "page": {
                    "size": "A4",
                    "width_cm": 21.0,
                    "height_cm": 29.7,
                    "margins_cm": {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0},
                    "one_sided": True
                },
                "typography": {
                    "body_font": "Arial",
                    "body_size_pt": 11.0
                },
                "header_footer": {
                    "footer_right": "page_number",
                    "page_number_start": 1
                },
                "front_matter": ["cover"],
                "section_blueprint": [
                    {"title": "Introduction", "level": 1, "required": True}
                ]
            }
        }
    }
    prof_file = tmp_path / "custom_pstart.json"
    prof_file.write_text(json.dumps(custom_prof), encoding="utf-8")
    prof = load_profile(prof_file)
    
    doc_path = tmp_path / "pstart_doc.docx"
    scaffold_document(doc_path, prof, "default", title="Test Doc")
    d = Document(doc_path)
    for p in d.paragraphs:
        if "[Document content]" in p.text or "[Section content...]" in p.text:
            p.text = "Valid body content."
    d.save(doc_path)
    
    val_res = validate_document(doc_path, prof, "default", verify_mode=True)
    assert val_res.valid is True
