from __future__ import annotations
from pathlib import Path
from docx import Document
from wordassistant.document.structure import inspect_document
from wordassistant.document.scaffold import scaffold_document
from wordassistant.profiles.loader import resolve_profile

def test_literal_page_text_does_not_trigger_page_field(tmp_path: Path):
    doc_path = tmp_path / "plain_page.docx"
    d = Document()
    d.add_heading("Section 1", 1)
    d.add_paragraph("This paragraph mentions the word PAGE and TOC in plain text.")
    d.save(doc_path)
    
    inv = inspect_document(doc_path)
    assert inv["fields"]["page_field"] is False
    assert inv["fields"]["toc_field"] is False
    assert inv["fields"].get("page_fields", 0) == 0
    assert inv["fields"].get("toc_fields", 0) == 0

def test_hcmunre_report_has_structural_page_and_toc_fields(tmp_path: Path):
    doc_path = tmp_path / "report.docx"
    prof = resolve_profile("hcmunre_dacn_2026_2027")
    scaffold_document(doc_path, prof, "report", title="TEST TITLE", student="STUDENT A", supervisor="SUPERVISOR B")
    
    inv = inspect_document(doc_path)
    assert inv["fields"]["page_field"] is True
    assert inv["fields"]["toc_field"] is True
    assert inv["fields"]["page_fields"] >= 1
    assert inv["fields"]["toc_fields"] >= 1
