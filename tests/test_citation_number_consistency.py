from docx import Document
import pytest

from wordassistant.document.compose import compose_document
from wordassistant.document.scaffold import scaffold_document
from wordassistant.profiles.loader import load_profile
from pathlib import Path
from wordassistant.evidence.registry import EvidenceRegistry
from wordassistant.operations.citations import (
    generate_bibliography, insert_citation, list_citations, remove_citation,
)


def _composed(tmp_path):
    source = tmp_path / 'source.docx'
    document = Document()
    document.add_heading('Body', 1)
    document.add_heading('References', 1)
    document.save(source)
    evidence = tmp_path / 'evidence.json'
    registry = EvidenceRegistry(evidence)
    for key, author, language in [('z', 'Zulu', 'en'), ('b', 'Nguyễn Bình', 'vi'), ('a', 'Nguyễn An', 'vi')]:
        registry.add(dict(id=key, authors=[author], language=language, author_sort_key=author.split()[-1], title=f'Source {key}', year=2024, url=f'https://example.org/{key}', verified_by='fixture', synthetic=False))
    output = tmp_path / 'composed.docx'
    compose_document(source, output_path=output, evidence_registry=evidence, content={
        'version': 1, 'citation_style': 'hcmunre_numeric',
        'bibliography_order': 'vietnamese_given_name_then_english_family_name',
        'bibliography_heading': 'References',
        'sections': [{'heading': 'Body', 'paragraphs': [{'text': 'Claim', 'citations': ['z', 'b', 'a']}]}],
    })
    assert [c['text'] for c in list_citations(output)] == ['[3]', '[2]', '[1]']
    return output, evidence


def _assert_bibliography(path, ids):
    document = Document(path)
    entries = [p for p in document.paragraphs if p.text.startswith('[')]
    assert [p.text.split('Source ')[1][0] for p in entries] == ids
    assert [p.text.split(']')[0] for p in entries] == [f'[{n}' for n in range(1, len(ids) + 1)]
    assert all(any(r.bold for r in p.runs) and any(r.italic for r in p.runs) for p in entries)
    return document


def test_sorted_repeat_insert_and_default_regeneration_preserve_policy(tmp_path):
    source, evidence = _composed(tmp_path)
    inserted = insert_citation(source, output_path=tmp_path / 'inserted.docx', evidence_registry=evidence, evidence_id='z')
    assert [c['text'] for c in list_citations(inserted.output_path)] == ['[3]', '[2]', '[1]', '[3]']
    result = generate_bibliography(inserted.output_path, output_path=tmp_path / 'regen.docx', evidence_registry=evidence)
    document = _assert_bibliography(result.output_path, ['a', 'b', 'z'])
    assert [p.text for p in document.paragraphs].count('Tiếng Việt:') == 1
    assert [p.text for p in document.paragraphs].count('Tiếng Anh:') == 1


def test_sorted_removal_updates_bibliography_and_last_source(tmp_path):
    source, _ = _composed(tmp_path)
    result = remove_citation(source, output_path=tmp_path / 'removed.docx', citation_index=2)
    assert [c['text'] for c in list_citations(result.output_path)] == ['[2]', '[1]']
    _assert_bibliography(result.output_path, ['b', 'z'])
    for n in range(2):
        result = remove_citation(result.output_path, output_path=tmp_path / f'remove-{n}.docx', citation_index=0)
    assert not list_citations(result.output_path)
    assert not any(p.text.startswith('[') or p.text in ('Tiếng Việt:', 'Tiếng Anh:') for p in Document(result.output_path).paragraphs)


def test_compose_again_inherits_policy_and_includes_existing_sources(tmp_path):
    source, evidence = _composed(tmp_path)
    result = compose_document(source, output_path=tmp_path / 'again.docx', evidence_registry=evidence, content={
        'version': 1, 'sections': [{'heading': 'Body', 'paragraphs': [{'text': 'Another claim', 'citations': ['b']}]}],
    })
    assert result['citations'] == 3
    assert [c['text'] for c in list_citations(result['output'])] == ['[2]', '[3]', '[2]', '[1]']
    _assert_bibliography(result['output'], ['a', 'b', 'z'])


def test_conflicting_existing_number_fails_without_output(tmp_path):
    source, evidence = _composed(tmp_path)
    document = Document(source)
    nodes = document.element.body.xpath('.//w:sdt//w:t')
    nodes[0].text = '[1]'
    document.save(source)
    output = tmp_path / 'bad.docx'
    with pytest.raises(ValueError, match='Conflicting citation'):
        insert_citation(source, output_path=output, evidence_registry=evidence, evidence_id='z')
    assert not output.exists()


def test_new_sorted_source_renumbers_existing_citations_and_bibliography(tmp_path):
    source, evidence = _composed(tmp_path)
    EvidenceRegistry(evidence).add(dict(id='c', authors=['Nguyễn Cường'], language='vi', author_sort_key='Cường', title='Source c', year=2024, url='https://example.org/c', verified_by='fixture', synthetic=False))
    result = insert_citation(source, output_path=tmp_path / 'new.docx', evidence_registry=evidence, evidence_id='c')
    assert [c['text'] for c in list_citations(result.output_path)] == ['[4]', '[2]', '[1]', '[3]']
    _assert_bibliography(result.output_path, ['a', 'b', 'c', 'z'])


def test_legacy_sorted_document_requires_explicit_adoption(tmp_path):
    source, evidence = _composed(tmp_path)
    document = Document(source)
    for node in document.settings.element.xpath('./w:docVars'):
        document.settings.element.remove(node)
    document.save(source)
    with pytest.raises(ValueError, match='adopt'):
        insert_citation(source, output_path=tmp_path / 'implicit.docx', evidence_registry=evidence, evidence_id='z')
    with pytest.raises(ValueError, match='adopt'):
        remove_citation(source, output_path=tmp_path / 'removed.docx', citation_index=0)
    result = generate_bibliography(source, output_path=tmp_path / 'adopted.docx', evidence_registry=evidence, style='hcmunre_numeric', heading='References')
    assert [c['text'] for c in list_citations(result.output_path)] == ['[1]', '[2]', '[3]']
    _assert_bibliography(result.output_path, ['z', 'b', 'a'])


def test_explicit_legacy_regeneration_repairs_conflicting_labels(tmp_path):
    source, evidence = _composed(tmp_path)
    document = Document(source)
    for node in document.settings.element.xpath('./w:docVars'):
        document.settings.element.remove(node)
    document.element.body.xpath('.//w:sdt//w:t')[0].text = '[1]'
    document.save(source)
    result = generate_bibliography(source, output_path=tmp_path / 'repaired.docx', evidence_registry=evidence, style='hcmunre_numeric', heading='References')
    assert [c['text'] for c in list_citations(result.output_path)] == ['[1]', '[2]', '[3]']
    _assert_bibliography(result.output_path, ['z', 'b', 'a'])


@pytest.mark.parametrize('operation', ['compose', 'insert'])
def test_source_backed_scaffold_requires_review_without_repeated_profile_flag(tmp_path, operation):
    _, evidence = _composed(tmp_path)
    profile = load_profile(Path(__file__).resolve().parents[1] / 'profiles/hcm-university-dacn-2026-2027.json')
    source = tmp_path / 'strict.docx'
    scaffold_document(source, profile, 'report', title='Profile policy fixture')
    output = tmp_path / 'should-not-exist.docx'
    with pytest.raises(ValueError, match='reviewed'):
        if operation == 'compose':
            compose_document(source, output_path=output, evidence_registry=evidence, content={
                'version': 1, 'sections': [{'heading': 'MỞ ĐẦU', 'paragraphs': [{'text': 'Legacy source claim', 'citations': ['z']}]}],
            })
        else:
            insert_citation(source, output_path=output, evidence_registry=evidence, evidence_id='z')
    assert not output.exists()


@pytest.mark.parametrize('operation', ['regenerate', 'remove'])
def test_reference_mutations_preserve_following_custom_outline_section(tmp_path, operation):
    source, evidence = _composed(tmp_path)
    document = Document(source)
    style = document.styles.add_style('Custom Appendix', 1)
    style.base_style = document.styles['Heading 1']
    document.add_paragraph('Appendix', style=style)
    document.add_paragraph('Appendix content must survive reference edits.')
    document.save(source)
    output = tmp_path / 'bounded-references.docx'
    if operation == 'regenerate':
        generate_bibliography(source, output_path=output, evidence_registry=evidence)
    else:
        remove_citation(source, output_path=output, citation_index=0)
    text = [p.text for p in Document(output).paragraphs]
    assert 'Appendix' in text
    assert 'Appendix content must survive reference edits.' in text


def test_source_backed_scaffold_allows_uncited_content_without_registry(tmp_path):
    profile = load_profile(Path(__file__).resolve().parents[1] / 'profiles/hcm-university-dacn-2026-2027.json')
    source = tmp_path / 'strict.docx'
    scaffold_document(source, profile, 'report', title='Profile policy fixture')
    output = tmp_path / 'authored.docx'
    result = compose_document(source, output_path=output, content={
        'version': 1, 'sections': [{'heading': 'MỞ ĐẦU', 'paragraphs': ['The project scope is stated here.']}],
    })
    assert result['citations'] == 0
    assert 'The project scope is stated here.' in [p.text for p in Document(output).paragraphs]
    # Source policy remains present for future source-linked content.
    assert 'require_reviewed' in Document(output).settings.element.xml


@pytest.mark.parametrize('anchor', ['References', 'Tiếng Việt:', '[1] Nguyễn An. Source a. 2024. https://example.org/a.'])
def test_citation_anchor_in_managed_bibliography_is_rejected(tmp_path, anchor):
    source, evidence = _composed(tmp_path)
    output = tmp_path / 'invalid-anchor.docx'
    before = source.read_bytes()
    with pytest.raises(ValueError, match='bibliography'):
        insert_citation(source, output_path=output, evidence_registry=evidence,
                        evidence_id='z', selector={'type': 'text', 'text': anchor})
    assert not output.exists()
    assert source.read_bytes() == before


@pytest.mark.parametrize('heading', ['References', 'Tiếng Việt:', 'Reference notes'])
@pytest.mark.parametrize('dry_run', [False, True])
def test_composition_target_in_managed_bibliography_is_rejected(tmp_path, heading, dry_run):
    source, evidence = _composed(tmp_path)
    document = Document(source)
    document.add_heading('Reference notes', 2)
    document.save(source)
    before = source.read_bytes()
    output = tmp_path / 'invalid-composition.docx'
    with pytest.raises(ValueError, match='bibliography'):
        compose_document(source, output_path=output, evidence_registry=evidence,
                         dry_run=dry_run, content={
            'version': 1, 'sections': [{'heading': heading, 'paragraphs': [
                {'text': 'This authored content must not disappear.', 'citations': ['z']}
            ]}],
        })
    assert not output.exists()
    assert source.read_bytes() == before


@pytest.mark.parametrize('already_authored', [False, True])
def test_composition_can_create_new_bibliography_heading(tmp_path, already_authored):
    _, evidence = _composed(tmp_path)
    source = tmp_path / 'body-only.docx'
    document = Document()
    document.add_heading('Body', 1)
    document.save(source)
    if already_authored:
        compose_document(source, output_path=source, content={
            'version': 1, 'sections': [{'heading': 'Body', 'paragraphs': ['Initial prose.']}],
        })
    output = tmp_path / 'new-bibliography.docx'
    compose_document(source, output_path=output, evidence_registry=evidence, content={
        'version': 1, 'bibliography_heading': 'References',
        'sections': [{'heading': 'Body', 'paragraphs': [{'text': 'Source-backed claim.', 'citations': ['z']}]}],
    })
    text = [p.text for p in Document(output).paragraphs]
    assert 'References' in text
    assert any(p.startswith('[1]') and 'Source z' in p for p in text)
    assert list_citations(output)[0]['text'] == '[1]'
