from __future__ import annotations

import csv
import hashlib
import json
import sys
import tempfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import check_course_essay_readiness as readiness  # noqa: E402


MATRIX_FIELDS = [
    "requirement_id", "rubric_requirement", "rubric_location", "report_section",
    "evidence_files", "verification", "status", "notes", "waived_by",
]
CLAIM_FIELDS = [
    "claim_id", "report_location", "claim_summary", "evidence_id", "source_file",
    "source_sha256", "source_location", "review_status", "support_confirmed",
    "reviewed_by", "review_date",
]
DISPOSITION_FIELDS = [
    "gate", "finding_index", "decision", "reason", "evidence", "reviewed_by", "review_date",
]
COVER_METADATA_FIELDS = [
    "field", "value", "authority", "authority_type", "status", "reviewed_by", "review_date", "notes",
]
STRUCTURE_PARITY_FIELDS = [
    "reference_id", "reference_level", "reference_heading", "target_heading", "target_level", "target_order", "decision",
    "coverage_note", "evidence_files", "status", "reviewed_by", "review_date",
]
ASSET_PARITY_FIELDS = [
    "reference_id", "item_type", "reference_index", "reference_caption", "reference_description",
    "target_section", "target_caption", "decision", "coverage_note", "evidence_files",
    "status", "reviewed_by", "review_date",
]


def _write_csv(path: Path, fields: list[str], rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _write_docx(path: Path, outline_rows: list[dict[str, str]]) -> None:
    document = Document()
    for row in outline_rows:
        document.add_paragraph(
            row["reference_heading"],
            style=f"Heading {row['reference_level']}",
        )
        document.add_paragraph("Synthetic section body long enough to preserve every heading in the outline.")
    document.save(path)


def _case(root: Path) -> tuple[Path, dict, Path, Path]:
    project_dir = root / "essay"
    project_dir.mkdir()
    qa_dir = project_dir / "qa"
    qa_dir.mkdir()
    (project_dir / "assignment.md").write_text("Assignment", encoding="utf-8")
    (project_dir / "rubric.md").write_text("Rubric", encoding="utf-8")
    source = project_dir / "source.txt"
    source.write_text("Reviewed evidence passage", encoding="utf-8")
    output = project_dir / "output" / "final.docx"
    output.parent.mkdir()
    baseline_rows = readiness._rows(readiness._STRONGWINE_STRUCTURE_BASELINE)
    _write_docx(output, baseline_rows)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    source_digest = hashlib.sha256(source.read_bytes()).hexdigest()

    project = {
        "schema_version": 1,
        "project_title": "Demo essay",
        "institution_profile": None,
        "profile_variant": "default",
        "assignment_file": "assignment.md",
        "assignment_summary": "",
        "rubric_file": "rubric.md",
        "source_project_root": str(project_dir),
        "output_docx": "output/final.docx",
        "cover_metadata_file": "cover-metadata.csv",
        "structure_parity_matrix_file": "structure-parity.csv",
        "asset_parity_matrix_file": "asset-parity.csv",
        "requirements_matrix_file": "requirements-matrix.csv",
        "claim_evidence_file": "claim-evidence.csv",
        "evidence_registry_file": "evidence.json",
        "integrity_warning_dispositions_file": "integrity-warning-dispositions.csv",
        "qa_reports": {
            "profile_validation": "qa/profile-validation.json",
            "word_verify": "qa/word-verify.json",
            "integrity_audit": "qa/integrity-audit.json",
        },
        "required_acceptance_checks": [
            "assignment_and_rubric_checked",
            "cover_metadata_fields_verified",
            "strongwine_structure_parity_reviewed",
            "strongwine_asset_parity_reviewed",
            "all_required_claims_mapped_to_reviewed_evidence",
            "all_integrity_warnings_reviewed_and_dispositioned",
            "source_routes_and_templates_checked",
            "wireframes_and_runtime_captures_distinguished",
            "figure_list_rebuilt_from_captions",
            "word_fields_updated_and_saved",
            "all_rendered_pages_visually_reviewed",
        ],
        "acceptance_checks": {
            "assignment_and_rubric_checked": True,
            "cover_metadata_fields_verified": True,
            "strongwine_structure_parity_reviewed": True,
            "strongwine_asset_parity_reviewed": True,
            "all_required_claims_mapped_to_reviewed_evidence": True,
            "all_integrity_warnings_reviewed_and_dispositioned": True,
            "source_routes_and_templates_checked": True,
            "wireframes_and_runtime_captures_distinguished": True,
            "figure_list_rebuilt_from_captions": True,
            "word_fields_updated_and_saved": True,
            "all_rendered_pages_visually_reviewed": True,
        },
    }
    (project_dir / "project.json").write_text(json.dumps(project), encoding="utf-8")
    _write_csv(
        project_dir / "requirements-matrix.csv",
        MATRIX_FIELDS,
        [{
            "requirement_id": "R1", "rubric_requirement": "Discuss the system",
            "rubric_location": "rubric.md", "report_section": "1",
        "evidence_files": "source.txt", "verification": "Opened source and checked passage",
        "status": "PASS", "notes": "", "waived_by": "",
        }],
    )
    metadata_rows = [{
        "field": field, "value": "", "authority": "", "authority_type": "NOT_PROVIDED",
        "status": "PLACEHOLDER", "reviewed_by": "Reviewer", "review_date": "2026-10-20",
        "notes": "Not supplied; leave blank in the document.",
    } for field in readiness._COVER_METADATA_FIELDS]
    _write_csv(project_dir / "cover-metadata.csv", COVER_METADATA_FIELDS, metadata_rows)
    baseline_rows = readiness._rows(readiness._STRONGWINE_STRUCTURE_BASELINE)
    parity_rows = [{
        "reference_id": row["reference_id"],
        "reference_level": row["reference_level"],
        "reference_heading": row["reference_heading"],
        "target_heading": row["reference_heading"],
        "target_level": row["reference_level"],
        "target_order": str(index),
        "decision": "MATCHED",
        "coverage_note": "Synthetic mapping for readiness test fixture.",
        "evidence_files": "source.txt",
        "status": "PASS",
        "reviewed_by": "Reviewer",
        "review_date": "2026-10-20",
    } for index, row in enumerate(baseline_rows, start=1)]
    _write_csv(project_dir / "structure-parity.csv", STRUCTURE_PARITY_FIELDS, parity_rows)
    asset_baseline_rows = readiness._rows(readiness._STRONGWINE_ASSET_BASELINE)
    asset_parity_rows = [{
        "reference_id": row["reference_id"],
        "item_type": row["item_type"],
        "reference_index": row["reference_index"],
        "reference_caption": row["reference_caption"],
        "reference_description": row["reference_description"],
        "target_section": "Synthetic test section",
        "target_caption": f"Synthetic {row['item_type'].lower()} target",
        "decision": "ADAPTED",
        "coverage_note": "Synthetic mapping for readiness test fixture.",
        "evidence_files": "source.txt",
        "status": "PASS",
        "reviewed_by": "Reviewer",
        "review_date": "2026-10-20",
    } for row in asset_baseline_rows]
    _write_csv(project_dir / "asset-parity.csv", ASSET_PARITY_FIELDS, asset_parity_rows)
    _write_csv(
        project_dir / "claim-evidence.csv",
        CLAIM_FIELDS,
        [{
            "claim_id": "C1", "report_location": "Section 1", "claim_summary": "Verified claim",
            "evidence_id": "S1", "source_file": "source.txt", "source_sha256": source_digest,
            "source_location": "paragraph 1", "review_status": "reviewed", "support_confirmed": "yes",
            "reviewed_by": "Reviewer", "review_date": "2026-10-20",
        }],
    )
    (project_dir / "evidence.json").write_text(json.dumps({
        "S1": {
            "id": "S1", "review_status": "reviewed", "source_sha256": source_digest,
            "support_location": "paragraph 1",
        }
    }), encoding="utf-8")
    _write_csv(project_dir / "integrity-warning-dispositions.csv", DISPOSITION_FIELDS, [])
    registry_digest = hashlib.sha256((project_dir / "evidence.json").read_bytes()).hexdigest()
    (qa_dir / "word-verify.json").write_text(json.dumps({
        "overall_status": "PASS", "delivery_status": "PASS", "render_status": "PASS", "visual_status": "PASS",
        "render": {"document_sha256": digest},
    }), encoding="utf-8")
    (qa_dir / "integrity-audit.json").write_text(json.dumps({
        "status": "PASS", "document": {"sha256": digest}, "gates": {
            "source_evidence": {
                "status": "PASS", "finding_count": 0, "findings": [],
                "linked_evidence_id_count": 1, "verified_record_count": 1,
                "registry": {"sha256": registry_digest},
            },
            "unresolved_placeholders": {"status": "PASS", "finding_count": 0, "findings": []},
        },
    }), encoding="utf-8")
    return project_dir, project, output, source


def _save_project(project_dir: Path, project: dict) -> None:
    (project_dir / "project.json").write_text(json.dumps(project), encoding="utf-8")


def test_all_required_evidence_and_word_gates_produce_ready_status() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))

        report = readiness.assess(project_dir)

    assert report["status"] == "PASS"
    assert report["gates"]["word_delivery"]["status"] == "PASS"
    assert report["gates"]["content_and_sources"]["status"] == "PASS"
    assert report["gates"]["submission_readiness"]["status"] == "PASS"


def test_structure_parity_requires_every_canonical_heading() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "structure-parity.csv")
        _write_csv(project_dir / "structure-parity.csv", STRUCTURE_PARITY_FIELDS, rows[:-1])

        report = readiness.assess(project_dir)

    parity = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "strongwine_structure_parity"
    )
    assert parity["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_structure_parity_rejects_omitted_or_not_applicable_sections() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "structure-parity.csv")
        rows[0]["decision"] = "OMIT"
        _write_csv(project_dir / "structure-parity.csv", STRUCTURE_PARITY_FIELDS, rows)

        report = readiness.assess(project_dir)

    parity = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "strongwine_structure_parity"
    )
    assert parity["status"] == "FAIL"
    assert report["gates"]["submission_readiness"]["status"] == "FAIL"


def test_structure_parity_rejects_changed_heading_level() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "structure-parity.csv")
        rows[0]["target_level"] = "2"
        _write_csv(project_dir / "structure-parity.csv", STRUCTURE_PARITY_FIELDS, rows)

        report = readiness.assess(project_dir)

    parity = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "strongwine_structure_parity"
    )
    assert parity["status"] == "FAIL"


def test_structure_parity_rejects_reordered_headings() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "structure-parity.csv")
        rows[0]["target_order"] = "2"
        _write_csv(project_dir / "structure-parity.csv", STRUCTURE_PARITY_FIELDS, rows)

        report = readiness.assess(project_dir)

    parity = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "strongwine_structure_parity"
    )
    assert parity["status"] == "FAIL"


def test_structure_parity_checks_the_actual_docx_outline() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, output, _source = _case(Path(temp))
        baseline_rows = readiness._rows(readiness._STRONGWINE_STRUCTURE_BASELINE)
        _write_docx(output, baseline_rows[:-1])

        report = readiness.assess(project_dir)

    outline_check = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "strongwine_structure_in_docx"
    )
    assert outline_check["status"] == "FAIL"


def test_structure_parity_rejects_reordered_headings_in_the_actual_docx() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, output, _source = _case(Path(temp))
        baseline_rows = readiness._rows(readiness._STRONGWINE_STRUCTURE_BASELINE)
        baseline_rows[0], baseline_rows[1] = baseline_rows[1], baseline_rows[0]
        _write_docx(output, baseline_rows)

        report = readiness.assess(project_dir)

    outline_check = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "strongwine_structure_in_docx"
    )
    assert outline_check["status"] == "FAIL"


def test_structure_parity_needs_existing_evidence_for_each_mapped_heading() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "structure-parity.csv")
        rows[0]["evidence_files"] = "missing-evidence.md"
        _write_csv(project_dir / "structure-parity.csv", STRUCTURE_PARITY_FIELDS, rows)

        report = readiness.assess(project_dir)

    parity = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "strongwine_structure_parity"
    )
    assert parity["status"] == "INCOMPLETE"


def test_asset_parity_requires_every_reference_image_and_table_slot() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "asset-parity.csv")
        _write_csv(project_dir / "asset-parity.csv", ASSET_PARITY_FIELDS, rows[:-1])

        report = readiness.assess(project_dir)

    parity = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "strongwine_asset_parity"
    )
    assert parity["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_asset_parity_rejects_dropped_visual_slots() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "asset-parity.csv")
        rows[0]["decision"] = "NOT_APPLICABLE"
        _write_csv(project_dir / "asset-parity.csv", ASSET_PARITY_FIELDS, rows)

        report = readiness.assess(project_dir)

    parity = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "strongwine_asset_parity"
    )
    assert parity["status"] == "FAIL"


def test_asset_parity_needs_evidence_for_each_replacement_artifact() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "asset-parity.csv")
        rows[0]["evidence_files"] = "missing-image-source.png"
        _write_csv(project_dir / "asset-parity.csv", ASSET_PARITY_FIELDS, rows)

        report = readiness.assess(project_dir)

    parity = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "strongwine_asset_parity"
    )
    assert parity["status"] == "INCOMPLETE"


def test_cover_metadata_cannot_use_the_reference_sample_as_authority() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "cover-metadata.csv")
        supervisor = next(row for row in rows if row["field"] == "supervisor_name")
        supervisor.update({
            "value": "Sample Supervisor", "authority": "StrongWine sample",
            "authority_type": "REFERENCE_SAMPLE", "status": "VERIFIED",
        })
        _write_csv(project_dir / "cover-metadata.csv", COVER_METADATA_FIELDS, rows)

        report = readiness.assess(project_dir)

    metadata = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "cover_metadata"
    )
    assert metadata["status"] == "FAIL"
    assert report["gates"]["submission_readiness"]["status"] == "FAIL"


def test_cover_metadata_requires_every_identity_field_to_be_resolved() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "cover-metadata.csv")
        rows = [row for row in rows if row["field"] != "supervisor_name"]
        _write_csv(project_dir / "cover-metadata.csv", COVER_METADATA_FIELDS, rows)

        report = readiness.assess(project_dir)

    metadata = next(
        item for item in report["gates"]["content_and_sources"]["checks"]
        if item["check"] == "cover_metadata"
    )
    assert metadata["status"] == "INCOMPLETE"


def test_word_delivery_pass_with_missing_rubric_is_not_ready() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, project, _output, _source = _case(Path(temp))
        project["rubric_file"] = ""
        _save_project(project_dir, project)

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "PASS"
    assert report["gates"]["content_and_sources"]["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_rubric_absence_cannot_be_waived_into_submission_readiness() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, project, _output, _source = _case(Path(temp))
        project["rubric_file"] = ""
        project["rubric_waiver"] = {"authorized_by": "Instructor", "reason": "No rubric supplied"}
        _save_project(project_dir, project)

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "PASS"
    assert report["gates"]["content_and_sources"]["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_all_core_human_acceptance_checks_are_required_even_if_config_is_shortened() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, project, _output, _source = _case(Path(temp))
        project["required_acceptance_checks"] = ["assignment_and_rubric_checked"]
        project["acceptance_checks"] = {"assignment_and_rubric_checked": True}
        _save_project(project_dir, project)

        report = readiness.assess(project_dir)

    assert report["gates"]["content_and_sources"]["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_unreviewed_claim_evidence_blocks_content_readiness() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        rows = readiness._rows(project_dir / "claim-evidence.csv")
        rows[0]["review_status"] = "candidate"
        _write_csv(project_dir / "claim-evidence.csv", CLAIM_FIELDS, rows)

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "PASS"
    assert report["gates"]["content_and_sources"]["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_integrity_warnings_need_one_complete_disposition_each() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        audit_path = project_dir / "qa" / "integrity-audit.json"
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        audit["status"] = "WARN"
        audit["gates"]["style_review"] = {
            "status": "WARN", "finding_count": 1, "findings": [{"kind": "repeated_phrase"}],
        }
        audit_path.write_text(json.dumps(audit), encoding="utf-8")

        incomplete = readiness.assess(project_dir)
        _write_csv(project_dir / "integrity-warning-dispositions.csv", DISPOSITION_FIELDS, [{
            "gate": "style_review", "finding_index": "1", "decision": "not_applicable",
            "reason": "Reviewed as a section title repeated in the contents page.",
            "evidence": "qa/integrity-audit.json:style_review:1", "reviewed_by": "Reviewer",
            "review_date": "2026-10-20",
        }])
        complete = readiness.assess(project_dir)

    assert incomplete["gates"]["submission_readiness"]["status"] == "INCOMPLETE"
    assert complete["gates"]["submission_readiness"]["status"] == "PASS"


def test_changed_docx_hash_fails_delivery_and_stale_audit_binding() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, output, _source = _case(Path(temp))
        output.write_bytes(b"changed after QA")

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "FAIL"
    assert report["gates"]["content_and_sources"]["status"] == "FAIL"
    assert report["gates"]["submission_readiness"]["status"] == "FAIL"


def test_reused_render_manifest_binds_word_delivery_to_the_docx_hash() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, output, _source = _case(Path(temp))
        verify_path = project_dir / "qa" / "word-verify.json"
        verify = json.loads(verify_path.read_text(encoding="utf-8"))
        document_digest = hashlib.sha256(output.read_bytes()).hexdigest()
        verify["render"].pop("document_sha256")
        verify["render_reuse_manifest"] = {"document_sha256": document_digest}
        verify_path.write_text(json.dumps(verify), encoding="utf-8")

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "PASS"
    assert report["gates"]["content_and_sources"]["status"] == "PASS"


def test_selected_profile_must_have_a_passing_format_gate() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, project, _output, _source = _case(Path(temp))
        project["institution_profile"] = "course-profile"
        _save_project(project_dir, project)

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_selected_profile_requires_matching_saved_validate_report() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, project, _output, _source = _case(Path(temp))
        project["institution_profile"] = "course-profile"
        _save_project(project_dir, project)
        verify_path = project_dir / "qa" / "word-verify.json"
        verify = json.loads(verify_path.read_text(encoding="utf-8"))
        verify["format_status"] = "PASS"
        verify_path.write_text(json.dumps(verify), encoding="utf-8")
        (project_dir / "qa" / "profile-validation.json").write_text(json.dumps({
            "file": "output/final.docx", "profile": "course-profile", "variant": "default", "valid": True,
        }), encoding="utf-8")

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "PASS"
    assert report["gates"]["submission_readiness"]["status"] == "PASS"


def test_vacuous_generic_source_gate_does_not_cover_claim_evidence_ids() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        audit_path = project_dir / "qa" / "integrity-audit.json"
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        audit["gates"]["source_evidence"]["linked_evidence_id_count"] = 0
        audit["gates"]["source_evidence"]["verified_record_count"] = 0
        audit_path.write_text(json.dumps(audit), encoding="utf-8")

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "PASS"
    assert report["gates"]["content_and_sources"]["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_report_citation_registry_is_separate_from_project_claim_registry() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, project, _output, _source = _case(Path(temp))
        citation_registry = project_dir / "report-evidence.json"
        citation_registry.write_text(json.dumps({"[1]": "reviewed paper"}), encoding="utf-8")
        project["report_citation_registry_file"] = "report-evidence.json"
        _save_project(project_dir, project)

        audit_path = project_dir / "qa" / "integrity-audit.json"
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        audit["gates"]["source_evidence"]["linked_evidence_id_count"] = 3
        audit["gates"]["source_evidence"]["verified_record_count"] = 3
        audit["gates"]["source_evidence"]["registry"]["sha256"] = hashlib.sha256(
            citation_registry.read_bytes()
        ).hexdigest()
        audit["gates"]["numeric_citation_bibliography_consistency"] = {
            "status": "PASS",
            "finding_count": 0,
            "findings": [],
            "citation_numbers": [1, 2, 3],
            "bibliography_numbers": [1, 2, 3],
        }
        audit_path.write_text(json.dumps(audit), encoding="utf-8")

        report = readiness.assess(project_dir)

    assert report["gates"]["content_and_sources"]["status"] == "PASS"
    assert report["gates"]["submission_readiness"]["status"] == "PASS"
    checks = report["gates"]["content_and_sources"]["checks"]
    registry_check = next(item for item in checks if item["check"] == "integrity_registry_hash")
    source_check = next(item for item in checks if item["check"] == "integrity_source_evidence")
    assert registry_check["status"] == "PASS"
    assert "3 unique in-text citation numbers" in source_check["detail"]


def test_integrity_evidence_counts_must_be_integers_not_booleans() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        audit_path = project_dir / "qa" / "integrity-audit.json"
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        source_gate = audit["gates"]["source_evidence"]
        source_gate["linked_evidence_id_count"] = True
        source_gate["verified_record_count"] = True
        audit_path.write_text(json.dumps(audit), encoding="utf-8")

        report = readiness.assess(project_dir)

    assert report["gates"]["content_and_sources"]["status"] == "INCOMPLETE"


def test_integrity_warning_count_must_be_an_integer_not_boolean() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        audit_path = project_dir / "qa" / "integrity-audit.json"
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        audit["status"] = "WARN"
        audit["gates"]["style_review"] = {
            "status": "WARN", "finding_count": True, "findings": [{"kind": "review"}],
        }
        audit_path.write_text(json.dumps(audit), encoding="utf-8")
        _write_csv(project_dir / "integrity-warning-dispositions.csv", DISPOSITION_FIELDS, [{
            "gate": "style_review", "finding_index": "1", "decision": "accepted",
            "reason": "Reviewed", "evidence": "qa/integrity-audit.json:style_review:1",
            "reviewed_by": "Reviewer", "review_date": "2026-10-20",
        }])

        report = readiness.assess(project_dir)

    assert report["gates"]["content_and_sources"]["status"] == "FAIL"


def test_integrity_audit_with_unknown_gate_status_is_not_ready() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        audit_path = project_dir / "qa" / "integrity-audit.json"
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        audit["gates"]["runtime_capture"] = {"status": "NOT_RUN"}
        audit_path.write_text(json.dumps(audit), encoding="utf-8")

        report = readiness.assess(project_dir)

    assert report["gates"]["content_and_sources"]["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_claim_review_date_requires_exact_iso_format() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        claims = readiness._rows(project_dir / "claim-evidence.csv")
        claims[0]["review_date"] = "20261020"
        _write_csv(project_dir / "claim-evidence.csv", CLAIM_FIELDS, claims)

        report = readiness.assess(project_dir)

    assert report["gates"]["content_and_sources"]["status"] == "INCOMPLETE"


def test_duplicate_claim_ids_fail_claim_mapping() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        claims = readiness._rows(project_dir / "claim-evidence.csv")
        duplicate = dict(claims[0], evidence_id="S1")
        claims.append(duplicate)
        _write_csv(project_dir / "claim-evidence.csv", CLAIM_FIELDS, claims)

        report = readiness.assess(project_dir)

    assert report["gates"]["content_and_sources"]["status"] == "FAIL"
    assert report["gates"]["submission_readiness"]["status"] == "FAIL"


def test_claim_id_must_exist_as_the_same_id_in_reviewed_registry() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, source = _case(Path(temp))
        source_digest = hashlib.sha256(source.read_bytes()).hexdigest()
        (project_dir / "evidence.json").write_text(json.dumps({
            "S2": {
                "id": "S2", "review_status": "reviewed", "source_sha256": source_digest,
                "support_location": "paragraph 1",
            }
        }), encoding="utf-8")
        audit_path = project_dir / "qa" / "integrity-audit.json"
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        registry_digest = hashlib.sha256((project_dir / "evidence.json").read_bytes()).hexdigest()
        audit["gates"]["source_evidence"]["registry"]["sha256"] = registry_digest
        audit_path.write_text(json.dumps(audit), encoding="utf-8")

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "PASS"
    assert report["gates"]["content_and_sources"]["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_evidence_registry_change_after_audit_fails_hash_binding() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        registry_path = project_dir / "evidence.json"
        registry = json.loads(registry_path.read_text(encoding="utf-8"))
        registry["S2"] = dict(registry["S1"], id="S2")
        registry_path.write_text(json.dumps(registry), encoding="utf-8")

        report = readiness.assess(project_dir)

    assert report["gates"]["content_and_sources"]["status"] == "FAIL"
    assert report["gates"]["submission_readiness"]["status"] == "FAIL"


def test_requirement_pass_needs_existing_evidence_files() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        matrix = readiness._rows(project_dir / "requirements-matrix.csv")
        matrix[0]["evidence_files"] = "missing-source.txt"
        _write_csv(project_dir / "requirements-matrix.csv", MATRIX_FIELDS, matrix)

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "PASS"
    assert report["gates"]["content_and_sources"]["status"] == "INCOMPLETE"
    assert report["gates"]["submission_readiness"]["status"] == "INCOMPLETE"


def test_duplicate_requirement_ids_fail_the_requirements_matrix() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        matrix = readiness._rows(project_dir / "requirements-matrix.csv")
        matrix.append(dict(matrix[0]))
        _write_csv(project_dir / "requirements-matrix.csv", MATRIX_FIELDS, matrix)

        report = readiness.assess(project_dir)

    assert report["gates"]["content_and_sources"]["status"] == "FAIL"
    assert report["gates"]["submission_readiness"]["status"] == "FAIL"


def test_failed_word_delivery_is_never_submission_ready() -> None:
    with tempfile.TemporaryDirectory() as temp:
        project_dir, _project, _output, _source = _case(Path(temp))
        verify_path = project_dir / "qa" / "word-verify.json"
        verify = json.loads(verify_path.read_text(encoding="utf-8"))
        verify["delivery_status"] = "FAIL"
        verify_path.write_text(json.dumps(verify), encoding="utf-8")

        report = readiness.assess(project_dir)

    assert report["gates"]["word_delivery"]["status"] == "FAIL"
    assert report["gates"]["submission_readiness"]["status"] == "FAIL"
