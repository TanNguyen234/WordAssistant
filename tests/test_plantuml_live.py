from __future__ import annotations

import json
from pathlib import Path
import pytest
from PIL import Image
from docx import Document

from wordassistant.diagrams.plantuml import PlantUMLRenderer, doctor, render_plantuml
from wordassistant.cli import run
from wordassistant.operations.images import insert_image

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS_DIR = ROOT / "artifacts" / "plantuml_test_renders"


DIAGRAM_SAMPLES = {
    "sequence": """@startuml
autonumber
actor User
participant "Client App" as App
participant "API Gateway" as Gateway
database "Database" as DB

User -> App: Submit Order
activate App
App -> Gateway: POST /orders
activate Gateway
Gateway -> DB: Insert record
DB --> Gateway: OK (id=42)
Gateway --> App: 201 Created
deactivate Gateway
App --> User: Order Confirmation
deactivate App
@enduml""",

    "class": """@startuml
abstract class Shape {
  # x : int
  # y : int
  + {abstract} area() : double
}
class Circle extends Shape {
  - radius : double
  + area() : double
}
class Rectangle extends Shape {
  - width : double
  - height : double
  + area() : double
}
@enduml""",

    "usecase": """@startuml
left to right direction
actor Student
actor Teacher
rectangle "Learning Management System" {
  usecase "Enroll in Course" as UC1
  usecase "Submit Assignment" as UC2
  usecase "Grade Assignment" as UC3
}
Student --> UC1
Student --> UC2
Teacher --> UC3
@enduml""",

    "activity": """@startuml
start
:Receive Request;
if (Validate Input?) then (yes)
  :Process Payment;
  :Generate Invoice;
else (no)
  :Log Error;
  :Return 400 Bad Request;
  stop
endif
:Send Email Notification;
stop
@enduml""",

    "component": """@startuml
package "Web Tier" {
  [React Frontend] as Frontend
}
package "App Tier" {
  [REST API Service] as API
  [Auth Service] as Auth
}
database "PostgreSQL" as DB
Frontend ..> API : HTTPS / JSON
API ..> Auth : Verify JWT
API --> DB : SQL Queries
@enduml""",

    "state": """@startuml
[*] --> Draft
Draft --> PendingReview : Submit
PendingReview --> Approved : Approve
PendingReview --> Rejected : Reject
Rejected --> Draft : Revise
Approved --> Published : Publish
Published --> [*]
@enduml""",

    "object": """@startuml
object user1 {
  id = 101
  name = "Alice Nguyen"
  role = "Administrator"
}
object session1 {
  token = "xyz-789"
  expiresAt = "2026-12-31"
}
user1 -- session1 : possesses
@enduml""",

    "deployment": """@startuml
node "Kubernetes Cluster" {
  node "Worker Node 1" {
    artifact "WordAssistant Pod" as Pod1
  }
  node "Worker Node 2" {
    artifact "WordAssistant Pod" as Pod2
  }
}
cloud "Cloudflare CDN" as CDN
database "Shared Storage" as Storage
CDN --> Pod1
CDN --> Pod2
Pod1 --> Storage
Pod2 --> Storage
@enduml""",

    "timing": """@startuml
robust "Web Browser" as WB
concise "API Server" as AS

@0
WB is Idle
AS is Idle

@100
WB is Waiting
AS is Processing

@300
WB is Rendering
AS is Idle

@400
WB is Idle
@enduml""",

    "mindmap": """@startmindmap
* WordAssistant System
** Core Engine
*** DOCX Parser
*** Layout Calculator
*** Style Manager
** Diagram Engine
*** PlantUML Standalone
*** Smetana Layout
** Agent Integration
*** Bounded Context
*** Audit & Integrity
@endmindmap""",

    "wbs": """@startwbs
* Project Delivery
** Phase 1: Engine Vendoring
*** Download Jar
*** Wrapper Scripts
** Phase 2: Service Layer
*** Python Renderer
*** Health Doctor
** Phase 3: Live Verification
*** 14+ UML Tests
*** DOCX Embedding Test
@endwbs""",

    "gantt": """@startgantt
[Specification] lasts 5 days
[Implementation] lasts 10 days
[Testing] lasts 4 days
[Specification] -> [Implementation]
[Implementation] -> [Testing]
@endgantt""",

    "erd": """@startuml
entity "Users" as users {
  * id : UUID <<PK>>
  --
  * email : VARCHAR(255)
  * full_name : VARCHAR(100)
  created_at : TIMESTAMP
}
entity "Documents" as docs {
  * id : UUID <<PK>>
  --
  * user_id : UUID <<FK>>
  * title : VARCHAR(200)
  file_path : VARCHAR(500)
  updated_at : TIMESTAMP
}
users ||..o{ docs : owns
@enduml""",

    "json": """@startjson
{
  "system": "WordAssistant",
  "version": "2.0.0",
  "features": ["DOCX", "PlantUML", "Integrity", "Academic"],
  "engine": {
    "type": "Standalone PlantUML",
    "format": "PNG",
    "status": "OPERATIONAL"
  }
}
@endjson""",

    "yaml": """@startyaml
name: WordAssistant
architecture:
  layers:
    - presentation: CLI
    - business: Document Engine
    - diagram: PlantUML
active: true
@endyaml""",

    "salt": """@startsalt
{+
  WordAssistant Export Settings
  [X] Render PlantUML diagrams
  () PNG | () SVG | () PDF
  [ Browse... ] | "output/report.docx"
  [  Generate DOCX  ] | [ Cancel ]
}
@endsalt""",

    "nwdiag": """@startnwdiag
nwdiag {
  network internal {
    address = "192.168.1.x/24"
    web01 [address = "192.168.1.10"];
    db01 [address = "192.168.1.20"];
  }
  network dmz {
    address = "10.0.0.x/24"
    web01 [address = "10.0.0.10"];
    lb01 [address = "10.0.0.1"];
  }
}
@endnwdiag""",
}


def test_plantuml_doctor_detects_java_and_engine():
    res = doctor()
    assert res["status"] == "SUCCESS"
    assert res["operational"] is True
    assert res["java_executable"] is not None
    assert "java version" in res["java_version"].lower() or "openjdk" in res["java_version"].lower()
    assert res["plantuml_jar"] is not None
    assert "PlantUML version" in res["plantuml_version"]


def test_plantuml_cli_doctor(capsys):
    code = run(["plantuml", "doctor", "--json"])
    assert code == 0
    out = capsys.readouterr().out
    data = json.loads(out)
    assert data["status"] == "SUCCESS"
    assert data["operational"] is True


@pytest.mark.parametrize("diagram_name,diagram_source", list(DIAGRAM_SAMPLES.items()))
def test_live_render_all_diagram_types(tmp_path: Path, diagram_name: str, diagram_source: str):
    """
    Renders all 17 UML and specialized diagram types to PNG,
    verifying clean exit, non-empty file, valid PNG magic bytes,
    and positive dimensions. Also archives a copy to artifacts/
    for live visual inspection.
    """
    ARTIFACTS_DIR.mkdir(parents=True, exist_ok=True)
    out_file = tmp_path / f"{diagram_name}.png"
    persistent_copy = ARTIFACTS_DIR / f"{diagram_name}.png"

    result = render_plantuml(diagram_source, output_path=out_file, format="png")

    assert result["status"] == "SUCCESS"
    assert result["format"] == "png"
    assert out_file.is_file()
    assert result["size_bytes"] > 0
    assert result["width"] > 0
    assert result["height"] > 0

    # Verify PNG magic header
    raw_bytes = out_file.read_bytes()
    assert raw_bytes.startswith(b"\x89PNG\r\n\x1a\n"), f"Invalid PNG magic bytes for {diagram_name}"

    # Verify PIL can open and decode the image
    with Image.open(out_file) as img:
        assert img.format == "PNG"
        assert img.width == result["width"]
        assert img.height == result["height"]

    # Archive to artifacts directory for visual inspection
    persistent_copy.write_bytes(raw_bytes)


def test_plantuml_multi_format_rendering(tmp_path: Path):
    sample = DIAGRAM_SAMPLES["sequence"]

    # SVG format
    svg_out = tmp_path / "diagram.svg"
    svg_res = render_plantuml(sample, output_path=svg_out, format="svg")
    assert svg_res["status"] == "SUCCESS"
    assert svg_out.is_file()
    svg_text = svg_out.read_text(encoding="utf-8")
    assert "<svg" in svg_text

    # TXT (ASCII/Unicode art) format
    txt_out = tmp_path / "diagram.txt"
    txt_res = render_plantuml(sample, output_path=txt_out, format="txt")
    assert txt_res["status"] == "SUCCESS"
    assert txt_out.is_file()
    txt_content = txt_out.read_text(encoding="utf-8")
    assert len(txt_content) > 20


def test_plantuml_cli_render(tmp_path: Path, capsys):
    out_png = tmp_path / "cli_rendered.png"
    code = run([
        "plantuml", "render",
        "--text", "@startuml\nBob -> Alice : Hello\n@enduml",
        "--output", str(out_png),
        "--format", "png",
        "--json",
    ])
    assert code == 0
    out = capsys.readouterr().out
    data = json.loads(out)
    assert data["status"] == "SUCCESS"
    assert out_png.is_file()
    assert out_png.stat().st_size > 500


def test_plantuml_embed_in_docx(tmp_path: Path):
    # 1. Render sequence diagram to PNG
    diagram_png = tmp_path / "architecture_sequence.png"
    render_plantuml(DIAGRAM_SAMPLES["sequence"], output_path=diagram_png)
    assert diagram_png.is_file()

    # 2. Create blank DOCX
    base_docx = tmp_path / "report.docx"
    doc = Document()
    doc.add_heading("Hệ thống xử lý đơn hàng", level=1)
    doc.add_paragraph("Dưới đây là sơ đồ tương tác giữa các thành phần:")
    doc.save(base_docx)

    # 3. Insert diagram into DOCX using WordAssistant operation
    output_docx = tmp_path / "report_with_diagram.docx"
    res = insert_image(
        base_docx,
        output_path=output_docx,
        image_path=diagram_png,
        caption="Sơ đồ tương tác đặt hàng",
        caption_label="Figure",
        title="Sơ đồ tương tác đặt hàng",
    )
    assert res.status == "SUCCESS"
    assert output_docx.is_file()

    # 4. Verify embedded image in DOCX
    result_doc = Document(output_docx)
    # Check paragraphs: caption must be present
    all_text = [p.text for p in result_doc.paragraphs]
    assert any("Sơ đồ tương tác đặt hàng" in t for t in all_text)

    # Check relationships: image must be packaged in docx part
    has_image_rel = any("image" in rel.target_ref.lower() for rel in result_doc.part.rels.values())
    assert has_image_rel is True


def test_plantuml_negative_invalid_syntax(tmp_path: Path):
    invalid_code = "@startuml\nthis is invalid syntax !!! @@@ ::: %%%\n@enduml"
    out_file = tmp_path / "invalid.png"

    with pytest.raises(RuntimeError) as exc_info:
        render_plantuml(invalid_code, output_path=out_file)

    assert "compilation failed" in str(exc_info.value).lower() or "error" in str(exc_info.value).lower()
