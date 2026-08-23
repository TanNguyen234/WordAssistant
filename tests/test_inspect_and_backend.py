from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from wordassistant.document.structure import inspect_document
from wordassistant.backends.status import get_backend_status, advanced_word_backend_available

def test_backend_status_reporting():
    status = get_backend_status()
    assert 'installed' in status
    assert 'adapter_supported' in status
    assert 'operational' in status
    assert 'core_capabilities' in status
    assert 'advanced_capabilities' in status
    assert 'engine' in status
    assert status['engine'] in ('python-docx-core', 'docx-mcp-advanced')

def test_inspect_document_rich_inventory(tmp_path):
    doc_path = tmp_path / 'rich.docx'
    d = Document()
    p1 = d.add_paragraph("Paragraph 1 with ")
    # Add a field
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p1._p.append(fld)
    
    # Add a table
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).text = "Cell 0,0"
    
    d.save(doc_path)
    
    inv = inspect_document(doc_path)
    assert inv['paragraphs'] >= 1
    assert inv['tables'] == 1
    assert inv['sections'] == 1
    assert inv['fields']['page_field'] is True
    assert 'advanced_backend_recommended' in inv
