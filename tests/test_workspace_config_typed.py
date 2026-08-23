from __future__ import annotations
import json
from pathlib import Path
from wordassistant.cli import run

def test_workspace_config_set_and_context_typed(tmp_path: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    
    # wordassistant init .
    assert run(["init", "."]) == 0
    
    # wordassistant config set context_budget 8000
    assert run(["config", "set", "context_budget", "8000"]) == 0
    
    # Check project.json has integer 8000
    cfg = json.loads((tmp_path / ".wordassistant" / "project.json").read_text(encoding="utf-8"))
    assert cfg["context_budget"] == 8000
    assert isinstance(cfg["context_budget"], int)
    
    # wordassistant config set unknown_key foo => MUST FAIL
    assert run(["config", "set", "unknown_key", "foo"]) != 0
    
    # wordassistant config set context_budget 50 => MUST FAIL (< 100)
    assert run(["config", "set", "context_budget", "50"]) != 0

def test_context_command_uses_typed_workspace_budget(tmp_path: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert run(["init", "."]) == 0
    assert run(["config", "set", "context_budget", "8000"]) == 0
    
    # Create a small valid docx
    from docx import Document
    d = Document()
    d.add_heading("Section 1", 1)
    d.add_paragraph("Content of section 1")
    doc_path = tmp_path / "test.docx"
    d.save(doc_path)
    
    # Run wordassistant context
    assert run(["context", str(doc_path), "--query", "section", "--json"]) == 0
