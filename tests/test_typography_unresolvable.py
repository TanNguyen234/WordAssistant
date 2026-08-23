from __future__ import annotations
from pathlib import Path
from docx import Document
from wordassistant.document.validate import validate_document
from wordassistant.profiles.loader import resolve_profile

def test_unresolvable_formatting_produces_warning_not_false_pass(tmp_path: Path):
    doc_path = tmp_path / "unresolvable.docx"
    d = Document()
    s = d.sections[0]
    s.page_width = Document().sections[0].page_width
    s.page_height = Document().sections[0].page_height
    
    # Add a normal paragraph whose run has no font or size specified, and style has no font name set
    p1 = d.add_paragraph("Chapter 1 Title")
    p1.style = d.styles["Heading 1"]
    
    p2 = d.add_paragraph("Body paragraph without explicit font override.")
    # Clear Normal style font name to simulate unresolvable font
    try:
        d.styles["Normal"].font.name = None
        d.styles["Normal"].font.size = None
        d.styles["Normal"].paragraph_format.line_spacing = None
    except Exception:
        pass
    d.save(doc_path)
    
    prof = resolve_profile("generic")
    res = validate_document(doc_path, prof, "default", verify_mode=False)
    # Warnings for unresolvable formatting should be present, not false pass or false error
    warning_codes = [i.code for i in res.issues if i.severity == "warning"]
    assert "unresolvable_formatting" in warning_codes
