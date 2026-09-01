from docx import Document
from pathlib import Path
import pytest

from wordassistant.backends.word_native.backend import WordNativeBackend, _run_native_worker
from wordassistant.operations.models import OperationRequest
from wordassistant.operations.fields import inspect_fields, insert_toc
from wordassistant.execution import ExecutionService
from wordassistant.document.verify import verify_document
from wordassistant.operations.lists import apply_heading_numbering
from wordassistant.operations.page_numbers import configure_page_numbers


def _backend_or_skip():
    backend = WordNativeBackend(timeout_seconds=90)
    health = backend.health()
    if not health.operational:
        pytest.skip(f"Word native unavailable: {health.reason}")
    return backend


def _document(path, text="Alpha beta"):
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    return path


def test_live_comment_create_list_delete(tmp_path):
    backend = _backend_or_skip()
    source = _document(tmp_path / "source.docx")
    created = tmp_path / "commented.docx"
    backend.execute(
        OperationRequest(
            "comment.create",
            input_path=source,
            output_path=created,
            options={"start": 0, "end": 5, "text": "Review this"},
        )
    )

    listed = backend.execute(OperationRequest("comment.list", input_path=created))
    assert listed.data["comments"][0]["text"].strip() == "Review this"

    cleaned = tmp_path / "cleaned.docx"
    backend.execute(
        OperationRequest(
            "comment.delete",
            input_path=created,
            output_path=cleaned,
            options={"index": 1},
        )
    )
    assert backend.execute(OperationRequest("comment.list", input_path=cleaned)).data["comments"] == []


def test_live_revision_toggle_compare_accept_and_native_render(tmp_path):
    backend = _backend_or_skip()
    original = _document(tmp_path / "original.docx", "Original text")
    revised = _document(tmp_path / "revised.docx", "Revised text")

    tracking = tmp_path / "tracking.docx"
    backend.execute(
        OperationRequest("revision.enable", input_path=original, output_path=tracking)
    )
    disabled = tmp_path / "disabled.docx"
    backend.execute(
        OperationRequest("revision.disable", input_path=tracking, output_path=disabled)
    )

    compared = tmp_path / "compared.docx"
    comparison = backend.execute(
        OperationRequest(
            "compare.documents",
            input_path=original,
            output_path=compared,
            options={"other_path": str(revised)},
        )
    )
    assert comparison.data["revisions"] > 0
    accepted = tmp_path / "accepted.docx"
    backend.execute(
        OperationRequest(
            "revision.accept_all", input_path=compared, output_path=accepted
        )
    )
    assert backend.execute(OperationRequest("revision.list", input_path=accepted)).data["revisions"] == []

    pdf = tmp_path / "accepted.pdf"
    backend.execute(
        OperationRequest("render.pdf", input_path=accepted, output_path=pdf)
    )
    assert pdf.read_bytes().startswith(b"%PDF-")


def test_live_field_update_survives_reopen(tmp_path):
    backend = _backend_or_skip()
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("Giới thiệu", level=1)
    document.add_paragraph("Nội dung tiếng Việt phải đi qua worker bằng UTF-8.")
    document.save(source)
    with_toc = insert_toc(source, output_path=tmp_path / "with-toc.docx")
    updated = tmp_path / "updated.docx"

    result = backend.execute(
        OperationRequest(
            "field.update_all",
            input_path=with_toc.output_path,
            output_path=updated,
        )
    )

    reopened = inspect_fields(updated)
    assert result.data["field_results_freshness"] == "UPDATED_NATIVE"
    assert reopened
    assert all(field["cached_result"] for field in reopened)
    assert all("update field" not in field["cached_result"].casefold() for field in reopened)
    assert any("Giới thiệu" in field["cached_result"] for field in reopened)


def test_live_page_and_numpages_are_evaluated_and_persisted(tmp_path):
    _backend_or_skip()
    source = tmp_path / "source.docx"
    document = Document()
    for page in range(3):
        document.add_paragraph(f"Page body {page + 1}")
        if page < 2:
            document.add_page_break()
    document.save(source)
    numbered = configure_page_numbers(
        source,
        output_path=tmp_path / "numbered.docx",
        section_index=0,
        page_x_of_y=True,
    )
    updated = tmp_path / "updated.docx"
    result = ExecutionService.default().execute(
        OperationRequest(
            "field.update_all",
            input_path=numbered.output_path,
            output_path=updated,
        )
    )

    native_values = {
        field["instruction"].split()[0]: field["result"]
        for field in result.data["fields"]
    }
    assert int(native_values["NUMPAGES"]) >= 3
    report = verify_document(updated, native_fields=True)
    assert report["field_freshness_status"] == "PASS"
    assert report["field_evaluation_status"] == "PASS"
    assert report["native_word_status"] == "PASS"
    assert report["overall_status"] == "PASS"
    verified_values = {
        field["instruction"].split()[0]: field["result"]
        for field in report["native_field_results"]
    }
    assert int(verified_values["NUMPAGES"]) >= 3


def test_live_heading_numbering_is_recognized_by_word_list_engine(tmp_path):
    _backend_or_skip()
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("One", level=1)
    document.add_heading("One One", level=2)
    document.add_heading("Two", level=1)
    document.save(source)
    numbered = apply_heading_numbering(
        source, output_path=tmp_path / "numbered.docx"
    )

    completed = _run_native_worker(
        {
            "action": "inspect_heading_numbering",
            "input_path": str(numbered.output_path.resolve()),
        },
        90,
    )
    import json

    payload = json.loads(completed.stdout.splitlines()[-1])
    assert completed.returncode == 0, payload
    observed = [
        (item["list_type"], item["list_string"])
        for item in payload["paragraphs"][:3]
    ]

    assert [value for _, value in observed] == ["1", "1.1", "2"]
    assert all(list_type != 0 for list_type, _ in observed)


def test_live_execution_service_stages_native_dry_run(tmp_path):
    _backend_or_skip()
    source = _document(tmp_path / "source.docx")
    original = source.read_bytes()
    output = tmp_path / "commented.docx"

    result = ExecutionService.default().execute(
        OperationRequest(
            "comment.create",
            input_path=source,
            output_path=output,
            options={"start": 0, "end": 5, "text": "review"},
            dry_run=True,
        )
    )

    assert result.status == "DRY_RUN"
    assert result.output_path is None
    assert source.read_bytes() == original
    assert not output.exists()


def test_live_execution_service_persists_field_evidence(tmp_path):
    _backend_or_skip()
    source = _document(tmp_path / "source.docx")
    document = Document(source)
    document.paragraphs[0].style = "Heading 1"
    document.save(source)
    with_toc = insert_toc(source, output_path=tmp_path / "with-toc.docx")
    output = tmp_path / "updated.docx"

    result = ExecutionService.default().execute(
        OperationRequest(
            "field.update_all",
            input_path=with_toc.output_path,
            output_path=output,
        )
    )

    assert Path(result.data["evidence_path"]).is_file()
    report = verify_document(output, native_fields=True)
    assert report["field_freshness_status"] == "PASS"
    assert report["field_evaluation_status"] == "PASS"
    assert report["native_word_status"] == "PASS"
    assert report["overall_status"] == "PASS"


def test_live_field_update_rejects_locked_dynamic_field_without_publish(tmp_path):
    _backend_or_skip()
    source = _document(tmp_path / "source.docx")
    document = Document(source)
    from wordassistant.operations.fields import append_field

    append_field(document.paragraphs[0], "PAGE", "999", locked=True)
    document.save(source)
    output = tmp_path / "updated.docx"

    with pytest.raises((ValueError, RuntimeError), match="locked|valid"):
        ExecutionService.default().execute(
            OperationRequest(
                "field.update_all",
                input_path=source,
                output_path=output,
            )
        )

    assert not output.exists()
