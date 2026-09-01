from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wordassistant.document.render import render_pdf_pages
from wordassistant.document.verify import verify_document
from wordassistant.operations.content_controls import create_content_control
from wordassistant.operations.fields import append_field, inspect_fields, remove_fields
from wordassistant.operations.references import (
    create_bookmark,
    delete_bookmark,
    insert_cross_reference,
    insert_hyperlink,
    rename_bookmark,
)
from wordassistant.operations.tables import create_table, set_table_layout
from wordassistant.operations.lists import apply_heading_numbering


def _source(path: Path, text="Target") -> Path:
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    return path


def test_autofit_true_survives_save_reopen(tmp_path):
    source = _source(tmp_path / "source.docx")
    created = create_table(source, output_path=tmp_path / "table.docx", rows=2, columns=2)
    result = set_table_layout(
        created.output_path,
        output_path=tmp_path / "autofit.docx",
        table_index=0,
        autofit=True,
    )

    table = Document(result.output_path).tables[0]
    layout = table._tbl.tblPr.find(qn("w:tblLayout"))
    assert table.autofit is True
    assert layout is None or layout.get(qn("w:type")) == "autofit"


def test_table_width_plus_indent_stays_inside_printable_width(tmp_path):
    source = _source(tmp_path / "source.docx")
    result = create_table(source, output_path=tmp_path / "table.docx", rows=2, columns=3)
    document = Document(result.output_path)
    table = document.tables[0]
    width = int(table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")))
    indent = int(table._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")))
    section = document.sections[-1]
    printable = int(section.page_width.twips - section.left_margin.twips - section.right_margin.twips)
    assert width + indent <= printable


def test_verify_resolves_numbering_in_heading_style(tmp_path):
    source = _source(tmp_path / "source.docx")
    document = Document(source)
    document.paragraphs[0].style = "Heading 1"
    document.save(source)
    numbered = apply_heading_numbering(source, output_path=tmp_path / "numbered.docx")

    report = verify_document(numbered.output_path)

    assert report["numbering_status"] == "PASS"


def test_advanced_content_controls_fail_closed(tmp_path):
    source = _source(tmp_path / "source.docx")
    with pytest.raises(ValueError, match="plain_text"):
        create_content_control(
            source,
            output_path=tmp_path / "checkbox.docx",
            control_type="checkbox",
            title="Approval",
            tag="approval",
        )


def test_complex_field_removal_removes_complete_field_span(tmp_path):
    source = _source(tmp_path / "source.docx")
    document = Document(source)
    paragraph = document.paragraphs[0]
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        run = paragraph.add_run()._r
        if kind:
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), kind)
            run.append(field)
        else:
            element = OxmlElement("w:instrText" if text.strip() == "PAGE" else "w:t")
            element.text = text
            run.append(element)
    document.save(source)

    result = remove_fields(source, output_path=tmp_path / "removed.docx", field_type="PAGE")

    assert result.data["removed"] == 1
    assert inspect_fields(result.output_path) == []


def test_bookmark_rename_updates_fields_and_hyperlink_anchor(tmp_path):
    source = _source(tmp_path / "source.docx")
    bookmarked = create_bookmark(
        source,
        output_path=tmp_path / "bookmarked.docx",
        name="old_name",
        selector={"type": "paragraph", "index": 0},
    )
    referenced = insert_cross_reference(
        bookmarked.output_path,
        output_path=tmp_path / "referenced.docx",
        bookmark="old_name",
        include_page=True,
    )
    linked = insert_hyperlink(
        referenced.output_path,
        output_path=tmp_path / "linked.docx",
        text="jump",
        bookmark="old_name",
    )

    renamed = rename_bookmark(
        linked.output_path,
        output_path=tmp_path / "renamed.docx",
        old_name="old_name",
        new_name="new_name",
    )

    with ZipFile(renamed.output_path) as package:
        xml = package.read("word/document.xml").decode("utf-8")
    assert "old_name" not in xml
    assert "new_name" in xml


def test_cross_reference_rejects_missing_bookmark(tmp_path):
    source = _source(tmp_path / "source.docx")
    with pytest.raises(ValueError, match="does not exist"):
        insert_cross_reference(
            source,
            output_path=tmp_path / "missing.docx",
            bookmark="missing",
        )


def test_stale_fields_and_unreviewed_render_are_incomplete(tmp_path, monkeypatch):
    source = _source(tmp_path / "source.docx")
    document = Document(source)
    append_field(document.paragraphs[0], "PAGE", "1")
    document.save(source)

    report = verify_document(source)
    assert report["field_freshness_status"] == "STALE_OR_UNKNOWN"
    assert report["overall_status"] == "INCOMPLETE"


def test_field_structure_failure_always_fails_even_with_forged_sidecar(tmp_path):
    source = _source(tmp_path / "source.docx")
    document = Document(source)
    paragraph = document.paragraphs[0]
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    paragraph.add_run()._r.append(begin)
    instruction = OxmlElement("w:instrText")
    instruction.text = " REF missing "
    paragraph.add_run()._r.append(instruction)
    document.save(source)
    import hashlib, json
    source.with_suffix(".docx.wa-evidence.json").write_text(
        json.dumps(
            {
                "version": 1,
                "operation": "field.update_all",
                "backend": "word-native",
                "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            }
        ),
        encoding="utf-8",
    )

    report = verify_document(source)

    assert report["field_structure_status"] == "FAIL"
    assert report["overall_status"] == "FAIL"
    assert report["native_word_status"] != "PASS"


def test_matching_sidecar_is_informational_not_native_attestation(tmp_path):
    source = _source(tmp_path / "source.docx")
    document = Document(source)
    append_field(document.paragraphs[0], "PAGE", "7")
    document.save(source)
    import hashlib, json
    source.with_suffix(".docx.wa-evidence.json").write_text(
        json.dumps(
            {
                "version": 1,
                "operation": "field.update_all",
                "backend": "word-native",
                "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            }
        ),
        encoding="utf-8",
    )

    report = verify_document(source)

    assert report["field_evidence_status"] == "INFORMATIONAL_UNTRUSTED"
    assert report["field_freshness_status"] == "STALE_OR_UNKNOWN"
    assert report["native_word_status"] != "PASS"
    assert report["overall_status"] == "INCOMPLETE"


def test_bookmark_delete_blocks_live_references_and_missing_internal_link(tmp_path):
    source = _source(tmp_path / "source.docx")
    bookmarked = create_bookmark(
        source,
        output_path=tmp_path / "bookmarked.docx",
        name="target",
        selector={"type": "paragraph", "index": 0},
    )
    referenced = insert_cross_reference(
        bookmarked.output_path,
        output_path=tmp_path / "referenced.docx",
        bookmark="target",
        include_page=True,
    )
    with pytest.raises(ValueError, match="referenced"):
        delete_bookmark(
            referenced.output_path,
            output_path=tmp_path / "deleted.docx",
            name="target",
        )
    with pytest.raises(ValueError, match="does not exist"):
        insert_hyperlink(
            source,
            output_path=tmp_path / "missing-link.docx",
            text="jump",
            bookmark="missing",
        )


def test_bookmark_reference_safety_covers_footer_stories(tmp_path):
    source = _source(tmp_path / "source.docx")
    bookmarked = create_bookmark(
        source,
        output_path=tmp_path / "bookmarked.docx",
        name="target",
        selector={"type": "paragraph", "index": 0},
    )
    document = Document(bookmarked.output_path)
    append_field(document.sections[0].footer.paragraphs[0], "REF target", "Target")
    document.save(bookmarked.output_path)

    renamed = rename_bookmark(
        bookmarked.output_path,
        output_path=tmp_path / "renamed.docx",
        old_name="target",
        new_name="renamed_target",
    )
    footer_fields = [
        field
        for field in inspect_fields(renamed.output_path)
        if field["story"].endswith("footer")
    ]
    assert footer_fields[0]["instruction"].startswith("REF renamed_target")
    with pytest.raises(ValueError, match="referenced"):
        delete_bookmark(
            renamed.output_path,
            output_path=tmp_path / "deleted.docx",
            name="renamed_target",
        )


def test_render_page_inventory_is_numeric_and_isolates_stale_pages(tmp_path, monkeypatch):
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(b"%PDF")
    pages = tmp_path / "pages"
    pages.mkdir()
    (pages / "page-99.png").write_bytes(b"stale")

    monkeypatch.setattr("wordassistant.document.render.shutil.which", lambda _: "pdftoppm")

    def fake_run(command, **kwargs):
        from PIL import Image

        prefix = Path(command[-1])
        for number in (1, 2, 10):
            image = Image.new("RGB", (2, 2), color=(number, number, number))
            image.save(prefix.with_name(f"{prefix.name}-{number}.png"), "PNG")
            image.close()
        return type("Completed", (), {"returncode": 0, "stderr": "", "stdout": ""})()

    monkeypatch.setattr("wordassistant.document.render.subprocess.run", fake_run)
    result = render_pdf_pages(pdf, pages)

    assert [path.name for path in result] == ["page-1.png", "page-2.png", "page-10.png"]
    assert (pages / "page-99.png").exists()
    assert all(path.parent != pages for path in result)
    assert "page-99.png" not in {path.name for path in result}
