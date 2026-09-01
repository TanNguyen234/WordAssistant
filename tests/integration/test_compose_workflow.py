from __future__ import annotations

from docx import Document
import pytest

from wordassistant.document.compose import compose_document
from wordassistant.document.integrity import audit_integrity
from wordassistant.document.scaffold import scaffold_document
from wordassistant.document.validate import validate_document
from wordassistant.evidence.registry import EvidenceRegistry
from wordassistant.operations.citations import list_citations
from wordassistant.operations.lists import apply_heading_numbering
from wordassistant.profiles.loader import resolve_profile


def _evidence(path, count: int = 10):
    registry = EvidenceRegistry(path)
    for number in range(1, count + 1):
        registry.add(
            {
                "id": f"source-{number}",
                "title": f"Verified Source {number}",
                "authors": [f"Author {number}"],
                "year": 2024,
                "url": f"https://example.com/source-{number}",
                "verified_by": "integration test",
                "synthetic": False,
                "tags": ["primary"],
            }
        )
    return path


def test_compose_populates_real_headings_citations_and_existing_bibliography(tmp_path):
    profile = resolve_profile("hcmunre_dacn_2026_2027")
    scaffold = tmp_path / "scaffold.docx"
    scaffold_document(
        scaffold,
        profile,
        "report",
        title="KIỂM THỬ WORDASSISTANT",
        student="NGUYỄN VĂN A - 2351010001",
        supervisor="TS. NGUYỄN VĂN B",
    )
    evidence = _evidence(tmp_path / "evidence.json")
    content = {
        "version": 1,
        "citation_style": "hcmunre_numeric",
        "bibliography_heading": "DANH MỤC TÀI LIỆU THAM KHẢO",
        "sections": [
            {
                "heading": "MỞ ĐẦU",
                "paragraphs": [
                    {
                        "text": "Nghiên cứu sử dụng nguồn đã kiểm chứng.",
                        "citations": [f"source-{number}" for number in range(1, 11)],
                    }
                ],
            },
            {
                "heading": "CHƯƠNG 1: TỔNG QUAN TÀI LIỆU",
                "paragraphs": ["Nội dung tổng quan có chủ đích."],
                "subsections": [
                    {
                        "title": "Ngữ nghĩa định dạng Word",
                        "level": 2,
                        "paragraphs": ["Style là cấu trúc Word, không phải chữ số gõ tay."],
                    },
                    {
                        "title": "Trường động",
                        "level": 3,
                        "paragraphs": ["TOC và tham chiếu phải là field thật."],
                    },
                ],
            },
        ],
    }

    result = compose_document(
        scaffold,
        output_path=tmp_path / "composed.docx",
        content=content,
        evidence_registry=evidence,
    )

    assert result["citations"] == 10
    document = Document(result["output"])
    assert sum(
        p.text == "DANH MỤC TÀI LIỆU THAM KHẢO" for p in document.paragraphs
    ) == 1
    assert next(p for p in document.paragraphs if p.text == "Ngữ nghĩa định dạng Word").style.name == "Heading 2"
    assert next(p for p in document.paragraphs if p.text == "Trường động").style.name == "Heading 3"
    assert len(list_citations(result["output"])) == 10
    assert sum(p.text.startswith("[") and "Verified Source" in p.text for p in document.paragraphs) == 10
    validation = validate_document(result["output"], profile, "report")
    assert not {
        issue.code for issue in validation.issues
    } & {"missing_references", "insufficient_references"}
    integrity = audit_integrity(
        result["output"],
        evidence_registry_path=evidence,
        profile=profile,
        variant_name="report",
    )
    assert integrity["gates"]["source_evidence"]["status"] == "PASS"
    assert integrity["gates"]["numeric_citation_bibliography_consistency"]["status"] == "PASS"


def test_scoped_heading_numbering_excludes_front_matter_and_terminal_sections(tmp_path):
    source = tmp_path / "headings.docx"
    document = Document()
    document.add_heading("MỤC LỤC", level=1)
    document.add_heading("CHƯƠNG 1: TỔNG QUAN", level=1)
    document.add_heading("Cơ sở", level=2)
    document.add_heading("Chi tiết", level=3)
    document.add_heading("KẾT LUẬN", level=1)
    document.save(source)

    result = apply_heading_numbering(
        source,
        output_path=tmp_path / "numbered.docx",
        levels=3,
        start_text="CHƯƠNG 1:",
        end_text="KẾT LUẬN",
    )
    paragraphs = Document(result.output_path).paragraphs

    assert paragraphs[0]._p.pPr.numPr is None
    assert int(paragraphs[1]._p.pPr.numPr.ilvl.val) == 0
    assert int(paragraphs[2]._p.pPr.numPr.ilvl.val) == 1
    assert int(paragraphs[3]._p.pPr.numPr.ilvl.val) == 2
    assert paragraphs[4]._p.pPr.numPr is None


def test_compose_fails_closed_before_writing_when_visible_budget_is_exceeded(tmp_path):
    profile = resolve_profile("hcmunre_dacn_2026_2027")
    scaffold = tmp_path / "scaffold.docx"
    output = tmp_path / "composed.docx"
    scaffold_document(scaffold, profile, "report", title="BUDGET TEST")
    content = {
        "version": 1,
        "sections": [
            {
                "heading": "MỞ ĐẦU",
                "paragraphs": ["Nội dung vượt ngân sách " * 100],
            }
        ],
    }

    with pytest.raises(ValueError, match="exceeds the allowed maximum"):
        compose_document(
            scaffold,
            output_path=output,
            content=content,
            max_visible_characters=100,
        )

    assert not output.exists()
