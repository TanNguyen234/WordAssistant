import hashlib
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from wordassistant import cli
from wordassistant.document.captions import audit_captions
from wordassistant.operations.fields import append_field


def _image(path):
    image = Image.new("RGB", (16, 16), "white")
    image.save(path, format="PNG")
    image.close()


def _add_picture(document, image_path):
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(str(image_path))
    return paragraph


def _set_table_caption_metadata(table, value):
    properties = table._tbl.tblPr
    caption = OxmlElement("w:tblCaption")
    caption.set(qn("w:val"), value)
    properties.append(caption)


def test_caption_audit_reports_adjacent_evidence_and_metadata_without_guessing(tmp_path):
    image_path = tmp_path / "figure.png"
    _image(image_path)
    source = tmp_path / "captions.docx"
    document = Document()

    _add_picture(document, image_path)
    document.add_paragraph("Hình 1.1: Ảnh minh họa", style="Caption")

    document.add_paragraph("Bảng 1.1: Dữ liệu mẫu", style="Caption")
    document.add_table(rows=1, cols=1)

    metadata_table = document.add_table(rows=1, cols=1)
    _set_table_caption_metadata(metadata_table, "Bảng thông tin")

    _add_picture(document, image_path)
    document.save(source)
    before_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    report = audit_captions(source)

    assert report["status"] == "REVIEW_REQUIRED"
    assert report["summary"]["figures"] == {
        "total": 2, "present": 1, "missing": 1, "empty": 0,
        "ambiguous": 0, "type_mismatch": 0, "metadata_only": 0,
    }
    assert report["summary"]["tables"]["present"] == 1
    assert report["tables"][1]["caption_status"] == "METADATA_ONLY"
    assert report["tables"][1]["caption_metadata"] == "Bảng thông tin"
    assert report["figures"][0]["captions"][0]["text"] == "Hình 1.1: Ảnh minh họa"
    assert report["scope"]["mutates_document"] is False
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before_hash


def test_caption_audit_marks_multiple_adjacent_captions_ambiguous(tmp_path):
    image_path = tmp_path / "figure.png"
    _image(image_path)
    source = tmp_path / "ambiguous.docx"
    document = Document()
    document.add_paragraph("Figure 1: first possible caption", style="Caption")
    _add_picture(document, image_path)
    document.add_paragraph("Figure 2: second possible caption", style="Caption")
    document.save(source)

    report = audit_captions(source)

    assert report["figures"][0]["caption_status"] == "AMBIGUOUS"
    assert any(issue["code"] == "caption_ambiguous" for issue in report["issues"])


def test_caption_audit_does_not_mistake_equation_sequence_for_caption(tmp_path):
    image_path = tmp_path / "figure.png"
    _image(image_path)
    source = tmp_path / "equation.docx"
    document = Document()
    _add_picture(document, image_path)
    equation = document.add_paragraph()
    append_field(equation, "SEQ Equation \\* ARABIC", "1")
    document.save(source)

    report = audit_captions(source)

    assert report["figures"][0]["caption_status"] == "MISSING"


def test_caption_audit_reports_caption_type_mismatches(tmp_path):
    image_path = tmp_path / "figure.png"
    _image(image_path)
    source = tmp_path / "type-mismatch.docx"
    document = Document()
    _add_picture(document, image_path)
    document.add_paragraph("Table 1: Data", style="Caption")
    document.add_paragraph("Figure 2: Diagram", style="Caption")
    document.add_table(rows=1, cols=1)
    document.save(source)

    report = audit_captions(source)

    assert report["figures"][0]["caption_status"] == "TYPE_MISMATCH"
    assert report["figures"][0]["mismatched_captions"][0]["caption_kind"] == "table"
    assert report["tables"][0]["caption_status"] == "TYPE_MISMATCH"
    assert report["tables"][0]["mismatched_captions"][0]["caption_kind"] == "figure"
    assert sum(issue["code"] == "caption_type_mismatch" for issue in report["issues"]) == 2


def test_caption_audit_reassembles_split_complex_seq_instruction(tmp_path):
    image_path = tmp_path / "figure.png"
    _image(image_path)
    source = tmp_path / "split-seq.docx"
    document = Document()
    _add_picture(document, image_path)
    paragraph = document.add_paragraph(style="Caption")

    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)
    for fragment in ("SE", "Q Figure \\\\* ARABIC"):
        instruction_run = paragraph.add_run()
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = fragment
        instruction_run._r.append(instruction)
    separator_run = paragraph.add_run()
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    separator_run._r.append(separator)
    paragraph.add_run("1")
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)
    document.save(source)

    report = audit_captions(source)

    assert report["figures"][0]["caption_status"] == "PRESENT"
    assert report["figures"][0]["captions"][0]["sequence_ids"] == ["Figure"]


def test_caption_audit_cli_is_report_only_and_returns_report_for_missing_caption(
    tmp_path, monkeypatch, capsys
):
    image_path = tmp_path / "figure.png"
    _image(image_path)
    source = tmp_path / "missing.docx"
    document = Document()
    _add_picture(document, image_path)
    document.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()

    result = cli.run(["caption", "audit", str(source), "--json"])

    report = json.loads(capsys.readouterr().out)
    assert result == 0
    assert report["figures"][0]["caption_status"] == "MISSING"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before

    rejected_output = tmp_path / "must-not-write.json"
    assert cli.run(
        ["caption", "audit", str(source), "--output", str(rejected_output), "--json"]
    ) == 1
    assert not rejected_output.exists()
