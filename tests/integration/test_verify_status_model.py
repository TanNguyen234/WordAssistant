from docx import Document

from wordassistant.document.render import RenderResult
from wordassistant.document.verify import verify_document
from wordassistant.operations.fields import insert_toc
from wordassistant.operations.tables import create_table


def test_verify_reports_independent_statuses_without_profile(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_heading("Report", level=1)
    doc.add_paragraph("Complete content")
    doc.save(source)

    report = verify_document(source, render=False)

    assert set(report) >= {
        "overall_status",
        "structural_status",
        "format_status",
        "numbering_status",
        "table_geometry_status",
        "field_structure_status",
        "field_freshness_status",
        "field_status",
        "render_status",
        "native_word_status",
        "manual_review_required",
    }
    assert report["overall_status"] == "PASS"
    assert report["delivery_status"] == "INCOMPLETE"
    assert report["structural_status"] == "PASS"
    assert report["format_status"] == "NOT_REQUESTED"
    assert report["render_status"] == "NOT_REQUESTED"


def test_verify_rejects_literal_heading_numbers_and_stale_toc(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_heading("1 Introduction", level=1)
    doc.save(source)
    with_toc = insert_toc(source, output_path=tmp_path / "toc.docx")

    report = verify_document(with_toc.output_path, render=False)

    assert report["overall_status"] == "FAIL"
    assert report["numbering_status"] == "FAIL"
    assert report["field_structure_status"] == "PASS"
    assert report["field_freshness_status"] == "FAIL"
    assert report["manual_review_required"] is True


def test_verify_reports_exact_table_geometry(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    result = create_table(
        source,
        output_path=tmp_path / "table.docx",
        rows=2,
        columns=3,
    )

    report = verify_document(result.output_path, render=False)

    assert report["table_geometry_status"] == "PASS"
    assert report["overall_status"] == "PASS"
    assert report["delivery_status"] == "INCOMPLETE"


def test_verify_fails_structurally_for_non_docx(tmp_path):
    source = tmp_path / "bad.docx"
    source.write_text("not a zip", encoding="utf-8")

    report = verify_document(source, render=False)

    assert report["overall_status"] == "FAIL"
    assert report["structural_status"] == "FAIL"
    assert report["issues"][0]["severity"] == "ERROR"


def test_verify_can_pass_after_explicit_visual_review(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Complete content with enough visible text for review.")
    document.save(source)
    pdf = tmp_path / "source.pdf"
    pdf.write_bytes(b"%PDF-test")
    page = tmp_path / "page-1.png"
    page.write_bytes(b"png")
    monkeypatch.setattr("wordassistant.document.verify.render_pdf_pages", lambda _: (page,))
    monkeypatch.setattr("wordassistant.document.verify._pdf_page_count", lambda _: 1)
    monkeypatch.setattr(
        "wordassistant.document.verify.audit_pdf_page_content",
        lambda *args, **kwargs: {
            "status": "PASS",
            "pages": [{"page": 1, "alphanumeric_characters": 50, "near_blank": False}],
            "near_blank_pages": [],
        },
    )

    challenge_report = verify_document(
        source,
        render=True,
        renderer=lambda *args, **kwargs: RenderResult(
            pdf, pdf.stat().st_size, engine="word", page_count=1
        ),
    )
    report = verify_document(
        source,
        render=True,
        renderer=lambda *args, **kwargs: RenderResult(
            pdf, pdf.stat().st_size, engine="word", page_count=1
        ),
        visual_reviewed=True,
        visual_review_token=challenge_report["visual_review_challenge"]["token"],
    )

    assert report["visual_status"] == "PASS"
    assert report["manual_review_required"] is False
    assert report["overall_status"] == "PASS"
    assert report["delivery_status"] == "PASS"


def test_visual_review_cannot_override_near_blank_failure(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)
    pdf = tmp_path / "source.pdf"
    pdf.write_bytes(b"%PDF-test")
    page = tmp_path / "page-1.png"
    page.write_bytes(b"png")
    monkeypatch.setattr("wordassistant.document.verify.render_pdf_pages", lambda _: (page,))
    monkeypatch.setattr("wordassistant.document.verify._pdf_page_count", lambda _: 1)
    monkeypatch.setattr(
        "wordassistant.document.verify.audit_pdf_page_content",
        lambda *args, **kwargs: {
            "status": "FAIL",
            "pages": [],
            "near_blank_pages": [1],
        },
    )

    report = verify_document(
        source,
        render=True,
        renderer=lambda *args, **kwargs: RenderResult(
            pdf, pdf.stat().st_size, engine="word", page_count=1
        ),
        visual_reviewed=True,
        visual_review_token="invalid",
    )

    assert report["visual_status"] == "FAIL"
    assert report["overall_status"] == "FAIL"


def test_visual_review_requires_token_bound_to_current_page_evidence(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)
    pdf = tmp_path / "source.pdf"
    pdf.write_bytes(b"%PDF-test")
    page = tmp_path / "page-1.png"
    page.write_bytes(b"first-render")
    monkeypatch.setattr("wordassistant.document.verify.render_pdf_pages", lambda _: (page,))
    monkeypatch.setattr("wordassistant.document.verify._pdf_page_count", lambda _: 1)
    monkeypatch.setattr(
        "wordassistant.document.verify.audit_pdf_page_content",
        lambda *args, **kwargs: {"status": "PASS", "pages": [], "near_blank_pages": []},
    )
    renderer = lambda *args, **kwargs: RenderResult(
        pdf, pdf.stat().st_size, engine="word", page_count=1
    )

    first = verify_document(source, render=True, renderer=renderer)
    page.write_bytes(b"changed-render")
    final = verify_document(
        source,
        render=True,
        renderer=renderer,
        visual_reviewed=True,
        visual_review_token=first["visual_review_challenge"]["token"],
    )

    assert final["visual_status"] == "FAIL"
    assert any(issue["code"] == "visual_review_attestation_invalid" for issue in final["issues"])


def test_verify_rejects_pdf_and_page_image_count_mismatch(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)
    pdf = tmp_path / "source.pdf"
    pdf.write_bytes(b"%PDF-test")
    page = tmp_path / "page-1.png"
    page.write_bytes(b"png")
    monkeypatch.setattr("wordassistant.document.verify.render_pdf_pages", lambda _: (page,))
    monkeypatch.setattr("wordassistant.document.verify._pdf_page_count", lambda _: 1)

    report = verify_document(
        source,
        render=True,
        renderer=lambda *args, **kwargs: RenderResult(
            pdf, pdf.stat().st_size, engine="word", page_count=2
        ),
        visual_reviewed=True,
        visual_review_token="irrelevant",
    )

    assert report["visual_status"] == "FAIL"
    assert any(issue["code"] == "rendered_page_count_mismatch" for issue in report["issues"])


def test_visual_review_token_ignores_nonvisual_pdf_byte_churn(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)
    pdf = tmp_path / "source.pdf"
    pdf.write_bytes(b"%PDF-first")
    page = tmp_path / "page-1.png"
    page.write_bytes(b"png")
    monkeypatch.setattr("wordassistant.document.verify.render_pdf_pages", lambda _: (page,))
    monkeypatch.setattr("wordassistant.document.verify._pdf_page_count", lambda _: 1)
    monkeypatch.setattr(
        "wordassistant.document.verify.audit_pdf_page_content",
        lambda *args, **kwargs: {"status": "PASS", "pages": [], "near_blank_pages": []},
    )
    renderer = lambda *args, **kwargs: RenderResult(
        pdf, pdf.stat().st_size, engine="word", page_count=None
    )

    first = verify_document(source, render=True, renderer=renderer)
    pdf.write_bytes(b"%PDF-second")
    final = verify_document(
        source,
        render=True,
        renderer=renderer,
        visual_reviewed=True,
        visual_review_token=first["visual_review_challenge"]["token"],
    )

    assert first["visual_review_challenge"]["pdf_sha256"] != final["visual_review_challenge"]["pdf_sha256"]
    assert final["visual_status"] == "PASS"
    assert not any(issue["code"] == "visual_review_attestation_invalid" for issue in final["issues"])


def test_missing_renderer_count_cannot_hide_pdf_image_mismatch(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)
    pdf = tmp_path / "source.pdf"
    pdf.write_bytes(b"%PDF-test")
    page = tmp_path / "page-1.png"
    page.write_bytes(b"png")
    monkeypatch.setattr("wordassistant.document.verify.render_pdf_pages", lambda _: (page,))
    monkeypatch.setattr("wordassistant.document.verify._pdf_page_count", lambda _: 2)

    report = verify_document(
        source,
        render=True,
        renderer=lambda *args, **kwargs: RenderResult(
            pdf, pdf.stat().st_size, engine="word", page_count=None
        ),
    )

    assert report["visual_status"] == "FAIL"
    assert any(issue["code"] == "rendered_page_count_mismatch" for issue in report["issues"])
