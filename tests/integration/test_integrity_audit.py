from __future__ import annotations

from hashlib import sha256
import json

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wordassistant.document.integrity import audit_integrity, build_provenance_manifest
from wordassistant.evidence.registry import EvidenceRegistry
from wordassistant.cli import run
from wordassistant.profiles.loader import resolve_profile


def _save_registry(path, records):
    registry = EvidenceRegistry(path)
    for record in records:
        registry.add(record)
    return path


def _real_source(record_id, *, synthetic=False, tags=None):
    return {
        "id": record_id,
        "title": f"Source {record_id}",
        "authors": ["Researcher A"],
        "year": 2026,
        "url": f"https://example.org/{record_id}",
        "verified_by": "human reviewer",
        "synthetic": synthetic,
        "tags": list(tags or []),
    }


def test_real_docx_audits_body_table_header_and_footer(tmp_path):
    path = tmp_path / "integrity-stories.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header TODO"
    document.sections[0].footer.paragraphs[0].text = "Footer\u200btext"
    document.add_heading("Analysis", level=1)
    document.add_paragraph("The measured value was 38.5% without a source.")
    document.add_paragraph(
        'The participant stated, "This quotation contains enough separate words to require an adjacent academic citation in the final report."'
    )
    long_quote = " ".join(f"word{index}" for index in range(45))
    document.add_paragraph(f'"{long_quote}"')
    document.add_paragraph("(Nguyen, 2024)")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "A styled remark — with an ellipsis…"
    document.add_heading("References", level=1)
    document.add_paragraph("[1] Verified source. https://example.org/one")
    document.save(path)

    report = audit_integrity(path)

    json.dumps(report, ensure_ascii=False)
    assert report["ai_detection_performed"] is False
    assert report["detector_evasion_assistance"] is False
    assert report["coverage"]["nonempty_story_counts"]["body"] >= 4
    assert report["coverage"]["nonempty_story_counts"]["table"] == 1
    assert report["coverage"]["nonempty_story_counts"]["header"] == 1
    assert report["coverage"]["nonempty_story_counts"]["footer"] == 1
    assert report["gates"]["invisible_unicode"]["status"] == "FAIL"
    assert report["gates"]["unresolved_placeholders"]["status"] == "FAIL"
    punctuation_gate = report["gates"]["typography_punctuation_review"]
    assert punctuation_gate["classification"] == "STYLE_REVIEW_NOT_AI_EVIDENCE"
    assert {item["punctuation"] for item in punctuation_gate["findings"]} >= {"em_dash", "ellipsis"}
    assert report["gates"]["long_verbatim_quotes"]["status"] == "WARN"
    assert report["gates"]["quotes_lacking_adjacent_citation"]["status"] == "WARN"
    assert report["gates"]["quotes_lacking_adjacent_citation"]["finding_count"] == 1
    assert report["gates"]["quantitative_claims_lacking_citation"]["status"] == "WARN"


def test_real_docx_cross_checks_numeric_references_and_registry_policy(tmp_path):
    path = tmp_path / "citation-audit.docx"
    document = Document()
    document.add_heading("Findings", level=1)
    document.add_paragraph("The verified finding is supported by evidence [1].")
    document.add_paragraph("A second claim points to a missing bibliography item [3].")
    document.add_heading("References", level=1)
    document.add_paragraph("[1] First source. https://example.org/one")
    document.add_paragraph("[2] Unused source. https://example.org/two")
    document.save(path)

    registry_path = _save_registry(
        tmp_path / "evidence.json",
        [
            _real_source("source-1", tags=["primary"]),
            _real_source("source-2"),
            _real_source("source-3", synthetic=True),
        ],
    )
    profile = {
        "id": "integrity-test",
        "version": 1,
        "citations": {"minimum_references": 3, "prefer_primary_source": True},
    }

    report = audit_integrity(path, registry_path, profile)

    citation_gate = report["gates"]["numeric_citation_bibliography_consistency"]
    assert citation_gate["status"] == "FAIL"
    assert citation_gate["citation_numbers"] == [1, 3]
    assert citation_gate["bibliography_numbers"] == [1, 2]
    evidence_gate = report["gates"]["source_evidence"]
    assert evidence_gate["status"] == "FAIL"
    assert evidence_gate["document_reference_count"] == 2
    assert evidence_gate["verified_record_count"] == 3
    assert evidence_gate["primary_record_count"] == 1
    assert evidence_gate["synthetic_record_count"] == 1
    kinds = {finding["kind"] for finding in evidence_gate["findings"]}
    assert {"minimum_references_not_met", "synthetic_evidence_present"} <= kinds


def test_repetition_is_warn_only_and_never_ai_detection(tmp_path):
    path = tmp_path / "repetition.docx"
    document = Document()
    for ending in ("alpha", "beta", "gamma"):
        document.add_paragraph(f"In this study we evaluate the same repeated phrase for {ending}.")
    document.add_paragraph("AI-use disclosure: no generative AI was used to author the assessed claims.")
    document.save(path)

    report = audit_integrity(
        path,
        profile={
            "integrity": {
                "repeated_opening_words": 3,
                "repeated_opening_minimum": 3,
                "ngram_size": 5,
                "repeated_ngram_minimum": 3,
                "review_punctuation": [],
            }
        },
    )

    gate = report["gates"]["repetitive_prose_review"]
    assert gate["status"] == "WARN"
    assert gate["classification"] == "HUMAN_REVIEW_HEURISTIC_NOT_AI_DETECTION"
    assert report["gates"]["ai_use_disclosure_provenance"]["status"] == "PASS"


def test_vietnamese_ai_assistance_disclosure_is_recognized(tmp_path):
    path = tmp_path / "vietnamese-ai-disclosure.docx"
    document = Document()
    document.add_paragraph(
        "WordAssistant hỗ trợ định dạng; công cụ AI hỗ trợ soạn thảo và rà soát diễn đạt."
    )
    document.save(path)

    report = audit_integrity(path)

    gate = report["gates"]["ai_use_disclosure_provenance"]
    assert gate["status"] == "PASS"
    assert gate["declared"] is True


def test_quantitative_review_ignores_figure_labels_and_software_versions(tmp_path):
    path = tmp_path / "quantitative-signal.docx"
    document = Document()
    document.add_paragraph("Hình 1.2: Giao diện trang đăng nhập", style="Caption")
    document.add_paragraph("Môi trường chạy sử dụng Java 21 và Spring Boot 3.")
    document.add_paragraph("Tỷ lệ giao dịch thành công đạt 98,7%.")
    document.save(path)

    report = audit_integrity(path)

    findings = report["gates"]["quantitative_claims_lacking_citation"]["findings"]
    assert len(findings) == 1
    assert "98,7%" in findings[0]["excerpt"]


def test_repetition_review_suppresses_common_function_word_phrases(tmp_path):
    path = tmp_path / "common-phrase.docx"
    document = Document()
    document.add_paragraph("For checkout, this is part of the same workflow.")
    document.add_paragraph("For review, this is part of the same workflow.")
    document.add_paragraph("For admin, this is part of the same workflow.")
    document.save(path)

    report = audit_integrity(
        path,
        profile={
            "integrity": {
                "repeated_opening_words": 4,
                "repeated_opening_minimum": 3,
                "ngram_size": 5,
                "repeated_ngram_minimum": 3,
                "review_punctuation": [],
            }
        },
    )

    gate = report["gates"]["repetitive_prose_review"]
    assert gate["status"] == "PASS"
    assert gate["findings"] == []


def test_provenance_manifest_binds_real_docx_registry_and_audit(tmp_path):
    path = tmp_path / "manifest.docx"
    document = Document()
    document.add_paragraph("AI-use disclosure: language tools were used only for spelling review.")
    document.save(path)
    registry_path = _save_registry(
        tmp_path / "evidence.json",
        [_real_source("source-1", tags=["primary"])],
    )
    audit = audit_integrity(path, registry_path)

    manifest = build_provenance_manifest(
        path,
        registry_path,
        disclosure="Language tools were used only for spelling review.",
        audit_report=audit,
        source_registry={"owner": "student"},
    )

    json.dumps(manifest, ensure_ascii=False)
    expected_doc_hash = sha256(path.read_bytes()).hexdigest()
    expected_registry_hash = sha256(registry_path.read_bytes()).hexdigest()
    assert manifest["document"]["sha256"] == expected_doc_hash
    assert manifest["source_registry"]["evidence"]["sha256"] == expected_registry_hash
    assert manifest["source_registry"]["metadata"] == {"owner": "student"}
    assert manifest["disclosure"]["source"] == "caller"
    assert manifest["audit_summary"]["document_sha256"] == expected_doc_hash
    assert manifest["audit_summary"]["ai_detection_performed"] is False
    assert manifest["generated_at_utc"].endswith("+00:00")


def test_manifest_rejects_stale_audit_after_docx_changes(tmp_path):
    path = tmp_path / "stale.docx"
    document = Document()
    document.add_paragraph("Original content")
    document.save(path)
    audit = audit_integrity(path)

    changed = Document(path)
    changed.add_paragraph("Changed content")
    changed.save(path)

    with pytest.raises(ValueError, match="does not match"):
        build_provenance_manifest(path, audit_report=audit)


def test_missing_requested_registry_fails_closed_on_real_docx(tmp_path):
    path = tmp_path / "missing-registry.docx"
    document = Document()
    document.add_paragraph("A supported claim [1].")
    document.add_heading("References", level=1)
    document.add_paragraph("[1] Source")
    document.save(path)

    report = audit_integrity(path, tmp_path / "does-not-exist.json")

    gate = report["gates"]["source_evidence"]
    assert gate["status"] == "FAIL"
    assert gate["findings"][0]["kind"] == "missing_evidence_registry"


def test_report_integrity_audit_rejects_unreviewed_evidence_linked_citation(tmp_path):
    path = tmp_path / "unreviewed-linked-citation.docx"
    document = Document()
    paragraph = document.add_paragraph("The claim is supported ")
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "wordassistant:citation:source-1")
    properties.append(tag)
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "[1]"
    run.append(text)
    content.append(run)
    sdt.extend([properties, content])
    paragraph._p.append(sdt)
    document.add_heading("References", level=1)
    document.add_paragraph("[1] Legacy record without source review.")
    document.save(path)
    registry_path = _save_registry(
        tmp_path / "evidence.json", [_real_source("source-1")]
    )

    report = audit_integrity(
        path,
        registry_path,
        profile={"citations": {"require_reviewed_sources": True}},
    )

    gate = report["gates"]["source_evidence"]
    assert gate["status"] == "FAIL"
    assert gate["linked_evidence_id_count"] == 1
    assert any(
        item["kind"] == "cited_source_not_reviewed"
        for item in gate["findings"]
    )


def test_manifest_rejects_registry_changed_after_audit(tmp_path):
    path = tmp_path / "registry-change.docx"
    document = Document()
    document.add_paragraph("AI-use disclosure: no generative AI was used.")
    document.save(path)
    registry_path = _save_registry(tmp_path / "evidence.json", [_real_source("source-1")])
    audit = audit_integrity(path, registry_path)
    EvidenceRegistry(registry_path).add(_real_source("source-2"))

    with pytest.raises(ValueError, match="evidence registry"):
        build_provenance_manifest(path, registry_path, audit_report=audit)


def test_hcmunre_report_variant_applies_integrity_and_reference_policy(tmp_path):
    path = tmp_path / "hcmunre-policy.docx"
    document = Document()
    document.add_paragraph("Một câu cần rà soát — nhưng không phải bằng chứng phát hiện AI.")
    document.save(path)
    profile = resolve_profile("hcmunre_dacn_2026_2027", start_path=tmp_path)

    report = audit_integrity(path, profile=profile, variant_name="report")

    assert report["profile"]["profile_id"] == "hcmunre_dacn_2026_2027"
    assert report["profile"]["minimum_references"] == 10
    assert report["profile"]["repeated_opening_minimum"] == 4
    assert report["gates"]["source_evidence"]["status"] == "FAIL"
    assert report["gates"]["typography_punctuation_review"]["status"] == "WARN"


def test_public_integrity_audit_and_atomic_provenance_cli(tmp_path, capsys):
    path = tmp_path / "public-integrity.docx"
    document = Document()
    document.add_paragraph("AI-use disclosure: no generative AI was used to author this test.")
    document.save(path)

    audit_exit = run(["integrity", "audit", str(path), "--json"])
    audit_payload = json.loads(capsys.readouterr().out)
    assert audit_exit == 0
    assert audit_payload["audit_kind"] == "ETHICAL_ACADEMIC_INTEGRITY"
    assert audit_payload["ai_detection_performed"] is False

    output = tmp_path / "provenance.json"
    manifest_exit = run(
        [
            "integrity",
            "provenance",
            str(path),
            "--output",
            str(output),
            "--disclosure",
            "No generative AI was used to author this test.",
            "--json",
        ]
    )
    manifest_payload = json.loads(capsys.readouterr().out)
    manifest = json.loads(output.read_text(encoding="utf-8"))
    assert manifest_exit == 0
    assert manifest_payload["status"] == "SUCCESS"
    assert manifest_payload["output"] == str(output.resolve())
    assert manifest["document"]["sha256"] == sha256(path.read_bytes()).hexdigest()
    assert manifest["audit_summary"]["ai_detection_performed"] is False
