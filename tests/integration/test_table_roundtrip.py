import pytest
from docx import Document

from wordassistant.operations.tables import create_table
from wordassistant.document.verify import verify_document


def test_nested_table_creation_fails_closed_until_recursive_geometry_is_verified(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    with pytest.raises(ValueError, match="Nested table"):
        create_table(
            source,
            output_path=tmp_path / "created.docx",
            rows=1,
            columns=1,
            nested={"row": 0, "column": 0, "rows": 2, "columns": 2},
        )


def test_verifier_audits_existing_nested_table_geometry(tmp_path):
    source = tmp_path / "nested.docx"
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).add_table(rows=1, cols=1)
    document.save(source)

    report = verify_document(source)

    assert report["table_geometry_status"] == "FAIL"
    assert report["overall_status"] == "FAIL"
