from docx import Document

from wordassistant.cli import run
from wordassistant.document.verify import verify_document


def test_agent_surface_creates_table_and_page_number_without_source_inspection(tmp_path, capsys):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("Report", 1)
    document.add_paragraph("Body content")
    document.save(source)

    tabled = tmp_path / "tabled.docx"
    assert run([
        "table", "create", str(source), "--output", str(tabled),
        "--options", '{"rows":2,"columns":3}', "--json",
    ]) == 0
    capsys.readouterr()

    numbered = tmp_path / "numbered.docx"
    assert run([
        "page-number", "apply", str(tabled), "--output", str(numbered),
        "--options", '{"section_index":0,"number_format":"decimal","start":1}',
        "--json",
    ]) == 0
    capsys.readouterr()

    report = verify_document(numbered)
    assert report["structural_status"] == "PASS"
    assert report["table_geometry_status"] == "PASS"
