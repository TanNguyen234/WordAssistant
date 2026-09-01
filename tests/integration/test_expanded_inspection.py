import base64

from docx import Document

from wordassistant.document.structure import inspect_document
from wordassistant.operations.fields import insert_caption
from wordassistant.operations.lists import create_list
from wordassistant.operations.references import create_bookmark
from wordassistant.operations.tables import create_table


def test_expanded_inspection_reports_object_details(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Measured values")
    doc.save(source)
    tabled = create_table(source, output_path=tmp_path / "table.docx", rows=2, columns=2, caption="Results table")
    listed = create_list(tabled.output_path, output_path=tmp_path / "list.docx", items=["One", "Two"], kind="numbered")
    bookmarked = create_bookmark(listed.output_path, output_path=tmp_path / "bookmark.docx", name="results", selector={"type": "heading", "text": "Results"})
    captioned = insert_caption(bookmarked.output_path, output_path=tmp_path / "final.docx", label="Table", title="Results")

    inventory = inspect_document(captioned.output_path)

    assert inventory["tables_detail"][0]["rows"] == 2
    assert inventory["tables_detail"][0]["caption"] == "Results table"
    assert inventory["lists_detail"][0]["text"] == "One"
    assert inventory["bookmarks_detail"][0]["name"] == "results"
    assert any(field["type"] == "SEQ" for field in inventory["fields_detail"])
    assert inventory["sections_detail"][0]["orientation"] == "portrait"
    assert inventory["paragraphs_detail"][0]["style"] == "Heading 1"

