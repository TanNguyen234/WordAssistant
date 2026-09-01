from docx import Document
import sys
import types

from wordassistant.document.verify import audit_pdf_page_content, visual_preflight


def test_visual_preflight_flags_explicit_blank_page_pattern(tmp_path):
    source = tmp_path / "blank.docx"
    doc = Document()
    doc.add_paragraph("Page one")
    doc.add_page_break()
    doc.add_page_break()
    doc.add_paragraph("Page three")
    doc.save(source)

    result = visual_preflight(source)

    assert result["status"] == "WARNING"
    assert "consecutive_page_breaks" in result["warnings"]
    assert result["manual_review_required"] is True


def test_visual_preflight_accepts_contentful_document(tmp_path):
    source = tmp_path / "normal.docx"
    doc = Document()
    doc.add_paragraph("A complete paragraph with enough content for structural preflight.")
    doc.save(source)

    result = visual_preflight(source)

    assert result["status"] == "PASS"
    assert result["warnings"] == []


def test_pdf_page_content_audit_flags_near_blank_pages(tmp_path, monkeypatch):
    pdf = tmp_path / "rendered.pdf"
    pdf.write_bytes(b"%PDF-test")
    outputs = {
        1: "Substantive report content with enough visible text.",
        2: "2",
        3: "Another substantive page with a meaningful conclusion.",
    }

    monkeypatch.setattr("wordassistant.document.verify.shutil.which", lambda _: "pdftotext")

    def fake_run(command, **kwargs):
        page = int(command[command.index("-f") + 1])
        return type("Completed", (), {"returncode": 0, "stdout": outputs[page], "stderr": ""})()

    monkeypatch.setattr("wordassistant.document.verify.subprocess.run", fake_run)

    result = audit_pdf_page_content(pdf, page_count=3, minimum_alphanumeric=20)

    assert result["status"] == "FAIL"
    assert result["near_blank_pages"] == [2]
    assert result["pages"][1]["alphanumeric_characters"] == 1


def test_pdf_page_content_audit_is_explicitly_unavailable_without_extractor(tmp_path, monkeypatch):
    pdf = tmp_path / "rendered.pdf"
    pdf.write_bytes(b"%PDF-test")
    monkeypatch.setattr("wordassistant.document.verify.shutil.which", lambda _: None)

    result = audit_pdf_page_content(pdf, page_count=1)

    assert result["status"] == "UNAVAILABLE"
    assert result["near_blank_pages"] == []


def test_pdf_page_content_audit_falls_back_to_packaged_pdfium(tmp_path, monkeypatch):
    pdf = tmp_path / "rendered.pdf"
    pdf.write_bytes(b"%PDF-test")
    monkeypatch.setattr("wordassistant.document.verify.shutil.which", lambda _: None)

    class TextPage:
        def get_text_bounded(self):
            return "Substantive page content extracted by packaged PDFium."

        def close(self):
            return None

    class Page:
        def get_textpage(self):
            return TextPage()

        def close(self):
            return None

    class Pdf:
        def __len__(self):
            return 1

        def __getitem__(self, index):
            assert index == 0
            return Page()

        def close(self):
            return None

    monkeypatch.setitem(
        sys.modules,
        "pypdfium2",
        types.SimpleNamespace(PdfDocument=lambda path: Pdf()),
    )

    result = audit_pdf_page_content(pdf, page_count=1)

    assert result["status"] == "PASS"
    assert result["near_blank_pages"] == []


def test_failed_pdftotext_falls_back_to_packaged_pdfium(tmp_path, monkeypatch):
    pdf = tmp_path / "rendered.pdf"
    pdf.write_bytes(b"%PDF-test")
    monkeypatch.setattr("wordassistant.document.verify.shutil.which", lambda _: "pdftotext")
    monkeypatch.setattr(
        "wordassistant.document.verify.subprocess.run",
        lambda *args, **kwargs: types.SimpleNamespace(
            returncode=1, stdout="", stderr="poppler extraction failed"
        ),
    )

    class TextPage:
        def get_text_bounded(self):
            return "Substantive fallback text extracted after Poppler failed."

        def close(self):
            return None

    class Page:
        def get_textpage(self):
            return TextPage()

        def close(self):
            return None

    class Pdf:
        def __len__(self):
            return 1

        def __getitem__(self, index):
            return Page()

        def close(self):
            return None

    monkeypatch.setitem(
        sys.modules, "pypdfium2", types.SimpleNamespace(PdfDocument=lambda path: Pdf())
    )

    result = audit_pdf_page_content(pdf, page_count=1)

    assert result["status"] == "PASS"


def test_pdftotext_exception_falls_back_to_packaged_pdfium(tmp_path, monkeypatch):
    pdf = tmp_path / "rendered.pdf"
    pdf.write_bytes(b"%PDF-test")
    monkeypatch.setattr("wordassistant.document.verify.shutil.which", lambda _: "pdftotext")

    def timeout(*args, **kwargs):
        raise __import__("subprocess").TimeoutExpired(args[0], 30)

    monkeypatch.setattr("wordassistant.document.verify.subprocess.run", timeout)

    class TextPage:
        def get_text_bounded(self):
            return "Substantive PDFium text after a pdftotext timeout."

        def close(self):
            return None

    class Page:
        def get_textpage(self):
            return TextPage()

        def close(self):
            return None

    class Pdf:
        def __len__(self):
            return 1

        def __getitem__(self, index):
            return Page()

        def close(self):
            return None

    monkeypatch.setitem(
        sys.modules, "pypdfium2", types.SimpleNamespace(PdfDocument=lambda path: Pdf())
    )

    assert audit_pdf_page_content(pdf, page_count=1)["status"] == "PASS"
