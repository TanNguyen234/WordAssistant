from pathlib import Path
import json

from docx import Document
from wordassistant.document.render import RenderResult

from wordassistant.document.verify import verify_document
from wordassistant.operations.fields import append_field


def test_native_field_verification_uses_copy_and_preserves_source(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    document = Document()
    paragraph = document.add_paragraph("Page ")
    append_field(paragraph, "PAGE", "1")
    document.save(source)
    original = source.read_bytes()
    observed_path = None

    def fake_worker(payload, timeout):
        nonlocal observed_path
        observed_path = Path(payload["input_path"])
        assert observed_path != source
        assert observed_path.read_bytes() == original
        observed_path.write_bytes(observed_path.read_bytes() + b"worker mutation")
        body = (
            '{"status":"SUCCESS","data":{'
            '"fields":[{"instruction":"PAGE","result":"1","locked":false}],'
            '"source_results_were_current":true}}'
        )
        return type("Completed", (), {"returncode": 0, "stdout": body, "stderr": ""})()

    monkeypatch.setattr(
        "wordassistant.backends.word_native.backend._run_native_worker", fake_worker
    )

    report = verify_document(source, native_fields=True)

    assert report["field_freshness_status"] == "PASS"
    assert source.read_bytes() == original
    assert observed_path is not None
    assert not observed_path.exists()


def _field_payload():
    return {
        "fields": [{"instruction": "PAGE", "result": "1", "locked": False}],
        "source_results_were_current": True,
    }


def test_combined_native_verify_renders_and_checks_fields_once_on_private_copy(
    tmp_path, monkeypatch
):
    source = tmp_path / "combined.docx"
    doc = Document()
    append_field(doc.add_paragraph("Page "), "PAGE", "1")
    doc.save(source)
    original = source.read_bytes()
    destination = tmp_path / "renders"
    captured = {}

    def fake_worker(payload, timeout):
        captured.update(payload)
        private_docx = Path(payload["input_path"])
        private_pdf = Path(payload["output_path"])
        assert private_docx != source and private_docx.read_bytes() == original
        private_pdf.write_bytes(b"%PDF-combined")
        return type("Completed", (), {
            "returncode": 0,
            "stdout": json.dumps({"status": "SUCCESS", "data": {
                "render": {"status": "SUCCESS", "data": {"page_count": 1}},
                "fields": {"status": "SUCCESS", "data": _field_payload()},
            }}),
            "stderr": "",
        })()

    page_image = tmp_path / "page-1.png"
    page_image.write_bytes(b"png")
    monkeypatch.setattr("wordassistant.backends.word_native.backend._run_native_worker", fake_worker)
    monkeypatch.setattr("wordassistant.document.verify.render_pdf_pages", lambda *args, **kwargs: (page_image,))
    monkeypatch.setattr("wordassistant.document.verify._pdf_page_count", lambda _: 1)
    monkeypatch.setattr("wordassistant.document.verify.audit_pdf_page_content", lambda *args, **kwargs: {
        "status": "PASS", "pages": [{"page": 1, "near_blank": False}], "near_blank_pages": []
    })

    report = verify_document(
        source, render=True, native_fields=True, render_output_dir=destination
    )

    assert captured["action"] == "verify_fields_and_render"
    assert report["field_evaluation_status"] == "PASS"
    assert report["render_status"] == "PASS"
    assert report["native_word_status"] == "PASS"
    assert (destination / "combined.pdf").read_bytes() == b"%PDF-combined"
    assert source.read_bytes() == original


def test_combined_field_failure_is_not_erased_by_successful_render(
    tmp_path, monkeypatch
):
    source = tmp_path / "failed-field.docx"
    doc = Document()
    append_field(doc.add_paragraph("Page "), "PAGE", "1")
    doc.save(source)
    pdf_image = tmp_path / "page-1.png"
    pdf_image.write_bytes(b"png")

    def fake_worker(payload, timeout):
        Path(payload["output_path"]).write_bytes(b"%PDF-output")
        return type("Completed", (), {
            "returncode": 0,
            "stdout": json.dumps({"status": "SUCCESS", "data": {
                "render": {"status": "SUCCESS", "data": {"page_count": 1}},
                "fields": {"status": "ERROR", "message": "field stage failed"},
            }}),
            "stderr": "",
        })()

    monkeypatch.setattr("wordassistant.backends.word_native.backend._run_native_worker", fake_worker)
    monkeypatch.setattr("wordassistant.document.verify.render_pdf_pages", lambda *args, **kwargs: (pdf_image,))
    monkeypatch.setattr("wordassistant.document.verify._pdf_page_count", lambda _: 1)
    monkeypatch.setattr("wordassistant.document.verify.audit_pdf_page_content", lambda *args, **kwargs: {
        "status": "PASS", "pages": [{"page": 1, "near_blank": False}], "near_blank_pages": []
    })

    report = verify_document(source, render=True, native_fields=True)

    assert report["render_status"] == "PASS"
    assert report["field_evaluation_status"] == "FAIL"
    assert report["native_word_status"] == "FAIL"
    assert report["overall_status"] == "FAIL"


def test_custom_renderer_and_reuse_keep_legacy_field_worker_path(tmp_path, monkeypatch):
    source = tmp_path / "compat.docx"
    doc = Document()
    append_field(doc.add_paragraph("Page "), "PAGE", "1")
    doc.save(source)
    actions = []
    pdf = tmp_path / "custom.pdf"
    pdf.write_bytes(b"%PDF-custom")
    image = tmp_path / "page-1.png"
    image.write_bytes(b"png")

    def fake_worker(payload, timeout):
        actions.append(payload["action"])
        return type("Completed", (), {
            "returncode": 0,
            "stdout": json.dumps({"status": "SUCCESS", "data": _field_payload()}),
            "stderr": "",
        })()

    monkeypatch.setattr("wordassistant.backends.word_native.backend._run_native_worker", fake_worker)
    monkeypatch.setattr("wordassistant.document.verify.render_pdf_pages", lambda *args, **kwargs: (image,))
    monkeypatch.setattr("wordassistant.document.verify._pdf_page_count", lambda _: 1)
    monkeypatch.setattr("wordassistant.document.verify.audit_pdf_page_content", lambda *args, **kwargs: {
        "status": "PASS", "pages": [{"page": 1, "near_blank": False}], "near_blank_pages": []
    })

    custom = verify_document(
        source,
        render=True,
        native_fields=True,
        renderer=lambda *_args, **_kwargs: RenderResult(pdf, pdf.stat().st_size, "libreoffice", 1),
    )
    reuse_source = tmp_path / "reuse.json"
    reuse_source.write_text("{}", encoding="utf-8")
    reuse = verify_document(
        source,
        native_fields=True,
        reuse_render_report=reuse_source,
        visual_reviewed=True,
        visual_review_token="bad-token",
    )

    assert custom["render_status"] == "PASS"
    assert reuse["render_status"] == "FAIL"
    assert actions == ["verify_fields", "verify_fields"]


def test_malformed_combined_render_metadata_falls_back_without_losing_field_result(
    tmp_path, monkeypatch
):
    source = tmp_path / "malformed.docx"
    doc = Document()
    append_field(doc.add_paragraph("Page "), "PAGE", "1")
    doc.save(source)
    page_image = tmp_path / "page-1.png"
    page_image.write_bytes(b"png")
    fallback_pdf = tmp_path / "fallback.pdf"
    fallback_pdf.write_bytes(b"%PDF-fallback")
    actions = []

    def fake_worker(payload, timeout):
        actions.append(payload["action"])
        return type("Completed", (), {
            "returncode": 0,
            "stdout": json.dumps({"status": "SUCCESS", "data": {
                "render": {"status": "SUCCESS", "data": {"page_count": "unknown"}},
                "fields": {"status": "SUCCESS", "data": _field_payload()},
            }}),
            "stderr": "",
        })()

    monkeypatch.setattr("wordassistant.backends.word_native.backend._run_native_worker", fake_worker)
    monkeypatch.setattr("wordassistant.document.verify._render_via_execution_service", lambda *args, **kwargs: RenderResult(fallback_pdf, fallback_pdf.stat().st_size, "libreoffice", 1))
    monkeypatch.setattr("wordassistant.document.verify.render_pdf_pages", lambda *args, **kwargs: (page_image,))
    monkeypatch.setattr("wordassistant.document.verify._pdf_page_count", lambda _: 1)
    monkeypatch.setattr("wordassistant.document.verify.audit_pdf_page_content", lambda *args, **kwargs: {
        "status": "PASS", "pages": [{"page": 1, "near_blank": False}], "near_blank_pages": []
    })

    report = verify_document(source, render=True, native_fields=True)

    assert actions == ["verify_fields_and_render"]
    assert report["field_evaluation_status"] == "PASS"
    assert report["render_status"] == "PASS"
    assert report["render"]["engine"] == "libreoffice"
