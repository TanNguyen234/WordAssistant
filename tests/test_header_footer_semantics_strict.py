from __future__ import annotations
import json
from pathlib import Path
from docx import Document
from wordassistant.document.validate import validate_document
from wordassistant.profiles.loader import resolve_profile, load_profile
from wordassistant.document.scaffold import scaffold_document

def _create_clean_report(path: Path):
    prof = resolve_profile("hcmunre_dacn_2026_2027")
    scaffold_document(
        path,
        prof,
        "report",
        title="XAY DUNG HE THONG RAG CHO VAN BAN PHAP LUAT",
        student="NGUYEN VAN A - 2351010001",
        supervisor="TS. NGUYEN VAN B"
    )
    d = Document(path)
    for p in d.paragraphs:
        if "[Front matter content]" in p.text:
            p.text = "Lời cảm ơn chi tiết."
        elif "[Section content...]" in p.text:
            p.text = "Nội dung chương hoàn chỉnh."
        elif "[List of figures]" in p.text:
            p.text = "Hình 1.1: Sơ đồ."
        elif "[List of tables]" in p.text:
            p.text = "Bảng 1.1: Bảng dữ liệu."
    for i in range(1, 12):
        d.add_paragraph(f"[{i}] Tac gia {i}, Sach {i}, 2024.")
    d.save(path)
    return prof

def test_header_supervisor_same_surname_mismatch_fails(tmp_path: Path):
    doc_path = tmp_path / "same_surname_sup.docx"
    prof = _create_clean_report(doc_path)
    
    d = Document(doc_path)
    body_sec = d.sections[-1]
    # Header has correct topic but different person with same surname "NGUYEN"
    body_sec.header.paragraphs[0].text = "XAY DUNG HE THONG RAG CHO VAN BAN PHAP LUAT\tTS. NGUYEN VAN C"
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues if i.severity == "error"]
    assert "missing_header_supervisor" in codes or "header_supervisor_mismatch" in codes

def test_header_supervisor_only_single_word_fails(tmp_path: Path):
    doc_path = tmp_path / "single_word_sup.docx"
    prof = _create_clean_report(doc_path)
    
    d = Document(doc_path)
    body_sec = d.sections[-1]
    # Header has topic but only "NGUYEN"
    body_sec.header.paragraphs[0].text = "XAY DUNG HE THONG RAG CHO VAN BAN PHAP LUAT\tNGUYEN"
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues if i.severity == "error"]
    assert "missing_header_supervisor" in codes or "header_supervisor_mismatch" in codes

def test_footer_student_only_single_word_fails(tmp_path: Path):
    doc_path = tmp_path / "single_word_stu.docx"
    prof = _create_clean_report(doc_path)
    
    d = Document(doc_path)
    body_sec = d.sections[-1]
    for p in body_sec.footer.paragraphs:
        if "NGUYEN VAN A" in p.text:
            p.text = "NGUYEN"
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues if i.severity == "error"]
    assert "footer_student_mismatch" in codes or "missing_footer_student" in codes

def test_generic_english_profile_semantic_header_footer_validation(tmp_path: Path):
    custom_prof = {
        "id": "english_academic",
        "name": "English Academic Report",
        "version": 1,
        "variants": {
            "default": {
                "page": {
                    "size": "A4",
                    "width_cm": 21.0,
                    "height_cm": 29.7,
                    "margins_cm": {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0},
                    "one_sided": True
                },
                "typography": {
                    "body_font": "Arial",
                    "body_size_pt": 11.0
                },
                "cover": {
                    "institution_lines": ["HARVARD UNIVERSITY", "DEPARTMENT OF COMPUTER SCIENCE"],
                    "document_type": "FINAL THESIS REPORT",
                    "metadata_fields": [
                        {"label": "Candidate:", "value_key": "student", "required": True},
                        {"label": "Advisor:", "value_key": "supervisor", "required": True}
                    ]
                },
                "header_footer": {
                    "header_left": "topic",
                    "header_right": "supervisor",
                    "footer_left": "student",
                    "footer_right": "page_number"
                },
                "front_matter": ["cover"],
                "section_blueprint": [
                    {"title": "Introduction", "level": 1, "required": True}
                ],
                "citations": {
                    "minimum_references": 1
                }
            }
        }
    }
    prof_file = tmp_path / "custom.json"
    prof_file.write_text(json.dumps(custom_prof), encoding="utf-8")
    prof = load_profile(prof_file)
    
    doc_path = tmp_path / "english_report.docx"
    scaffold_document(
        doc_path,
        prof,
        "default",
        title="ADVANCED AI AGENTS RESEARCH",
        student="John Doe - ID12345",
        supervisor="Prof. Jane Smith"
    )
    
    # Fill in placeholders
    d = Document(doc_path)
    for p in d.paragraphs:
        if "[Section content...]" in p.text:
            p.text = "This is complete research content."
        elif "[Document content]" in p.text:
            p.text = "This is thesis content."
    d.add_heading("References", 1)
    d.add_paragraph("[1] J. Doe, AI Systems, 2025.")
    d.save(doc_path)
    
    # Verify should succeed without needing any Vietnamese labels
    val_res = validate_document(doc_path, prof, "default", verify_mode=True)
    assert val_res.valid is True
