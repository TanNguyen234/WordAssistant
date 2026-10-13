from pathlib import Path
import json
from docx import Document
from wordassistant.cli import run
from wordassistant.capabilities import capabilities

ROOT = Path(__file__).resolve().parents[1]
GUIDE = ROOT / 'use_cases/hcmunre_dacn_2026_2027/assets/huong_dan_trinh_bay_dacn_2026_2027.docx'

def test_doctor_json_is_core_only_and_succeeds(capsys):
    code = run(['doctor', '--json'])
    out = capsys.readouterr().out
    assert code == 0
    assert 'python-docx' in out

def test_status_command(tmp_path, capsys):
    run(['init', str(tmp_path), '--json'])
    capsys.readouterr()
    code = run(['status', '--json'])
    out = capsys.readouterr().out
    assert code == 0
    assert 'is_workspace' in out

def test_context_cli_needs_only_file_and_query(capsys):
    code = run(['context', str(GUIDE), '--query', 'trích dẫn tài liệu gốc', '--max-tokens', '1200', '--json'])
    out = capsys.readouterr().out.casefold()
    assert code == 0
    assert 'trích dẫn' in out

def test_init_creates_small_agent_workspace(tmp_path, capsys):
    target = tmp_path / 'proj'
    code = run(['init', str(target), '--json'])
    assert code == 0
    assert (target / '.wordassistant/project.json').exists()
    assert (target / '.wordassistant/evidence.json').exists()

def test_capabilities_explain_core_and_advanced(capsys):
    assert run(['capabilities', '--json']) == 0
    out = capsys.readouterr().out
    assert 'replace' in out
    assert 'tracked_changes' in out
    selector = json.loads(out)['selector_schema']
    assert selector['allowed_keys'] == ['type', 'text', 'index', 'style', 'case_sensitive']
    assert selector['types']['heading']['example'] == {'type': 'heading', 'text': 'Introduction'}
    payload = json.loads(out)
    assert {'protect', 'properties', 'mail-merge'} <= set(payload['core']['commands'])
    toc = payload['operation_contracts']['toc.insert']
    assert toc['cli_routes'][0]['family'] == 'toc'
    assert toc['cli_routes'][0]['action'] == 'insert'
    assert toc['selector_supported'] is True
    assert any(option['name'] == 'heading_depth' for option in toc['options'])
    comment = payload['operation_contracts']['comment.create']
    assert any(option['name'] == 'text' and option['required'] for option in comment['options'])
    assert 'end' not in comment['options_example']
    assert payload['operation_contracts']['mail-merge.execute']['output_required'] is True
    assert 'locked' not in payload['operation_contracts']['field.unlock']['options_example']
    assert {'font_family', 'font_size_pt', 'bold'} <= {
        option['name'] for option in payload['operation_contracts']['format.run']['options']
    }
    assert payload['operation_contracts']['format.run']['options_example'] == {}
    assert payload['operation_contracts']['format.paragraph']['options_example'] == {}
    assert {'title', 'author', 'custom'} <= {
        option['name'] for option in payload['operation_contracts']['properties.set']['options']
    }
    edit_routes = payload['operation_contracts']['text.edit']['cli_routes']
    replace_route = next(route for route in edit_routes if route['action'] == 'replace')
    assert 'action' not in replace_route['options_example']
    assert replace_route['options_example']['pattern'] == 'value'
    assert payload['operation_contracts']['content-control.create']['options_example']['control_type'] == 'plain_text'
    assert payload['operation_contracts']['protection.protect']['options_example']['mode'] == 'read_only'
    assert payload['operation_contracts']['reference.hyperlink']['options_example']['url'] == 'https://example.com'
    table_example = payload['operation_contracts']['table.create']['options_example']
    assert table_example == {'rows': 2, 'columns': 2}
    assert payload['operation_contracts']['image.insert']['options_example']['image_path'] == 'IMAGE.png'
    caption_separator = next(
        option for option in payload['operation_contracts']['image.insert']['options']
        if option['name'] == 'caption_separator'
    )
    assert caption_separator['default'] == '.'
    border_inset = next(
        option for option in payload['operation_contracts']['section.layout']['options']
        if option['name'] == 'page_border_space_pt'
    )
    assert border_inset['type'] == 'number'
    style_options = {
        option['name']: option
        for option in payload['operation_contracts']['style.update']['options']
    }
    assert style_options['font_color']['type'] == 'string'
    assert style_options['left_indent_pt']['type'] == 'number'
    assert payload['operation_contracts']['image.resize']['options_example'] == {
        'image_index': 0, 'width_cm': 5.0
    }
    assert payload['operation_contracts']['bookmark.rename']['options_example'] == {
        'old_name': 'OldBookmark', 'new_name': 'NewBookmark'
    }
    for operation in ('comment.delete', 'comment.reply', 'comment.resolve', 'comment.reopen', 'revision.accept', 'revision.reject'):
        assert payload['operation_contracts'][operation]['options_example']['index'] == 1
    normalized_types = {'string', 'integer', 'number', 'boolean', 'object', 'path', 'any'}
    for contract in payload['operation_contracts'].values():
        for option in contract['options']:
            assert option['type'] in normalized_types or option['type'].startswith('array[')
        assert all(
            not (isinstance(value, str) and value.startswith('<'))
            for value in contract['options_example'].values()
        )
    assert all(contract['cli_routes'] for contract in payload['operation_contracts'].values())


def test_agent_command_help_explains_selector_shape(capsys):
    try:
        run(['image', '--help'])
    except SystemExit as exc:
        assert exc.code == 0
    out = capsys.readouterr().out
    assert '{"type":"heading","text":"Introduction"}' in out
    assert 'type, text, index, style, case_sensitive' in ' '.join(out.split())


def test_capabilities_returns_an_independent_selector_schema():
    first = capabilities()
    first['selector_schema']['allowed_keys'].append('caller_mutation')

    assert 'caller_mutation' not in capabilities()['selector_schema']['allowed_keys']

def test_profiles_command_lists_variants(capsys):
    assert run(['profiles', '--json']) == 0
    out = capsys.readouterr().out
    assert 'hcmunre_dacn_2026_2027' in out
    assert 'report' in out and 'proposal' in out

def test_profile_show_and_validate(capsys):
    code = run(['profile', 'show', 'generic', '--json'])
    out = capsys.readouterr().out
    assert code == 0
    assert 'typography' in out
    
    code = run(['profile', 'validate', str(ROOT / 'profiles/generic.json'), '--json'])
    out = capsys.readouterr().out
    assert code == 0
    assert '"valid": true' in out

def test_inspect_cli(capsys):
    code = run(['inspect', str(GUIDE), '--json'])
    out = capsys.readouterr().out
    assert code == 0
    assert 'paragraphs' in out
    assert 'sections' in out


def test_inspect_summary_omits_document_text_and_large_inventories(tmp_path, capsys):
    path = tmp_path / 'long_report.docx'
    document = Document()
    document.add_heading('Methods', level=1)
    document.add_paragraph('SENSITIVE BODY TEXT THAT MUST NOT BE IN SUMMARY')
    document.add_heading('Results', level=1)
    document.add_paragraph('More body text.')
    document.save(path)

    code = run(['inspect', str(path), '--summary', '--json'])
    payload = json.loads(capsys.readouterr().out)

    assert code == 0
    assert payload['paragraphs'] == 4
    assert payload['outline_count'] == 2
    assert [item['title'] for item in payload['outline_preview']] == ['Methods', 'Results']
    assert 'paragraphs_detail' not in payload
    assert 'tables_detail' not in payload
    assert 'SENSITIVE BODY TEXT' not in json.dumps(payload)


def test_inspect_summary_bounds_heading_preview(tmp_path, capsys):
    path = tmp_path / 'many_sections.docx'
    document = Document()
    for index in range(130):
        document.add_heading(f'Section {index}', level=1)
        document.add_paragraph(f'Content {index}')
    document.save(path)

    code = run(['inspect', str(path), '--summary', '--json'])
    payload = json.loads(capsys.readouterr().out)

    assert code == 0
    assert payload['outline_count'] == 130
    assert len(payload['outline_preview']) == 100
    assert payload['outline_preview_truncated'] is True


def test_context_only_section_returns_no_neighboring_text(tmp_path, capsys):
    path = tmp_path / 'two_chapters.docx'
    document = Document()
    document.add_heading('Methods', level=1)
    document.add_paragraph('The target method and only its details.')
    document.add_heading('Results', level=1)
    document.add_paragraph('Neighboring results details.')
    document.save(path)

    code = run([
        'context', str(path), '--query', 'details', '--only-section', 'methods',
        '--max-tokens', '4000', '--json',
    ])
    payload = json.loads(capsys.readouterr().out)

    assert code == 0
    assert [item['id'] for item in payload['sections']] == ['methods']
    assert 'The target method' in payload['text']
    assert 'Neighboring results' not in payload['text']


def test_context_cli_rejects_overfeed_budget(tmp_path, capsys):
    path = tmp_path / 'report.docx'
    document = Document()
    document.add_heading('Methods', level=1)
    document.add_paragraph('Some content.')
    document.save(path)

    code = run([
        'context', str(path), '--query', 'content', '--max-tokens', '12001', '--json',
    ])
    captured = capsys.readouterr()

    assert code == 1
    assert 'hard maximum of 12000' in (captured.out + captured.err)

def test_evidence_cli_crud(tmp_path, capsys):
    reg = tmp_path / 'evidence.json'
    
    # 1. Add
    code = run([
        'evidence', 'add', str(reg),
        '--id', 'ev1',
        '--title', 'Machine Learning Paper',
        '--authors', 'John Doe, Jane Smith',
        '--year', '2023',
        '--url', 'https://doi.org/10.1234/example',
        '--verified-by', 'Alice',
        '--verified-real-source',
        '--json'
    ])
    assert code == 0
    capsys.readouterr()
    
    # 2. Get
    code = run(['evidence', 'get', str(reg), 'ev1', '--json'])
    out = capsys.readouterr().out
    assert code == 0
    assert 'Machine Learning Paper' in out
    
    # 3. List
    code = run(['evidence', 'list', str(reg), '--json'])
    out = capsys.readouterr().out
    assert code == 0
    assert 'ev1' in out
    
    # 4. Validate
    code = run(['evidence', 'validate', str(reg), '--json'])
    out = capsys.readouterr().out
    assert code == 0
    assert '"valid": true' in out
    
    # 5. Remove
    code = run(['evidence', 'remove', str(reg), 'ev1', '--json'])
    assert code == 0

def test_verify_orchestrator_cli(tmp_path, capsys):
    out = tmp_path / 'test_doc.docx'
    run(['scaffold', str(out), '--profile', 'generic', '--variant', 'default', '--title', 'Final Report', '--json'])
    capsys.readouterr()
    
    d = Document(out)
    for p in d.paragraphs:
        if '[Document content]' in p.text or '[Section content...]' in p.text or '[Front matter content]' in p.text:
            p.text = 'Completed report content without draft placeholders.'
    d.save(out)
    
    code = run(['verify', str(out), '--profile', 'generic', '--variant', 'default', '--json'])
    out_txt = capsys.readouterr().out
    assert code == 0
    assert '"valid": true' in out_txt
    assert 'inventory' in out_txt
