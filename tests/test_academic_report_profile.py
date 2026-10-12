from __future__ import annotations

from pathlib import Path
import hashlib

import pytest

from docx import Document
from docx.oxml.ns import qn

from wordassistant.document.scaffold import scaffold_document
from wordassistant.document.compose import compose_document
from wordassistant.document.integrity import audit_integrity
from wordassistant.document.validate import validate_document
from wordassistant.document.structure import inspect_document
from wordassistant.operations.citations import list_citations
from wordassistant.evidence.registry import EvidenceRegistry
from wordassistant.profiles.loader import load_profile


ROOT = Path(__file__).resolve().parents[1]
PROFILE_PATH = ROOT / "profiles" / "hcm-university-dacn-2026-2027.json"


def test_report_profile_is_guide_only_and_encodes_required_order():
    profile = load_profile(PROFILE_PATH)
    report = profile.variant("report")

    assert profile.source_assets == (
        "use_cases/hcmunre_dacn_2026_2027/assets/huong_dan_trinh_bay_dacn_2026_2027.docx",
    )
    assert report.page.size == "A4"
    assert (
        report.page.margins.top_cm,
        report.page.margins.bottom_cm,
        report.page.margins.left_cm,
        report.page.margins.right_cm,
    ) == (2, 2, 3, 2)
    assert report.typography.body_font == "Times New Roman"
    assert report.typography.body_size_pt == 13
    assert report.typography.line_spacing == 1.5
    assert report.header_footer["front_matter_page_number_format"] == "none"
    assert report.citations["minimum_references"] == 10
    assert report.citations["require_reviewed_sources"] is True
    assert report.citations["bibliography_order"] == (
        "vietnamese_given_name_then_english_family_name"
    )
    assert [item["display"] for item in report.captions["figure_labels"]] == [
        "HÌNH",
        "SƠ ĐỒ",
    ]
    assert report.captions["table_label"]["display"] == "BẢNG"
    titles = [item["title"] for item in report.section_blueprint]
    assert titles == [
        "MỞ ĐẦU",
        "TỔNG QUAN TÀI LIỆU",
        "PHƯƠNG PHÁP THỰC HIỆN",
        "CÀI ĐẶT THỬ NGHIỆM, KẾT QUẢ VÀ ĐÁNH GIÁ",
        "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN",
        "PHỤ LỤC",
        "DANH MỤC TÀI LIỆU THAM KHẢO",
    ]


def test_scaffold_persists_word_styles_sections_and_cover_border(tmp_path):
    profile = load_profile(PROFILE_PATH)
    output = tmp_path / "report.docx"
    scaffold_document(
        output,
        profile,
        "report",
        title="MẪU ĐỀ TÀI ĐỒ ÁN",
        student="Nguyễn Văn A",
        supervisor="TS. Trần Văn B",
        cover_metadata={
            "student_id": "123456",
            "class": "DHTH",
            "cohort": "2023–2027",
            "month": "6",
            "year": "2027",
        },
    )

    document = Document(output)
    assert len(document.sections) == 3
    assert len(document.tables) >= 2  # schedule and abbreviations; covers stay table-free
    assert all(
        item["geometry_status"] == "PASS"
        for item in inspect_document(output)["tables_detail"]
    )

    normal = document.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size.pt == 13
    assert normal.paragraph_format.line_spacing == 1.5
    assert document.styles["Heading 1"].font.size.pt == 16
    for name in (
        "WA Cover University",
        "WA Cover Faculty",
        "WA Cover Document",
        "WA Cover Title",
        "WA Cover Metadata Line",
        "WA Cover Location Year",
        "WA Front Matter Title",
        "WA References",
        "WA Reference Language Heading",
    ):
        style = document.styles[name]
        assert style.hidden is False
        assert style.quick_style is True

    border = document.sections[0]._sectPr.find(qn("w:pgBorders"))
    assert border is not None
    assert border.get(qn("w:offsetFrom")) == "page"
    assert border.get(qn("w:display")) == "firstPage"
    assert border.find(qn("w:top")).get(qn("w:val")) == "thinThickSmallGap"
    assert border.find(qn("w:bottom")).get(qn("w:val")) == "thickThinSmallGap"
    assert all(
        border.find(qn(f"w:{edge}")).get(qn("w:sz")) == "24"
        and border.find(qn(f"w:{edge}")).get(qn("w:space")) == "24"
        for edge in ("top", "left", "bottom", "right")
    )
    assert document.sections[1]._sectPr.find(qn("w:pgBorders")) is None
    assert document.sections[2]._sectPr.find(qn("w:pgBorders")) is None

    # No layout table appears before the first section break (the two covers).
    body = document.element.body
    for child in body:
        if child.tag == qn("w:tbl"):
            raise AssertionError("Cover layout must not contain a Word table")
        if child.tag == qn("w:p") and child.find(f".//{qn('w:sectPr')}") is not None:
            break

    front_footer = document.sections[1].footer._element.xml
    body_footer = document.sections[2].footer._element.xml
    assert "PAGE" not in front_footer
    assert "PAGE" in body_footer
    assert "TOC" in document._element.xml
    assert '&quot;Hinh&quot;' in document._element.xml
    assert '&quot;SoDo&quot;' in document._element.xml
    assert '&quot;Bang&quot;' in document._element.xml
    assert "[MSSV]" not in "\n".join(p.text for p in document.paragraphs)
    assert "123456" in "\n".join(p.text for p in document.paragraphs)

    heading_text = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert "TỔNG QUAN TÀI LIỆU" in heading_text
    assert heading_text[-1] == "DANH MỤC TÀI LIỆU THAM KHẢO"
    assert document.settings.element.find(qn("w:updateFields")) is not None
    validation = validate_document(output, profile, "report", verify_mode=False)
    assert not any(issue.code == "missing_front_matter" for issue in validation.issues)
    audit = audit_integrity(output, profile=profile, variant_name="report")
    assert not any(
        finding.get("text") == "hồ chí minh"
        for finding in audit["gates"]["repetitive_prose_review"]["findings"]
    )


def test_profile_composition_sorts_bibliography_and_matches_numeric_citations(tmp_path):
    profile = load_profile(PROFILE_PATH)
    scaffold = tmp_path / "scaffold.docx"
    output = tmp_path / "report.docx"
    registry_path = tmp_path / ".wordassistant" / "evidence.json"
    scaffold_document(scaffold, profile, "report", title="Mẫu")
    registry = EvidenceRegistry(registry_path)
    for evidence_id, author, sort_key in (
        ("english", "Aho A.V.", "Aho"),
        ("viet_binh", "Nguyễn Bảo Bình", "Bình"),
        ("viet_an", "Nguyễn Văn An", "An"),
    ):
        source_bytes = f"%PDF-1.4\n{evidence_id} source excerpt".encode()
        source_path = registry_path.parent / "sources" / f"{evidence_id}.pdf"
        source_path.parent.mkdir(parents=True, exist_ok=True)
        source_path.write_bytes(source_bytes)
        registry.add({
            "id": evidence_id,
            "title": f"Source {evidence_id}",
            "authors": [author],
            "author_sort_key": sort_key,
            "language": "vi" if evidence_id.startswith("viet_") else "en",
            "container_title": "Journal of Tested Sources",
            "volume": "2",
            "issue": "1",
            "pages": "10-20",
            "publisher": "Example Publisher",
            "year": 2024,
            "url": f"https://example.org/{evidence_id}",
            "verified_by": "test fixture",
            "synthetic": False,
            "review_status": "reviewed",
            "reviewed_by": "test reviewer",
            "reviewed_at": "2026-10-12T00:00:00+00:00",
            "support_location": "Page 1",
            "support_excerpt": f"{evidence_id} source excerpt",
            "source_file": f"{evidence_id}.pdf",
            "source_sha256": hashlib.sha256(source_bytes).hexdigest(),
            "source_license": "CC-BY-4.0",
            "source_url": f"https://example.org/{evidence_id}.pdf",
            "retrieved_at": "2026-10-12T00:00:00+00:00",
            "claim": "Fixture claim supported by exact source excerpt.",
        })

    compose_document(
        scaffold,
        output_path=output,
        content={
            "version": 1,
            "sections": [{
                "heading": "MỞ ĐẦU",
                "paragraphs": [
                    {"text": "English source claim.", "citations": ["english"]},
                    {"text": "Vietnamese source claim.", "citations": ["viet_binh"]},
                    {"text": "Another Vietnamese source claim.", "citations": ["viet_an"]},
                ],
            }],
        },
        evidence_registry=registry_path,
        profile=PROFILE_PATH,
        variant_name="report",
    )

    citations = list_citations(output)
    assert [item["text"] for item in citations] == ["[3]", "[2]", "[1]"]
    document = Document(output)
    reference_heading = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text == "DANH MỤC TÀI LIỆU THAM KHẢO"
    )
    references = [paragraph.text for paragraph in document.paragraphs[reference_heading + 1:]]
    assert references[0] == "Tiếng Việt:"
    assert references[1].startswith("[1] Nguyễn Văn An. ")
    assert references[1].endswith("https://example.org/viet_an.")
    assert references[2].startswith("[2] Nguyễn Bảo Bình")
    assert references[3] == "Tiếng Anh:"
    assert references[4].startswith("[3] Aho A.V.")
    group_heading = document.paragraphs[reference_heading + 1]
    assert group_heading.text == "Tiếng Việt:"
    assert group_heading.style.name == "WA Reference Language Heading"
    reference_paragraph = document.paragraphs[reference_heading + 2]
    assert any(run.text == "Nguyễn Văn An" and run.bold is True for run in reference_paragraph.runs)
    assert any(run.text == "Source viet_an" and run.italic is True for run in reference_paragraph.runs)
    assert "In: Journal of Tested Sources" in references[1]
    assert "Vol. 2" in references[1]
    assert "No. 1" in references[1]
    assert "pp. 10-20" in references[1]
    assert "Example Publisher" in references[1]


def test_report_composition_rejects_records_without_reviewed_source(tmp_path):
    profile = load_profile(PROFILE_PATH)
    scaffold = tmp_path / "scaffold.docx"
    registry_path = tmp_path / ".wordassistant" / "evidence.json"
    scaffold_document(scaffold, profile, "report", title="Mẫu")
    EvidenceRegistry(registry_path).add({
        "id": "legacy-source",
        "title": "Source title",
        "authors": ["Author Name"],
        "year": 2024,
        "url": "https://example.org/source",
        "verified_by": "old workflow",
        "synthetic": False,
    })

    with pytest.raises(ValueError, match="no reviewed, downloaded source"):
        compose_document(
            scaffold,
            output_path=tmp_path / "report.docx",
            content={
                "version": 1,
                "sections": [{
                    "heading": "MỞ ĐẦU",
                    "paragraphs": [{
                        "text": "Claim requiring a reviewed source.",
                        "citations": ["legacy-source"],
                    }],
                }],
            },
            evidence_registry=registry_path,
            profile=PROFILE_PATH,
            variant_name="report",
        )


@pytest.mark.parametrize(
    "override",
    [
        {"citation_style": "numeric"},
        {"bibliography_order": "citation_order"},
    ],
)
def test_report_composition_does_not_allow_citation_policy_overrides(tmp_path, override):
    profile = load_profile(PROFILE_PATH)
    scaffold = tmp_path / "scaffold.docx"
    scaffold_document(scaffold, profile, "report", title="Mẫu")
    content = {
        "version": 1,
        "sections": [{
            "heading": "MỞ ĐẦU",
            "paragraphs": ["Text without invented claims."],
        }],
        **override,
    }

    with pytest.raises(ValueError, match="cannot override"):
        compose_document(
            scaffold,
            output_path=tmp_path / "report.docx",
            content=content,
            profile=PROFILE_PATH,
            variant_name="report",
        )
