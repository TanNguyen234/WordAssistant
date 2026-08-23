from __future__ import annotations
from pathlib import Path
from docx import Document
from wordassistant.document.validate import validate_document
from wordassistant.profiles.loader import resolve_profile
from wordassistant.document.scaffold import scaffold_document

def test_blueprint_heading_level_mismatch_fails_verify(tmp_path: Path):
    doc_path = tmp_path / "level_mismatch.docx"
    prof = resolve_profile("hcmunre_dacn_2026_2027")
    scaffold_document(
        doc_path,
        prof,
        "report",
        title="XAY DUNG HE THONG RAG CHO VAN BAN PHAP LUAT",
        student="NGUYEN VAN A - 2351010001",
        supervisor="TS. NGUYEN VAN B"
    )
    d = Document(doc_path)
    # Replace placeholder content
    for p in d.paragraphs:
        if "[Front matter content]" in p.text:
            p.text = "Lời cảm ơn."
        elif "[Section content...]" in p.text:
            p.text = "Nội dung chương."
        elif "[List of figures]" in p.text:
            p.text = "Hình 1.1."
        elif "[List of tables]" in p.text:
            p.text = "Bảng 1.1."
            
    # Change CHƯƠNG 1 (level 1 in blueprint) to style Heading 2 (level 2)
    for p in d.paragraphs:
        if "CHƯƠNG 1" in p.text.upper():
            p.style = d.styles["Heading 2"]
            
    for i in range(1, 12):
        d.add_paragraph(f"[{i}] Tac gia {i}, Sach {i}, 2024.")
    d.save(doc_path)
    
    res = validate_document(doc_path, prof, "report", verify_mode=True)
    assert res.valid is False
    codes = [i.code for i in res.issues if i.severity == "error"]
    assert "blueprint_heading_level_mismatch" in codes or any("level" in c for c in codes)
