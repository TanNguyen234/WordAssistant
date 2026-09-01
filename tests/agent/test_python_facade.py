from docx import Document

from wordassistant import WordDocument
from wordassistant.execution import ExecutionService


def test_facade_table_operation_uses_public_namespace(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    document = WordDocument.open(source)

    result = document.tables.create(
        rows=2, columns=3, output=tmp_path / "table.docx"
    )

    assert result.backend_used == "python-docx"
    assert len(Document(result.output_path).tables[0].columns) == 3


def test_facade_inspect_is_bounded_public_read(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_heading("Report", 1)
    doc.save(source)

    assert WordDocument.open(source).inspect()["paragraphs"] == 1


def test_facade_mutation_uses_execution_service(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)
    called = []

    class Service:
        def execute(self, request):
            called.append(request)
            return type("Result", (), {"output_path": None})()

    monkeypatch.setattr(ExecutionService, "default", classmethod(lambda cls: Service()))
    WordDocument.open(source).execute(
        "table.create",
        output=tmp_path / "output.docx",
        rows=1,
        columns=1,
    )

    assert called[0].name == "table.create"
