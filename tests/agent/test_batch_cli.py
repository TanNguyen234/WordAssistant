from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys

from docx import Document


def test_public_apply_command_runs_atomic_plan_on_a_real_docx(tmp_path: Path):
    source = tmp_path / "input.docx"
    output = tmp_path / "final.docx"
    plan = tmp_path / "operations.json"
    document = Document()
    document.add_paragraph("Public CLI")
    document.save(source)
    plan.write_text(
        json.dumps(
            {
                "version": 1,
                "steps": [
                    {
                        "operation": "text.edit",
                        "selector": {"type": "paragraph", "index": 0},
                        "options": {"action": "append", "text": " verified"},
                    },
                    {
                        "operation": "format.paragraph",
                        "selector": {"type": "paragraph", "index": 0},
                        "options": {"keep_with_next": True},
                    },
                ],
            }
        ),
        encoding="utf-8",
    )

    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "wordassistant.cli",
            "apply",
            str(source),
            "--output",
            str(output),
            "--plan",
            str(plan),
            "--json",
        ],
        cwd=tmp_path,
        text=True,
        encoding="utf-8",
        capture_output=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr or completed.stdout
    payload = json.loads(completed.stdout)
    result = Document(output)
    assert payload["status"] == "SUCCESS"
    assert len(payload["steps"]) == 2
    assert result.paragraphs[0].text == "Public CLI verified"
    assert result.paragraphs[0].paragraph_format.keep_with_next is True
    assert sorted(path.name for path in tmp_path.glob("*.docx")) == [
        "final.docx",
        "input.docx",
    ]
