import json

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from wordassistant.cli import run


@pytest.mark.parametrize("batch", [False, True])
def test_public_image_caption_accepts_localized_sequence(tmp_path, capsys, batch):
    source = tmp_path / "source.docx"
    output = tmp_path / "captioned.docx"
    picture = tmp_path / "figure.png"
    Document().save(source)
    Image.new("RGB", (20, 20), "white").save(picture)
    options = {
        "image_path": str(picture),
        "caption": "Kiến trúc",
        "caption_label": "HÌNH",
        "caption_sequence": "Hinh",
    }
    if batch:
        plan = tmp_path / "plan.json"
        plan.write_text(json.dumps({"version": 1, "steps": [
            {"operation": "image.insert", "options": options}
        ]}), encoding="utf-8")
        args = ["apply", str(source), "--output", str(output), "--plan", str(plan)]
    else:
        args = ["image", "insert", str(source), "--output", str(output),
                "--options", json.dumps(options)]

    code = run([*args, "--json"])
    payload = json.loads(capsys.readouterr().out)
    assert code == 0, payload
    document = Document(output)
    assert len(document.inline_shapes) == 1
    instructions = [node.text or "" for node in document.element.body.iter(qn("w:instrText"))]
    instructions.extend(node.get(qn("w:instr"), "")
                        for node in document.element.body.iter(qn("w:fldSimple")))
    assert any("SEQ Hinh" in instruction for instruction in instructions)
    assert any("HÌNH" in paragraph.text for paragraph in document.paragraphs)
