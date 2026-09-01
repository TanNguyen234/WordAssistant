from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from docx import Document

from wordassistant.cli import run
from wordassistant.execution import ExecutionService
from wordassistant.operations.models import BatchPlan, BatchStep, OperationUnavailable
from wordassistant.transactions import CheckpointStore


def _document(path: Path, text: str) -> Path:
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    return path


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_plan(path: Path, steps: list[dict]) -> Path:
    path.write_text(
        json.dumps({"version": 1, "steps": steps}, ensure_ascii=False),
        encoding="utf-8",
    )
    return path


def test_apply_cli_executes_real_docx_steps_and_publishes_only_final(
    tmp_path, capsys
):
    source = _document(tmp_path / "input.docx", "alpha")
    output = _document(tmp_path / "final.docx", "pre-existing")
    source_before = source.read_bytes()
    plan = _write_plan(
        tmp_path / "operations.json",
        [
            {
                "operation": "text.edit",
                "selector": {"type": "paragraph", "index": 0},
                "options": {"action": "append", "text": " beta"},
            },
            {
                "operation": "format.run",
                "selector": {"type": "paragraph", "index": 0},
                "options": {"bold": True},
            },
        ],
    )

    exit_code = run(
        [
            "apply",
            str(source),
            "--output",
            str(output),
            "--plan",
            str(plan),
            "--json",
        ]
    )

    payload = json.loads(capsys.readouterr().out)
    result = Document(output)
    assert exit_code == 0
    assert source.read_bytes() == source_before
    assert result.paragraphs[0].text == "alpha beta"
    assert result.paragraphs[0].runs[0].bold is True
    assert payload["status"] == "SUCCESS"
    assert payload["operation"] == "batch.apply"
    assert [step["operation"] for step in payload["steps"]] == [
        "text.edit",
        "format.run",
    ]
    assert payload["sha256"] == _sha256(output)
    assert "staging" not in json.dumps(payload).casefold()
    assert sorted(path.name for path in tmp_path.glob("*.docx")) == [
        "final.docx",
        "input.docx",
    ]


def test_apply_failure_rolls_back_real_docx_chain_and_removes_staging(
    tmp_path, capsys
):
    source = _document(tmp_path / "input.docx", "alpha")
    output = _document(tmp_path / "final.docx", "keep me")
    source_before = source.read_bytes()
    output_before = output.read_bytes()
    plan = _write_plan(
        tmp_path / "operations.json",
        [
            {
                "operation": "text.edit",
                "selector": {"type": "paragraph", "index": 0},
                "options": {"action": "append", "text": " changed"},
            },
            {
                "operation": "format.run",
                "selector": {"type": "paragraph", "index": 0},
                "options": {"highlight": "not-a-word-color"},
            },
        ],
    )

    exit_code = run(
        [
            "apply",
            str(source),
            "--output",
            str(output),
            "--plan",
            str(plan),
            "--json",
        ]
    )

    payload = json.loads(capsys.readouterr().out)
    assert exit_code == 1
    assert payload["status"] == "ERROR"
    assert payload["error"]["code"] == "BATCH_STEP_FAILED"
    assert payload["error"]["step_index"] == 1
    assert len(payload["error"]["completed_steps"]) == 1
    assert source.read_bytes() == source_before
    assert output.read_bytes() == output_before
    assert sorted(path.name for path in tmp_path.glob("*.docx")) == [
        "final.docx",
        "input.docx",
    ]


@pytest.mark.parametrize(
    "operation",
    [
        "batch.apply",
        "render.pdf",
        "compare.documents",
        "merge.documents",
        "mail-merge.execute",
        "field.inspect",
    ],
)
def test_batch_rejects_non_composable_operation_before_checkpoint(
    tmp_path, operation
):
    source = _document(tmp_path / "input.docx", "alpha")
    output = _document(tmp_path / "final.docx", "keep me")
    output_before = output.read_bytes()
    checkpoint_store = CheckpointStore(tmp_path / "checkpoints")
    service = ExecutionService.default()
    service.checkpoint_store = checkpoint_store
    plan = BatchPlan.from_dict(
        {"version": 1, "steps": [{"operation": operation, "options": {}}]}
    )

    with pytest.raises((KeyError, ValueError), match="batch|composable|Unknown"):
        service.execute_batch(
            source,
            output,
            plan,
            checkpoint="auto",
        )

    assert output.read_bytes() == output_before
    assert not (tmp_path / "checkpoints").exists()
    assert sorted(path.name for path in tmp_path.glob("*.docx")) == [
        "final.docx",
        "input.docx",
    ]


def test_batch_rejects_unproven_operation_before_checkpoint(tmp_path):
    source = _document(tmp_path / "input.docx", "alpha")
    output = _document(tmp_path / "final.docx", "keep me")
    output_before = output.read_bytes()
    checkpoint_store = CheckpointStore(tmp_path / "checkpoints")
    service = ExecutionService.default()
    service.checkpoint_store = checkpoint_store
    plan = BatchPlan.from_dict(
        {
            "version": 1,
            "steps": [
                    {
                        "operation": "style.clone",
                        "options": {"source_name": "Normal", "name": "Copy"},
                    }
            ],
        }
    )

    with pytest.raises(OperationUnavailable) as raised:
        service.execute_batch(source, output, plan, checkpoint=True)

    assert raised.value.error.code == "EVIDENCE_UNVERIFIED"
    assert output.read_bytes() == output_before
    assert not (tmp_path / "checkpoints").exists()


def test_batch_validates_all_step_options_before_checkpoint_or_staging(tmp_path):
    source = _document(tmp_path / "input.docx", "alpha")
    output = _document(tmp_path / "final.docx", "keep me")
    output_before = output.read_bytes()
    checkpoint_store = CheckpointStore(tmp_path / "checkpoints")
    service = ExecutionService.default()
    service.checkpoint_store = checkpoint_store
    plan = BatchPlan.from_dict(
        {
            "version": 1,
            "steps": [
                {
                    "operation": "text.edit",
                    "options": {"action": "append", "text": " would mutate"},
                },
                {"operation": "table.create", "options": {"rows": 2}},
            ],
        }
    )

    with pytest.raises(ValueError, match="Missing required options.*columns"):
        service.execute_batch(source, output, plan, checkpoint=True)

    assert output.read_bytes() == output_before
    assert not (tmp_path / "checkpoints").exists()
    assert sorted(path.name for path in tmp_path.glob("*.docx")) == [
        "final.docx",
        "input.docx",
    ]


def test_batch_plan_rejects_per_step_control_fields():
    with pytest.raises(ValueError, match="Unsupported batch step keys"):
        BatchPlan.from_dict(
            {
                "version": 1,
                "steps": [
                    {
                        "operation": "text.edit",
                        "output_path": "intermediate.docx",
                        "options": {"action": "append", "text": "x"},
                    }
                ],
            }
        )

    with pytest.raises(ValueError, match="reserved control options"):
        BatchPlan(
            version=1,
            steps=(
                BatchStep(
                    operation="text.edit",
                    options={"action": "append", "text": "x", "checkpoint": True},
                ),
            ),
        )
    with pytest.raises(ValueError, match="reserved control options"):
        BatchPlan.from_dict(
            {
                "version": 1,
                "steps": [
                    {
                        "operation": "text.edit",
                        "options": {
                            "action": "append",
                            "text": "x",
                            "dry_run": True,
                        },
                    }
                ],
            }
        )


def test_batch_creates_exactly_one_optional_outer_checkpoint(tmp_path):
    source = _document(tmp_path / "input.docx", "alpha")
    output = tmp_path / "final.docx"
    checkpoint_store = CheckpointStore(tmp_path / "checkpoints")
    service = ExecutionService.default()
    service.checkpoint_store = checkpoint_store
    plan = BatchPlan.from_dict(
        {
            "version": 1,
            "steps": [
                {
                    "operation": "text.edit",
                    "options": {"action": "append", "text": " beta"},
                },
                {
                    "operation": "format.paragraph",
                    "options": {"alignment": "center"},
                },
            ],
        }
    )

    result = service.execute_batch(source, output, plan, checkpoint="auto")

    checkpoints = checkpoint_store.list()
    assert len(checkpoints) == 1
    assert checkpoints[0].operation == "batch.apply"
    assert result.checkpoint_created is True
    assert Document(output).paragraphs[0].text == "alpha beta"
