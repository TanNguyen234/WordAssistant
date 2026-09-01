from pathlib import Path

import pytest
from docx import Document

from wordassistant.backends.base import BackendHealth, MappedOperationBackend
from wordassistant.backends.router import BackendRouter
from wordassistant.execution import ExecutionService
from wordassistant.operations.models import (
    Fidelity,
    OperationRequest,
    OperationResult,
    OperationUnavailable,
)
from wordassistant.operations.registry import OperationRegistry, OperationSpec
from wordassistant.transactions import CheckpointStore


def _source(path: Path) -> Path:
    document = Document()
    document.add_paragraph("original")
    document.save(path)
    return path


class _Backend(MappedOperationBackend):
    fidelity = Fidelity.HIGH

    def health(self):
        return BackendHealth(self.name, True, True, "test backend")


def _service(tmp_path, handler, *, mutating=True):
    registry = OperationRegistry()
    registry.register(
        OperationSpec(
            "text.edit",
            mutating,
            "text.edit",
            ("python-docx",),
            Fidelity.HIGH,
        )
    )
    router = BackendRouter(
        registry, [_Backend("python-docx", {"text.edit": handler})]
    )
    return ExecutionService(
        registry, router, CheckpointStore(tmp_path / "checkpoints")
    )


def test_mutation_requires_distinct_output(tmp_path):
    source = _source(tmp_path / "source.docx")
    service = _service(tmp_path, lambda request: None)

    with pytest.raises(ValueError, match="distinct output"):
        service.execute(OperationRequest("text.edit", input_path=source))
    with pytest.raises(ValueError, match="distinct output"):
        service.execute(
            OperationRequest("text.edit", input_path=source, output_path=source)
        )


def test_dry_run_is_a_side_effect_free_plan(tmp_path):
    source = _source(tmp_path / "source.docx")
    output = tmp_path / "new-parent" / "output.docx"
    original = source.read_bytes()
    backend_calls = 0

    def mutate(request):
        nonlocal backend_calls
        backend_calls += 1
        document = Document(request.input_path)
        document.paragraphs[0].text = "changed"
        document.save(request.output_path)
        return OperationResult.success(
            request.name,
            "python-docx",
            Fidelity.HIGH,
            output_path=request.output_path,
        )

    result = _service(tmp_path, mutate).execute(
        OperationRequest(
            "text.edit",
            input_path=source,
            output_path=output,
            dry_run=True,
        )
    )

    assert result.status == "DRY_RUN"
    assert result.output_path is None
    assert source.read_bytes() == original
    assert not output.exists()
    assert not output.parent.exists()
    assert backend_calls == 0
    assert not (tmp_path / "checkpoints").exists()


def test_unsupported_selector_is_rejected_before_dry_run(tmp_path):
    source = _source(tmp_path / "source.docx")
    registry = OperationRegistry()
    registry.register(
        OperationSpec(
            "comment.create", True, "comment.create", ("word-native",), Fidelity.NATIVE
        )
    )
    router = BackendRouter(
        registry, [_Backend("word-native", {"comment.create": lambda request: None})]
    )
    service = ExecutionService(
        registry, router, CheckpointStore(tmp_path / "checkpoints")
    )

    with pytest.raises(ValueError, match="does not support --selector"):
        service.execute(
            OperationRequest(
                "comment.create",
                input_path=source,
                output_path=tmp_path / "output.docx",
                selector={"type": "document"},
                dry_run=True,
            )
        )


def test_explicit_kwargs_contract_rejects_typo_before_dry_run(tmp_path):
    source = _source(tmp_path / "source.docx")

    with pytest.raises(ValueError, match="font_szie_pt"):
        ExecutionService.default().execute(
            OperationRequest(
                "format.run",
                input_path=source,
                output_path=tmp_path / "output.docx",
                options={"font_szie_pt": 12},
                dry_run=True,
            )
        )

    with pytest.raises(ValueError, match="must be number"):
        ExecutionService.default().execute(
            OperationRequest(
                "format.run",
                input_path=source,
                output_path=tmp_path / "output.docx",
                options={"font_size_pt": "12"},
                dry_run=True,
            )
        )


def test_transaction_rolls_back_existing_output_on_backend_failure(tmp_path):
    source = _source(tmp_path / "source.docx")
    output = _source(tmp_path / "output.docx")
    original_source = source.read_bytes()
    original_output = output.read_bytes()

    def fail(request):
        Path(request.output_path).write_bytes(b"corrupt")
        raise RuntimeError("backend failed")

    with pytest.raises(RuntimeError, match="backend failed"):
        _service(tmp_path, fail).execute(
            OperationRequest("text.edit", input_path=source, output_path=output)
        )

    assert source.read_bytes() == original_source
    assert output.read_bytes() == original_output


def test_output_artifact_is_atomically_published(tmp_path):
    source = _source(tmp_path / "source.docx")
    output = _source(tmp_path / "output.docx")
    original_output = output.read_bytes()
    registry = OperationRegistry()
    registry.register(
        OperationSpec(
            "compare.documents",
            True,
            "compare.create",
            ("word-native",),
            Fidelity.NATIVE,
        )
    )

    class Backend(_Backend):
        fidelity = Fidelity.NATIVE

    def fail(request):
        Path(request.output_path).write_bytes(b"partial")
        raise RuntimeError("compare failed")

    router = BackendRouter(
        registry, [Backend("word-native", {"compare.documents": fail})]
    )
    service = ExecutionService(
        registry, router, CheckpointStore(tmp_path / "checkpoints")
    )

    with pytest.raises(RuntimeError, match="compare failed"):
        service.execute(
            OperationRequest(
                "compare.documents",
                input_path=source,
                output_path=output,
            )
        )

    assert output.read_bytes() == original_output


def test_unverified_core_operation_is_not_executable(tmp_path):
    source = _source(tmp_path / "source.docx")
    registry = OperationRegistry()
    registry.register(
        OperationSpec(
            "style.clone",
            True,
            "style.clone",
            ("python-docx",),
            Fidelity.HIGH,
        )
    )
    router = BackendRouter(
        registry,
        [_Backend("python-docx", {"style.clone": lambda request: None})],
    )

    with pytest.raises(OperationUnavailable) as raised:
        ExecutionService(registry, router).execute(
            OperationRequest(
                "style.clone",
                input_path=source,
                output_path=tmp_path / "output.docx",
            )
        )

    assert getattr(raised.value, "error").code == "EVIDENCE_UNVERIFIED"


def test_compare_output_must_differ_from_both_inputs(tmp_path):
    source = _source(tmp_path / "source.docx")
    other = _source(tmp_path / "other.docx")
    service = ExecutionService.default()

    with pytest.raises(ValueError, match="input artifact"):
        service.execute(
            OperationRequest(
                "compare.documents",
                input_path=source,
                output_path=other,
                options={"other_path": str(other)},
            )
        )


def test_document_mutation_rejects_non_docx_output(tmp_path):
    source = _source(tmp_path / "source.docx")
    service = _service(tmp_path, lambda request: None)
    with pytest.raises(ValueError, match=".docx"):
        service.execute(
            OperationRequest(
                "text.edit", input_path=source, output_path=tmp_path / "output.bin"
            )
        )


def test_field_sidecar_failure_is_a_warning_after_successful_commit(tmp_path, monkeypatch):
    source = _source(tmp_path / "source.docx")
    output = _source(tmp_path / "output.docx")
    registry = OperationRegistry()
    registry.register(
        OperationSpec(
            "field.update_all",
            True,
            "fields.update",
            ("word-native",),
            Fidelity.NATIVE,
        )
    )

    class NativeBackend(_Backend):
        fidelity = Fidelity.NATIVE

    def update(request):
        document = Document(request.input_path)
        document.paragraphs[0].text = "updated"
        document.save(request.output_path)
        return OperationResult.success(
            request.name, "word-native", Fidelity.NATIVE, output_path=request.output_path
        )

    service = ExecutionService(
        registry,
        BackendRouter(registry, [NativeBackend("word-native", {"field.update_all": update})]),
        CheckpointStore(tmp_path / "checkpoints"),
    )
    real_replace = __import__("os").replace

    def fail_sidecar(source_path, destination_path):
        if str(destination_path).endswith(".wa-evidence.json"):
            raise PermissionError("sidecar denied")
        return real_replace(source_path, destination_path)

    monkeypatch.setattr("wordassistant.execution.os.replace", fail_sidecar)

    result = service.execute(
        OperationRequest("field.update_all", input_path=source, output_path=output)
    )

    assert Document(output).paragraphs[0].text == "updated"
    assert result.status == "SUCCESS"
    assert any("sidecar" in warning for warning in result.warnings)
