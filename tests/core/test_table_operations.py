from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn

from wordassistant.operations.tables import (
    create_table,
    delete_column,
    insert_column,
    insert_row,
    merge_cells,
    set_cell_text,
    set_table_layout,
)


def _source(path):
    doc = Document()
    doc.add_heading("Data", level=1)
    doc.save(path)
    return path


def _geometry(table):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    layout = tbl_pr.find(qn("w:tblLayout"))
    width = int(tbl_w.get(qn("w:w")))
    grid = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.gridCol_lst]
    row_widths = []
    for row in table._tbl.tr_lst:
        row_widths.append(
            sum(int(cell.tcPr.tcW.get(qn("w:w"))) for cell in row.tc_lst)
        )
    return {
        "width_type": tbl_w.get(qn("w:type")),
        "width": width,
        "indent": int(tbl_ind.get(qn("w:w"))),
        "layout": layout.get(qn("w:type")),
        "grid": grid,
        "rows": row_widths,
    }


def test_table_create_edit_layout_and_merge_roundtrip(tmp_path):
    result = create_table(
        _source(tmp_path / "source.docx"),
        output_path=tmp_path / "created.docx",
        rows=2,
        columns=2,
        data=[["A", "B"], ["C", "D"]],
        caption="Table 1. Results",
    )
    result = insert_row(result.output_path, output_path=tmp_path / "row.docx", table_index=0, index=1, values=["E", "F"])
    result = insert_column(result.output_path, output_path=tmp_path / "column.docx", table_index=0, index=1, values=["X", "Y", "Z"])
    result = set_cell_text(result.output_path, output_path=tmp_path / "cell.docx", table_index=0, row=2, column=2, text="Final")
    result = merge_cells(result.output_path, output_path=tmp_path / "merged.docx", table_index=0, start=(0, 0), end=(0, 1))
    result = set_table_layout(result.output_path, output_path=tmp_path / "layout.docx", table_index=0, repeat_header=True, alignment="center", autofit=False)

    document = Document(result.output_path)
    table = document.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 3
    assert table.cell(2, 2).text == "Final"
    assert "A" in table.cell(0, 0).text and "X" in table.cell(0, 0).text
    assert "w:tblHeader" in table.rows[0]._tr.xml
    assert "w:tblCaption" in table._tbl.tblPr.xml
    assert table.autofit is False
    geometry = _geometry(table)
    assert geometry["width_type"] == "dxa"
    assert geometry["indent"] == 0
    assert geometry["layout"] == "fixed"
    assert sum(geometry["grid"]) == geometry["width"]
    assert geometry["rows"] == [geometry["width"]] * 3


def test_delete_column_updates_every_row(tmp_path):
    result = create_table(
        _source(tmp_path / "source.docx"),
        output_path=tmp_path / "created.docx",
        rows=2,
        columns=3,
        data=[["A", "B", "C"], ["D", "E", "F"]],
    )

    delete_column(result.output_path, output_path=tmp_path / "deleted.docx", table_index=0, index=1)

    table = Document(tmp_path / "deleted.docx").tables[0]
    assert [[cell.text for cell in row.cells] for row in table.rows] == [["A", "C"], ["D", "F"]]
    geometry = _geometry(table)
    assert sum(geometry["grid"]) == geometry["width"]
    assert geometry["rows"] == [geometry["width"]] * 2


def test_table_width_uses_the_section_that_contains_the_table(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    first = document.sections[0]
    first.left_margin = first.right_margin = 914400
    document.add_table(rows=1, cols=2)
    second = document.add_section(WD_SECTION_START.NEW_PAGE)
    second.orientation = 1
    second.page_width, second.page_height = second.page_height, second.page_width
    document.save(source)

    result = set_table_layout(
        source,
        output_path=tmp_path / "layout.docx",
        table_index=0,
        autofit=False,
    )
    reopened = Document(result.output_path)
    geometry = _geometry(reopened.tables[0])
    expected = int(
        reopened.sections[0].page_width.twips
        - reopened.sections[0].left_margin.twips
        - reopened.sections[0].right_margin.twips
    )
    assert geometry["width"] <= expected


def test_table_can_be_placed_after_heading_boundary(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("Data", level=1)
    document.add_paragraph("Text after table")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Existing"
    document.save(source)

    result = create_table(
        source,
        output_path=tmp_path / "placed.docx",
        rows=1,
        columns=1,
        data=[["Placed"]],
        selector={"type": "heading", "text": "Data"},
    )

    reopened = Document(result.output_path)
    body = list(reopened.element.body.iterchildren())
    assert [element.tag.rsplit("}", 1)[-1] for element in body[:3]] == [
        "p",
        "tbl",
        "p",
    ]
    assert reopened.tables[0].cell(0, 0).text == "Placed"
    assert result.data["table_index"] == 0
