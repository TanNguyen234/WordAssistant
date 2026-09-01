from docx import Document

from wordassistant.evidence.registry import EvidenceRegistry
from wordassistant.operations.citations import (
    generate_bibliography,
    insert_citation,
    list_citations,
    remove_citation,
)


def _registry(path):
    registry = EvidenceRegistry(path)
    for index in (1, 2):
        registry.add(
            {
                "id": f"src-{index}",
                "title": f"Verified Source {index}",
                "authors": [f"Author {index}"],
                "year": 2024,
                "url": f"https://example.com/{index}",
                "verified_by": "test",
                "synthetic": False,
            }
        )
    return path


def test_evidence_linked_numeric_citations_and_bibliography(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    evidence = _registry(tmp_path / "evidence.json")
    first = insert_citation(source, output_path=tmp_path / "one.docx", evidence_registry=evidence, evidence_id="src-1", style="numeric")
    second = insert_citation(first.output_path, output_path=tmp_path / "two.docx", evidence_registry=evidence, evidence_id="src-2", style="numeric")
    bibliography = generate_bibliography(second.output_path, output_path=tmp_path / "bibliography.docx", evidence_registry=evidence, style="IEEE")

    text = "\n".join(paragraph.text for paragraph in Document(bibliography.output_path).paragraphs)
    assert "[1]" in text and "[2]" in text
    citations = list_citations(bibliography.output_path)
    assert [item["evidence_id"] for item in citations] == ["src-1", "src-2"]
    assert "Verified Source 1" in text and "https://example.com/2" in text

    removed = remove_citation(bibliography.output_path, output_path=tmp_path / "removed.docx", citation_index=0)
    remaining = list_citations(removed.output_path)
    assert [item["evidence_id"] for item in remaining] == ["src-2"]
    assert remaining[0]["text"] == "[1]"


def test_citation_can_be_anchored_to_an_existing_paragraph(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Claim supported by evidence")
    document.add_paragraph("Following paragraph")
    document.save(source)
    evidence = _registry(tmp_path / "evidence.json")

    result = insert_citation(
        source,
        output_path=tmp_path / "anchored.docx",
        evidence_registry=evidence,
        evidence_id="src-1",
        style="numeric",
        selector={"type": "text", "text": "Claim supported by evidence"},
    )

    reopened = Document(result.output_path)
    first_xml = reopened.paragraphs[0]._p.xml
    assert "wordassistant:citation:src-1" in first_xml
    assert "[1]" in first_xml
    assert "wordassistant:citation" not in reopened.paragraphs[1]._p.xml


def test_synthetic_evidence_is_visibly_labeled(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Synthetic QA claim")
    document.save(source)
    evidence_path = tmp_path / "evidence.json"
    registry = EvidenceRegistry(evidence_path)
    registry.add(
        {
            "id": "synthetic-1",
            "title": "Synthetic QA Source",
            "authors": ["QA Fixture"],
            "year": 2026,
            "url": "https://example.invalid/synthetic-fixture",
            "verified_by": "test harness",
            "synthetic": True,
        }
    )

    cited = insert_citation(
        source,
        output_path=tmp_path / "cited.docx",
        evidence_registry=evidence_path,
        evidence_id="synthetic-1",
        selector={"type": "paragraph", "index": 0},
    )
    bibliography = generate_bibliography(
        cited.output_path,
        output_path=tmp_path / "bibliography.docx",
        evidence_registry=evidence_path,
    )

    xml = Document(cited.output_path).paragraphs[0]._p.xml
    text = "\n".join(p.text for p in Document(bibliography.output_path).paragraphs)
    assert "[SYNTHETIC-1]" in xml
    assert "[SYNTHETIC TEST DATA]" in text


def test_repeated_source_reuses_number_and_removal_renumbers_unique_sources(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("First claim")
    document.add_paragraph("Second claim")
    document.add_paragraph("Repeated first claim")
    document.save(source)
    evidence = _registry(tmp_path / "evidence.json")

    current = source
    for index, evidence_id in enumerate(("src-1", "src-2", "src-1")):
        result = insert_citation(
            current,
            output_path=tmp_path / f"cited-{index}.docx",
            evidence_registry=evidence,
            evidence_id=evidence_id,
            selector={"type": "paragraph", "index": index},
        )
        current = result.output_path

    assert [item["text"] for item in list_citations(current)] == [
        "[1]",
        "[2]",
        "[1]",
    ]
    removed = remove_citation(
        current,
        output_path=tmp_path / "removed.docx",
        citation_index=0,
    )
    assert [
        (item["evidence_id"], item["text"])
        for item in list_citations(removed.output_path)
    ] == [("src-2", "[1]"), ("src-1", "[2]")]
