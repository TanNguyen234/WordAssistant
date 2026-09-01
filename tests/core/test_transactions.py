from pathlib import Path

import pytest
from docx import Document

from wordassistant.operations.models import OperationRequest, OperationResult
from wordassistant.transactions import CheckpointStore, TransactionCoordinator


def _source(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("original")
    doc.save(path)
    return path


def test_failed_mutation_leaves_input_and_existing_output_untouched(tmp_path):
    source = _source(tmp_path / "source.docx")
    output = _source(tmp_path / "output.docx")
    original_source = source.read_bytes()
    original_output = output.read_bytes()
    coordinator = TransactionCoordinator(CheckpointStore(tmp_path / "checkpoints"))

    def fail(staging: Path) -> OperationResult:
        staging.write_bytes(b"corrupt")
        raise RuntimeError("mutation failed")

    request = OperationRequest(
        name="text.replace", input_path=source, output_path=output
    )
    with pytest.raises(RuntimeError, match="mutation failed"):
        coordinator.execute(request, fail)

    assert source.read_bytes() == original_source
    assert output.read_bytes() == original_output


def test_dry_run_does_not_create_output(tmp_path):
    source = _source(tmp_path / "source.docx")
    original = source.read_bytes()
    output = tmp_path / "output.docx"
    coordinator = TransactionCoordinator(CheckpointStore(tmp_path / "checkpoints"))

    def simulate(staging: Path) -> OperationResult:
        staging.write_bytes(b"simulated mutation")
        return OperationResult.success(
            "text.replace", "python-docx", "high", {"replacements": 1}
        )

    result = coordinator.execute(
        OperationRequest(
            name="text.replace",
            input_path=source,
            output_path=output,
            dry_run=True,
        ),
        simulate,
    )

    assert result.status == "DRY_RUN"
    assert source.read_bytes() == original
    assert not output.exists()


def test_automatic_checkpoint_records_hash_and_can_restore(tmp_path):
    source = _source(tmp_path / "source.docx")
    store = CheckpointStore(tmp_path / "checkpoints")
    checkpoint = store.create(source, operation="paragraph.format")
    source.write_bytes(b"changed")

    restored = store.restore(checkpoint.checkpoint_id, output_path=source)

    assert Document(restored).paragraphs[0].text == "original"
    assert store.list()[0].sha256 == checkpoint.sha256
