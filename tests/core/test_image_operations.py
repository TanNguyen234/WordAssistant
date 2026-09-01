import base64
from zipfile import ZipFile

from docx import Document

from wordassistant.operations.images import insert_image, remove_image, resize_image
from wordassistant.capabilities import capabilities


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def test_image_insert_resize_alt_text_and_remove_preserve_package(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "pixel.png"
    image.write_bytes(PNG_1X1)
    doc = Document()
    doc.add_paragraph("Before")
    doc.save(source)

    inserted = insert_image(
        source,
        output_path=tmp_path / "inserted.docx",
        image_path=image,
        width_cm=3,
        alt_text="Accessible pixel",
        title="Architecture image",
    )
    resized = resize_image(
        inserted.output_path,
        output_path=tmp_path / "resized.docx",
        image_index=0,
        width_cm=4,
        preserve_aspect_ratio=True,
    )

    with ZipFile(resized.output_path) as package:
        xml = package.read("word/document.xml").decode("utf-8")
        media = [name for name in package.namelist() if name.startswith("word/media/")]
    assert len(media) == 1
    assert 'descr="Accessible pixel"' in xml
    assert 'title="Architecture image"' in xml
    assert "1440000" in xml  # four centimetres in EMU

    removed = remove_image(
        resized.output_path,
        output_path=tmp_path / "removed.docx",
        image_index=0,
    )
    with ZipFile(removed.output_path) as package:
        xml = package.read("word/document.xml").decode("utf-8")
    assert "<w:drawing" not in xml


def test_image_single_target_accepts_heading_selector(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "pixel.png"
    image.write_bytes(PNG_1X1)
    document = Document()
    document.add_heading("Architecture", level=1)
    document.add_paragraph("Description")
    document.add_paragraph().add_run().add_picture(str(image))
    document.save(source)

    result = insert_image(
        source,
        output_path=tmp_path / "inserted.docx",
        image_path=image,
        selector={"type": "heading", "text": "Architecture"},
    )

    reopened = Document(result.output_path)
    assert "w:drawing" not in reopened.paragraphs[0]._p.xml
    assert "w:drawing" in reopened.paragraphs[1]._p.xml
    assert result.data["image_index"] == 0


def test_image_resize_capability_example_is_executable(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "pixel.png"
    image.write_bytes(PNG_1X1)
    document = Document()
    document.add_paragraph().add_run().add_picture(str(image))
    document.save(source)
    example = capabilities()["operation_contracts"]["image.resize"]["options_example"]

    result = resize_image(source, output_path=tmp_path / "resized.docx", **example)

    assert result.output_path.exists()
