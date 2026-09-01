import pytest
from docx import Document

from wordassistant.selectors import (
    Selector,
    SelectorAmbiguous,
    SelectorNotFound,
    resolve_paragraphs,
    resolve_single_paragraph,
)


def _document():
    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Opening text")
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Target result")
    doc.add_paragraph("Target result")
    return doc


def test_selector_from_json_and_heading_scope_resolves_bounded_paragraphs():
    selector = Selector.from_value('{"type":"heading","text":"Results"}')

    matches = resolve_paragraphs(_document(), selector)

    assert [paragraph.text for paragraph in matches] == [
        "Results",
        "Target result",
        "Target result",
    ]


def test_exact_text_selector_fails_closed_when_ambiguous():
    selector = Selector(type="text", text="Target result")

    with pytest.raises(SelectorAmbiguous) as exc:
        resolve_paragraphs(_document(), selector, require_single=True)

    assert exc.value.match_count == 2


def test_missing_selector_reports_structured_failure():
    with pytest.raises(SelectorNotFound) as exc:
        resolve_paragraphs(_document(), Selector(type="heading", text="Appendix"))

    assert exc.value.selector.type == "heading"


def test_duplicate_heading_selector_fails_closed():
    document = Document()
    document.add_heading("Results", level=1)
    document.add_paragraph("First")
    document.add_heading("Results", level=1)
    document.add_paragraph("Second")

    with pytest.raises(SelectorAmbiguous):
        resolve_paragraphs(document, {"type": "heading", "text": "Results"})


@pytest.mark.parametrize(
    "selector",
    [
        {"type": "heading"},
        {"type": "heading", "text": "Results", "index": 0},
        {"type": "paragraph", "index": "0"},
        {"type": "text", "text": "Result", "case_sensitive": "false"},
        {"type": "document", "text": "ignored"},
    ],
)
def test_selector_schema_is_enforced(selector):
    with pytest.raises((TypeError, ValueError)):
        Selector.from_value(selector)


def test_single_paragraph_resolution_uses_heading_itself_not_heading_scope():
    match = resolve_single_paragraph(
        _document(), Selector(type="heading", text="Results")
    )

    assert match.text == "Results"
