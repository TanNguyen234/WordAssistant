from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from wordassistant.operations.paragraph import format_paragraphs, format_runs
from wordassistant.operations.styles import create_style, apply_style, delete_style


def _source(path):
    doc = Document()
    doc.add_paragraph("Target")
    doc.save(path)
    return path


def test_run_and_paragraph_formatting_roundtrip(tmp_path):
    source = _source(tmp_path / "source.docx")
    run_result = format_runs(
        source,
        output_path=tmp_path / "runs.docx",
        selector={"type": "paragraph", "index": 0},
        font_family="Arial",
        font_size_pt=14,
        bold=True,
        italic=True,
        font_color="336699",
        small_caps=True,
    )
    format_paragraphs(
        run_result.output_path,
        output_path=tmp_path / "paragraph.docx",
        selector={"type": "paragraph", "index": 0},
        alignment="center",
        left_indent_cm=1.2,
        first_line_indent_cm=0.5,
        space_before_pt=3,
        space_after_pt=6,
        line_spacing=1.5,
        keep_with_next=True,
    )

    paragraph = Document(tmp_path / "paragraph.docx").paragraphs[0]
    run = paragraph.runs[0]
    assert run.font.name == "Arial"
    assert run.font.size.pt == 14
    assert run.bold is True and run.italic is True
    assert str(run.font.color.rgb) == "336699"
    assert run.font.small_caps is True
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert round(paragraph.paragraph_format.left_indent.cm, 1) == 1.2
    assert paragraph.paragraph_format.space_after.pt == 6
    assert paragraph.paragraph_format.keep_with_next is True


def test_custom_style_create_apply_and_safe_delete(tmp_path):
    source = _source(tmp_path / "source.docx")
    created = create_style(source, output_path=tmp_path / "styled.docx", name="Analysis Note", font_family="Calibri", font_size_pt=11, italic=True)
    applied = apply_style(created.output_path, output_path=tmp_path / "applied.docx", name="Analysis Note", selector={"type": "paragraph", "index": 0})

    paragraph = Document(applied.output_path).paragraphs[0]
    assert paragraph.style.name == "Analysis Note"

    blocked = delete_style(applied.output_path, output_path=tmp_path / "blocked.docx", name="Analysis Note")
    assert blocked.status == "UNSAFE_EDIT_BLOCKED"
    assert not (tmp_path / "blocked.docx").exists()

