from docx import Document

from wordassistant.operations.text import edit_text
from wordassistant.selectors import Selector


def _source(path):
    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("old1 outside")
    doc.add_heading("Results", level=1)
    p = doc.add_paragraph()
    p.add_run("old").bold = True
    p.add_run("2 and old3")
    doc.save(path)
    return path


def test_regex_replace_is_scoped_to_selected_heading_and_preserves_unmatched_runs(tmp_path):
    source = _source(tmp_path / "source.docx")
    output = tmp_path / "output.docx"

    result = edit_text(
        source,
        output_path=output,
        action="replace",
        pattern=r"old\d",
        replacement="new",
        regex=True,
        selector=Selector(type="heading", text="Results"),
    )

    reopened = Document(output)
    assert result.data["replacements"] == 2
    assert reopened.paragraphs[1].text == "old1 outside"
    assert reopened.paragraphs[3].text == "new and new"
    assert reopened.paragraphs[3].runs[0].bold is True


def test_whole_word_case_insensitive_replace_does_not_replace_substrings(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Cat cat category concatenate")
    doc.save(source)

    edit_text(
        source,
        output_path=tmp_path / "output.docx",
        action="replace",
        pattern="cat",
        replacement="dog",
        whole_word=True,
        case_sensitive=False,
    )

    assert Document(tmp_path / "output.docx").paragraphs[0].text == "dog dog category concatenate"


def test_insert_append_and_prepend_use_paragraph_selector(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("middle")
    doc.save(source)

    first = edit_text(source, output_path=tmp_path / "a.docx", action="prepend", text="before ", selector={"type": "paragraph", "index": 0})
    second = edit_text(first.output_path, output_path=tmp_path / "b.docx", action="append", text=" after", selector={"type": "paragraph", "index": 0})

    assert Document(second.output_path).paragraphs[0].text == "before middle after"

